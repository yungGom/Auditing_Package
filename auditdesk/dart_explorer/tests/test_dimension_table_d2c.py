"""D-2c 게이트: 차원표 편집기 3열 충전 — 금감원 배포 엑셀 조인.

G2: 렌더 산출의 표본 행 3열 값 ↔ 배포 엑셀 원본 독립 재파싱 일치.
확장요소는 '확장' 표시(빈칸 금지), 시트 상단 충전율 문구.
"""
import os
import random
import sys

import pytest
from openpyxl import load_workbook

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__))))
if os.path.join(_ROOT, "dsd_workbench") not in sys.path:
    sys.path.insert(0, os.path.join(_ROOT, "dsd_workbench"))

_PKG = os.path.join(_ROOT, "dart_explorer", "cache", "xbrl", "00136925",
                    "20260323000596_11011")
_ASSET_DIR = os.path.join(_ROOT, "assets", "taxonomy")

_ready = os.path.isdir(_PKG) and any(
    n.endswith(".xlsx") for n in
    (os.listdir(_ASSET_DIR) if os.path.isdir(_ASSET_DIR) else []))

pytestmark = pytest.mark.skipif(
    not _ready, reason="신원 패키지 또는 배포 엑셀 자산 없음")


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    from dart_explorer.xbrl.dimension_table import render_dimension_tables
    out = str(tmp_path_factory.mktemp("d2c") / "차원표.xlsx")
    return render_dimension_tables(_PKG, out_path=out)


def _raw_label_map():
    """배포 엑셀 독립 재파싱 — 리졸버를 거치지 않는 원본 대조 기준."""
    from dsd_tool.taxonomy_labels import default_asset_path
    wb = load_workbook(default_asset_path(), read_only=True)
    names, ko = {}, {}
    rows = wb["Concepts"].iter_rows(values_only=True)
    next(rows)
    for row in rows:
        if row and row[1] and row[2]:
            names[f"{row[1]}_{row[2]}"] = (str(row[1]), str(row[2]))
    for i, row in enumerate(wb["Label Link"].iter_rows(values_only=True), 1):
        if i <= 4 or not row or not row[1] or not row[2]:
            continue
        lab = row[3] or row[10]                 # label | terseLabel
        if lab:
            ko.setdefault(f"{row[1]}_{row[2]}", str(lab).strip())
    wb.close()
    return names, ko


def test_g2_sample_matches_distribution_xlsx(rendered):
    names, ko = _raw_label_map()
    wb = load_workbook(rendered["out_path"], read_only=False)
    samples = []                                # (sheet, row, ko, en, id)
    for ws in wb:
        # 3열 위치: 헤더 행2에서 '한글 표준레이블' 탐색
        col0 = next((c.column for c in ws[2]
                     if c.value == "한글 표준레이블"), None)
        if col0 is None:
            continue
        for row in ws.iter_rows(min_row=3):
            v = row[col0 - 1].value if len(row) >= col0 else None
            qn = row[col0 + 1].value if len(row) > col0 + 1 else None
            if v and qn and v != "확장":
                samples.append((ws.title, row[0].row, v,
                                row[col0].value, qn))
    assert len(samples) >= 500                  # 전 시트 충전 확인
    rng = random.Random(42)
    picked = rng.sample(samples, 10)
    for sheet, r, ko_v, en_v, qn in picked:
        cid = str(qn).replace(":", "_", 1)
        assert cid in names, f"{sheet}!R{r}: {qn} 배포 엑셀 미수록"
        prefix, name = names[cid]
        assert en_v == name, f"{sheet}!R{r}: 영문명 {en_v} != {name}"
        expect_ko = ko.get(cid) or name
        assert ko_v == expect_ko, \
            f"{sheet}!R{r}: 한글 {ko_v!r} != {expect_ko!r}"
    print(f"\n[G2] 표본 10행 3열 = 배포 엑셀 원본 일치 "
          f"(모집단 {len(samples):,}행)")


def test_extension_marked_not_blank(rendered):
    """확장요소 = '확장' 표시 — 미매칭≠0, 빈칸 금지."""
    wb = load_workbook(rendered["out_path"], read_only=False)
    n_ext = 0
    for ws in wb:
        col0 = next((c.column for c in ws[2]
                     if c.value == "한글 표준레이블"), None)
        if col0 is None:
            continue
        for row in ws.iter_rows(min_row=3):
            if len(row) > col0 + 1 and row[col0 + 1].value and \
                    str(row[col0 + 1].value).startswith("entity"):
                # 회사 확장요소 행 — 한글레이블 열은 반드시 '확장'
                assert row[col0 - 1].value == "확장", \
                    f"{ws.title}!R{row[0].row}"
                n_ext += 1
    assert n_ext >= 50                          # 신원 실측 확장 367
    print(f"\n[G2] 확장 표시 행 {n_ext} — 빈칸 0")


def test_fill_rate_line_and_summary(rendered):
    wb = load_workbook(rendered["out_path"], read_only=True)
    ws = wb["BS"]
    line = ws.cell(2, 1).value
    assert line and "편집기 열 충전율" in line and "표준" in line
    fills = [r.get("label_fill") for r in rendered["roles"]]
    assert all(f and (f["std"] + f["ext"]) > 0 for f in fills)

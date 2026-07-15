"""D-1 게이트: 택사노미 트리 뷰.

- entity00136925 패키지: Role 트리 정확 재구성, pre.xml 아크 133개 전수 반영
- 확장 element 빨간 표시: entity00136925(3개) + 삼성전자 B-3 캐시(598개)
- 금감원 택사노미 xlsm도 같은 뷰어로 열람
"""
import glob
import os

import pytest
from openpyxl import load_workbook

from dart_explorer.xbrl.taxonomy import (TaxonomyPackage, build_tree_view,
                                         fss_xlsm_role_rows)

_XBRL_DIR = r"C:\Users\moonyong\OneDrive - 서현회계법인\바탕 화면\XBRL"
_SHINWON = os.path.join(
    _XBRL_DIR, "신원종합개발",
    "[신원종합개발]반기보고서_IFRS(원문XBRL)(2025.08.14)")
_FSS_XLSM = os.path.join(
    _XBRL_DIR, "Taxonomy_매크로",
    "금융감독원 DART XBRL Taxonomy Search_배포용(20251513).xlsm")
_SAMSUNG_PKG = os.path.join(
    os.path.dirname(__file__), "..", "cache", "xbrl", "00126380",
    "20260310002820_11011")


# ---------------------------------------------------------------------------
# 게이트 1: entity00136925 — Role 트리 정확 재구성 (아크 133개 전수)
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not os.path.isdir(_SHINWON), reason="신원 패키지 없음")
def test_gate_shinwon_tree(tmp_path):
    pkg = TaxonomyPackage(_SHINWON)
    roles = pkg.roles()
    codes = sorted(d.split("]")[0].strip("[") for _, d, _ in roles)
    assert codes == ["D210005", "D431415", "D520005", "D610005"]
    assert pkg.arc_count() == 133              # 스펙 명시 수치

    role_rows = pkg.all_role_rows()
    total_rows = sum(len(r) for _, r in role_rows)
    # 전수 반영: 행 수 = 아크 133 + Role별 루트 각 1개
    assert total_rows == 133 + len(roles) == 137

    # 확장 element (실측 3개 — 스펙 기재 0개와 달리 존재)
    ext = [r for _, rows in role_rows for r in rows if r["ext"]]
    assert len(ext) == 3
    assert all(r["prefix"] == "entity00136925" for r in ext)

    # 트리 내용 스팟체크: BS 루트와 유동자산 계층
    bs = next(rows for d, rows in role_rows if "D210005" in d)
    assert bs[0]["depth"] == 0 and "재무상태표" in bs[0]["ko"]
    cur = next(r for r in bs if r["id"] == "CurrentAssets")
    assert cur["depth"] >= 2 and cur["ko"]     # 들여쓰기 + 한글 라벨

    out = build_tree_view(_SHINWON, str(tmp_path / "신원.xlsx"))
    assert out["roles"] == 4 and out["arcs"] == 133
    wb = load_workbook(out["out_path"])
    assert len(wb.sheetnames) == 4
    # 빨간 글씨 확인
    reds = [c.value for ws in wb.worksheets for row in ws.iter_rows()
            for c in row
            if c.font and c.font.color and c.font.color.rgb and
            str(c.font.color.rgb).endswith("CC0000") and c.column == 4]
    assert len(reds) == 3


# ---------------------------------------------------------------------------
# 게이트 2: 삼성전자 (B-3 캐시) — 확장 element 표시
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not os.path.isdir(_SAMSUNG_PKG),
                    reason="삼성 XBRL 캐시 없음 (B-3 파이프라인 선실행 필요)")
def test_gate_samsung_extensions(tmp_path):
    pkg = TaxonomyPackage(_SAMSUNG_PKG)
    assert len(pkg.extensions) >= 500          # 실측 598개 확장 선언
    out = build_tree_view(_SAMSUNG_PKG, str(tmp_path / "삼성.xlsx"))
    assert out["roles"] >= 50
    assert out["extensions"] >= 1000           # 트리 위치 기준 빨간 행


# ---------------------------------------------------------------------------
# 금감원 택사노미 xlsm — 같은 뷰어로 열람 (D-1-5)
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not os.path.exists(_FSS_XLSM), reason="금감원 xlsm 없음")
def test_fss_xlsm_viewer(tmp_path):
    role_rows = fss_xlsm_role_rows(_FSS_XLSM, role_filter="D210000")
    assert role_rows
    definition, rows = role_rows[0]
    assert "D210000" in definition
    assert any(r["id"] == "CurrentAssets" for r in rows)
    assert all(r["ext"] is False for r in rows)   # 표준 택사노미
    # depth 컬럼 그대로 들여쓰기
    root = rows[0]
    assert root["depth"] == 0
    out = build_tree_view(_FSS_XLSM, str(tmp_path / "fss.xlsx"),
                          role_filter="D210000")
    assert out["rows"] == sum(len(r) for _, r in role_rows)

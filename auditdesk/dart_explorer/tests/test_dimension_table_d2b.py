"""D-2b 게이트: 차원 표 렌더러 보완 — 육안 검증 피드백 2건.

1-1. 기간 블록 라벨: 실제 컨텍스트 일자로 당기/전기 + 3개월/누적 판정
1-2. 시트명: DSD 편집용 엑셀과 동일 체계 (BS/PL/PL1/CE/CF, 주석 1,2,3...)
"""
import datetime
import os

import pytest
from openpyxl import load_workbook

from dart_explorer.xbrl.dimension_table import (HypercubeDef, RoleTable,
                                                XbrlInstance, _classify_period,
                                                _duration_desc, _fs_abbrev,
                                                _is_note_role,
                                                render_dimension_tables)
from dart_explorer.xbrl.taxonomy import TaxonomyPackage

_SHINWON = os.path.join(
    r"C:\Users\moonyong\OneDrive - 서현회계법인\바탕 화면\XBRL",
    "신원종합개발", "[신원종합개발]반기보고서_IFRS(원문XBRL)(2025.08.14)")
_SAMSUNG = os.path.join(
    os.path.dirname(__file__), "..", "cache", "xbrl", "00126380",
    "20260310002820_11011")


# ---------------------------------------------------------------------------
# 1-1. 기간 분류 단위 테스트
# ---------------------------------------------------------------------------

def test_duration_desc_rules():
    assert _duration_desc(datetime.date(2025, 4, 1),
                          datetime.date(2025, 6, 30)) == "3개월"
    assert _duration_desc(datetime.date(2025, 1, 1),
                          datetime.date(2025, 6, 30)) == "누적"
    assert _duration_desc(datetime.date(2025, 1, 1),
                          datetime.date(2025, 9, 30)) == "9개월 누적"
    assert _duration_desc(datetime.date(2025, 1, 1),
                          datetime.date(2025, 12, 31)) is None   # 연간


def test_classify_period_duration_and_instant():
    ref = datetime.date(2025, 6, 30)
    p = _classify_period(ref, ("duration", "2025-04-01", "2025-06-30"))
    assert p == {"name": "당기", "desc": "3개월", "kind": "duration"}
    p2 = _classify_period(ref, ("duration", "2024-01-01", "2024-06-30"))
    assert p2["name"] == "전기" and p2["desc"] == "누적"
    p3 = _classify_period(ref, ("instant", "2024-01-01", "2024-01-01"))
    assert p3["name"] == "전기" and p3["desc"] == "초"
    p4 = _classify_period(ref, ("instant", "2025-06-30", "2025-06-30"))
    assert p4["name"] == "당기" and p4["desc"] == "말"


def test_fs_abbrev_keywords():
    assert _fs_abbrev("[D210000] 재무상태표, 유동/비유동법 - 연결") == "BS"
    assert _fs_abbrev("[D310000] 손익계산서, 기능별 분류 - 연결") == "PL"
    assert _fs_abbrev("[D410000] 포괄손익계산서, 세후 - 연결") == "PL1"
    assert _fs_abbrev("[D610000] 자본변동표 - 연결") == "CE"
    assert _fs_abbrev("[D520000] 현금흐름표, 간접법 - 연결") == "CF"
    assert _fs_abbrev("[D822380] 28. 재무위험관리") is None


def test_is_note_role():
    assert _is_note_role("http://x/role-D822380") is True
    assert _is_note_role("http://x/role-D210000") is False


# ---------------------------------------------------------------------------
# 게이트 1: 신원(반기) 재렌더 — 당기 3개월/누적 + 전기 3개월/누적 + 일자 병기
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not os.path.isdir(_SHINWON), reason="신원 패키지 없음")
def test_gate_shinwon_half_year_labels(tmp_path):
    res = render_dimension_tables(_SHINWON, str(tmp_path / "신원.xlsx"))
    wb = load_workbook(res["out_path"])

    # 1-2 부수 확인: 단일 포괄손익계산서만 있으므로 PL(=PL1 아님), 시트명 체계
    assert set(wb.sheetnames) == {"BS", "PL", "CF", "CE"}

    ce = wb["CE"]
    labels = [ce.cell(r, 1).value for r in range(1, ce.max_row + 1)
             if ce.cell(r, 1).value]
    assert any("당기 누적 (2025-01-01 ~ 2025-06-30)" in s for s in labels)
    assert any("전기 누적 (2024-01-01 ~ 2024-06-30)" in s for s in labels)

    pl = wb["PL"]
    header = [pl.cell(3, c).value for c in range(2, 6)]
    assert header[0] == "당기 누적\n2025-01-01~2025-06-30"
    assert header[1] == "당기 3개월\n2025-04-01~2025-06-30"
    assert header[2] == "전기 누적\n2024-01-01~2024-06-30"
    assert header[3] == "전기 3개월\n2024-04-01~2024-06-30"

    # role 코드 + 정의 전문은 1행에 그대로 보존
    assert pl.cell(1, 1).value.startswith("[D431415]")


# ---------------------------------------------------------------------------
# 게이트 2: 삼성전자(연간) 재렌더 — 시트명 BS/PL/.../1,2,3.. + role 코드 보존
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not os.path.isdir(_SAMSUNG), reason="삼성 캐시 없음")
def test_gate_samsung_sheet_naming(tmp_path):
    res = render_dimension_tables(_SAMSUNG, str(tmp_path / "삼성.xlsx"))
    wb = load_workbook(res["out_path"])
    names = set(wb.sheetnames)

    # 연결이 기본, 별도는 접미사 — 5개 재무제표 계열 전부
    for base in ("BS", "PL", "PL1", "CE", "CF"):
        assert base in names, base
        assert f"{base}_별도" in names, base

    # 주석은 role 코드 오름차순 순번 (숫자 시트명)
    note_sheets = [n for n in wb.sheetnames if n.isdigit()]
    assert len(note_sheets) >= 20
    nums = sorted(int(n) for n in note_sheets)
    assert nums == list(range(1, len(nums) + 1))   # 연속 결번 없음

    # 추적성: 각 주석 시트 1행에 원래 role 코드 보존
    for n in note_sheets[:5]:
        cell = wb[n].cell(1, 1).value
        assert cell.startswith("[D8"), (n, cell)

    # 연간 보고서는 3개월/누적 접미사 없음 — 당기/전기만
    bs = wb["BS"]
    header = [bs.cell(3, c).value for c in range(2, 4)]
    assert "당기말" in header[0] and "3개월" not in header[0]
    assert "전기말" in header[1]


@pytest.mark.skipif(not os.path.isdir(_SAMSUNG), reason="삼성 캐시 없음")
def test_gate_samsung_annual_pl_duration_no_suffix(tmp_path):
    res = render_dimension_tables(_SAMSUNG, str(tmp_path / "삼성.xlsx"),
                                  role_filter="D310000")
    wb = load_workbook(res["out_path"])
    pl = wb["PL"]
    header = [pl.cell(3, c).value for c in range(2, 4)]
    assert header[0].startswith("당기\n") and "개월" not in header[0]
    assert header[1].startswith("전기\n")

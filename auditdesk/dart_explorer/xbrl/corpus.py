"""D-3a: 계정과목 매핑 코퍼스 구축 (전업종, 주기 실행).

상장사 전체 목록(corpCode) → 사업보고서 XBRL 일괄 수신(B-3 재사용, 영구 캐시)
→ 인스턴스별 element 사용 실태를 mapping_corpus.sqlite에 적재.

- usages: (corp_code, element_id, 회사 라벨(lab-ko), Role 코드들, 차원 문맥,
  팩트 수, 확장 여부) — 사용 회사수는 조회 시 집계
- extensions: 확장 element 별도 테이블 (신규 계정과목 정의 참고자료)
- standard_labels: 금감원 택사노미 xlsm에서 표준 한/영 라벨 적재
- companies: 처리 상태(ok/no_annual/no_xbrl/error) — 재개(resume) 지원
- 요청 한도 보호: 캐시 히트 시 재다운로드 없음 (일 20,000건)

LLM/외부 AI 불사용 — 실증 빈도 기반 (보안·재현성·감사조서 근거).
"""
import datetime
import os
import sqlite3

from lxml import etree

from ..client.opendart import OpenDartClient, OpenDartError
from .pipeline import (REPRT, XbrlNotAvailable, download_xbrl,
                       find_periodic_rcept, unpack_xbrl)
from .taxonomy import LB, XL, TaxonomyPackage, _pkg_glob
from .dimension_table import XbrlInstance

DEFAULT_DB = os.path.join(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__))), "corpus", "mapping_corpus.sqlite")

_SCHEMA = """
CREATE TABLE IF NOT EXISTS companies(
  corp_code TEXT PRIMARY KEY, corp_name TEXT, stock_code TEXT,
  induty_code TEXT, year INTEGER, rcept_no TEXT, report_nm TEXT,
  status TEXT NOT NULL, error TEXT, fetched_at TEXT);
CREATE TABLE IF NOT EXISTS usages(
  corp_code TEXT NOT NULL, element_id TEXT NOT NULL, is_ext INTEGER,
  label_ko TEXT, roles TEXT, dims TEXT, n_facts INTEGER,
  PRIMARY KEY(corp_code, element_id));
CREATE TABLE IF NOT EXISTS extensions(
  corp_code TEXT NOT NULL, element_id TEXT NOT NULL,
  label_ko TEXT, roles TEXT, dims TEXT, n_facts INTEGER,
  PRIMARY KEY(corp_code, element_id));
CREATE TABLE IF NOT EXISTS standard_labels(
  element_id TEXT PRIMARY KEY, label_ko TEXT, label_en TEXT);
CREATE INDEX IF NOT EXISTS idx_usages_element ON usages(element_id);
CREATE INDEX IF NOT EXISTS idx_usages_label ON usages(label_ko);
"""


def connect(db_path=None):
    path = db_path or DEFAULT_DB
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    con = sqlite3.connect(path)
    con.executescript(_SCHEMA)
    # D-4a: standard_labels에 택사노미 세대 태그 (기존 DB 마이그레이션)
    try:
        con.execute("ALTER TABLE standard_labels ADD COLUMN version TEXT")
    except sqlite3.OperationalError:
        pass
    return con


# ---------------------------------------------------------------------------
# 패키지 1개 → 사용 실태 추출
# ---------------------------------------------------------------------------

def _element_roles(folder):
    """pre.xml에서 element → {role 코드} (loc 기반, 트리 순회 없이)."""
    out = {}
    for path in _pkg_glob(folder, "*_pre.xml"):
        root = etree.parse(path).getroot()
        for pl in root.findall(f"{{{LB}}}presentationLink"):
            role = pl.get(f"{{{XL}}}role") or ""
            code = role.rsplit("role-", 1)[-1]
            for loc in pl.findall(f"{{{LB}}}loc"):
                concept = (loc.get(f"{{{XL}}}href") or "").split("#")[-1]
                out.setdefault(concept, set()).add(code)
    return out


def extract_package(folder):
    """XBRL 패키지 → [usage dict]. (표준+확장, 팩트가 있는 element만)"""
    pkg = TaxonomyPackage(folder)
    inst = XbrlInstance(folder)
    roles = _element_roles(folder)
    rows = []
    for element_id, facts in inst.facts.items():
        dims = sorted({"|".join(sorted(f["ctx"]["dims"])) or "(없음)"
                       for f in facts})
        rows.append({
            "element_id": element_id,
            "is_ext": int(element_id in pkg.extensions),
            "label_ko": pkg._label(pkg.labels_ko, element_id),
            "roles": ",".join(sorted(roles.get(element_id, set()))),
            "dims": ";".join(dims),
            "n_facts": len(facts),
        })
    return rows


# ---------------------------------------------------------------------------
# 배치 빌드 (재개 가능)
# ---------------------------------------------------------------------------

def build(year, limit=None, db_path=None, cli=None, listed_only=True,
          progress=print):
    """상장사 전체(또는 limit)를 순회하며 코퍼스 적재. 요약 dict 반환."""
    cli = cli or OpenDartClient()
    con = connect(db_path)
    corps = [c for c in cli.corp_codes()
             if (c["stock_code"] if listed_only else True)]
    corps.sort(key=lambda c: c["corp_code"])
    done = {r[0] for r in con.execute(
        "SELECT corp_code FROM companies WHERE year = ?", (year,))}
    todo = [c for c in corps if c["corp_code"] not in done]
    if limit is not None:
        # limit = 목표 총 회사 수 (기존 처리분 포함) — 재개 시 잔여분만
        todo = todo[:max(0, limit - len(done))]
    counts = {"ok": 0, "no_annual": 0, "no_xbrl": 0, "error": 0}

    for i, corp in enumerate(todo, 1):
        code = corp["corp_code"]
        status, err, rcept, report_nm, induty = "ok", None, None, None, None
        try:
            doc = find_periodic_rcept(cli, code, year)
            if doc is None:
                status = "no_annual"
            else:
                rcept, report_nm = doc["rcept_no"], doc.get("report_nm")
                _, zpath = download_xbrl(cli, code, rcept, REPRT["annual"])
                folder = unpack_xbrl(zpath)
                rows = extract_package(folder)
                try:
                    induty = cli.company(code).get("induty_code")
                except OpenDartError:
                    induty = None
                con.execute("DELETE FROM usages WHERE corp_code = ?", (code,))
                con.execute("DELETE FROM extensions WHERE corp_code = ?",
                            (code,))
                for r in rows:
                    con.execute(
                        "INSERT OR REPLACE INTO usages VALUES(?,?,?,?,?,?,?)",
                        (code, r["element_id"], r["is_ext"], r["label_ko"],
                         r["roles"], r["dims"], r["n_facts"]))
                    if r["is_ext"]:
                        con.execute(
                            "INSERT OR REPLACE INTO extensions "
                            "VALUES(?,?,?,?,?,?)",
                            (code, r["element_id"], r["label_ko"],
                             r["roles"], r["dims"], r["n_facts"]))
        except XbrlNotAvailable:
            status = "no_xbrl"
        except OpenDartError as e:
            if "020" in str(e):                    # 일 한도 초과 → 중단
                progress(f"요청 한도 도달 — 중단 (재실행 시 이어서 진행): {e}")
                break
            status, err = "error", str(e)[:300]
        except Exception as e:                     # 손상 패키지 등
            status, err = "error", f"{type(e).__name__}: {e}"[:300]
        counts[status] += 1
        con.execute(
            "INSERT OR REPLACE INTO companies "
            "VALUES(?,?,?,?,?,?,?,?,?,?)",
            (code, corp["corp_name"], corp["stock_code"], induty, year,
             rcept, report_nm, status, err,
             datetime.datetime.now().isoformat(timespec="seconds")))
        con.commit()
        if i % 10 == 0 or i == len(todo):
            progress(f"  [{i}/{len(todo)}] {corp['corp_name']} → {status} "
                     f"(누적 ok {counts['ok']})")
    con.close()
    return counts


# ---------------------------------------------------------------------------
# 표준 라벨 적재 (금감원 xlsm)
# ---------------------------------------------------------------------------

def load_standard_labels(xlsm_path, db_path=None, version=None):
    """금감원 xlsm 표준 라벨 적재. version 미지정 시 systemid 열에서 자동 감지
    (dart_all_YYYY-MM-DD_pre 기준일 — D-4a). 기존 element는 갱신하지 않는다
    (INSERT OR IGNORE: 최초 적재 세대로 고정, version 열은 그 세대의 감사조서
    근거로만 사용)."""
    from .taxonomy import fss_xlsm_role_rows
    if version is None:
        from .taxonomy_diff import detect_version
        version = detect_version(xlsm_path)
    con = connect(db_path)
    for _, rows in fss_xlsm_role_rows(xlsm_path):
        for r in rows:
            eid = f"{r['prefix']}_{r['id']}" if r["prefix"] else r["id"]
            con.execute(
                "INSERT OR IGNORE INTO standard_labels VALUES(?,?,?,?)",
                (eid, r["ko_std"], r["en_std"], version))
    con.commit()
    n = con.execute("SELECT COUNT(*) FROM standard_labels").fetchone()[0]
    by_version = dict(con.execute(
        "SELECT version, COUNT(*) FROM standard_labels GROUP BY version"))
    con.close()
    return {"total": n, "version": version, "by_version": by_version}


def export_references(xlsm_path, out_json=None):
    """금감원 xlsm Reference Link → element별 기준서 참조 JSON.

    (전량 빌드 중 sqlite 쓰기 경합을 피하려고 DB가 아닌 JSON으로 저장 —
    D-3b 조회기가 읽는다.)
    """
    import json
    from openpyxl import load_workbook
    wb = load_workbook(xlsm_path, read_only=True)
    ws = wb["Reference Link"]
    refs = {}
    for row in ws.iter_rows(values_only=True):
        prefix, name = row[1], row[2]
        if not name or prefix in (None, "prefix"):
            continue
        parts = [str(v).strip() for v in (row[6], row[7], row[8])
                 if v is not None and str(v).strip()]
        if not parts:
            continue
        eid = f"{prefix}_{name}"
        ref = " ".join(parts)
        refs.setdefault(eid, [])
        if ref not in refs[eid]:
            refs[eid].append(ref)
    wb.close()
    if out_json is None:
        out_json = os.path.join(os.path.dirname(DEFAULT_DB),
                                "standard_refs.json")
    os.makedirs(os.path.dirname(out_json), exist_ok=True)
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(refs, f, ensure_ascii=False)
    return out_json, len(refs)


def export_note_assets(xlsx_path, out_dir=None):
    """F-2용 표준 자산 추출 → JSON 2종 (sqlite 잠금 없는 파일 교환).

    - standard_roles.json: 주석(D8xxxxx) role 정의 목록
      [{code, sector, consol, definition}]
    - standard_hypercubes.json: role 코드 → 축·member
      {code: [{axis, axis_label, members: [[id, label], ...]}]}
    열 배치는 헤더 행에서 탐지 (D-4b 교훈 — 위치 하드코딩 금지).
    """
    import json
    import re as _re
    from openpyxl import load_workbook

    out_dir = out_dir or os.path.dirname(DEFAULT_DB)
    os.makedirs(out_dir, exist_ok=True)
    wb = load_workbook(xlsx_path, read_only=True)

    # --- Role 일람표 → 주석 role 정의 ---------------------------------------
    roles = []
    ws = wb["Role 일람표"]
    for row in ws.iter_rows(values_only=True):
        cells = [str(v) if v is not None else "" for v in row]
        definition = next((c for c in cells
                           if _re.match(r"^\[D\d{6}", c)), None)
        if not definition:
            continue
        code = _re.match(r"^\[(D\d{6})", definition).group(1)
        sector = next((c for c in cells if c in
                       ("비금융업", "금융업", "보험업", "증권업")), "")
        consol = "별도" if "별도" in " ".join(cells) else "연결"
        roles.append({"code": code, "sector": sector, "consol": consol,
                      "definition": definition})
    with open(os.path.join(out_dir, "standard_roles.json"), "w",
              encoding="utf-8") as f:
        json.dump(roles, f, ensure_ascii=False)

    # --- Definition Link → role별 하이퍼큐브 축·member ----------------------
    cubes = {}
    ws = wb["Definition Link"]
    code = None
    col = {}
    cur_axes = []
    for row in ws.iter_rows(values_only=True):
        if row[0] == "LinkRole":
            if code and cur_axes:
                cubes.setdefault(code, []).extend(cur_axes)
            uri = str(row[1] or "")
            m = _re.search(r"role[-/](D\d{6})", uri) or \
                _re.search(r"(D\d{6})", uri)
            code = m.group(1) if m else None
            cur_axes = []
            col = {}
            continue
        if row[0] == "prefix" and "arcrole" in row:
            col = {name: i for i, name in enumerate(row) if name}
            continue
        if not code or not col or row[col.get("name", 1)] is None:
            continue
        prefix = row[col["prefix"]] or ""
        name = row[col["name"]]
        label = row[col.get("label", 2)] or ""
        arc = str(row[col["arcrole"]] or "")
        eid = f"{prefix}_{name}"
        if arc.endswith("hypercube-dimension"):
            cur_axes.append({"axis": eid, "axis_label": label,
                             "members": []})
        elif ("domain-member" in arc or "dimension-domain" in arc) \
                and cur_axes:
            cur_axes[-1]["members"].append([eid, label])
    if code and cur_axes:
        cubes.setdefault(code, []).extend(cur_axes)
    wb.close()
    with open(os.path.join(out_dir, "standard_hypercubes.json"), "w",
              encoding="utf-8") as f:
        json.dump(cubes, f, ensure_ascii=False)
    return {"roles": len(roles), "hypercube_roles": len(cubes)}


# ---------------------------------------------------------------------------
# 통계
# ---------------------------------------------------------------------------

def stats(db_path=None):
    con = connect(db_path)
    out = {}
    out["companies"] = dict(con.execute(
        "SELECT status, COUNT(*) FROM companies GROUP BY status"))
    out["usages"] = con.execute("SELECT COUNT(*) FROM usages").fetchone()[0]
    out["distinct_elements"] = con.execute(
        "SELECT COUNT(DISTINCT element_id) FROM usages").fetchone()[0]
    out["extensions"] = con.execute(
        "SELECT COUNT(*) FROM extensions").fetchone()[0]
    out["standard_labels"] = con.execute(
        "SELECT COUNT(*) FROM standard_labels").fetchone()[0]
    out["top_elements"] = con.execute(
        "SELECT element_id, COUNT(DISTINCT corp_code) AS n, MAX(label_ko) "
        "FROM usages WHERE is_ext = 0 GROUP BY element_id "
        "ORDER BY n DESC LIMIT 10").fetchall()
    con.close()
    return out


# ---------------------------------------------------------------------------
# F-3: 승계 자산 내보내기 (자기 기말 인스턴스 → JSON, 파일 교환 경계)
# ---------------------------------------------------------------------------

def export_succession_assets(folder, out_json=None, against=None,
                             progress=None):
    """자기 기말 XBRL 패키지 → 승계 자산 JSON (dsd_workbench가 읽음).

    - elements: 팩트 있는 element 전수 (extract_package 스키마 —
      element 선택·확장 정의·role 구성 그대로 승계의 원천)
    - roles: pre.xml role 정의 목록 [{code, definition}]
    - members: 인스턴스 문맥 explicitMember 전수 [{id, label_ko}]
    - taxcheck: against(신버전 세대) 지정 시 D-4c 상태 —
      {element_id: {status: 녹색|노랑|파랑|확장, detail}} (폐지 대체
      후보는 dsd_workbench가 D-3b로 산출)
    """
    import glob as _glob
    import json
    import re as _re

    pkg = TaxonomyPackage(folder)
    rows = extract_package(folder)

    roles = []
    for _uri, definition, _pl in pkg.roles():
        m = _re.match(r"^\[(D\d{6})\]", definition or "")
        if m:
            roles.append({"code": m.group(1), "definition": definition})

    xbrldi = "http://xbrl.org/2006/xbrldi"
    used_members = set()
    for path in _glob.glob(os.path.join(folder, "*.xbrl")):
        root = etree.parse(path).getroot()
        for m in root.iter(f"{{{xbrldi}}}explicitMember"):
            used_members.add((m.text or "").strip().replace(":", "_"))
    members = [{"id": mid, "label_ko": pkg._label(pkg.labels_ko, mid)}
               for mid in sorted(used_members)]

    taxcheck = {}
    if against:
        from .taxonomy_diff import (parse_presentation_concepts,
                                    resolve_version_dir)
        new_concepts = parse_presentation_concepts(
            resolve_version_dir(against), progress)
        for r in rows:
            eid = r["element_id"]
            if eid.startswith("dart-gcd"):
                continue
            if r["is_ext"]:
                taxcheck[eid] = {"status": "확장",
                                 "detail": "당사 확장 element — 승계 유지"}
            elif eid in new_concepts:
                nl = new_concepts[eid]["label_ko"]
                if (nl or "").strip() == (r["label_ko"] or "").strip():
                    taxcheck[eid] = {"status": "녹색",
                                     "detail": "그대로 사용 가능"}
                else:
                    taxcheck[eid] = {
                        "status": "파랑",
                        "detail": f"라벨 변경: '{r['label_ko']}' → '{nl}'"}
            else:
                taxcheck[eid] = {"status": "노랑",
                                 "detail": "신버전에서 폐지 — 대체 후보 확인"}

    data = {"source": os.path.abspath(folder), "against": against,
            "elements": rows, "roles": roles, "members": members,
            "taxcheck": taxcheck}
    if out_json is None:
        out_json = os.path.join(folder, "succession_assets.json")
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    if progress:
        progress(f"승계 자산: elements {len(rows)} / roles {len(roles)} / "
                 f"members {len(members)} / taxcheck {len(taxcheck)}")
    return {"out_json": out_json, "n_elements": len(rows),
            "n_roles": len(roles), "n_members": len(members),
            "taxcheck": {s: sum(1 for t in taxcheck.values()
                                if t["status"] == s)
                         for s in ("녹색", "노랑", "파랑", "확장")}}

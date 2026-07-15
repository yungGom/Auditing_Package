"""D-1: 택사노미 트리 뷰 (이미지2 형식).

- XBRL 패키지(인스턴스+pre+lab-ko/en+xsd) → Role별 presentation 트리 재구성
  (parentChildArc order 정렬, preferredLabel 반영, 회사 확장 element 빨간 표시)
- 금감원 택사노미 xlsm(Presentation Link 시트)도 같은 형식으로 열람
  → 표준 트리 vs 회사 인스턴스 트리 비교의 기초

출력 열 (이미지2 그대로):
표시 한글명(들여쓰기 트리) | Prefix | 표시 영어명 | ID | 기본 한글명 | 기본 영어명
"""
import glob
import os
import re

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

LB = "http://www.xbrl.org/2003/linkbase"
XL = "http://www.w3.org/1999/xlink"
XS = "http://www.w3.org/2001/XMLSchema"
STD_LABEL = "http://www.xbrl.org/2003/role/label"

_RED = Font(color="FFCC0000")
_BOLD = Font(bold=True)
_HDR_FILL = PatternFill("solid", start_color="D9E1F2")

HEADER = ["표시 한글명", "Prefix", "표시 영어명", "ID",
          "기본 한글명", "기본 영어명"]


def _pkg_glob(folder, pattern):
    return sorted(glob.glob(os.path.join(glob.escape(folder), pattern)))


def _concept_of(href):
    return href.split("#")[-1]


def _split_id(concept_id):
    """'ifrs-full_Assets' → ('ifrs-full', 'Assets')"""
    if "_" in concept_id:
        p, local = concept_id.split("_", 1)
        return p, local
    return "", concept_id


class TaxonomyPackage:
    """XBRL 패키지 폴더 (B-3 해제 산출물 또는 수동 패키지)."""

    def __init__(self, folder):
        self.folder = folder
        pre = _pkg_glob(folder, "*_pre.xml")
        if not pre:
            raise FileNotFoundError(f"pre.xml 없음: {folder}")
        self.pre_root = etree.parse(pre[0]).getroot()
        self.labels_ko = self._load_labels(_pkg_glob(folder, "*_lab-ko.xml"))
        self.labels_en = self._load_labels(_pkg_glob(folder, "*_lab-en.xml"))
        self.role_defs = {}
        self.extensions = set()
        for xsd in _pkg_glob(folder, "*.xsd"):
            root = etree.parse(xsd).getroot()
            for rt in root.iter(f"{{{LB}}}roleType"):
                d = rt.find(f"{{{LB}}}definition")
                self.role_defs[rt.get("roleURI")] = \
                    (d.text or "").strip() if d is not None else ""
            for el in root.iter(f"{{{XS}}}element"):
                if el.get("id"):
                    self.extensions.add(el.get("id"))

    @staticmethod
    def _load_labels(paths):
        """{concept_id: {label_role_uri: text}}"""
        out = {}
        for path in paths:
            root = etree.parse(path).getroot()
            loc_map, arcs, labels = {}, {}, {}
            for loc in root.iter(f"{{{LB}}}loc"):
                loc_map[loc.get(f"{{{XL}}}label")] = \
                    _concept_of(loc.get(f"{{{XL}}}href") or "")
            for arc in root.iter(f"{{{LB}}}labelArc"):
                arcs.setdefault(arc.get(f"{{{XL}}}from"), []).append(
                    arc.get(f"{{{XL}}}to"))
            for lab in root.iter(f"{{{LB}}}label"):
                labels.setdefault(lab.get(f"{{{XL}}}label"), []).append(
                    (lab.get(f"{{{XL}}}role") or STD_LABEL, lab.text or ""))
            for lloc, concept in loc_map.items():
                for to in arcs.get(lloc, []):
                    for role, text in labels.get(to, []):
                        out.setdefault(concept, {})[role] = text
        return out

    def _label(self, table, concept, preferred=None):
        d = table.get(concept, {})
        if preferred and preferred in d:
            return d[preferred]
        if STD_LABEL in d:
            return d[STD_LABEL]
        return next(iter(d.values()), _split_id(concept)[1])

    # ------------------------------------------------------------------
    def roles(self):
        """[(roleURI, definition, presentationLink element)]"""
        out = []
        for pl in self.pre_root.findall(f"{{{LB}}}presentationLink"):
            uri = pl.get(f"{{{XL}}}role")
            out.append((uri, self.role_defs.get(uri, uri), pl))
        return out

    def arc_count(self):
        return sum(len(pl.findall(f"{{{LB}}}presentationArc"))
                   for _, _, pl in self.roles())

    def tree_rows(self, pl):
        """presentationLink → 트리 순회 행 목록 (아크 전수 반영)."""
        loc_map = {}
        for loc in pl.findall(f"{{{LB}}}loc"):
            loc_map[loc.get(f"{{{XL}}}label")] = \
                _concept_of(loc.get(f"{{{XL}}}href") or "")
        children = {}
        has_parent = set()
        seen_arcs = set()
        for arc in pl.findall(f"{{{LB}}}presentationArc"):
            src = loc_map.get(arc.get(f"{{{XL}}}from"))
            dst = loc_map.get(arc.get(f"{{{XL}}}to"))
            if src is None or dst is None:
                continue
            pref = arc.get("preferredLabel")
            # 동일 (부모,자식,표시라벨) 중복 아크는 1회만 (서브트리 중복 방지).
            # preferredLabel이 다른 중복(기초/기말 등)은 정당한 표시 행.
            if (src, dst, pref) in seen_arcs:
                continue
            seen_arcs.add((src, dst, pref))
            order = float(arc.get("order", "0") or 0)
            children.setdefault(src, []).append((order, dst, pref))
            has_parent.add(dst)
        roots = [c for c in dict.fromkeys(loc_map.values())
                 if c not in has_parent]
        rows = []

        def walk(concept, depth, preferred, path):
            p, local = _split_id(concept)
            rows.append({
                "depth": depth,
                "ko": self._label(self.labels_ko, concept, preferred),
                "prefix": p,
                "en": self._label(self.labels_en, concept, preferred),
                "id": local,
                "ko_std": self._label(self.labels_ko, concept),
                "en_std": self._label(self.labels_en, concept),
                "ext": concept in self.extensions,
            })
            if concept in path:                   # 순환 가드
                return
            for order, child, pref in sorted(children.get(concept, []),
                                             key=lambda t: t[0]):
                walk(child, depth + 1, pref, path | {concept})

        for r in roots:
            walk(r, 0, None, frozenset())
        return rows

    def all_role_rows(self):
        """[(정의, rows)] — Role별."""
        return [(definition, self.tree_rows(pl))
                for _, definition, pl in self.roles()]


# ---------------------------------------------------------------------------
# 금감원 택사노미 xlsm (Presentation Link 시트)
# ---------------------------------------------------------------------------

def fss_xlsm_role_rows(xlsm_path, role_filter=None):
    """xlsm의 Presentation Link 시트 → 같은 행 모델. [(정의, rows)] 반환.

    depth 열이 이미 계산돼 있어 그대로 사용. 표준 택사노미이므로 확장 없음.
    role_filter: 'D210000'처럼 코드 부분 문자열 필터 (None=전체).
    """
    wb = load_workbook(xlsm_path, read_only=True)
    ws = wb["Presentation Link"]
    blocks = []
    cur_def, cur_sector, cur_rows = None, "", []
    for row in ws.iter_rows(values_only=True):
        a, b = row[0], row[3]
        if b == "LinkRole":                       # 블록 시작
            if cur_def and cur_rows:
                blocks.append((cur_sector, cur_def, cur_rows))
            cur_def, cur_rows = None, []
            continue
        if b == "Definition":
            cur_def = str(row[4] or "").strip()
            continue
        if b == "prefix":                          # 컬럼 헤더
            continue
        if cur_def and row[4]:                     # 데이터 행 (name 존재)
            cur_sector = str(a or cur_sector or "").strip()
            depth = int(row[7] or 0)
            cur_rows.append({
                "depth": depth,
                "ko": str(row[6] or ""), "prefix": str(row[3] or ""),
                "en": str(row[5] or ""), "id": str(row[4] or ""),
                "ko_std": str(row[6] or ""), "en_std": str(row[5] or ""),
                "ext": False,
            })
    if cur_def and cur_rows:
        blocks.append((cur_sector, cur_def, cur_rows))
    wb.close()
    out = []
    for sector, definition, rows in blocks:
        if role_filter and role_filter not in definition:
            continue
        name = f"{sector} {definition}".strip()
        out.append((name, rows))
    return out


# ---------------------------------------------------------------------------
# 엑셀 출력 (이미지2 레이아웃)
# ---------------------------------------------------------------------------

def _sheet_name(definition, used):
    m = re.search(r"\[(D?\w+)\]", definition)
    base = m.group(1) if m else re.sub(r"[\\/*?:\[\]]", "", definition)[:20]
    prefix = definition.split("[")[0].strip()[:8]
    name = f"{prefix}{base}"[:31] if prefix else base[:31]
    n, i = name, 2
    while n in used:
        n = f"{name[:28]}_{i}"
        i += 1
    used.add(n)
    return n


def write_tree_excel(role_rows, out_path):
    """[(정의, rows)] → Role별 시트 엑셀. 확장 element는 빨간 글씨."""
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for definition, rows in role_rows:
        ws = wb.create_sheet(_sheet_name(definition, used))
        ws.append([definition])
        ws.cell(row=1, column=1).font = _BOLD
        ws.append(HEADER)
        for c in range(1, len(HEADER) + 1):
            cell = ws.cell(row=2, column=c)
            cell.font = _BOLD
            cell.fill = _HDR_FILL
        for r in rows:
            ws.append(["    " * r["depth"] + r["ko"], r["prefix"], r["en"],
                       r["id"], r["ko_std"], r["en_std"]])
            if r["ext"]:
                for c in range(1, len(HEADER) + 1):
                    ws.cell(row=ws.max_row, column=c).font = _RED
        ws.freeze_panes = "A3"
        widths = [56, 12, 44, 40, 32, 32]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(out_path)
    return out_path


def build_tree_view(source, out_path=None, role_filter=None):
    """패키지 폴더 또는 금감원 xlsm → 트리 뷰 엑셀. 요약 dict 반환."""
    if os.path.isdir(source):
        pkg = TaxonomyPackage(source)
        role_rows = pkg.all_role_rows()
        if role_filter:
            role_rows = [(d, r) for d, r in role_rows if role_filter in d]
        arc_total = pkg.arc_count()
        ext_total = sum(1 for _, rows in role_rows
                        for r in rows if r["ext"])
    else:
        role_rows = fss_xlsm_role_rows(source, role_filter)
        arc_total = None
        ext_total = 0
    if out_path is None:
        base = os.path.basename(source.rstrip("/\\"))
        out_path = os.path.join(
            os.path.dirname(source) if os.path.isfile(source) else source,
            re.sub(r"\.(xlsm|xlsx)$", "", base) + "_트리뷰.xlsx")
    write_tree_excel(role_rows, out_path)
    return {"out_path": out_path, "roles": len(role_rows),
            "rows": sum(len(r) for _, r in role_rows),
            "arcs": arc_total, "extensions": ext_total}

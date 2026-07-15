"""D-2: 차원 표 렌더러 (이미지1 형식 — DART 뷰어식 배치).

- 열: 축(axis)의 member 조합 (축 2개면 2단 중첩 헤더) + 마지막 [합계]
  = 차원 미지정(default member) 팩트
- 행: LineItems (pre.xml order, 들여쓰기)
- 기간별 블록 분리 "(기간: 당기)/(전기)", 음수는 빨간 괄호 (숫자서식)
- def.xml: all → hypercube-dimension → dimension-domain → domain-member
  아크롤 체인 (dimension-member 표기 수용)

DART 구조 반영 (삼성전자 실측):
- 연결/별도 축(ConsolidatedAndSeparateFinancialStatementsAxis)은 열이 아니라
  role별 필터 — 정의 문자열의 "- 연결/- 별도"로 member를 선택해 팩트를 거른다
- 주석 role 하나가 def에서 a/b/c… 서브롤(하이퍼큐브 단위)로 분할됨
  → 한 시트에 하이퍼큐브별 섹션 테이블을 연속 렌더
스코프 제외: typed dimension.
"""
import collections
import datetime
import os
import re

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .taxonomy import LB, XL, TaxonomyPackage, _pkg_glob

XI = "http://www.xbrl.org/2003/instance"
XD = "http://xbrl.org/2006/xbrldi"

FILTER_AXIS = "ifrs-full_ConsolidatedAndSeparateFinancialStatementsAxis"
_MEMBER_CONsolidated = "ifrs-full_ConsolidatedMember"
_MEMBER_SEPARATE = "ifrs-full_SeparateMember"

_BOLD = Font(bold=True)
_HDR_FILL = PatternFill("solid", start_color="D9E1F2")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_NUMFMT = "#,##0;[RED](#,##0)"              # 음수 = 빨간 괄호
_NUMFMT_DEC = "#,##0.00;[RED](#,##0.00)"

_PERIOD_NAMES = ["당기", "전기", "전전기"]

# 재무제표 시트 약칭 판정 (role 정의 텍스트 키워드 매칭 — 코드 하드코딩 대신
# "재무상태표" 등 명칭 우선). "포괄손익"을 "손익계산서"보다 먼저 검사해야
# "포괄손익계산서"가 PL로 잘못 분류되지 않는다.
_FS_KEYWORDS = [
    ("포괄손익", "PL1"), ("comprehensive income", "PL1"),
    ("재무상태표", "BS"), ("financial position", "BS"),
    ("손익계산서", "PL"), ("income statement", "PL"),
    ("자본변동표", "CE"), ("changes in equity", "CE"),
    ("현금흐름표", "CF"), ("cash flows", "CF"),
]


def _qid(text):
    return (text or "").strip().replace(":", "_")


def _role_code(uri):
    m = re.search(r"role[-_]([A-Za-z0-9]+)$", uri or "")
    return m.group(1) if m else (uri or "")


def _base_code(code):
    return re.sub(r"[a-z]+$", "", code)


def _try_num(val):
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def _parse_date(s):
    try:
        return datetime.date.fromisoformat(s)
    except (TypeError, ValueError):
        return None


def _is_note_role(role_uri):
    """role 코드가 D8xxxxx(주석) 형태인지 — 재무제표 본문(D2~D6)과 구분."""
    m = re.match(r"D(\d)", _role_code(role_uri))
    return bool(m) and m.group(1) == "8"


def _is_separate(definition):
    return "별도" in definition or \
        ("Separate" in definition and "연결" not in definition)


def _fs_abbrev(definition):
    for kw, code in _FS_KEYWORDS:
        if kw.lower() in definition.lower():
            return code
    return None


def _tier_name(years_back):
    if 0 <= years_back < len(_PERIOD_NAMES):
        return _PERIOD_NAMES[years_back]
    return f"{years_back + 1}기 전"


def _duration_desc(start, end):
    """기간 길이 설명: 3개월/누적/N개월 누적/None(연간, 접미사 없음)."""
    days = (end - start).days + 1
    months = round(days / 30.44)
    if days >= 350:
        return None                          # 연간 — 사업보고서 등
    if 80 <= days <= 100:
        return "3개월"
    if start.month == 1 and start.day == 1:  # 회계연도 개시일부터 누적
        if 175 <= days <= 190:
            return "누적"                     # 반기 누적(관용 표기)
        return f"{months}개월 누적"
    return f"{months}개월"


def _classify_period(ref_date, block):
    """block=(type,start,end) → {name, desc, kind}. 라벨은 실측 일자 기반."""
    btype, bstart, bend = block
    end_date = _parse_date(bend)
    ref_year = (ref_date or end_date).year if (ref_date or end_date) else None
    years_back = max(0, ref_year - end_date.year) \
        if (ref_year and end_date) else 0
    name = _tier_name(years_back)
    if btype == "duration":
        start_date = _parse_date(bstart)
        desc = _duration_desc(start_date, end_date) \
            if start_date and end_date else None
        return {"name": name, "desc": desc, "kind": "duration"}
    suffix = "초" if end_date and (end_date.month, end_date.day) == (1, 1) \
        else "말"
    return {"name": name, "desc": suffix, "kind": "instant"}


# ---------------------------------------------------------------------------
# 인스턴스
# ---------------------------------------------------------------------------

class XbrlInstance:
    """팩트·컨텍스트 (explicit dimension만)."""

    def __init__(self, folder):
        path = _pkg_glob(folder, "*.xbrl")
        if not path:
            raise FileNotFoundError(f".xbrl 인스턴스 없음: {folder}")
        root = etree.parse(path[0]).getroot()
        self.contexts = {}
        for c in root.findall(f"{{{XI}}}context"):
            p = c.find(f"{{{XI}}}period")
            inst = p.find(f"{{{XI}}}instant")
            if inst is not None:
                period = ("instant", inst.text, inst.text)
            else:
                s = p.find(f"{{{XI}}}startDate")
                e = p.find(f"{{{XI}}}endDate")
                period = ("duration",
                          s.text if s is not None else "",
                          e.text if e is not None else "")
            dims = {}
            for m in c.findall(f".//{{{XD}}}explicitMember"):
                dims[_qid(m.get("dimension"))] = _qid(m.text)
            self.contexts[c.get("id")] = {
                "type": period[0], "start": period[1], "end": period[2],
                "dims": dims,
            }
        self.facts = collections.defaultdict(list)
        for el in root:
            if not isinstance(el.tag, str):
                continue
            cref = el.get("contextRef")
            if cref is None or cref not in self.contexts:
                continue
            val = (el.text or "").strip()
            if not val:
                continue
            qn = etree.QName(el.tag)
            cid = f"{el.prefix}_{qn.localname}" if el.prefix else qn.localname
            self.facts[cid].append({
                "ctx": self.contexts[cref], "value": val,
                "decimals": el.get("decimals", ""),
            })

        # 보고기준일 (dart-gcd_DocumentPeriodEndDate) — 당기/전기 판정 기준.
        # 없으면 instant 컨텍스트 중 최신 종료일로 대체.
        self.doc_period_end = None
        ref = self.facts.get("dart-gcd_DocumentPeriodEndDate")
        if ref:
            self.doc_period_end = _parse_date(ref[0]["value"])
        if self.doc_period_end is None:
            ends = [_parse_date(c["end"]) for c in self.contexts.values()
                    if c["type"] == "instant"]
            ends = [d for d in ends if d]
            if ends:
                self.doc_period_end = max(ends)


# ---------------------------------------------------------------------------
# def.xml 하이퍼큐브
# ---------------------------------------------------------------------------

class HypercubeDef:
    """definitionLink별 하이퍼큐브. 주석 서브롤(a/b/c…)을 기본 코드로 묶는다."""

    def __init__(self, folder):
        self.by_base = collections.defaultdict(list)
        paths = _pkg_glob(folder, "*_def.xml")
        if not paths:
            return
        root = etree.parse(paths[0]).getroot()
        for dl in root.findall(f"{{{LB}}}definitionLink"):
            role_uri = dl.get(f"{{{XL}}}role")
            loc = {l.get(f"{{{XL}}}label"):
                   (l.get(f"{{{XL}}}href") or "").split("#")[-1]
                   for l in dl.findall(f"{{{LB}}}loc")}
            arcs = collections.defaultdict(list)
            for a in dl.findall(f"{{{LB}}}definitionArc"):
                kind = (a.get(f"{{{XL}}}arcrole") or "").rsplit("/", 1)[-1]
                if kind == "dimension-member":
                    kind = "domain-member"
                arcs[kind].append((loc.get(a.get(f"{{{XL}}}from")),
                                   loc.get(a.get(f"{{{XL}}}to")),
                                   float(a.get("order", "0") or 0)))
            if not arcs["all"]:
                continue
            children = collections.defaultdict(list)
            for src, dst, order in arcs["domain-member"]:
                children[src].append((order, dst))

            def closure(node):
                out, stack = [], [node]
                seen = {node}
                while stack:
                    cur = stack.pop(0)
                    out.append(cur)
                    for _, ch in sorted(children.get(cur, [])):
                        if ch not in seen:
                            seen.add(ch)
                            stack.append(ch)
                return out

            for li_node, table, _ in sorted(arcs["all"], key=lambda t: t[2]):
                axes = []
                for src, axis, order in sorted(arcs["hypercube-dimension"],
                                               key=lambda t: t[2]):
                    if src != table:
                        continue
                    domains = [d for s, d, _ in arcs["dimension-domain"]
                               if s == axis]
                    default = next((d for s, d, _ in
                                    arcs["dimension-default"] if s == axis),
                                   None)
                    members = []
                    for dom in domains:
                        members += [m for m in closure(dom)[1:]
                                    if m not in members]
                    axes.append({"axis": axis, "domains": domains,
                                 "default": default, "members": members})
                code = _role_code(role_uri)
                self.by_base[_base_code(code)].append({
                    "role_uri": role_uri, "code": code, "axes": axes,
                    "primary": set(closure(li_node)),
                })


# ---------------------------------------------------------------------------
# Role 테이블 모델 (섹션 = 하이퍼큐브 1개)
# ---------------------------------------------------------------------------

def _structural_concepts(cubes):
    out = set()
    for cube in cubes:
        for ax in cube["axes"]:
            out.add(ax["axis"])
            out.update(ax["domains"])
            out.update(ax["members"])
            if ax["default"]:
                out.add(ax["default"])
    return out


class RoleTable:
    def __init__(self, pkg: TaxonomyPackage, inst: XbrlInstance,
                 cube: HypercubeDef, role_uri, definition, pl):
        self.pkg = pkg
        self.definition = definition
        self.ref_date = inst.doc_period_end
        self.cubes = cube.by_base.get(_base_code(_role_code(role_uri)), [])
        structural = _structural_concepts(self.cubes)

        pre_rows = []
        for r in pkg.tree_rows(pl):
            cid = f"{r['prefix']}_{r['id']}" if r["prefix"] else r["id"]
            if cid.endswith(("Table", "Axis")) or cid in structural:
                continue
            pre_rows.append((r, cid))

        # 연결/별도 필터 member (role 정의 문자열 기준)
        self.filter_member = None
        if "별도" in definition or "Separate" in definition:
            self.filter_member = _MEMBER_SEPARATE
        elif "연결" in definition or "Consolidated" in definition:
            self.filter_member = _MEMBER_CONsolidated

        def keep(f):
            m = f["ctx"]["dims"].get(FILTER_AXIS)
            return m is None or self.filter_member is None or \
                m == self.filter_member

        facts = {cid: [f for f in inst.facts.get(cid, []) if keep(f)]
                 for _, cid in pre_rows}

        # 섹션 구성: 하이퍼큐브별 (없으면 전체 1섹션·축0)
        self.sections = []
        if self.cubes:
            for cb in self.cubes:
                rows = [(r, cid) for r, cid in pre_rows
                        if cid in cb["primary"]]
                if not rows:
                    continue
                sec = _Section(self, rows, facts, cb)
                if sec.has_facts:
                    self.sections.append(sec)
        if not self.sections:
            sec = _Section(self, pre_rows, facts, None)
            if sec.has_facts:
                self.sections.append(sec)

    @property
    def has_facts(self):
        return bool(self.sections)

    def member_label(self, member):
        return self.pkg._label(self.pkg.labels_ko, member)


class _Section:
    """하이퍼큐브 1개 = 표 1개."""

    def __init__(self, parent: RoleTable, rows, all_facts, cube):
        self.parent = parent
        self.rows = rows
        self.facts = {cid: all_facts.get(cid, []) for _, cid in rows}
        self.has_facts = any(self.facts.values())

        axes_def = [ax for ax in (cube["axes"] if cube else [])
                    if ax["axis"] != FILTER_AXIS]
        self.axes = []
        for ax in axes_def:
            observed = {f["ctx"]["dims"].get(ax["axis"])
                        for fl in self.facts.values() for f in fl}
            observed.discard(None)
            used = [m for m in ax["members"] if m in observed]
            if not ax["members"]:
                # 큐브가 member 미선언 시에만 관측값 사용
                # (선언된 경우 타 하이퍼큐브 member 유입 차단)
                used = sorted(observed)
            if used:
                self.axes.append({**ax, "members": used})
        self.axis_ids = [ax["axis"] for ax in self.axes]

        periods = {}
        for fl in self.facts.values():
            for f in fl:
                c = f["ctx"]
                periods[(c["type"], c["start"], c["end"])] = True
        durs = sorted([p for p in periods if p[0] == "duration"],
                      key=lambda p: p[1])
        durs.sort(key=lambda p: p[2], reverse=True)
        insts = sorted([p for p in periods if p[0] == "instant"],
                       key=lambda p: p[2], reverse=True)
        self.block_kind = "duration" if durs else "instant"
        self.blocks = durs if durs else insts
        self.inst_cols = insts
        if self.axes:
            def has_member_fact(block):
                return any(
                    (f["ctx"]["type"], f["ctx"]["start"],
                     f["ctx"]["end"]) == block
                    and any(f["ctx"]["dims"].get(a) for a in self.axis_ids)
                    for fl in self.facts.values() for f in fl)
            self.blocks = [b for b in self.blocks if has_member_fact(b)]

    # --- 값 조회 -----------------------------------------------------------
    def _dims_match(self, ctx, combo):
        proj = {a: ctx["dims"].get(a) for a in self.axis_ids}
        want = {a: None for a in self.axis_ids}
        want.update(combo)
        return proj == want

    def _in_block(self, ctx, block):
        btype, bstart, bend = block
        if ctx["type"] == btype and (ctx["start"], ctx["end"]) == \
                (bstart, bend):
            return True
        if btype == "duration" and ctx["type"] == "instant":
            return ctx["end"] == bend
        return False

    def cell(self, concept, block, combo):
        for f in self.facts.get(concept, []):
            if self._in_block(f["ctx"], block) and \
                    self._dims_match(f["ctx"], combo):
                return f
        return None

    def block_label(self, block):
        """실측 일자 기반 라벨. 예: '당기 3개월 (2025-04-01 ~ 2025-06-30)'."""
        p = _classify_period(self.parent.ref_date, block)
        _, bstart, bend = block
        if p["kind"] == "duration":
            head = f"{p['name']} {p['desc']}" if p["desc"] else p["name"]
            return f"{head} ({bstart} ~ {bend})"
        return f"{p['name']}{p['desc']} ({bend})"


# ---------------------------------------------------------------------------
# 엑셀 렌더링
# ---------------------------------------------------------------------------

def _write_cell(ws, r, c, fact):
    if fact is None:
        return False
    num = _try_num(fact["value"])
    cell = ws.cell(row=r, column=c)
    if num is not None:
        cell.value = num
        dec = fact.get("decimals", "")
        try:
            frac = dec not in ("", "0", "INF") and int(float(dec)) > 0
        except ValueError:
            frac = False
        cell.number_format = _NUMFMT_DEC if frac else _NUMFMT
    else:
        text = fact["value"]
        cell.value = text[:100] + ("…" if len(text) > 100 else "")
    return True


def _hdr(ws, r, c, value):
    cell = ws.cell(row=r, column=c, value=value)
    cell.font = _BOLD
    cell.fill = _HDR_FILL
    cell.alignment = _CENTER
    return cell


def _render_section_flat(ws, sec: _Section, r):
    """축 0개: 행=element, 열=기간 (실측 일자 기반 라벨)."""
    ref_date = sec.parent.ref_date
    cols = list(dict.fromkeys(
        ([b for b in sec.blocks] if sec.block_kind == "duration" else [])
        + sec.inst_cols))
    labels = []
    for c in cols:
        p = _classify_period(ref_date, c)
        if p["kind"] == "duration":
            head = f"{p['name']} {p['desc']}" if p["desc"] else p["name"]
            labels.append(f"{head}\n{c[1]}~{c[2]}")
        else:
            labels.append(f"{p['name']}{p['desc']}\n{c[2]}")
    _hdr(ws, r, 1, "계정과목")
    for j, lab in enumerate(labels):
        _hdr(ws, r, 2 + j, lab)
    r += 1
    for row, concept in sec.rows:
        ws.cell(row=r, column=1, value="    " * row["depth"] + row["ko"])
        for j, c in enumerate(cols):
            _write_cell(ws, r, 2 + j, sec.cell(concept, c, {}))
        r += 1
    return r + 1


def _render_section_dim(ws, sec: _Section, r):
    """축 1~2개: 기간 블록 × (member 조합 + [합계])."""
    for block in sec.blocks:
        ws.cell(row=r, column=1, value=sec.block_label(block)).font = _BOLD
        r += 1
        combos = [{}]
        for ax in sec.axes:
            combos = [dict(c, **{ax["axis"]: m}) for c in combos
                      for m in ax["members"]]
        combos = [c for c in combos
                  if any(sec.cell(concept, block, c)
                         for _, concept in sec.rows)]
        if not combos:
            r += 1
            continue
        nhdr = len(sec.axes)
        for ai, ax in enumerate(sec.axes):
            ws.cell(row=r + ai, column=1,
                    value="계정과목" if ai == nhdr - 1 else "")
            if ai == nhdr - 1:
                ws.cell(row=r + ai, column=1).font = _BOLD
            col = 2
            j = 0
            while j < len(combos):
                m = combos[j][ax["axis"]]
                span = 1
                while j + span < len(combos) and \
                        combos[j + span][ax["axis"]] == m and \
                        all(combos[j + span][a["axis"]] ==
                            combos[j][a["axis"]]
                            for a in sec.axes[:ai]):
                    span += 1
                _hdr(ws, r + ai, col, sec.parent.member_label(m))
                if span > 1:
                    ws.merge_cells(start_row=r + ai, start_column=col,
                                   end_row=r + ai, end_column=col + span - 1)
                col += span
                j += span
            if ai == nhdr - 1:
                _hdr(ws, r + ai, 2 + len(combos), "[합계]")
        r += nhdr
        for row, concept in sec.rows:
            ws.cell(row=r, column=1,
                    value="    " * row["depth"] + row["ko"])
            for j, combo in enumerate(combos):
                _write_cell(ws, r, 2 + j, sec.cell(concept, block, combo))
            _write_cell(ws, r, 2 + len(combos),
                        sec.cell(concept, block, {}))
            r += 1
        r += 1
    return r


def _assign_sheet_names(entries):
    """(role_uri, definition) 목록 → DSD 편집용 엑셀과 동일한 시트명 체계.

    - 재무제표 본문(D2~D6): 정의 텍스트 키워드로 BS/PL/PL1/CE/CF 판정.
      연결·별도가 "모두 렌더되는" 경우에만 별도에 `_별도` 접미사(연결이 기본).
    - 주석(D8xxxxx): role 코드 오름차순으로 1, 2, 3, ... 순번.
    - 판정 불가·충돌 시 role 코드로 폴백 + 31자 제한·중복 접미사 처리.
    """
    note_sorted = sorted(
        [(uri, d) for uri, d in entries if _is_note_role(uri)],
        key=lambda e: _role_code(e[0]))
    note_number = {uri: str(i + 1) for i, (uri, _) in enumerate(note_sorted)}

    fs_abbrev = {}
    for uri, definition in entries:
        if _is_note_role(uri):
            continue
        ab = _fs_abbrev(definition)
        if ab:
            fs_abbrev[uri] = ab
    # 손익계산서(PL)가 따로 없이 "단일 포괄손익계산서"만 있는 경우
    # (K-IFRS 단일 표시 방식) — 유일한 손익 시트이므로 PL1이 아닌 PL로 명명.
    # PL1은 손익계산서·포괄손익계산서가 별개 표로 둘 다 존재할 때만 쓴다.
    if "PL" not in fs_abbrev.values() and "PL1" in fs_abbrev.values():
        fs_abbrev = {u: ("PL" if a == "PL1" else a)
                     for u, a in fs_abbrev.items()}

    def_by_uri = dict(entries)
    variants = collections.defaultdict(set)
    for uri, ab in fs_abbrev.items():
        variants[ab].add(
            "별도" if _is_separate(def_by_uri[uri]) else "연결")

    used, names = set(), []
    for uri, definition in entries:
        if _is_note_role(uri):
            base = note_number.get(uri, _role_code(uri))
        else:
            ab = fs_abbrev.get(uri)
            if ab is None:
                base = _role_code(uri)               # 미분류 폴백
            elif _is_separate(definition) and len(variants[ab]) > 1:
                base = f"{ab}_별도"
            else:
                base = ab
        name, i = base[:31], 2
        while name in used:
            name = f"{base[:29]}_{i}"
            i += 1
        used.add(name)
        names.append(name)
    return names


def render_dimension_tables(folder, out_path=None, role_filter=None):
    """패키지 폴더 → Role별 차원 표 엑셀. 요약 dict 반환."""
    pkg = TaxonomyPackage(folder)
    inst = XbrlInstance(folder)
    cube = HypercubeDef(folder)

    entries = []
    for role_uri, definition, pl in pkg.roles():
        if role_filter and role_filter not in definition and \
                role_filter not in role_uri:
            continue
        table = RoleTable(pkg, inst, cube, role_uri, definition, pl)
        if table.has_facts:
            entries.append((role_uri, definition, table))
    if not entries:
        raise ValueError("렌더링할 Role이 없습니다 (필터 확인).")

    names = _assign_sheet_names([(u, d) for u, d, _ in entries])

    wb = Workbook()
    wb.remove(wb.active)
    rendered = []
    for (role_uri, definition, table), sheet_name in zip(entries, names):
        ws = wb.create_sheet(sheet_name)
        # 추적성 보존: role 코드 + 정의 전문 (definition에 "[Dxxxxxx] ..." 포함)
        ws.append([definition])
        ws.cell(row=1, column=1).font = _BOLD
        r = 3
        for sec in table.sections:
            if sec.axes:
                r = _render_section_dim(ws, sec, r)
            else:
                r = _render_section_flat(ws, sec, r)
        ws.column_dimensions["A"].width = 52
        for ci in range(2, 14):
            ws.column_dimensions[get_column_letter(ci)].width = 17
        ws.freeze_panes = "B2"
        rendered.append({
            "definition": definition,
            "sheet": sheet_name,
            "sections": len(table.sections),
            "axes": max((len(s.axes) for s in table.sections), default=0),
            "rows": sum(len(s.rows) for s in table.sections),
        })
    if out_path is None:
        out_path = os.path.join(folder, "차원표.xlsx")
    wb.save(out_path)
    return {"out_path": out_path, "roles": rendered}

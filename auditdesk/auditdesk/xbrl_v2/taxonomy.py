"""Offline DART spreadsheet preview parser; not a DTS/DRS validator.

Namespace and release applicability declarations are caller-supplied evidence,
never inferred from filenames, prefixes, presentation paths, or prior sources.
"""
from __future__ import annotations
from dataclasses import dataclass, asdict
from datetime import date
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile
from openpyxl import load_workbook
from .model import ExpandedQName, stable_id

PARSER_VERSION = 'dart-workbook-preview/3'

@dataclass(frozen=True)
class TaxonomyMetadata:
    version: str
    source_role: str
    report_period: str
    applicable_from: str
    applicable_to: str
    applicability_evidence: str
    content_sha256: str

@dataclass(frozen=True)
class NamespaceBinding:
    prefix: str
    namespace_uri: str
    evidence: str

@dataclass(frozen=True)
class WorkbookLocation:
    snapshot_id: str
    sheet: str
    row: int
    column: int = 0

@dataclass(frozen=True)
class TaxonomyDiagnostic:
    code: str
    message: str
    severity: str = 'ERROR'
    source: WorkbookLocation | None = None

@dataclass(frozen=True)
class TaxonomyConcept:
    concept_id: str
    qname: ExpandedQName | None
    prefix: str
    name: str
    type_name: str
    period_type: str
    abstract: bool | None
    source: WorkbookLocation
    source_id: str = ''

@dataclass(frozen=True)
class TaxonomyRole:
    role_uri: str
    definition: str
    used_on: str
    source: WorkbookLocation
    source_id: str = ''

@dataclass(frozen=True)
class TaxonomyOccurrence:
    occurrence_id: str
    qname: ExpandedQName | None
    role_uri: str
    path: tuple[str, ...]
    preferred_label: str
    arcrole: str
    order: str
    source: WorkbookLocation

@dataclass(frozen=True)
class TaxonomyLabel:
    resource_id: str
    qname: ExpandedQName | None
    language: str
    role: str
    text: str
    source: WorkbookLocation

@dataclass(frozen=True)
class TaxonomyReference:
    resource_id: str
    qname: ExpandedQName | None
    role: str
    parts: tuple[tuple[str, str], ...]
    source: WorkbookLocation

@dataclass(frozen=True)
class TaxonomySnapshot:
    snapshot_id: str
    content_sha256: str
    logical_uri: str
    parser_version: str
    metadata: TaxonomyMetadata
    namespaces: tuple[NamespaceBinding, ...]
    concepts: tuple[TaxonomyConcept, ...]
    roles: tuple[TaxonomyRole, ...]
    occurrences: tuple[TaxonomyOccurrence, ...]
    labels: tuple[TaxonomyLabel, ...]
    references: tuple[TaxonomyReference, ...]
    diagnostics: tuple[TaxonomyDiagnostic, ...]
    current_applicable: bool
    capabilities: str = 'PARTIAL'
    dimensions_known: bool = False


def parse_taxonomy_workbook(data: bytes, *, logical_uri: str,
                            metadata: TaxonomyMetadata,
                            namespaces: tuple[NamespaceBinding, ...] = ()) -> TaxonomySnapshot:
    """Parse bounded xlsx preview; structural or applicability errors fail closed.

    ``current_applicable`` means the supplied release evidence is consistent with
    these exact bytes and report date. It is not independent verification of the
    evidence, dimensional validity, or an accounting approval.
    """
    bindings = tuple(sorted(namespaces, key=lambda n: (n.prefix, n.namespace_uri, n.evidence)))
    digest = sha256(data).hexdigest()
    sid = stable_id('taxonomy', (digest, logical_uri, PARSER_VERSION, asdict(metadata), [asdict(n) for n in bindings]))
    diagnostics=[]; concepts=[]; roles=[]; occurrences=[]; labels=[]; references=[]
    def loc(sheet, row, column=0): return WorkbookLocation(sid, sheet, row, column)
    def error(code, message, source=None): diagnostics.append(TaxonomyDiagnostic(code, message, 'ERROR', source))
    def result():
        return TaxonomySnapshot(sid,digest,logical_uri,PARSER_VERSION,metadata,bindings,tuple(concepts),tuple(roles),tuple(occurrences),tuple(labels),tuple(references),tuple(diagnostics),not any(d.severity=='ERROR' for d in diagnostics))
    if not logical_uri.strip() or not metadata.version.strip() or not metadata.applicability_evidence.strip():
        error('MISSING_APPLICABILITY', 'Logical source, version and release applicability evidence required')
    if metadata.source_role != 'CURRENT_TAXONOMY':
        error('NON_CURRENT_SOURCE', 'Prior/reference sources cannot be promoted to current taxonomy')
    if metadata.content_sha256 != digest:
        error('SOURCE_HASH_MISMATCH','Applicability declaration is not bound to these workbook bytes')
    try:
        if not date.fromisoformat(metadata.applicable_from) <= date.fromisoformat(metadata.report_period) <= date.fromisoformat(metadata.applicable_to):
            error('OUTSIDE_APPLICABILITY','Report date lies outside supplied release applicability')
    except (ValueError, TypeError): error('INVALID_APPLICABILITY_DATE','Explicit ISO report and release range dates required')
    ns={}; invalid_prefixes=set()
    for n in bindings:
        if n.prefix in ns or not n.prefix.strip() or not n.namespace_uri.strip() or not n.evidence.strip():
            invalid_prefixes.add(n.prefix); error('AMBIGUOUS_NAMESPACE','Unique namespace mapping with provenance required')
        ns[n.prefix]=n.namespace_uri
    def qname(prefix, name, source):
        if not prefix or not name:
            error('MISSING_CONCEPT_IDENTITY', 'Nonempty data row requires both prefix and name', source)
            return None
        if prefix not in ns or prefix in invalid_prefixes:
            error('UNRESOLVED_NAMESPACE', f'Cannot expand {prefix}:{name} without unambiguous declaration', source); return None
        return ExpandedQName(ns[prefix],name)
    def text(v): return '' if v is None else str(v).strip()
    try:
        if len(data)>32*1024*1024: raise ValueError('compressed workbook limit exceeded')
        with ZipFile(BytesIO(data)) as z:
            if len(z.infolist())>1000 or sum(i.file_size for i in z.infolist())>128*1024*1024:
                raise ValueError('expanded workbook limit exceeded')
        wb=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    except Exception as exc:
        error('INVALID_WORKBOOK',f'Workbook could not be opened ({type(exc).__name__})'); return result()
    sheets={}
    try:
        for ws in wb:
            if ws.title not in {'Concepts', 'RoleTypes', 'Presentation Link', 'Label Link', 'Reference Link'}:
                continue
            if ws.max_row is None or ws.max_column is None or ws.max_row>200000 or ws.max_column>256:
                error('WORKBOOK_LIMIT','Worksheet dimensions missing or exceed preview bounds'); return result()
            rows=[]
            for row in ws.iter_rows():
                if any(c.data_type=='f' for c in row): error('FORMULA_UNSUPPORTED','Formula cells are not authoritative taxonomy data',loc(ws.title,row[0].row))
                rows.append(tuple(text(c.value) for c in row))
            sheets[ws.title]=rows
    except Exception as exc:
        error('INVALID_WORKBOOK',f'Worksheet could not be read ({type(exc).__name__})'); return result()
    finally: wb.close()
    def header(sheet, required):
        rows=sheets.get(sheet,[])
        for index,row in enumerate(rows[:8]):
            if set(required).issubset(row):
                if len([v for v in row if v]) != len(set(v for v in row if v)):
                    error('DUPLICATE_HEADER','Duplicate header names',loc(sheet,index+1)); return None
                return index,{v:i for i,v in enumerate(row) if v}
        error('MISSING_HEADER',f'{sheet} requires explicit headers {required}'); return None
    def get(row,h,key):
        index=h.get(key); return row[index] if index is not None and index<len(row) else ''
    def source_identifier(row, h, source, seen, reported):
        # The workbook identifies record kinds, not XSD document scopes. Keep
        # supplied IDs and reject collisions within a kind without guessing
        # that different prefixes represent separate schema documents.
        identifier = get(row, h, 'id')
        if identifier:
            if identifier in seen:
                if identifier not in reported:
                    error('DUPLICATE_SOURCE_ID', 'Repeated source identifier has unresolved schema scope', seen[identifier])
                    reported.add(identifier)
                error('DUPLICATE_SOURCE_ID', 'Repeated source identifier has unresolved schema scope', source)
            else:
                seen[identifier] = source
        return identifier
    ch=header('Concepts',('prefix','name','type','periodType','abstract'))
    concept_keys={}; ambiguous=set(); concept_ids={}; duplicate_concept_ids=set()
    if ch:
        start,h=ch
        for index,row in enumerate(sheets['Concepts'][start+1:],start+2):
            if not any(row): continue
            source=loc('Concepts',index); prefix=get(row,h,'prefix'); name=get(row,h,'name')
            source_id=source_identifier(row,h,source,concept_ids,duplicate_concept_ids)
            q=qname(prefix,name,source); key=q.key if q else prefix+':'+name
            abstract={'true':True,'false':False,'1':True,'0':False}.get(get(row,h,'abstract').lower())
            period=get(row,h,'periodType'); typ=get(row,h,'type')
            if abstract is None:
                if get(row,h,'abstract'):
                    error('INVALID_CONCEPT','Invalid abstract flag',source)
                else:
                    diagnostics.append(TaxonomyDiagnostic('UNKNOWN_ABSTRACT','Workbook omits abstract flag; line-item eligibility is unknown','WARNING',source))
            if period not in {'instant','duration'} or not typ:
                error('INVALID_CONCEPT','Explicit type and periodType required',source)
            if key in concept_keys: ambiguous.add(key); error('DUPLICATE_CONCEPT','Repeated expanded QName is ambiguous',source)
            concept_keys[key]=q
            concepts.append(TaxonomyConcept(stable_id('concept',(sid,index,key)),q,prefix,name,typ,period,abstract,source,source_id))
    if not concepts: error('EMPTY_CONCEPTS','No concepts available')
    rh=header('RoleTypes',('roleURI','definition','usedOn'))
    role_keys=set(); role_ids={}; duplicate_role_ids=set()
    if rh:
        start,h=rh
        for index,row in enumerate(sheets['RoleTypes'][start+1:],start+2):
            if not any(row): continue
            uri=get(row,h,'roleURI'); source=loc('RoleTypes',index)
            source_id=source_identifier(row,h,source,role_ids,duplicate_role_ids)
            if not uri or uri in role_keys: error('AMBIGUOUS_ROLE','Missing or duplicate Role URI',source)
            role_keys.add(uri); roles.append(TaxonomyRole(uri,get(row,h,'definition'),get(row,h,'usedOn'),source,source_id))
    def linked(prefix,name,source):
        q=qname(prefix,name,source)
        key=q.key if q else prefix+':'+name
        if key not in concept_keys or key in ambiguous: error('UNRESOLVED_CONCEPT','Link references missing or ambiguous concept',source); return None
        return q
    rows=sheets.get('Presentation Link',[]); role=''; h=None; stack=[]; section=0
    section_start=0; section_header=False
    def presentation_metadata(row, marker):
        # Explicit layouts observed in official DART and the existing preview.
        # Extra populated cells are ambiguous, never guessed or discarded.
        for marker_col, value_col in ((0,1), (1,3)):
            if len(row)>marker_col and row[marker_col]==marker:
                if len(row)<=value_col or any(value for col,value in enumerate(row) if col not in {marker_col,value_col}):
                    return True, None
                return True, row[value_col]
        return False, None
    def close_presentation_section():
        if section_start and not section_header:
            error('MISSING_PRESENTATION_HEADER', 'Presentation section has no unambiguous header', loc('Presentation Link',section_start))
    if not rows: error('MISSING_PRESENTATION','Presentation Link sheet required for occurrence preview')
    for index,row in enumerate(rows,1):
        is_role, role_value=presentation_metadata(row,'LinkRole')
        if is_role:
            close_presentation_section()
            role=role_value or ''; section+=1; h=None; stack=[]
            section_start=index; section_header=False
            if role_value is None:
                error('MALFORMED_PRESENTATION_METADATA', 'Role metadata has ambiguous populated cells', loc('Presentation Link',index))
            if role not in role_keys: error('UNRESOLVED_ROLE','Presentation Role not declared',loc('Presentation Link',index))
            continue
        if 'prefix' in row and 'name' in row:
            if any(row.count(k)!=1 for k in ('prefix','name','depth','parent','order','arcrole','preferredLabel')):
                error('MALFORMED_PRESENTATION_HEADER','Missing/duplicate occurrence header',loc('Presentation Link',index)); h=None
            else:
                h={v:i for i,v in enumerate(row) if v}
                section_header=True
            continue
        if not any(row): continue
        if h is None:
            # Only the documented Definition metadata layout may precede a
            # section header; every other nonempty row gets a rejection locator.
            is_definition, definition_value=presentation_metadata(row,'Definition')
            if section_start and is_definition and definition_value is not None:
                continue
            error('REJECTED_PRESENTATION_ROW', 'Nonempty row cannot be interpreted without a presentation header', loc('Presentation Link',index))
            continue
        if not role or role not in role_keys:
            error('UNRESOLVED_ROLE','Occurrence lacks a declared Role section',loc('Presentation Link',index))
        source=loc('Presentation Link',index); q=linked(get(row,h,'prefix'),get(row,h,'name'),source)
        try:
            depth=int(get(row,h,'depth'))
            if depth<0 or depth>len(stack): raise ValueError()
        except ValueError:
            error('INVALID_PRESENTATION_DEPTH','Hierarchy depth cannot be resolved',source); stack=[]; continue
        key=q.key if q else get(row,h,'prefix')+':'+get(row,h,'name')
        stack=stack[:depth]; parent=get(row,h,'parent')
        if depth and not parent:
            error('AMBIGUOUS_PARENT','Nested occurrence has no explicit parent',source)
        if not depth and parent:
            error('AMBIGUOUS_PARENT','Root occurrence unexpectedly names a parent',source)
        if depth and parent:
            raw=parent.split(':',1)
            pq=qname(raw[0],raw[1],source) if len(raw)==2 else None
            if pq is None or pq.key!=stack[-1]: error('AMBIGUOUS_PARENT','Parent does not match depth path',source)
        stack.append(key)
        occurrences.append(TaxonomyOccurrence(stable_id('occurrence',(sid,section,index,role,stack)),q,role,tuple(stack),get(row,h,'preferredLabel'),get(row,h,'arcrole'),get(row,h,'order'),source))
    close_presentation_section()
    if not occurrences:
        error('EMPTY_PRESENTATION','No resolvable presentation occurrences')
    # Label headers repeat roles across language groups; carry language only from
    # the explicit language row, never infer it from text or workbook filename.
    rows=sheets.get('Label Link',[]); hrow=None; langs=[]
    if not rows: diagnostics.append(TaxonomyDiagnostic('MISSING_LABELS','Label Link sheet absent','WARNING'))
    for index,row in enumerate(rows,1):
        if 'prefix' in row and 'name' in row:
            hrow=row; previous=rows[index-2] if index>=2 else (); language=''; langs=[]
            for col in range(len(row)):
                value=previous[col] if col<len(previous) else ''
                if value: language=value
                langs.append(language)
            keys=[(langs[c],row[c]) for c in range(3,len(row)) if row[c]]
            if row[:3] != ('#','prefix','name') or len(keys)!=len(set(keys)):
                error('AMBIGUOUS_LABEL_HEADER','Unique language/role pairs and explicit identity columns required',loc('Label Link',index))
                hrow=None
            continue
        if hrow is None or not any(row): continue
        h={v:i for i,v in enumerate(hrow[:3]) if v}
        q=linked(get(row,h,'prefix'),get(row,h,'name'),loc('Label Link',index))
        for col in range(3,min(len(row),len(hrow))):
            if not row[col]: continue
            if not langs[col] or not hrow[col]: error('AMBIGUOUS_LABEL','Label has no explicit language/role',loc('Label Link',index,col+1)); continue
            source=loc('Label Link',index,col+1)
            labels.append(TaxonomyLabel(stable_id('label',(sid,index,col)),q,langs[col],hrow[col],row[col],source))
    if 'Label Link' in sheets and hrow is None:
        error('MISSING_LABEL_HEADER','Label Link has no unambiguous language/role header')
    rh=None
    if 'Reference Link' in sheets:
        rh=header('Reference Link',('prefix','name','role'))
    else:
        diagnostics.append(TaxonomyDiagnostic('MISSING_REFERENCES','Reference Link sheet absent','WARNING'))
    if rh:
        start,h=rh
        for index,row in enumerate(sheets['Reference Link'][start+1:],start+2):
            if not any(row): continue
            source=loc('Reference Link',index); q=linked(get(row,h,'prefix'),get(row,h,'name'),source)
            role=get(row,h,'role')
            parts=tuple((name,get(row,h,name)) for name in h if name not in {'#','prefix','name','label','verbose','role'} and get(row,h,name))
            if not role or not parts: error('AMBIGUOUS_REFERENCE','Reference needs role and named parts',source)
            references.append(TaxonomyReference(stable_id('reference',(sid,index)),q,role,parts,source))
    return result()

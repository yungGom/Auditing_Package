"""/api/workbench — dsd_workbench 함수만 import (네트워크 금지 경계).

라우터는 얇은 래퍼 — 로직은 전부 dsd_tool 기존 함수.
repack의 승인 게이트(★)는 서버가 강제: diff 미실행 → 409.
부분 승인은 백엔드(repack)가 전체 일괄 적용만 지원하므로 422로 안내
(백엔드 수정 금지 원칙 — 변경을 빼려면 엑셀에서 되돌린 뒤 재-diff).
"""
import datetime
import hashlib
import json
import os
import uuid

from fastapi import APIRouter, HTTPException

from .. import jobs

router = APIRouter()

_WORKDIR = os.path.join(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__))), "data", "sessions")

_CHECKLIST = [
    "DART 편집기에서 파일이 정상적으로 열리는지 확인",
    "수정한 셀의 표시가 의도대로인지 확인 (표·주석)",
    "편집기 자체 검증(오류 검사) 통과 확인",
]


def _sha1(path):
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _session(sid):
    with jobs.connect() as con:
        row = con.execute(
            "SELECT id, dsd_path, created, meta, state, xlsx_path, diff, "
            "diff_options, repack, foot, recon FROM sessions WHERE id=?",
            (sid,)).fetchone()
    if row is None:
        raise HTTPException(404, "세션 없음")
    meta = json.loads(row[3] or "{}")
    # editver 검증 여부는 저장 스냅샷이 아니라 현재 KNOWN_VERSIONS 표
    # 기준으로 재판정 (batch_validate로 등재가 갱신되면 즉시 반영)
    if meta.get("editver"):
        from dsd_tool.version import is_known
        meta["editver_known"] = bool(is_known(meta["editver"]))
    # UI-7: 구버전 상태값을 읽기 시점에 화면 용어로 정규화
    state = "반영완료" if row[4] == "repack" + "완료" else row[4]
    return {"session_id": row[0], "dsd_path": row[1], "created": row[2],
            "meta": meta, "state": state,
            "xlsx_path": row[5],
            "diff": json.loads(row[6]) if row[6] else None,
            "diff_options": json.loads(row[7]) if row[7] else None,
            "repack": json.loads(row[8]) if row[8] else None,
            "foot": json.loads(row[9]) if row[9] else None,
            "recon": json.loads(row[10]) if row[10] else None}


def _update(sid, **cols):
    sets = ", ".join(f"{k}=?" for k in cols)
    with jobs.connect() as con:
        con.execute(f"UPDATE sessions SET {sets} WHERE id=?",
                    (*cols.values(), sid))


@router.post("/sessions", status_code=201)
def create_session(body: dict):
    dsd_path = body.get("dsd_path") or ""
    if dsd_path.lower().endswith(".ixd"):
        # UI-7 델타 g: IXD 가드 — 침묵 실패 금지
        raise HTTPException(
            400, "IXD는 편집기 프로젝트 파일입니다. 편집기에서 생성한 "
                 "제출용 XBRL 패키지(또는 DSD 파일)를 투입하세요")
    if not os.path.isfile(dsd_path):
        raise HTTPException(400, f"DSD 파일이 없습니다: {dsd_path}")
    from dsd_tool.version import is_known, read_version_info
    with open(dsd_path, "rb") as f:
        data = f.read()
    v = read_version_info(data)
    meta = {
        "file": os.path.basename(dsd_path),
        "sha1": _sha1(dsd_path),
        "size": os.path.getsize(dsd_path),
        "editver": v.get("editver"), "docver": v.get("docver"),
        "editver_known": bool(is_known(v.get("editver"))),
        "company": os.path.basename(dsd_path).split("_")[0].strip("[]"),
        "cells": None, "notes": None, "cr_only": None,   # extract 후 채움
    }
    sid = uuid.uuid4().hex[:10]
    with jobs.connect() as con:
        con.execute(
            "INSERT INTO sessions "
            "(id, dsd_path, created, meta, state, xlsx_path, diff, "
            "diff_options, repack, foot, recon) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?)",
            (sid, os.path.abspath(dsd_path),
             datetime.datetime.now().isoformat(timespec="seconds"),
             json.dumps(meta, ensure_ascii=False), "생성됨",
             None, None, None, None, None, None))
    return {"session_id": sid, "meta": meta}


@router.get("/sessions")
def list_sessions():
    with jobs.connect() as con:
        rows = con.execute(
            "SELECT id FROM sessions ORDER BY created DESC LIMIT 50")
        ids = [r[0] for r in rows]
    return [
        {k: s[k] for k in ("session_id", "dsd_path", "created", "meta",
                           "state", "xlsx_path")}
        for s in (_session(i) for i in ids)]


@router.get("/sessions/{sid}")
def get_session(sid: str):
    s = _session(sid)
    # 파이프라인 단계 상태 (추출 → 엑셀 편집(외부) → 변경검토 → repack → DART 확인)
    s["pipeline"] = [
        {"key": "extract", "label": "추출", "done": bool(s["xlsx_path"])},
        {"key": "edit", "label": "엑셀에서 수정",
         "done": bool(s["diff"] or s["repack"])},
        {"key": "diff", "label": "수정 확인", "done": bool(s["diff"])},
        {"key": "repack", "label": "DSD에 반영", "done": bool(s["repack"])},
        {"key": "confirm", "label": "DART 확인", "done": False},
    ]
    return s


@router.post("/sessions/{sid}/extract", status_code=202)
def extract_session(sid: str):
    s = _session(sid)
    dsd = s["dsd_path"]
    os.makedirs(_WORKDIR, exist_ok=True)
    out = os.path.join(_WORKDIR, f"{sid}_{os.path.splitext(
        os.path.basename(dsd))[0]}.xlsx")

    def _run(progress):
        progress("DSD → Excel 추출 중…")
        from dsd_tool.excel_out import extract
        info = extract(dsd, out)
        meta = s["meta"]
        meta.update(cells=info["mapped_cells"], notes=info["note_count"],
                    cr_only=info["cr_only_cells"],
                    fs_sheets=info["fs_sheets"],
                    deduped_notes=info["deduped_notes"])
        _update(sid, xlsx_path=info["out_path"], state="추출됨",
                meta=json.dumps(meta, ensure_ascii=False))
        return {"xlsx_path": info["out_path"], "meta": meta}

    return {"job_id": jobs.submit("extract", _run)}


@router.get("/sessions/{sid}/sheets")
def list_sheets(sid: str):
    s = _session(sid)
    if not s["xlsx_path"] or not os.path.exists(s["xlsx_path"]):
        raise HTTPException(409, "extract 미실행 — 먼저 추출하세요")
    from openpyxl import load_workbook
    wb = load_workbook(s["xlsx_path"], read_only=True)
    names = [n for n in wb.sheetnames if not n.startswith("_")]
    wb.close()
    return {"sheets": names}


@router.get("/sessions/{sid}/sheets/{name}")
def sheet_grid(sid: str, name: str, max_rows: int = 400,
               max_cols: int = 12):
    s = _session(sid)
    if not s["xlsx_path"]:
        raise HTTPException(409, "extract 미실행")
    from openpyxl import load_workbook
    wb = load_workbook(s["xlsx_path"], read_only=True)
    if name not in wb.sheetnames:
        wb.close()
        raise HTTPException(404, "시트 없음")
    ws = wb[name]
    grid = []
    for r, row in enumerate(ws.iter_rows(max_row=max_rows,
                                         max_col=max_cols), 1):
        vals = [(c.value if c.value is not None else None) for c in row]
        if any(v is not None for v in vals):
            grid.append({"row": r, "cells": vals})
    wb.close()
    return {"name": name, "rows": grid, "truncated": len(grid) >= max_rows}


@router.post("/sessions/{sid}/diff")
def diff_session(sid: str, body: dict = None):
    """dry-run — 변경 목록. repack 승인 게이트의 입력."""
    s = _session(sid)
    if not s["xlsx_path"] or not os.path.exists(s["xlsx_path"]):
        raise HTTPException(409, "추출 미실행 — 수정 확인 대상 엑셀이 없습니다")
    options = (body or {}).get("options") or {}
    clean_cr = bool(options.get("clean_cr", True))
    from dsd_tool.repack import diff
    changes, _text, _data, stats = diff(s["xlsx_path"], s["dsd_path"],
                                        clean_cr=clean_cr)
    items = [{"id": i, "sheet": c["sheet"], "row": c["row"], "col": c["col"],
              "cell": f"R{c['row']}C{c['col']}",
              "before": c["old"], "after": c["new"], "reason": c["reason"]}
             for i, c in enumerate(changes)]
    counts = {
        "edit": sum(1 for c in items if c["reason"] == "edit"),
        "clean_cr": sum(1 for c in items if c["reason"] == "clean-cr"),
        "note_dedup": sum(1 for c in items if c["reason"] == "note-dedup"),
        "total": len(items),
    }
    payload = {"changes": items, "counts": counts, "stats": stats,
               "options": {"clean_cr": clean_cr},
               "ts": datetime.datetime.now().isoformat(timespec="seconds")}
    _update(sid, diff=json.dumps(payload, ensure_ascii=False),
            diff_options=json.dumps({"clean_cr": clean_cr}),
            state="수정중" if items else "추출됨")
    return payload


@router.post("/sessions/{sid}/repack", status_code=202)
def repack_session(sid: str, body: dict = None):
    s = _session(sid)
    body = body or {}
    if s["diff"] is None:
        # ★ 수정 확인이 반영 전 확인 관문 — 서버가 강제
        raise HTTPException(
            409, "수정 확인 미실행 — DSD에 반영하기 전에 수정 내용을 "
                 "확인·승인해야 합니다")
    options = body.get("options") or {}
    clean_cr = bool(options.get("clean_cr",
                                s["diff"]["options"]["clean_cr"]))
    if clean_cr != s["diff"]["options"]["clean_cr"]:
        raise HTTPException(409, "옵션이 수정 확인 시점과 다릅니다 — "
                                 "동일 옵션으로 다시 비교를 실행하세요")
    approved = body.get("approved_change_ids", "all")
    all_ids = [c["id"] for c in s["diff"]["changes"]]
    if approved != "all":
        if sorted(approved) != sorted(all_ids):
            raise HTTPException(
                422, "부분 승인 미지원 — 반영은 수정본 기준 전체 일괄 "
                     "적용입니다. 제외할 수정은 엑셀에서 되돌린 뒤 다시 "
                     "수정 확인을 실행하세요 (백엔드 수정 금지 원칙)")
    dsd, xlsx = s["dsd_path"], s["xlsx_path"]

    def _run(progress):
        progress("DSD에 반영 실행 중…")
        from dsd_tool.repack import repack
        info = repack(xlsx, dsd, clean_cr=clean_cr)
        result = {"output_path": info["out_path"],
                  "sha1": _sha1(info["out_path"]),
                  "n_changes": len(info["changes"]),
                  "checklist": _CHECKLIST}
        _update(sid, repack=json.dumps(result, ensure_ascii=False),
                state="반영완료")
        return result

    return {"job_id": jobs.submit("repack", _run)}


# --------------------------------------------------------------------------
# UI-2: Footing (A-4/A-5a) + 전기대사 (A-5b) — 기존 함수 얇은 래퍼
# --------------------------------------------------------------------------

def _run_foot(sid, xlsx, excel, limit, progress):
    """foot() 실행 + 세션 저장 (레벨 재검증과 공용)."""
    progress("Footing 검증 중…")
    from dsd_tool.foot import foot
    res = foot(xlsx, limit=limit)
    payload = {
        "summary": {"match": res["match"], "rounding": res["fuzzy"],
                    "mismatch": res["mismatch"],
                    "cross": res["note_missing"],
                    "note_found": res["note_found"],
                    "manual_overrides": res["manual_overrides"]},
        "findings": res["foot"], "levels": res["levels"],
        "notes": res["notes"],
        "ts": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    if excel:
        progress("AI_Footing 엑셀 생성 중…")
        from dsd_tool.foot_excel import write_ai_footing
        x = write_ai_footing(xlsx, res)
        payload["ai_excel_path"] = x["out_path"]
    _update(sid, foot=json.dumps(payload, ensure_ascii=False, default=str))
    return {"summary": payload["summary"],
            "ai_excel_path": payload.get("ai_excel_path")}


@router.post("/sessions/{sid}/foot", status_code=202)
def foot_session(sid: str, body: dict = None):
    s = _session(sid)
    if not s["xlsx_path"] or not os.path.exists(s["xlsx_path"]):
        raise HTTPException(409, "extract 미실행 — 먼저 추출하세요")
    body = body or {}
    excel = bool(body.get("excel"))
    limit = int(body.get("limit") or 2)
    xlsx = s["xlsx_path"]
    return {"job_id": jobs.submit(
        "foot", lambda p: _run_foot(sid, xlsx, excel, limit, p))}


@router.put("/sessions/{sid}/foot/levels", status_code=202)
def foot_levels(sid: str, body: dict):
    """레벨 오버라이드 → 재검증. 기존 관례 그대로 자동화 —
    _FOOT [레벨] 섹션의 '수동(수정 후 재실행)' 열에 기입 후 foot 재실행."""
    s = _session(sid)
    if not s["foot"]:
        raise HTTPException(409, "Footing 미실행 — 먼저 검증을 실행하세요")
    overrides = body.get("overrides") or []
    if not overrides:
        raise HTTPException(400, "overrides가 비었습니다")
    want = {(str(o["sheet"]), int(o["row"])):
            (int(o["level"]) if o.get("level") else None)
            for o in overrides}
    xlsx = s["xlsx_path"]

    def _run(progress):
        progress("레벨 오버라이드 기입…")
        from openpyxl import load_workbook

        from dsd_tool.foot import FOOT_SHEET
        wb = load_workbook(xlsx)
        ws = wb[FOOT_SHEET]
        in_sec = False
        applied = 0
        for row in ws.iter_rows():
            v0 = row[0].value
            if v0 == "[레벨]":
                in_sec = True
                continue
            if in_sec and isinstance(v0, str) and v0.startswith("["):
                break
            if in_sec and v0 and row[1].value is not None:
                try:
                    key = (str(v0), int(row[1].value))
                except (TypeError, ValueError):
                    continue
                if key in want:
                    ws.cell(row=row[0].row, column=5).value = want[key]
                    applied += 1
        wb.save(xlsx)
        progress(f"수동 레벨 {applied}건 기입 — 재검증…")
        out = _run_foot(sid, xlsx, False, 2, progress)
        out["applied"] = applied
        return out

    return {"job_id": jobs.submit("foot-levels", _run)}


@router.post("/sessions/{sid}/recon", status_code=202)
def recon_session(sid: str, body: dict):
    s = _session(sid)
    prior = body.get("prior_path") or body.get("prior_cache_ref") or ""
    if not os.path.isfile(prior):
        raise HTTPException(400, f"전기 파일 없음: {prior}")
    tolerance = float(body.get("tolerance") or 0)
    # 전기 소스 뱃지: dart_explorer 캐시 경유 여부 (파일 경로 기준)
    norm = os.path.normpath(prior).lower()
    source = "opendart-cache" if os.sep + "dart_explorer" + os.sep in norm \
        and os.sep + "cache" + os.sep in norm else "local"
    dsd = s["dsd_path"]
    os.makedirs(_WORKDIR, exist_ok=True)
    out = os.path.join(_WORKDIR, f"{sid}_전기대사.xlsx")

    def _run(progress):
        progress("전기대사 실행 중…")
        from dsd_tool.recon import GUIDE, recon
        res = recon(dsd, prior, out_path=out, tolerance=tolerance,
                    progress=lambda m: progress(m))
        body_rows = []
        for sheet, rows in (res.get("stmt_results") or {}).items():
            for r in rows:
                body_rows.append({
                    "sheet": sheet, "label": r.get("label"),
                    "cur": r.get("cur_prior"), "pri": r.get("pri_current"),
                    "true": bool(r.get("true")), "note": r.get("note", "")})
        note_rows = []
        for sheet, rows in (res.get("note_results") or {}).items():
            for r in rows:
                note_rows.append({
                    "sheet": sheet, "table": r.get("table"),
                    "label": r.get("label"), "cur": r.get("cur_prior"),
                    "pri": r.get("pri_current"),
                    "true": bool(r.get("true")), "note": r.get("note", "")})
        payload = {
            "verdict": (res["stmt"]["false"] + res["notes"]["false"]) == 0,
            "stmt": res["stmt"], "notes_summary": res["notes"],
            "note_matched": res["note_matched"],
            "note_total": res["note_total"],
            "guide": GUIDE, "excel_path": res["out_path"],
            "prior_path": prior, "source": source,
            "tolerance": tolerance, "body": body_rows, "notes": note_rows,
            "ts": datetime.datetime.now().isoformat(timespec="seconds"),
        }
        _update(sid, recon=json.dumps(payload, ensure_ascii=False,
                                      default=str))
        return {"verdict": payload["verdict"], "stmt": res["stmt"],
                "notes": res["notes"], "excel_path": res["out_path"]}

    return {"job_id": jobs.submit("recon", _run)}


@router.get("/sessions/{sid}/history")
def session_history(sid: str, limit: int = 20, changes: bool = False):
    s = _session(sid)
    from dsd_tool.history import query_changes, query_runs
    runs = query_runs(s["dsd_path"], limit=limit)
    if changes:
        for r in runs:
            r["changes"] = query_changes(r["id"])
    return {"runs": runs}


@router.get("/version-check")
def version_check(dsd_path: str):
    if not os.path.isfile(dsd_path):
        raise HTTPException(400, f"파일 없음: {dsd_path}")
    from dsd_tool.version import g2_smoke, is_known, read_version_info
    with open(dsd_path, "rb") as f:
        v = read_version_info(f.read())
    smoke = g2_smoke(dsd_path)
    return {"editver": v.get("editver"), "known": bool(is_known(
        v.get("editver"))),
        "g2_smoke": "pass" if (smoke["changes"] == 0 and
                               smoke["byte_identical"]) else "fail"}

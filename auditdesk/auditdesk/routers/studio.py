"""/api/studio — dart_explorer xbrl 모듈 + dsd_tool 매핑 (수신은 캐시 경유만).

라우터는 얇은 래퍼. 차원표·워크시트 화면 데이터는 기존 함수가 산출한
xlsx를 다시 읽어 그대로 서빙한다 — 화면과 파일이 어긋날 수 없는 구조.
"""
import datetime
import getpass
import json
import os
import sqlite3
import uuid

from fastapi import APIRouter, HTTPException

from .. import jobs

router = APIRouter()

_WORKDIR = os.path.join(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__))), "data", "studio")

_MAP_SCHEMA = """
CREATE TABLE IF NOT EXISTS mappings(
  id TEXT PRIMARY KEY, created TEXT, payload TEXT);
CREATE TABLE IF NOT EXISTS mapping_decisions(
  mapping_id TEXT, account_idx INTEGER, decided TEXT,
  decided_by TEXT, decided_at TEXT,
  PRIMARY KEY(mapping_id, account_idx));
"""


def _con():
    con = jobs.connect()
    con.executescript(_MAP_SCHEMA)
    return con


def _grid_from_xlsx(path, max_rows=400, max_cols=14):
    """산출 xlsx → 시트별 셀 그리드 (화면 = 파일 보장용 재독)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    out = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for r, row in enumerate(ws.iter_rows(max_row=max_rows,
                                             max_col=max_cols), 1):
            vals = [c.value for c in row]
            if any(v is not None for v in vals):
                rows.append({"row": r, "cells": vals})
        out.append({"sheet": name, "rows": rows})
    wb.close()
    return out


def _resolve_source(body):
    src = body.get("package_path")
    if not src and body.get("taxonomy_version"):
        from dart_explorer.xbrl.taxonomy_diff import resolve_version_dir
        src = resolve_version_dir(body["taxonomy_version"])
    if not src or not os.path.isdir(src):
        raise HTTPException(400, f"패키지 폴더 없음: {src}")
    return src


# --------------------------------------------------------------------------
# 트리 뷰 (D-1)
# --------------------------------------------------------------------------

@router.post("/taxtree", status_code=202)
def taxtree(body: dict):
    src = _resolve_source(body)
    role = body.get("role")
    os.makedirs(_WORKDIR, exist_ok=True)
    out = os.path.join(_WORKDIR, f"트리뷰_{uuid.uuid4().hex[:6]}.xlsx")

    def _run(progress):
        progress("택사노미 파싱 중…")
        from dart_explorer.xbrl.taxonomy import TaxonomyPackage, \
            build_tree_view
        res = build_tree_view(src, out_path=out, role_filter=role)
        pkg = TaxonomyPackage(src)
        roles = []
        for _uri, definition, pl in pkg.roles():
            entry = {"definition": definition}
            if role and role in definition:
                entry["rows"] = pkg.tree_rows(pl)[:600]
            roles.append(entry)
        if role is None and roles:
            # 필터 없으면 첫 role의 트리만 즉시 제공 (나머지는 재요청)
            for _uri, definition, pl in pkg.roles():
                roles[0]["rows"] = pkg.tree_rows(pl)[:600]
                break
        return {"xlsx_path": res["out_path"], "arcs": res["arcs"],
                "extensions": res["extensions"], "roles": roles}

    return {"job_id": jobs.submit("taxtree", _run)}


# --------------------------------------------------------------------------
# 차원 표 뷰어 (D-2/D-2b) — 화면 = 산출 xlsx 재독
# --------------------------------------------------------------------------

@router.post("/dimtable", status_code=202)
def dimtable(body: dict):
    src = _resolve_source(body)
    role = body.get("role")
    os.makedirs(_WORKDIR, exist_ok=True)
    out = os.path.join(_WORKDIR, f"차원표_{uuid.uuid4().hex[:6]}.xlsx")

    def _run(progress):
        progress("차원 표 렌더링 중…")
        from dart_explorer.xbrl.dimension_table import \
            render_dimension_tables
        res = render_dimension_tables(src, out_path=out, role_filter=role)
        progress("산출 xlsx 재독 (화면 = 파일)…")
        return {"xlsx_path": res["out_path"], "roles": res["roles"],
                "grids": _grid_from_xlsx(res["out_path"])}

    return {"job_id": jobs.submit("dimtable", _run)}


# --------------------------------------------------------------------------
# 매핑 확정 (D-3b + 상태 4종 + 확정 기록)
# --------------------------------------------------------------------------

@router.post("/mapping", status_code=202)
def mapping(body: dict):
    accounts = body.get("accounts") or []
    if not accounts:
        raise HTTPException(400, "accounts가 비었습니다 "
                                 "[{name, category?, value?}]")
    induty = body.get("induty")
    mapping_id = uuid.uuid4().hex[:10]

    def _run(progress):
        from dsd_tool.mapping import (STATE_EXTENSION, STATE_MANUAL,
                                      STATE_PREFER_STANDARD, STATE_STANDARD,
                                      MappingCorpus, classify_suggestion,
                                      suggest)
        state_map = {STATE_STANDARD: "normal",
                     STATE_PREFER_STANDARD: "standard_recommended",
                     STATE_EXTENSION: "extension_needed",
                     STATE_MANUAL: "manual"}
        corpus = MappingCorpus()
        items = []
        for i, a in enumerate(accounts):
            if i % 5 == 0:
                progress(f"추천 생성 {i}/{len(accounts)}…")
            res = suggest(corpus, a["name"], a.get("category"),
                          induty=induty)
            state, _top = classify_suggestion(res)
            items.append({
                "account": a["name"], "category": a.get("category"),
                "value": a.get("value"),
                "state": state_map.get(state, "manual"),
                "state_label": state,
                "verdict": res["verdict"],
                "candidates": [
                    {"element": c["element_id"], "label": c["std_label"],
                     "score": c["score"], "similarity": c["sim"],
                     "firms": c["n_companies"],
                     "industry_boost": c.get("n_same_induty", 0),
                     "evidence": c.get("evidence", "")}
                    for c in (res.get("candidates") or [])[:4]],
                "similar_extensions": [
                    {"label": e["label"], "firms": e["n_companies"],
                     "similarity": e["sim"]}
                    for e in (res.get("similar_extensions") or [])[:3]],
            })
        payload = {"mapping_id": mapping_id, "items": items,
                   "induty": induty,
                   "ts": datetime.datetime.now().isoformat(
                       timespec="seconds")}
        with _con() as con:
            con.execute("INSERT OR REPLACE INTO mappings VALUES(?,?,?)",
                        (mapping_id, payload["ts"],
                         json.dumps(payload, ensure_ascii=False)))
        return {"mapping_id": mapping_id, "n": len(items)}

    return {"job_id": jobs.submit("mapping", _run)}


@router.get("/mapping/{mapping_id}")
def get_mapping(mapping_id: str):
    with _con() as con:
        row = con.execute("SELECT payload FROM mappings WHERE id=?",
                          (mapping_id,)).fetchone()
        if row is None:
            raise HTTPException(404, "매핑 작업 없음")
        payload = json.loads(row[0])
        for idx, decided, by, at in con.execute(
                "SELECT account_idx, decided, decided_by, decided_at "
                "FROM mapping_decisions WHERE mapping_id=?", (mapping_id,)):
            if 0 <= idx < len(payload["items"]):
                payload["items"][idx]["decided"] = json.loads(decided)
                payload["items"][idx]["decided_by"] = by
                payload["items"][idx]["decided_at"] = at
    return payload


@router.put("/mapping/{mapping_id}/decide")
def decide(mapping_id: str, body: dict):
    """확정 기록 — 누가/언제. 자동 확정 없음(회계사 클릭이 유일 경로)."""
    idx = body.get("account_idx")
    if idx is None:
        raise HTTPException(400, "account_idx 필요")
    decided = {"element": body.get("element"),
               "extension": body.get("extension")}
    if not decided["element"] and not decided["extension"]:
        raise HTTPException(400, "element 또는 extension 필요")
    by = getpass.getuser()
    at = datetime.datetime.now().isoformat(timespec="seconds")
    with _con() as con:
        if con.execute("SELECT 1 FROM mappings WHERE id=?",
                       (mapping_id,)).fetchone() is None:
            raise HTTPException(404, "매핑 작업 없음")
        con.execute(
            "INSERT OR REPLACE INTO mapping_decisions VALUES(?,?,?,?,?)",
            (mapping_id, int(idx), json.dumps(decided, ensure_ascii=False),
             by, at))
    return {"account_idx": idx, "decided": decided,
            "decided_by": by, "decided_at": at}


# --------------------------------------------------------------------------
# 택사노미 체크 (D-4c)
# --------------------------------------------------------------------------

@router.post("/taxcheck", status_code=202)
def taxcheck(body: dict):
    prior = body.get("prior_package")
    against = body.get("against_version")
    if not prior or not os.path.isdir(prior):
        raise HTTPException(400, f"전기 패키지 폴더 없음: {prior}")
    if not against:
        raise HTTPException(400, "against_version 필요")
    # UI 계약: 승격 감지(수십 분)는 기본 생략 — 별도 실행 버튼으로 요청
    skip_promotions = bool(body.get("skip_promotions", True))
    os.makedirs(_WORKDIR, exist_ok=True)
    out = os.path.join(_WORKDIR, f"taxcheck_{uuid.uuid4().hex[:6]}.xlsx")

    def _run(progress):
        progress("신버전 택사노미 대조 중…")
        from dart_explorer.xbrl.taxonomy_diff import run_taxcheck
        res = run_taxcheck(prior, against, out_path=out,
                           skip_promotions=skip_promotions,
                           progress=lambda m: progress(m))
        groups = {"green": [], "yellow": [], "blue": [], "ext": []}
        for r in res["rows"]:
            key = {"녹색": "green", "노랑": "yellow", "파랑": "blue",
                   "확장(회사고유)": "ext"}.get(r["status"])
            if key:
                groups[key].append({
                    "element": r["element_id"], "label": r["label_ko"],
                    "detail": r["detail"],
                    "candidates": [
                        {"element": c["element_id"],
                         "label": c.get("std_label", ""),
                         "score": c.get("score"),
                         "evidence": c.get("evidence", "")}
                        for c in (r.get("candidates") or [])[:4]]})
        return {"xlsx_path": res["out_path"], "counts": {
            "green": res["green"], "yellow": res["yellow"],
            "blue": res["blue"], "ext": len(groups["ext"])},
            **groups,
            "skip_promotions": skip_promotions,
            "promotions": len(res.get("promotions") or [])
            if not skip_promotions else None}

    return {"job_id": jobs.submit("taxcheck", _run)}


# --------------------------------------------------------------------------
# 작성 워크시트 (F-1/F-2/F-3 — 전 모드 활성)
# --------------------------------------------------------------------------

@router.post("/worksheet", status_code=202)
def worksheet(body: dict):
    dsd = body.get("dsd_path") or ""
    if not os.path.isfile(dsd):
        raise HTTPException(400, f"DSD 파일 없음: {dsd}")
    report = body.get("report") or "annual"
    mode = body.get("mode") or "new"
    include_notes = bool(body.get("include_notes", True))
    induty = body.get("induty")
    inherit_pkg = body.get("inherit_package")
    if mode == "inherit":
        if not inherit_pkg or not os.path.isdir(inherit_pkg):
            raise HTTPException(400, "승계 모드는 자기 기말 XBRL 패키지 "
                                     f"폴더가 필요합니다: {inherit_pkg}")
    os.makedirs(_WORKDIR, exist_ok=True)
    out = os.path.join(
        _WORKDIR, f"워크시트_{uuid.uuid4().hex[:6]}_{report}_{mode}.xlsx")

    def _run(progress):
        succession = None
        if mode == "inherit":
            sj = os.path.join(inherit_pkg, "succession_assets.json")
            if not os.path.exists(sj):
                progress("승계 자산 내보내기 (기말 인스턴스 → JSON)…")
                from dart_explorer.xbrl.corpus import \
                    export_succession_assets
                from dart_explorer.xbrl.taxonomy_diff import TAXONOMY_DIR
                versions = sorted(
                    d for d in os.listdir(TAXONOMY_DIR)
                    if os.path.isdir(os.path.join(TAXONOMY_DIR, d))
                ) if os.path.isdir(TAXONOMY_DIR) else []
                export_succession_assets(
                    inherit_pkg, against=versions[-1] if versions else None)
            from dsd_tool.succession import load_assets
            succession = load_assets(sj)
        progress("워크시트 생성 중… (주석 포함 시 수 분)")
        from dsd_tool.worksheet import build_worksheet
        res = build_worksheet(dsd, out_path=out, report_type=report,
                              induty=induty, include_notes=include_notes,
                              succession=succession,
                              progress=lambda m: progress(m))
        st = res["stats"]
        notes = res.get("notes") or {}
        progress("산출 xlsx 재독 (화면 = 파일)…")
        return {
            "xlsx_path": res["out_path"],
            "summary": {
                "mapped": st["mapped"],
                "unmapped": st["rows"] - st["mapped"] - st["skipped"],
                "extension_candidates": st["extension"],
                "rows": st["rows"], "inherited": st.get("inherited", 0),
                "new_accounts": st.get("new", 0),
                "deprecated": st.get("deprecated", 0),
                "mode": mode, "report": report,
                "note_unassigned": len(notes.get("unassigned", []))
                if isinstance(notes, dict) and "unassigned" in notes
                else None,
            },
            "grids": _grid_from_xlsx(res["out_path"]),
        }

    return {"job_id": jobs.submit("worksheet", _run)}


# --------------------------------------------------------------------------
# 보유 택사노미 (D-4a) + taxdiff
# --------------------------------------------------------------------------

@router.get("/taxonomies")
def taxonomies():
    from dart_explorer.xbrl.taxonomy_diff import TAXONOMY_DIR
    if not os.path.isdir(TAXONOMY_DIR):
        return {"versions": []}
    return {"versions": sorted(
        d for d in os.listdir(TAXONOMY_DIR)
        if os.path.isdir(os.path.join(TAXONOMY_DIR, d)))}


@router.post("/taxdiff", status_code=202)
def taxdiff_route(body: dict):
    old, new = body.get("old"), body.get("new")
    if not old or not new:
        raise HTTPException(400, "old/new 버전 필요")

    def _run(progress):
        progress("세대 간 diff 중…")
        from dart_explorer.xbrl.taxonomy_diff import taxdiff
        res = taxdiff(old, new, progress=lambda m: progress(m))
        return {"xlsx_path": res["out_path"],
                "new_version": res["new_version"],
                "old_version": res["old_version"],
                "added": len(res.get("added", [])),
                "removed": len(res.get("removed", [])),
                "changed": len(res.get("changed", [])),
                "kept": res.get("kept")}

    return {"job_id": jobs.submit("taxdiff", _run)}

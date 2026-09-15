"""Product assembly for existing M-1, V-2 and F-3b cores."""
import os
import re
import tempfile
from zipfile import BadZipFile

from openpyxl import load_workbook


def input_file(path, label):
    if not isinstance(path, str) or not os.path.isfile(path):
        raise ValueError(f'{label}: 존재하는 입력 파일을 지정하세요')
    return path


def package_facts(package):
    if not isinstance(package, str) or not os.path.isdir(package):
        raise ValueError('제출용 XBRL 패키지 폴더를 지정하세요 (IXD 아님)')
    from dart_explorer.xbrl.dimension_table import XbrlInstance
    inst = XbrlInstance(package)
    if not inst.facts or inst.doc_period_end is None:
        raise ValueError('XBRL 팩트 또는 보고기간말이 없습니다 — 제출파일을 확인하세요')
    facts = {eid: [dict(value=f['value'], decimals=f.get('decimals'),
                       unit=f.get('unit', ''), **f['ctx']) for f in fl]
             for eid, fl in inst.facts.items()}
    return facts, str(inst.doc_period_end)


def validate_pairs(left, right, pairs):
    input_file(left, '좌측 원문 Excel'); input_file(right, '우측 공시화면 Excel')
    if not isinstance(pairs, list) or not pairs:
        raise ValueError('pairs: 좌·우 시트명과 조서 시트명 목록이 필요합니다')
    books = []
    try:
        for p in (left, right):
            books.append(load_workbook(p, read_only=True))
        names = {'총괄표'}
        for pair in pairs:
            if not isinstance(pair, dict):
                raise ValueError('pairs: 각 항목은 left/right/name 객체여야 합니다')
            name = pair.get('name')
            if (not isinstance(name, str) or not name.strip() or len(name) > 31
                    or re.search(r'[\\/*?:\[\]]', name) or name.casefold() in names):
                raise ValueError('pairs.name: 중복 없는 유효한 Excel 시트명(31자 이내)이 필요합니다')
            names.add(name.casefold())
            if not pair.get('left') and not pair.get('right'):
                raise ValueError('pairs: 적어도 한쪽 시트를 지정하세요')
            for side, wb in zip(('left', 'right'), books):
                sheet = pair.get(side)
                if sheet is not None and (not isinstance(sheet, str) or sheet not in wb.sheetnames):
                    raise ValueError(f'pairs.{side}: 입력 Excel에 없는 시트입니다')
    except BadZipFile:
        raise ValueError('Excel 파일을 읽을 수 없습니다 — 손상되지 않은 .xlsx 파일을 지정하세요') from None
    finally:
        for wb in books:
            wb.close()


def attr_workbook(xlsx, package, out, progress):
    input_file(xlsx, '원문 extract Excel')
    progress('제출파일 기간·unit·decimals·주석명 검사 중…')
    facts, end = package_facts(package)
    from dart_explorer.xbrl.taxonomy import TaxonomyPackage
    from dsd_tool.taxonomy_labels import get_resolver, normalize_concept
    from dsd_tool.attr_check import attr_check, collect_note_titles
    from dsd_tool.foot import FootingContext
    tp = TaxonomyPackage(package)
    resolver = get_resolver()
    ctx = FootingContext(xlsx)
    try:
        titles = collect_note_titles(ctx)
    finally:
        ctx.wb.close()
    result = attr_check(facts,
        lambda eid: resolver.attrs.get(normalize_concept(eid)) or tp.ext_attrs.get(normalize_concept(eid)),
        tp.role_defs, titles, end,
        std_label=lambda eid: resolver.resolve(eid)['ko'] or '(확장)',
        out_path=out, progress=progress)
    return result


def run_rollforward(half, package, year_end, out, current=None, progress=print):
    """DSD/extract inputs are adapted in a private directory; sources stay intact."""
    for path, label in ((half, '① 전기 반기 DSD/Excel'), (year_end, '③ 전기 기말 DSD/Excel')):
        input_file(path, label)
    if current:
        input_file(current, '갱신용 당기 DSD/Excel')
    if os.path.exists(out):
        raise ValueError('산출 경로에 파일이 있습니다 — 새 파일명을 지정하세요')
    if not os.path.isdir(os.path.dirname(os.path.abspath(out))):
        raise ValueError('산출 폴더가 없습니다 — 존재하는 폴더를 지정하세요')
    progress('② 기말 XBRL 팩트와 승계 자산 준비…')
    facts, end = package_facts(package)
    from dart_explorer.xbrl.corpus import export_succession_assets
    from dsd_tool.succession import load_assets
    from dsd_tool.excel_out import extract
    from dsd_tool.rollforward import rollforward
    from dsd_tool.taxonomy_labels import get_resolver
    corpus = None
    def recommend(label):
        nonlocal corpus
        from dsd_tool.mapping import MappingCorpus, suggest
        if corpus is None:
            corpus = MappingCorpus()
        return [c['element_id'] for c in suggest(corpus, label)['candidates']]
    with tempfile.TemporaryDirectory(prefix='auditdesk_rollforward_') as tmp:
        def excel(path, name):
            if path.lower().endswith('.dsd'):
                target = os.path.join(tmp, name + '.xlsx')
                extract(path, target)
                return target
            if not path.lower().endswith('.xlsx'):
                raise ValueError('입력은 DSD 또는 extract Excel(.xlsx)이어야 합니다')
            return path
        sj = os.path.join(package, 'succession_assets.json')
        if not os.path.isfile(sj):
            sj = os.path.join(tmp, 'succession_assets.json')
            export_succession_assets(package, out_json=sj)
        result = rollforward(excel(half, 'half'), facts, end, load_assets(sj),
            ye_xlsx=excel(year_end, 'year_end'), resolver=get_resolver(),
            current_xlsx=excel(current, 'current') if current else None,
            recommend=recommend if current else None,
            out_path=out, progress=progress,
            source_warning='원천 자료의 회사·기간·연결/별도 범위를 확인하세요. 미매칭은 수동 확인입니다.')
    return result

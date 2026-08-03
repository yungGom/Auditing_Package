"""V-1: DSD ↔ XBRL 인스턴스 대사 (인벡터 기능 ② 대응).

같은 회사의 DSD 본문 값 ↔ 인스턴스 팩트 값 자동 대조 — 태깅이 공시
본문과 일치하는지, 제출 직전 최종 검증. 조립 패치:
- 팩트는 호출자(라우터 조립층)가 XbrlInstance로 파싱해 dict로 전달
  (import 방향 불변 — 이 모듈은 dart_explorer를 모른다)
- element 매핑: F-1 확정(decided) → F-3 승계(SuccessionAssets) 순 재사용,
  신규 추론·강제 매칭 금지 — 없으면 '매핑 없음' 별도 집계 (미매칭≠0)
- 판정·엑셀: 전기대사(A-5) 규격 — TRUE/FALSE + 비고, FALSE 하이라이트

스코프: 본문 FS (BS/PL/PL1/CF — CE는 자본항목 축이 필요해 제외, 비고 명시).
"""
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .foot import FootingContext, _label
from .mapping import normalize as _norm

_BOLD = Font(bold=True)
_HDR_FILL = PatternFill("solid", start_color="D9E1F2")
_FALSE_FILL = PatternFill("solid", start_color="FFC7CE")
_WARN_FILL = PatternFill("solid", start_color="FFEB9C")
_LINK_FONT = Font(color="FF0563C1", underline="single")
_NUMFMT = "#,##0;(#,##0)"

CONSOL_AXIS = "ifrs-full_ConsolidatedAndSeparateFinancialStatementsAxis"
_MEMBER = {"연결": "ifrs-full_ConsolidatedMember",
           "별도": "ifrs-full_SeparateMember"}

V_MATCH, V_DIFF = "일치", "값 상이"
V_NOFACT, V_ONLYX, V_NOMAP = "태깅 누락", "제출파일에만 있음", "매핑 없음"

GUIDE = ("FALSE 존재 시 태깅 값·단위·문맥(연결/별도)을 확인하십시오 "
         "(판정은 기계, 해석은 회계사)")

def _resolver():
    """D-2c 리졸버 재사용 — 배포 엑셀 부재 시 표준레이블 열만 생략."""
    try:
        from .taxonomy_labels import get_resolver
        return get_resolver()
    except Exception:
        return None


def _std_label(resolver, eid):
    if not resolver or not eid:
        return ""
    hit = resolver.resolve(eid)
    return hit["ko"] if hit["standard"] else "(확장)"


_UNIT_RE = re.compile(r"단\s*위\s*[:：]\s*([^\)\s]+)")
_UNIT_SCALE = {"원": 1, "천원": 1_000, "천 원": 1_000,
               "백만원": 1_000_000, "백만 원": 1_000_000,
               "십억원": 1_000_000_000, "억원": 100_000_000}


def _detect_unit(ws, max_scan=8):
    """시트 상단의 '(단위 : 백만원)' 행 → 원 환산 배수."""
    for r in range(1, max_scan + 1):
        for c in range(1, 4):
            m = _UNIT_RE.search(str(ws.cell(r, c).value or ""))
            if m:
                unit = m.group(1).strip()
                for k, s in _UNIT_SCALE.items():
                    if unit.startswith(k.replace(" ", "")[0]) and \
                            k.replace(" ", "") in unit.replace(" ", ""):
                        return s, unit
                return 1, unit                  # 미인식 단위 → 원 가정
    return 1, "원(표기 없음)"


def _current_fact(facts, element_id, doc_end, member, kind):
    """당기 보고기간 순수 문맥 팩트 선별.

    kind: 'instant'(BS류) | 'duration'(PL/CF류 — end=doc_end 중 최장).
    연결/별도 축 member 일치 또는 무차원만 허용, 그 외 축 문맥 배제
    (D-2c 교훈 — 섹션 외 축 유입 차단).
    """
    best = None
    for f in facts.get(element_id, []):
        dims = f.get("dims") or {}
        extra = [a for a in dims if a != CONSOL_AXIS]
        if extra:
            continue
        if CONSOL_AXIS in dims and dims[CONSOL_AXIS] != member:
            continue
        if kind == "instant":
            if f.get("type") == "instant" and f.get("end") == doc_end:
                best = best or f
        else:
            if f.get("type") == "duration" and f.get("end") == doc_end:
                if best is None or (f.get("start") or "") < \
                        (best.get("start") or ""):
                    best = f                    # 최장 기간(연간 누적)
    return best


def _current_dim_fact(facts, element_id, doc_end, member):
    """당기 보고기간의 차원(멤버 조합) 팩트 — 무차원 부재 시 표기용.

    판정에는 쓰지 않는다(본문 대사는 순수 문맥만). '인스턴스에만 있음'
    집계에서 무차원/차원을 나눠 보여주기 위한 탐지 전용.
    """
    for f in facts.get(element_id, []):
        dims = f.get("dims") or {}
        extra = sorted(a for a in dims if a != CONSOL_AXIS)
        if not extra:
            continue
        if CONSOL_AXIS in dims and dims[CONSOL_AXIS] != member:
            continue
        if f.get("end") == doc_end:
            combo = " × ".join(
                str(dims[a]).rsplit("_", 1)[-1] for a in extra)
            return f, combo
    return None, None


def _tolerance(fact, scale, override):
    if override is not None:
        return float(override)
    tol = 0.5 * scale                           # DSD 표시단위 반올림
    dec = str(fact.get("decimals") or "") if fact else ""
    try:
        d = int(float(dec))
        if d < 0:
            tol = max(tol, 0.5 * (10 ** -d))    # 인스턴스 decimals 기반
    except (ValueError, OverflowError):
        # decimals 미표기/INF — 팩트 값 자체의 10^k 배수성으로 반올림
        # 단위 추정 (실측: 천원 반올림 값에 decimals 누락 관행)
        try:
            fv = int(float(fact["value"]))
            k = 0
            while k < 6 and fv and fv % (10 ** (k + 1)) == 0:
                k += 1
            if k:
                tol = max(tol, 0.5 * (10 ** k))
        except (ValueError, TypeError, KeyError):
            pass
    return tol


def xbrl_recon(xlsx_path, facts, doc_end, decided=None, succession=None,
               body_elements=None, tolerance=None, out_path=None,
               source_warning=None, progress=None):
    """DSD 편집용 xlsx ↔ 인스턴스 팩트 대사. 요약 dict 반환 (+엑셀).

    facts: {element_id: [{value, decimals, type, start, end, dims}]}
    decided: {정규화 계정명: element_id} — F-1 확정 기록
    succession: SuccessionAssets — F-3 승계 (없으면 None)
    body_elements: 본문 role(D2~D6) 소속 element 집합 —
      '인스턴스에만 있음' 판정 범위
    """
    doc_end = str(doc_end)                      # date 객체 → ISO 문자열
    ctx = FootingContext(xlsx_path)
    decided = decided or {}
    rows_out = {}
    used_elements = set()
    counts = {V_MATCH: 0, V_DIFF: 0, V_NOFACT: 0, V_NOMAP: 0}
    tol_used = set()

    for sheet in ctx.fs_sheets:
        if sheet.endswith("CE"):
            continue                            # 자본항목 축 필요 — 스코프 외
        ws = ctx.wb[sheet]
        scale, unit_txt = _detect_unit(ws)
        member = _MEMBER["연결" if "연결" in sheet else "별도"]
        kind = "instant" if sheet.endswith("BS") else "duration"
        periods, data_rows = ctx.fs_sequences(sheet)
        cur = dict(next((s for n, s in periods if n == "당기"), []))
        if not cur:                             # CE류 colN 구조 등
            continue
        out = []
        for r in data_rows:
            label = _label(ws, r)
            v = cur.get(r)
            if not label or v is None:
                continue
            qn = _norm(label)
            # ── element 매핑: F-1 확정 → F-3 승계 (신규 추론 금지) ──
            eid, src = decided.get(qn), "확정"
            if eid is None and succession is not None:
                # V-1 자동 판정은 워크시트 제안보다 엄격 — 유사 폴백
                # 0.85 미만은 매핑 없음으로 정직 분류 (강제 매칭 금지;
                # 실측: 0.73 '장기투자자산'→Assets 같은 오매핑이 값
                # 상이·태깅 누락 오탐을 만든다)
                inh = succession.inherit(label, min_sim=0.85)
                if inh is not None:
                    eid, src = inh["element_id"], "승계"
                    if inh.get("sim", 1.0) < 1.0:
                        src = f"승계(유사 {inh['sim']:.2f})"
            row = {"row": r, "label": label.strip(), "dsd": v,
                   "won": v * scale, "unit": unit_txt,
                   "element": eid, "map_src": src if eid else "",
                   "fact": None, "diff": None}
            if eid is None:
                row["verdict"] = V_NOMAP
                row["true"] = None              # 판정 불가 — 별도 집계
                counts[V_NOMAP] += 1
            else:
                used_elements.add(eid)
                f = _current_fact(facts, eid, doc_end, member, kind)
                if f is None:
                    row["verdict"] = V_NOFACT
                    row["true"] = False
                    counts[V_NOFACT] += 1
                else:
                    fv = float(f["value"])
                    tol = _tolerance(f, scale, tolerance)
                    tol_used.add(tol)
                    row["fact"] = fv
                    row["diff"] = row["won"] - fv
                    ok = abs(row["diff"]) <= tol
                    row["verdict"] = V_MATCH if ok else V_DIFF
                    row["true"] = ok
                    counts[V_MATCH if ok else V_DIFF] += 1
            out.append(row)
        if out:
            rows_out[sheet] = {"rows": out, "unit": unit_txt,
                               "member": member, "kind": kind}
        if progress:
            progress(f"  [{sheet}] {len(out)}행 대사")

    # ── 인스턴스에만 있음: 본문 role 소속 + 당기 팩트 + 미사용 ──
    # 무차원(순수 문맥) 우선, 없으면 차원(멤버 조합) 팩트로 분리 집계
    # (판정 로직 불변 — 본문 대사는 여전히 순수 문맥만, 표기 구분 전용)
    only_inst = []
    for eid in sorted(body_elements or []):
        if eid in used_elements:
            continue
        entry = None
        for member in set(_MEMBER.values()):
            for kind in ("instant", "duration"):
                f = _current_fact(facts, eid, doc_end, member, kind)
                if f is not None:
                    entry = {"element": eid, "fact": float(f["value"]),
                             "member": member.split("_")[-1],
                             "dim": None}
                    break
            if entry:
                break
        if entry is None:
            for member in set(_MEMBER.values()):
                f, combo = _current_dim_fact(facts, eid, doc_end, member)
                if f is not None:
                    entry = {"element": eid, "fact": float(f["value"]),
                             "member": member.split("_")[-1],
                             "dim": combo}
                    break
        if entry is not None:
            only_inst.append(entry)

    matched = counts[V_MATCH] + counts[V_DIFF] + counts[V_NOFACT]
    total = matched + counts[V_NOMAP]
    oi_dim = sum(1 for x in only_inst if x["dim"])
    summary = {
        "counts": dict(counts),
        "only_instance": len(only_inst) - oi_dim,   # 무차원(기존 집계)
        "only_instance_dim": oi_dim,                # 차원(멤버 조합)
        "matched": matched, "total": total,
        "match_rate": round(matched / total, 4) if total else None,
        "true": counts[V_MATCH],
        "false": counts[V_DIFF] + counts[V_NOFACT],
        "tolerance": ("수동 " + str(tolerance)) if tolerance is not None
        else "자동(decimals·표시단위): " + ", ".join(
            f"±{t:,.0f}" for t in sorted(tol_used)[:4]),
        "doc_end": doc_end, "source_warning": source_warning,
        "sheets": {s: len(d["rows"]) for s, d in rows_out.items()},
    }
    if out_path:
        _write_excel(out_path, rows_out, only_inst, summary)
        summary["out_path"] = out_path
    summary["rows"] = {s: d["rows"] for s, d in rows_out.items()}
    return summary


def _write_excel(out_path, rows_out, only_inst, summary):
    """대사표 엑셀 — A-5 규격 (요약 + 시트별 상세 + FALSE 하이라이트)."""
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "요약"
    ws0.append(["DSD ↔ XBRL 제출파일 태깅 대사"])
    ws0.cell(1, 1).font = Font(bold=True, size=14)
    if summary.get("source_warning"):
        ws0.append([summary["source_warning"]])
        ws0.cell(ws0.max_row, 1).fill = _WARN_FILL
    ws0.append([GUIDE])
    ws0.cell(ws0.max_row, 1).fill = _WARN_FILL
    ws0.cell(ws0.max_row, 1).font = _BOLD
    ws0.append([])
    c = summary["counts"]
    ws0.append([f"대조율: {summary['matched']}/{summary['total']} "
                f"({(summary['match_rate'] or 0):.1%}) · "
                f"허용오차 {summary['tolerance']} · "
                f"보고기간말 {summary['doc_end']}"])
    ws0.cell(ws0.max_row, 1).font = _BOLD
    ws0.append([f"판정 분포: 일치 {c[V_MATCH]} / 값 상이 {c[V_DIFF]} / "
                f"태깅 누락 {c[V_NOFACT]} / 제출파일에만 있음 "
                f"{summary['only_instance']}"
                f"(+차원 {summary.get('only_instance_dim', 0)})"
                f" / 매핑 없음 {c[V_NOMAP]}"])
    ws0.append([])
    ws0.append(["시트", "대사", "TRUE", "FALSE", "매핑 없음", "링크"])
    for cell in ws0[ws0.max_row]:
        cell.font = _BOLD
        cell.fill = _HDR_FILL
    for sheet, d in rows_out.items():
        rows = d["rows"]
        t = sum(1 for r in rows if r["true"] is True)
        f = sum(1 for r in rows if r["true"] is False)
        nm = sum(1 for r in rows if r["verdict"] == V_NOMAP)
        ws0.append([sheet, len(rows), t, f, nm,
                    f'=HYPERLINK("#\'{sheet}\'!A1","시트 바로가기")'])
        ws0.cell(ws0.max_row, 6).font = _LINK_FONT
        if f:
            ws0.cell(ws0.max_row, 4).fill = _FALSE_FILL
    ws0.column_dimensions["A"].width = 44
    for col in "BCDEF":
        ws0.column_dimensions[col].width = 13

    resolver = _resolver()
    headers = ["행 라벨", "DSD 값", "원 환산", "태깅 요소 ID",
               "한글 표준레이블", "매핑 근거", "제출파일 값", "차이",
               "판정", "비고"]
    for sheet, d in rows_out.items():
        ws = wb.create_sheet(sheet)
        ws.append([f"{sheet} — 표시단위 {d['unit']} · "
                   f"{d['member'].split('_')[-1]} · {d['kind']}"])
        ws.cell(1, 1).font = _BOLD
        ws.append(headers)
        for cell in ws[2]:
            cell.font = _BOLD
            cell.fill = _HDR_FILL
        for r in d["rows"]:
            note = "" if r["true"] else (
                "계정 매핑 확정·기말 승계 기록 모두 부재 — 판정 불가"
                if r["verdict"] == V_NOMAP else
                "당기 문맥 팩트 없음 — 태깅 확인"
                if r["verdict"] == V_NOFACT else
                f"차이 {r['diff']:,.0f}원"
                + (" · 가이드 5.Ⅱ.3(1)아 확인(유동/비유동 구분은"
                   " 축이 아닌 행)"
                   if ("유동" in r["label"] or
                       "Current" in (r["element"] or "")) else ""))
            ws.append([r["label"], r["dsd"], r["won"],
                       (r["element"] or "").replace("_", ":", 1),
                       _std_label(resolver, r["element"]),
                       r["map_src"], r["fact"], r["diff"],
                       r["verdict"] if r["true"] is None
                       else ("TRUE" if r["true"] else "FALSE"), note])
            for ci in (2, 3, 7, 8):
                ws.cell(ws.max_row, ci).number_format = _NUMFMT
            if r["true"] is False:
                for cell in ws[ws.max_row]:
                    cell.fill = _FALSE_FILL
            elif r["true"] is None:
                ws.cell(ws.max_row, 9).fill = _WARN_FILL
        ws.column_dimensions["A"].width = 40
        ws.column_dimensions["D"].width = 44
        for col in ("B", "C", "G", "H"):
            ws.column_dimensions[col].width = 16
        for col in ("E", "F", "I", "J"):
            ws.column_dimensions[col].width = 22
        ws.freeze_panes = "A3"

    ws = wb.create_sheet("제출파일에만_매핑없음")
    ws.append(["① 제출파일에만 있음 — 본문 구획 소속 당기 값 중 DSD "
               "미대응 (무차원/차원 구분 — 차원 값은 표기 전용,"
               " 판정 대상 아님)"])
    ws.cell(1, 1).font = _BOLD
    ws.append(["element", "한글 표준레이블", "구분", "연결/별도", "팩트 값"])
    for cell in ws[2]:
        cell.font = _BOLD
        cell.fill = _HDR_FILL
    for x in only_inst:
        ws.append([x["element"].replace("_", ":", 1),
                   _std_label(resolver, x["element"]),
                   f"차원({x['dim']})" if x.get("dim") else "무차원",
                   x["member"], x["fact"]])
        ws.cell(ws.max_row, 5).number_format = _NUMFMT
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["E"].width = 18
    wb.save(out_path)

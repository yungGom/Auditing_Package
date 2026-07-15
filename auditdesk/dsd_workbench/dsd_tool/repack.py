"""편집된 Excel → DSD 역변환.

_MAP과 원본 contents.xml을 비교해 변경된 셀만 오프셋 역순으로 교체한다.
거짓변경 방지 규칙: 스펙 3.4. 출력은 {원본명}_수정.dsd (원본 덮어쓰기 금지).
"""
import hashlib
import os
import re

from openpyxl import load_workbook

from .excel_out import MAP_SHEET, META_SHEET
from .textutil import (clean_text, dedup_note_number, enclosing_tag,
                       escape_for_dsd, format_number_like, is_cr_only_cell,
                       numbers_equal, try_number)
from .zipsplice import read_contents, replace_contents


class RepackError(Exception):
    pass


def _cell_value_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value)


def _normalize_for_compare(s: str) -> str:
    return s.replace("\r\n", "\n").replace("\r", "\n").strip()


def diff(xlsx_path: str, dsd_path: str, clean_cr: bool = True):
    """변경 셀 목록 계산. (changes, text, data, stats) 반환.

    change: dict(sheet,row,col,xml_start,xml_end,old,new,new_raw,reason)
      reason: "edit"(사용자 수정) | "clean-cr"(&cr;-only 셀 정리)
              | "note-dedup"(주석 번호 중복 정리)
    stats: {"cr_only_total": 원본의 &cr;-only 셀 수}

    clean_cr(기본 True): 내용이 &amp;cr; 반복뿐인 TD/TH/TE/TU 셀을 빈 문자열로
    정리한다 — 한 셀에는 한 줄이 정상값. 값과 &cr;이 섞인 셀(의도적 줄바꿈)과
    P(문단)는 절대 건드리지 않는다. 원문 보존이 필요하면 False(--keep-cr).
    사용자가 해당 셀에 값을 입력한 경우는 일반 수정(edit)이 우선한다.
    """
    with open(dsd_path, "rb") as f:
        data = f.read()
    contents = read_contents(data)
    text = contents.decode("utf-8")

    wb = load_workbook(xlsx_path, data_only=True)
    if MAP_SHEET not in wb.sheetnames:
        raise RepackError("_MAP 시트가 없습니다. 이 도구로 추출한 파일이 아닙니다.")

    # _META 검증: 다른 DSD와의 짝 맞춤 실수 방지
    if META_SHEET in wb.sheetnames:
        meta = {r[0]: r[1] for r in wb[META_SHEET].iter_rows(values_only=True)
                if r and r[0]}
        stored_sha = meta.get("contents_sha1")
        if stored_sha and stored_sha != hashlib.sha1(contents).hexdigest():
            raise RepackError(
                "원본 DSD가 추출 당시 파일과 다릅니다 (contents.xml 해시 불일치). "
                "추출에 사용한 원본 DSD를 지정하세요.")

    changes = []
    cr_only_total = 0
    rows = wb[MAP_SHEET].iter_rows(min_row=2, values_only=True)
    for entry in rows:
        if not entry or entry[0] is None:
            continue
        sheet, row, col, xml_start, xml_end, orig_stored = entry[:6]
        row, col = int(row), int(col)
        xml_start, xml_end = int(xml_start), int(xml_end)
        orig_stored = "" if orig_stored is None else str(orig_stored)

        orig_raw = text[xml_start:xml_end]
        if orig_raw[:len(orig_stored)] != orig_stored:
            raise RepackError(
                f"_MAP 무결성 오류: {sheet}!R{row}C{col} 위치의 원본이 "
                f"일치하지 않습니다. 원본 DSD가 변경되었을 수 있습니다.")

        if str(sheet) not in wb.sheetnames:
            raise RepackError(f"시트 '{sheet}'가 없습니다. 시트명을 되돌리세요.")
        cur = wb[str(sheet)].cell(row=row, column=col).value

        orig_display = clean_text(orig_raw)
        cur_num = try_number(cur)
        orig_num = try_number(orig_display)
        cr_only = is_cr_only_cell(text, xml_start, orig_raw)
        if cr_only:
            cr_only_total += 1

        # --- 변경 판정 (스펙 3.4) ---
        unchanged = False
        if cur is None and orig_display == "":
            unchanged = True
        elif cur_num is not None and orig_num is not None:
            unchanged = numbers_equal(cur_num, orig_num)
        else:
            cur_str = _normalize_for_compare(_cell_value_to_str(cur))
            unchanged = cur_str == _normalize_for_compare(orig_display)

        if unchanged:
            # 개행 엔티티만 있는 셀 → 빈 문자열 (기본 동작)
            if clean_cr and cr_only:
                changes.append({
                    "sheet": str(sheet), "row": row, "col": col,
                    "xml_start": xml_start, "xml_end": xml_end,
                    "old": orig_raw, "new": "", "new_raw": "",
                    "reason": "clean-cr",
                })
            continue

        # --- 새 XML 텍스트 생성 (스펙 3.5) ---
        if cur is None:
            new_text = ""
        elif cur_num is not None and not isinstance(cur, str):
            # 숫자 셀: 원본 표기 스타일(콤마/괄호음수)로 포맷
            new_text = format_number_like(orig_display, cur_num)
        elif cur_num is not None:
            # 문자열로 입력된 숫자("(250)", "4.35")는 입력 표기 그대로
            new_text = str(cur).strip()
        else:
            new_text = str(cur)
        new_raw = escape_for_dsd(new_text) if new_text else ""

        # 주석 헤더(SPAN 매핑)에서 정확히 번호 중복 정리에 해당하는 변경 구분
        reason = "edit"
        if new_text and enclosing_tag(text, xml_start) == "SPAN" and \
                dedup_note_number(orig_display) == new_text.strip() and \
                orig_display.strip() != new_text.strip():
            reason = "note-dedup"

        changes.append({
            "sheet": str(sheet), "row": row, "col": col,
            "xml_start": xml_start, "xml_end": xml_end,
            "old": orig_display, "new": new_text, "new_raw": new_raw,
            "reason": reason,
        })
    return changes, text, data, {"cr_only_total": cr_only_total}


def repack(xlsx_path: str, dsd_path: str, out_path: str = None,
           clean_cr: bool = True, record_history: bool = True) -> dict:
    """역변환 실행. 변경이 없으면 원본 바이트 그대로 복사(G2).

    clean_cr(기본 True): &cr;-only 셀 정리 포함. 원문 보존은 False(--keep-cr).
    record_history(기본 True): 실행 내역을 history/history.sqlite에 기록
    (감사조서 증빙). --dry-run(diff)은 기록되지 않는다.
    """
    changes, text, data, stats = diff(xlsx_path, dsd_path, clean_cr=clean_cr)

    if out_path is None:
        out_path = re.sub(r"\.dsd$", "", dsd_path, flags=re.I) + "_수정.dsd"
    if os.path.abspath(out_path) == os.path.abspath(dsd_path):
        raise RepackError("원본 덮어쓰기는 금지되어 있습니다. 다른 출력 경로를 지정하세요.")

    if not changes:
        with open(out_path, "wb") as f:
            f.write(data)
        if record_history:
            from .history import record_run
            record_run(dsd_path, xlsx_path, out_path, clean_cr, [])
        return {"out_path": out_path, "changes": [], "stats": stats}

    # 오프셋 큰 것부터 교체 → 앞쪽 오프셋 밀림 방지 (스펙 3.5)
    spans = sorted(changes, key=lambda c: c["xml_start"], reverse=True)
    for i in range(len(spans) - 1):
        if spans[i + 1]["xml_end"] > spans[i]["xml_start"]:
            raise RepackError("겹치는 _MAP 구간이 있습니다. 파일이 손상되었습니다.")
    new_text = text
    for ch in spans:
        new_text = (new_text[:ch["xml_start"]] + ch["new_raw"]
                    + new_text[ch["xml_end"]:])

    new_zip = replace_contents(data, new_text.encode("utf-8"))
    with open(out_path, "wb") as f:
        f.write(new_zip)
    if record_history:
        from .history import record_run
        record_run(dsd_path, xlsx_path, out_path, clean_cr, changes)
    return {"out_path": out_path, "changes": changes, "stats": stats}

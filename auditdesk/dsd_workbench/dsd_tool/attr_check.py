"""V-2: XBRL 제출파일 속성 검증 — 기간·단위·주석 명칭 (실검토 대응).

조립 원칙 (신규 파서 금지 — 기존 파서의 속성 판독 확장만 사용):
- 팩트(unit 포함)·확장 속성·role 정의는 호출자(조립층)가 dict로 전달
  (이 모듈은 dart_explorer를 모른다 — V-1 경계 그대로)
- 표준 element 속성 = 금감원 배포 엑셀 Concepts(type·balance·periodType,
  D-2c 리졸버 확장 로드), 확장 = 패키지 xsd 속성
- 주석 명칭 매칭 = 전기대사 제목 정규화(_norm_label) 재사용
- 판정·엑셀 = A-5 규격 (TRUE/FALSE + 비고, FALSE 하이라이트),
  판정 불가는 '판정 불가'로 정직 노출 (미판정≠0)
"""
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .recon import _norm_label

_BOLD = Font(bold=True)
_HDR_FILL = PatternFill("solid", start_color="D9E1F2")
_FALSE_FILL = PatternFill("solid", start_color="FFC7CE")
_WARN_FILL = PatternFill("solid", start_color="FFEB9C")

GUIDE = ("FALSE 존재 시 편집기에서 해당 요소의 속성·단위·주석 명칭을 "
         "확인하십시오 (판정은 기계, 해석은 회계사)")

# 관례적 decimals 값 (원=0/INF, 천원=-3, 백만원=-6, 십억=-9)
_DEC_OK = {"", "INF", "0", "-3", "-6", "-8", "-9"}


def _kind_of_type(t):
    """element type → 기대 단위 종류."""
    t = (t or "").lower()
    if "monetary" in t:
        return "화폐"
    if "pershare" in t:
        return "주당"
    if "shares" in t:
        return "주식수"
    if any(k in t for k in ("string", "text", "date", "anyuri", "block")):
        return "비수치"
    return "순수"


def _unit_ok(kind, unit):
    u = (unit or "").lower()
    if kind == "화폐":
        return "iso4217" in u and "share" not in u
    if kind == "주당":
        return "iso4217" in u and "share" in u
    if kind == "주식수":
        return "share" in u and "iso4217" not in u
    if kind == "비수치":
        return u == ""
    return True                                 # 순수(pure 등) — 관대


_NOTE_TITLE_RE = re.compile(r"^\s*(\d{1,2})\s*[.．]\s*([가-힣][^\n]{0,40})")


def collect_note_titles(ctx):
    """DSD 주석 제목 수집 — 시트 단위 + 원문 텍스트 패턴 병행 (D-1 ②).

    뭉침 내성: 시트가 안 갈라진 문서(수신본 3덩어리·뻗틀이 작성)에서도
    시트 본문 안의 'N. 제목' 평문을 제목으로 수집해 명칭 대사가 돌게
    한다. 반환: [(제목, 출처)] — 출처 ∈ {"시트", "원문"}. 시트 제목이
    같은 번호를 이미 가지면 원문 패턴은 중복 수집하지 않는다.
    """
    out, seen = [], set()
    for sname in ctx.note_sheets:
        ws = ctx.wb[sname]
        title = str(ws.cell(1, 1).value or "").strip()
        m = _NOTE_TITLE_RE.match(title)
        if m:
            seen.add(int(m.group(1)))
        out.append((title, "시트"))
        for row in ws.iter_rows(min_row=2, max_col=1):
            v = row[0].value
            if not isinstance(v, str):
                continue
            m = _NOTE_TITLE_RE.match(v.strip())
            if m and int(m.group(1)) not in seen:
                seen.add(int(m.group(1)))
                out.append((v.strip()[:60], "원문"))
    return out


def attr_check(facts, attrs_of, role_defs, note_titles, doc_end,
               std_label=None, out_path=None, source_warning=None,
               progress=None):
    """제출파일 속성 검증 3종. 요약 dict 반환 (+A-5 규격 엑셀).

    facts: {element_id: [{value, decimals, unit, type, start, end, dims}]}
    attrs_of: callable(element_id) -> {periodType, type, balance} | None
      (조립층에서 표준=배포 엑셀 → 확장=패키지 xsd 순으로 합성)
    role_defs: {role_uri: 한글 정의} — 주석 role 명칭 대사용
    note_titles: DSD 주석 제목 리스트 (["1. 일반사항", ...])
    doc_end: 보고기간말 (당기/전기 분류 기준 — F-3b 기간 체계와 동일)
    std_label: callable(element_id) -> 한글 표준레이블 (없으면 빈칸)
    """
    doc_end = str(doc_end)
    lab = std_label or (lambda e: "")

    # ── ① 기간 속성 ────────────────────────────────────────
    period_rows, pair_rows = [], []
    for eid in sorted(facts):
        fl = facts[eid]
        a = attrs_of(eid)
        expected = (a or {}).get("periodType") or ""
        used = sorted({f.get("type") or "" for f in fl})
        if not expected:
            verdict, true, note = "판정 불가", None, \
                "속성 원천 없음(배포 엑셀·패키지 xsd 모두 부재)"
        elif all(u == expected for u in used):
            verdict, true, note = "일치", True, ""
        else:
            verdict, true, note = "속성 위반", False, \
                f"periodType={expected}인데 {'/'.join(used)} 컨텍스트 사용"
        period_rows.append({
            "element": eid, "label": lab(eid), "expected": expected,
            "used": "/".join(used), "verdict": verdict, "true": true,
            "note": note,
        })
        # 비교 쌍 존재 검사 (판정 아님 — 노출): 당기=end==doc_end
        cur = any(f.get("end") == doc_end for f in fl)
        pri = any((f.get("end") or "") < doc_end for f in fl)
        if cur != pri:
            pair_rows.append({
                "element": eid, "label": lab(eid),
                "side": "당기만 사용" if cur else "전기만 사용",
            })

    # ── ② 단위 속성 ────────────────────────────────────────
    unit_rows = []
    # 문서 대표 decimals (화폐 팩트 최빈값) — 표시단위 정합 기준
    from collections import Counter
    dec_counter = Counter()
    for eid, fl in facts.items():
        a = attrs_of(eid)
        if _kind_of_type((a or {}).get("type")) == "화폐":
            for f in fl:
                dec_counter[str(f.get("decimals") or "")] += 1
    doc_dec = dec_counter.most_common(1)[0][0] if dec_counter else ""
    for eid in sorted(facts):
        fl = facts[eid]
        a = attrs_of(eid)
        if a is None:
            unit_rows.append({
                "element": eid, "label": lab(eid), "kind": "",
                "units": "/".join(sorted({f.get("unit") or "" for f in fl})),
                "decimals": "", "verdict": "판정 불가", "true": None,
                "note": "속성 원천 없음"})
            continue
        kind = _kind_of_type(a.get("type"))
        units = sorted({f.get("unit") or "" for f in fl})
        decs = sorted({str(f.get("decimals") or "") for f in fl})
        bad_unit = [u for u in units if not _unit_ok(kind, u)]
        bad_dec = [d for d in decs if kind == "화폐" and d not in _DEC_OK]
        notes = []
        if bad_unit:
            notes.append(f"단위 부적정: {'/'.join(bad_unit) or '(없음)'}")
        if bad_dec:
            notes.append(f"비관례 decimals: {'/'.join(bad_dec)}"
                         f" (문서 관례 {doc_dec or '미표기'})")
        if kind == "화폐" and not bad_dec and doc_dec and \
                any(d not in ("", "INF", "0", doc_dec) for d in decs):
            notes.append(f"표시단위 상이(문서 관례 {doc_dec})")
        ok = not bad_unit and not bad_dec
        unit_rows.append({
            "element": eid, "label": lab(eid), "kind": kind,
            "units": "/".join(units), "decimals": "/".join(decs),
            "verdict": "일치" if ok else "속성 위반", "true": ok,
            "note": " · ".join(notes),
        })

    # ── ③ 주석 명칭 대사 ───────────────────────────────────
    def _split_def(txt):
        """'[D8..] 한글 | 영문' → (한글, 영문)."""
        t = str(txt or "")
        ko, _, en = t.partition("|")
        return ko.strip(), en.strip()

    def _clean_def(txt):
        # '[D831150] 8. 수익 관련 공시 - 별도' 류 정리 (한글부만)
        t = _split_def(txt)[0]
        t = re.sub(r"\[[^\]]*\]", "", t)
        t = re.sub(r"-\s*(연결|별도)\s*$", "", t.strip())
        t = re.sub(r"^\d+\.\s*", "", t.strip())
        return _norm_label(t)

    note_roles = {}
    for uri, d in role_defs.items():
        code = uri.rsplit("/", 1)[-1]
        # 대표 주석 role만 (서브롤 a/b/c 제외 — 제목 단위 대사)
        if re.search(r"D8\d+$", code) or                 re.search(r"\[D8\d+\]", str(d or "")):
            ko, en = _split_def(d)
            note_roles[uri] = {"code": code, "def": ko,
                               "en": en, "norm": _clean_def(d)}
    name_rows = []
    used_uris = set()
    for item in note_titles:
        title, src = item if isinstance(item, tuple) else (item, "시트")
        t_norm = _clean_def(re.sub(r"^\d+\.\s*", "", str(title)))
        hit = next((r for u, r in note_roles.items()
                    if u not in used_uris and r["norm"] and
                    (r["norm"] == t_norm or t_norm in r["norm"] or
                     r["norm"] in t_norm)), None)
        if hit:
            used_uris.add(next(u for u, r in note_roles.items()
                               if r is hit))
        name_rows.append({
            "src": src,
            "title": str(title).strip(),
            "role": hit["code"] if hit else "",
            "role_def": hit["def"] if hit else "",
            "en": (hit["en"] or
                   hit["code"] + " (영문 정의 미제공 — role 코드 대체)")
            if hit else "",
            "verdict": "매칭" if hit else "미매칭", "true": bool(hit),
            "note": "" if hit else "제출파일 주석 role에 대응 없음",
        })
    only_roles = [r for u, r in sorted(note_roles.items())
                  if u not in used_uris]

    def _cnt(rows):
        t = sum(1 for r in rows if r["true"] is True)
        f = sum(1 for r in rows if r["true"] is False)
        na = sum(1 for r in rows if r["true"] is None)
        return {"true": t, "false": f, "na": na, "total": len(rows)}

    summary = {
        "period": _cnt(period_rows), "unit": _cnt(unit_rows),
        "name": _cnt(name_rows),
        "pair_only": len(pair_rows), "only_roles": len(only_roles),
        "doc_dec": doc_dec, "doc_end": doc_end,
        "source_warning": source_warning,
    }
    result = {"summary": summary, "period_rows": period_rows,
              "pair_rows": pair_rows, "unit_rows": unit_rows,
              "name_rows": name_rows, "only_roles": only_roles}
    if out_path:
        _write_excel(out_path, result)
        result["out_path"] = out_path
    if progress:
        s = summary
        progress(f"  기간 {s['period']['false']}건 위반 · 단위 "
                 f"{s['unit']['false']}건 위반 · 명칭 미매칭 "
                 f"{s['name']['false']}건")
    return result


def _write_excel(out_path, result):
    """A-5 규격 — 요약 + ①기간 ②단위 ③주석명칭 시트, FALSE 하이라이트."""
    s = result["summary"]
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "요약"
    ws0.append(["XBRL 제출파일 속성 검증 (기간·단위·주석 명칭)"])
    ws0.cell(1, 1).font = Font(bold=True, size=13)
    if s.get("source_warning"):
        ws0.append([s["source_warning"]])
        ws0.cell(ws0.max_row, 1).fill = _WARN_FILL
    ws0.append([GUIDE])
    ws0.cell(ws0.max_row, 1).fill = _WARN_FILL
    ws0.cell(ws0.max_row, 1).font = _BOLD
    ws0.append([])
    ws0.append(["구분", "대상", "일치", "위반", "판정 불가"])
    for c in ws0[ws0.max_row]:
        c.font = _BOLD
        c.fill = _HDR_FILL
    for name, key in (("① 기간 속성", "period"), ("② 단위 속성", "unit"),
                      ("③ 주석 명칭", "name")):
        c = s[key]
        ws0.append([name, c["total"], c["true"], c["false"], c["na"]])
        if c["false"]:
            ws0.cell(ws0.max_row, 4).fill = _FALSE_FILL
    ws0.append([])
    ws0.append([f"비교 쌍 한쪽만 사용(노출): {s['pair_only']}건 · "
                f"제출파일에만 있는 주석 role: {s['only_roles']}건 · "
                f"문서 관례 decimals: {s['doc_dec'] or '미표기'} · "
                f"보고기간말 {s['doc_end']}"])
    ws0.column_dimensions["A"].width = 46
    for col in "BCDE":
        ws0.column_dimensions[col].width = 12

    def sheet(name, headers, rows, cells, widths):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for c in ws[1]:
            c.font = _BOLD
            c.fill = _HDR_FILL
        for r in rows:
            ws.append(cells(r))
            if r.get("true") is False:
                for c in ws[ws.max_row]:
                    c.fill = _FALSE_FILL
            elif r.get("true") is None and "verdict" in r:
                ws.cell(ws.max_row, len(headers) - 1).fill = _WARN_FILL
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        return ws

    sheet("기간속성",
          ["태깅 요소 ID", "한글 표준레이블", "정의 periodType",
           "사용 컨텍스트", "판정", "비고"],
          result["period_rows"],
          lambda r: [r["element"].replace("_", ":", 1), r["label"],
                     r["expected"], r["used"], r["verdict"], r["note"]],
          {"A": 52, "B": 24, "C": 16, "D": 16, "E": 12, "F": 40})
    ws_p = wb["기간속성"]
    ws_p.append([])
    ws_p.append(["— 비교 쌍 한쪽만 사용 (판정 아님, 노출) —"])
    ws_p.cell(ws_p.max_row, 1).font = _BOLD
    for r in result["pair_rows"]:
        ws_p.append([r["element"].replace("_", ":", 1), r["label"],
                     "", "", r["side"], ""])
    sheet("단위속성",
          ["태깅 요소 ID", "한글 표준레이블", "유형", "단위", "decimals",
           "판정", "비고"],
          result["unit_rows"],
          lambda r: [r["element"].replace("_", ":", 1), r["label"],
                     r["kind"], r["units"], r["decimals"], r["verdict"],
                     r["note"]],
          {"A": 52, "B": 24, "C": 10, "D": 22, "E": 12, "F": 12, "G": 40})
    sheet("주석명칭",
          ["DSD 주석 제목", "출처", "role 코드", "role 한글명",
           "영문명(대체 명기)", "판정", "비고"],
          result["name_rows"],
          lambda r: [r["title"], r.get("src", "시트"), r["role"],
                     r["role_def"], r["en"], r["verdict"], r["note"]],
          {"A": 40, "B": 8, "C": 14, "D": 40, "E": 34, "F": 10,
           "G": 30})
    ws_n = wb["주석명칭"]
    ws_n.append([])
    ws_n.append(["— 제출파일에만 있는 주석 role (노출) —"])
    ws_n.cell(ws_n.max_row, 1).font = _BOLD
    for r in result["only_roles"]:
        ws_n.append(["", r["code"], r["def"], "", "", ""])
    wb.save(out_path)

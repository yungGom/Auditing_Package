"""D-2c: 금감원 택소노미 레이블 리졸버 (공용 — 차원표·F 워크시트·V-1).

자산: assets/taxonomy/*DART_Taxonomy*.xlsx (사용자 배치·교체 —
다운로드 코드 없음, 수신 허용·송신 금지. KNOWN_VERSIONS.md에 버전 기록).

파싱:
- Concepts 시트: prefix·name·id → 표준 concept 등록부
- Label Link 시트 ko 블록: label 열 (표준 우선, 없으면 terseLabel)
- 배포용에 영문 레이블 블록(en label)이 실측으론 존재하나, 지시대로
  영문명 = element name(CamelCase) 채택 — 편집기 영문 검색은 name 기준.

키 정규화: QName 'prefix:name' ↔ 'prefix_name' 콜론/언더스코어 동치.
표준 미수록(회사 확장요소)은 standard=False — 미매칭≠0, 빈칸 금지
(확장 표시 자체가 편집기 별도 생성 대상이라는 정보).
"""
import glob
import os

_ASSETS_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__)))), "assets", "taxonomy")

_KO_LABEL_COL = 4       # Label Link ko 블록: label
_KO_TERSE_COL = 11      # Label Link ko 블록: terseLabel


def default_asset_path():
    hits = sorted(glob.glob(os.path.join(_ASSETS_DIR,
                                         "*DART_Taxonomy*.xlsx")))
    return hits[-1] if hits else None


def normalize_concept(concept: str) -> str:
    """QName 'prefix:name' → 'prefix_name' (첫 콜론만)."""
    return str(concept).replace(":", "_", 1)


class LabelResolver:
    """배포 엑셀 1회 파싱 스냅샷. resolve(concept) → 3열 값."""

    def __init__(self, xlsx_path=None):
        path = xlsx_path or default_asset_path()
        if not path or not os.path.exists(path):
            raise FileNotFoundError(
                "금감원 택소노미 배포 엑셀이 없습니다 — "
                f"{_ASSETS_DIR} 에 *DART_Taxonomy*.xlsx 를 배치하세요 "
                "(KNOWN_VERSIONS.md 참조)")
        self.path = path
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)

        # Concepts: 표준 concept 등록부 (prefix, name)
        # V-2: 속성(type/balance/periodType)도 같은 행에서 함께 로드
        self.names = {}                       # id → (prefix, name)
        self.attrs = {}                       # id → {type,balance,periodType}
        ws = wb["Concepts"]
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows)]
        col = {h: i for i, h in enumerate(header)}
        for row in rows:
            if not row or not row[1] or not row[2]:
                continue
            prefix, name = str(row[1]).strip(), str(row[2]).strip()
            cid = f"{prefix}_{name}"
            self.names[cid] = (prefix, name)
            self.attrs[cid] = {
                k: str(row[col[k]] or "").strip()
                for k in ("type", "balance", "periodType") if k in col}

        # Label Link ko 블록: label (없으면 terseLabel)
        self.ko = {}                          # id → 한글 표준레이블
        ws = wb["Label Link"]
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i <= 4 or not row or not row[1] or not row[2]:
                continue                      # 1~4행 = LinkRole/lang/헤더
            cid = f"{str(row[1]).strip()}_{str(row[2]).strip()}"
            label = row[_KO_LABEL_COL - 1] or row[_KO_TERSE_COL - 1]
            if label:
                self.ko.setdefault(cid, str(label).strip())
        wb.close()

    def resolve(self, concept: str) -> dict:
        """concept(QName 또는 underscore id) → 편집기 3열 값.

        {ko: 한글 표준레이블('확장'이면 None), en: element name,
         qname: 'prefix:name', standard: bool}
        """
        cid = normalize_concept(concept)
        hit = self.names.get(cid)
        if hit:
            prefix, name = hit
            return {"ko": self.ko.get(cid) or name, "en": name,
                    "qname": f"{prefix}:{name}", "standard": True}
        # 표준 미수록 = 회사 확장요소 — id에서 prefix/name 분해
        prefix, _, name = cid.partition("_")
        return {"ko": None, "en": name or cid,
                "qname": f"{prefix}:{name}" if name else cid,
                "standard": False}


def attrs_of(resolver, concept):
    """V-2: 표준 concept 속성 — 없으면 None (확장은 패키지 xsd 속성으로)."""
    from .taxonomy_labels import normalize_concept
    return resolver.attrs.get(normalize_concept(concept))


_CACHED = None


def get_resolver(xlsx_path=None) -> LabelResolver:
    """모듈 싱글턴 (배포 엑셀 1회 파싱)."""
    global _CACHED
    if xlsx_path:
        return LabelResolver(xlsx_path)
    if _CACHED is None:
        _CACHED = LabelResolver()
    return _CACHED

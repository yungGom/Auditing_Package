"""H-2: 병합 셀 안전 기입 — 공용 헬퍼 (단일 소스).

배경: openpyxl에서 병합 범위 내부 좌표의 셀은 MergedCell(읽기 전용)
이라 .value/.comment 등 대입이 AttributeError로 터진다. 추출 시트에는
원문 표의 ROWSPAN/COLSPAN 병합이 그대로 있으므로, 기존 시트에 표시를
기입하는 조립 경로는 전부 이 헬퍼를 쓴다.

규약:
- 병합 범위 내부 좌표 → 좌상단 앵커 셀로 리다이렉트해 기입
- 기존 메모가 있으면 이어붙임 (덮어쓰기 금지)
- 한 셀 실패로 전체 생성을 중단하지 않음 — failures 목록에
  "기입 불가+사유"로 축적해 산출물에 노출 (침묵 금지)
"""
from openpyxl.comments import Comment

_UNSET = object()


def anchor(ws, row, col):
    """병합 범위 내부 좌표면 좌상단 앵커 셀, 아니면 해당 셀."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and \
                rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col)
    return ws.cell(row=row, column=col)


def put(ws, row, col, value=_UNSET, comment=None, fill=None, font=None,
        failures=None, what="", preserve_value=False):
    """안전 기입. 성공 시 실제 기입된 셀, 실패 시 None.

    comment: (text, author) 튜플 또는 Comment — 기존 메모엔 이어붙임.
    failures: list — 실패 시 {"sheet","cell","what","reason"} 추가.
    preserve_value: 링크 기입 시 기존 원문을 보존하고 표 오른쪽 빈 셀 사용.
    """
    try:
        c = anchor(ws, row, col)
        if value is not _UNSET:
            if preserve_value and c.value is not None and c.value != value:
                fallback_col = max(ws.max_column + 1, col + 1)
                if fallback_col > 16384:
                    raise ValueError("원문 보존: 링크를 기록할 빈 열이 없습니다")
                c = ws.cell(row=row, column=fallback_col)
            c.value = value
        if fill is not None:
            c.fill = fill
        if font is not None:
            c.font = font
        if comment is not None:
            if isinstance(comment, tuple):
                text, author = comment
            else:
                text, author = comment.text, comment.author
            if c.comment is not None and c.comment.text:
                text = c.comment.text + "\n― ― ―\n" + text
            c.comment = Comment(text, author)
        return c
    except Exception as e:                      # 한 셀 실패 ≠ 전체 중단
        if failures is not None:
            failures.append({
                "sheet": ws.title, "cell": f"R{row}C{col}",
                "what": what, "reason": f"{type(e).__name__}: {e}"})
        return None

"""Phase F: XBRL 작성 워크시트 생성기 — DSD → 기공시 형식 엑셀 (F-1: 본문).

배경: 금감원 XBRL 편집기는 수작업 입력만 가능(업로드 양식 없음) →
산출물은 "전사(轉寫) 가이드 워크시트". 입력자가 위→아래로 따라가며
편집기에 입력한다. XBRL 인스턴스 직접 생성은 스코프 제외(DSD 창작 금지와
동일 원칙 — 편집기가 유일한 공식 작성 경로).

완전 로컬: 고객사 DSD·계정명은 외부로 나가지 않는다. 코퍼스는 dart_explorer
산출 sqlite를 읽기 전용 파일로만 접근 (D-3b와 동일 경계).

열 구성 (스펙 공통 양식):
계정/항목(들여쓰기) | 값(기간 블록별) | element ID | 편집기 검색어 |
확신도 | 대안(2~4위) | 확정 ☐
- 확장 필요 판정: element ID 자리에 [확장 필요] + 유사 확장 실증 사례 병기
- 시트 맨 앞 "작성 개요": 시트 구성·미매핑·확장 후보 요약 + 미확정=미완성 경고
"""
import datetime
import os
import re
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .excel_out import extract
from .foot import FootingContext, _label
from .mapping import MappingCorpus, suggest

_BOLD = Font(bold=True)
_HDR_FILL = PatternFill("solid", start_color="D9E1F2")
_WARN_FILL = PatternFill("solid", start_color="FFEB9C")
_EXT_FILL = PatternFill("solid", start_color="FCE4D6")
_WRAP = Alignment(wrap_text=True, vertical="top")
_NUMFMT = "#,##0;[RED](#,##0)"

# 보고서 유형 → 기간 열 라벨 (D-2b 기간 블록 규칙과 동일 어휘).
# DSD 본문 표의 금액 열 수(보통 당기/전기 2개)에 맞춰 앞에서부터 소비한다.
PERIOD_LABELS = {
    "annual": ["당기", "전기", "전전기"],
    "half": ["당기 누적", "전기 누적", "당기 3개월", "전기 3개월"],
    "q1": ["당기 3개월", "전기 3개월"],
    "q3": ["당기 누적", "전기 누적", "당기 3개월", "전기 3개월"],
}
# BS류(시점 잔액) 열 라벨
PERIOD_LABELS_INSTANT = {
    "annual": ["당기말", "전기말"],
    "half": ["당기말", "전기말"],
    "q1": ["당기말", "전기말"],
    "q3": ["당기말", "전기말"],
}

_HEADER = ["계정/항목", None, "element ID", "편집기 검색어", "확신도",
           "대안(2~4위)", "확정 ☐"]


def _sheet_category(sheet_name):
    """extract 시트명(반기연결BS 등) → 매핑 구분(BS/PL/CE/CF)."""
    m = re.search(r"(BS|PL1?|CE|CF)$", sheet_name)
    if not m:
        return None
    code = m.group(1)
    return "PL" if code == "PL1" else code


def _confidence(cand, induty):
    parts = []
    if induty and cand.get("n_same_induty"):
        parts.append(f"동업종 {cand['n_same_induty']}사")
    parts.append(f"실증 {cand['n_companies']}사")
    parts.append(f"유사도 {cand['sim']:.2f}")
    return " · ".join(parts)


def _alternatives(cands):
    return "\n".join(
        f"{i}. {c['element_id']} ({c['std_label']}) — {c['score']:.2f}"
        for i, c in enumerate(cands[1:4], 2)) or ""


def _extension_note(res):
    exts = res.get("similar_extensions") or []
    if not exts:
        return "유사 확장 실증 사례 없음"
    return "\n".join(
        f"확장사례: '{e['label']}' ({e['n_companies']}사, 유사도 {e['sim']:.2f})"
        for e in exts[:3])


class _MemberMatcher:
    """CE 자본구성요소 열 → 표준 member 추천 (자본금→IssuedCapitalMember 등).

    코퍼스 usages에는 member 단위 실증이 없어(축 시그니처만 저장) 표준
    라벨 유사도로 추천한다 — 한계는 워크시트에 명시.
    """

    def __init__(self, corpus: MappingCorpus):
        self.members = {eid: label for eid, label
                        in corpus.std_labels.items()
                        if eid.endswith("Member") and label}

    def match(self, col_label):
        from .mapping import normalize, similarity
        qn = normalize(col_label)
        best = None
        for eid, label in self.members.items():
            sim = similarity(qn, col_label, label)
            if best is None or sim > best[1]:
                best = (eid, sim, label)
        if best and best[1] >= 0.4:
            return {"member_id": best[0], "label": best[2],
                    "sim": round(best[1], 2)}
        return None


def build_worksheet(dsd_path, out_path=None, report_type="annual",
                    induty=None, corpus_db=None, holdout_corp=None,
                    include_notes=True, succession=None, progress=None):
    """회사 DSD → XBRL 전사 가이드 워크시트 (본문 F-1 + 주석 F-2 + 승계 F-3).

    holdout_corp: 게이트 검증용 — 해당 회사를 코퍼스 집계에서 제외(누수 방지).
    include_notes: 주석 워크시트(role 배정·표→차원 매핑) 포함 여부.
    succession: SuccessionAssets(F-3) — 자기 기말 인스턴스 자산. 지정 시
      element는 기말 사용분을 그대로 승계하고 D-3b는 신규 계정에만 작동.
      D-4c 폐지(노랑) element는 경고 + 대체 후보 병기.
    """
    if report_type not in PERIOD_LABELS:
        raise ValueError(f"report_type은 {sorted(PERIOD_LABELS)} 중 하나")
    corpus = MappingCorpus(
        corpus_db, exclude_corps={holdout_corp} if holdout_corp else None)

    with tempfile.TemporaryDirectory(prefix="xbrl_ws_") as tmp:
        xlsx = os.path.join(tmp, "extract.xlsx")
        extract(dsd_path, xlsx)
        ctx = FootingContext(xlsx)

        wb = Workbook()
        overview = wb.active
        overview.title = "작성 개요"

        stats = {"rows": 0, "mapped": 0, "extension": 0, "skipped": 0,
                 "inherited": 0, "new": 0, "deprecated": 0}
        sheet_summaries = []
        results_by_sheet = {}

        for sheet in ctx.fs_sheets:
            category = _sheet_category(sheet)
            periods, data_rows = ctx.fs_sequences(sheet)
            if not data_rows:
                continue
            src = ctx.wb[sheet]
            if category == "CE":
                # CE의 값 열은 기간이 아니라 자본구성요소 — 원본 헤더 그대로
                hdr_row = min(r for r in ctx.rowmaps[sheet]
                              if len(ctx.rowmaps[sheet][r]) >= 5)
                labels = []
                for name, _ in periods:
                    m = re.match(r"col(\d+)", name)
                    txt = str(src.cell(row=hdr_row,
                                       column=int(m.group(1))).value or name
                              ).replace("\n", "") if m else name
                    labels.append(txt)
            else:
                period_names = (PERIOD_LABELS_INSTANT if category == "BS"
                                else PERIOD_LABELS)[report_type]
                labels = [period_names[i] if i < len(period_names)
                          else f"기간{i + 1}" for i in range(len(periods))]

            ws = wb.create_sheet(sheet)
            header = list(_HEADER)
            header[1:2] = [f"값: {lb}" for lb in labels]
            ws.append(header)
            for c in ws[1]:
                c.font = _BOLD
                c.fill = _HDR_FILL

            # CE: 자본구성요소 열 → member 매핑 블록 (F-1-4)
            if category == "CE":
                matcher = _MemberMatcher(corpus)
                ws.append(["[자본변동표 열 → member 매핑 — 표준 라벨 유사도 "
                           "기준 (코퍼스에 member 실증 미축적, 최종 판단 필요)]"])
                ws.cell(row=ws.max_row, column=1).font = _BOLD
                header_row = min(r for r, _ in
                                 [(r, None) for r in ctx.rowmaps[sheet]
                                  if len(ctx.rowmaps[sheet][r]) >= 5])
                for c in ctx.rowmaps[sheet][header_row][2:]:
                    col_label = str(src.cell(row=header_row, column=c).value
                                    or "").replace("\n", "")
                    if not col_label:
                        continue
                    m = matcher.match(col_label)
                    ws.append([f"  열 '{col_label}'",
                               *([None] * len(labels)),
                               m["member_id"] if m else "[매핑 실패]",
                               m["label"] if m else "",
                               f"유사도 {m['sim']}" if m else "", "", "☐"])
                ws.append([])

            sheet_rows = []
            values_by_row = {}
            for name, seq in periods:
                for r, v in seq:
                    values_by_row.setdefault(r, {})[name] = v

            for r in data_rows:
                label = _label(src, r)
                if not label:
                    continue
                vals = values_by_row.get(r, {})
                indent = len(str(src.cell(row=r, column=1).value or "")) - \
                    len(str(src.cell(row=r, column=1).value or "").lstrip())
                stats["rows"] += 1
                if not vals and not label.strip():
                    stats["skipped"] += 1
                    continue
                row_vals = [vals.get(n) for n, _ in periods]

                # F-3 승계: 기말 사용 element 그대로 — D-3b는 신규에만
                inh = succession.inherit(label) if succession else None
                if inh is not None:
                    stats["inherited"] += 1
                    stats["mapped"] += 1
                    tc = succession.taxcheck_of(inh["element_id"]) or {}
                    conf = f"승계 — 기말 사용 (팩트 {inh.get('n_facts', 0)})"
                    if inh.get("sim", 1.0) < 1.0:
                        conf += f" · 라벨 유사 {inh['sim']:.2f}"
                    alts = ""
                    if tc:
                        conf += f" · D-4c {tc['status']}"
                        if tc["status"] != "녹색":
                            alts = tc["detail"]
                    if tc.get("status") == "노랑":     # 폐지 → 대체 후보
                        stats["deprecated"] += 1
                        res = suggest(corpus, label, category, induty=induty)
                        alts += "\n대체 후보:\n" + "\n".join(
                            f"{i}. {c['element_id']} ({c['std_label']})"
                            for i, c in enumerate(
                                res.get("candidates", [])[:4], 1))
                    ws.append([" " * indent + label.strip(), *row_vals,
                               inh["element_id"].replace("_", ":", 1),
                               inh.get("label_ko") or label.strip(),
                               conf, alts, "☐"])
                    if tc.get("status") == "노랑":
                        for c in ws[ws.max_row]:
                            c.fill = _WARN_FILL
                    sheet_rows.append({"row": r, "label": label.strip(),
                                       "top1": inh["element_id"],
                                       "candidates": [], "route": "승계"})
                    for j in range(len(labels)):
                        cell = ws.cell(row=ws.max_row, column=2 + j)
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = _NUMFMT
                    continue

                res = suggest(corpus, label, category, induty=induty)
                if succession is not None:
                    stats["new"] += 1
                new_tag = "신규 계정 · " if succession is not None else ""
                if res["verdict"] == "후보":
                    top = res["candidates"][0]
                    stats["mapped"] += 1
                    ws.append([" " * indent + label.strip(), *row_vals,
                               top["element_id"].replace("_", ":", 1),
                               top["std_label"],
                               new_tag + _confidence(top, induty),
                               _alternatives(res["candidates"]), "☐"])
                    sheet_rows.append({"row": r, "label": label.strip(),
                                       "top1": top["element_id"],
                                       "candidates": res["candidates"]})
                else:
                    stats["extension"] += 1
                    ws.append([" " * indent + label.strip(), *row_vals,
                               "[확장 필요]", "", "",
                               _extension_note(res), "☐"])
                    for c in ws[ws.max_row]:
                        c.fill = _EXT_FILL
                    sheet_rows.append({"row": r, "label": label.strip(),
                                       "top1": None, "candidates": []})
                for j in range(len(labels)):
                    cell = ws.cell(row=ws.max_row, column=2 + j)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = _NUMFMT
                if progress and stats["rows"] % 20 == 0:
                    progress(f"  [{sheet}] {stats['rows']}행 매핑…")

            widths = [40] + [16] * len(labels) + [38, 30, 26, 46, 8]
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
            for row in ws.iter_rows(min_row=2):
                row[len(labels) + 4].alignment = _WRAP
            ws.freeze_panes = "B2"
            results_by_sheet[sheet] = sheet_rows
            sheet_summaries.append(
                (sheet, category, len(sheet_rows),
                 sum(1 for x in sheet_rows if x["top1"] is None)))

        # --- 주석 워크시트 (F-2) ---------------------------------------
        note_results = None
        if include_notes and ctx.note_sheets:
            try:
                from .note_worksheet import NoteAssets, build_note_sheets
                assets = NoteAssets(db_path=corpus_db)
                note_results = build_note_sheets(
                    wb, ctx, corpus, assets, induty=induty,
                    succession=succession, progress=progress)
                sheet_summaries.append(
                    ("주석 1~%d" % len(ctx.note_sheets), "주석",
                     len(ctx.note_sheets), len(note_results["unassigned"])))
            except FileNotFoundError as e:
                note_results = {"error": f"표준 자산 없음: {e} — "
                                "dart_explorer corpus export 필요"}

        _write_overview(overview, dsd_path, report_type, stats,
                        sheet_summaries, succession=succession)
        if out_path is None:
            out_path = re.sub(r"\.dsd$", "", dsd_path,
                              flags=re.I) + "_XBRL작성워크시트.xlsx"
        wb.save(out_path)
    return {"out_path": out_path, "stats": stats,
            "sheets": results_by_sheet, "notes": note_results}


def _write_overview(ws, dsd_path, report_type, stats, sheet_summaries,
                    succession=None):
    ws.append(["XBRL 작성 워크시트 (전사 가이드)"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append([f"원본 DSD: {os.path.basename(dsd_path)}"])
    ws.append([f"보고서 유형: {report_type} · 생성: "
               f"{datetime.date.today().isoformat()}"])
    if succession is not None:
        ws.append([f"승계 모드(F-3): 기말 자산 {os.path.basename(succession.source)}"
                   f" — 승계 {stats['inherited']} / 신규 계정 {stats['new']}"
                   f" / D-4c 폐지 경고 {stats['deprecated']}"
                   + (f" (신버전 {succession.against} 대조)"
                      if succession.against else " (택사노미 대조 생략)")])
        ws.cell(row=ws.max_row, column=1).font = _BOLD
    ws.append([])
    ws.append(["⚠ 확정 ☐ 열이 전부 체크되기 전까지 이 워크시트는 '완성'이 "
               "아닙니다. element 추천은 실증·유사도 기반 후보이며 최종 "
               "판단은 회계사가 합니다."])
    ws.cell(row=5, column=1).fill = _WARN_FILL
    ws.cell(row=5, column=1).font = _BOLD
    ws.append([])
    ws.append(["시트(role) 구성", "구분", "항목 수", "확장 후보"])
    for c in ws[7]:
        c.font = _BOLD
        c.fill = _HDR_FILL
    for name, cat, n, n_ext in sheet_summaries:
        ws.append([name, cat, n, n_ext])
    ws.append([])
    ws.append([f"합계: 항목 {stats['rows']} / 매핑 {stats['mapped']} / "
               f"확장 후보 {stats['extension']} / 제외 {stats['skipped']}"])
    ws.append(["사용법: 각 시트를 위에서 아래로 따라가며 '편집기 검색어'를 "
               "금감원 XBRL 편집기 검색창에 입력해 element를 찾고, 값을 "
               "전사한 뒤 확정 ☐에 체크하세요."])
    ws.column_dimensions["A"].width = 60
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 14

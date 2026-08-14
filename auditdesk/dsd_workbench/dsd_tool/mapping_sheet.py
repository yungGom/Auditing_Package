"""M-1: 대사 조서 자동 조립 — 좌(감사보고서 원문) ↔ 우(XBRL 공시화면).

사양(승인 스키마가 유일 — 양식 원본 파일 미반입):
- 시트 = 총괄표 + FS(BS/PL/CE/CF) + 주석 번호별
- 각 시트 = 좌(원문 extract) ↔ 우(인스턴스 차원표 + 편집기 3열
  레이블) 병렬 — **우 블록 시작열 = 좌 최대열 + 2** (승인 규칙)
- 판정 = 대조 지점 TRUE/FALSE (합계검증·태깅 검증 규약): 행 단위 —
  우 블록 수치가 좌 블록 값 집합에서 발견되는가. 동일 ±허용오차
  우선, 단위 리졸버(A-7) 환산 폴백(좌 표시단위 × 배율 ↔ 우 원 단위)
- 총괄표: 시트명 | 감사보고서명 | XBRL명 | 대응 | 수행 | 오류
  개수 `=COUNTIF(INDIRECT("'"&시트명&"'!A1:ZZ10000"),"FALSE")` |
  오류상세(비고 초안 — 확정은 회계사)
- 좌우 짝 = 조립층이 전달 (FS 제목·주석 명칭 대사 결과 — V-2 ③).
  짝 없는 쪽은 단독 시트 + 총괄표 "대응 없음" 자동 기재

경계: 이 모듈은 xlsx 두 부와 짝 목록만 받는다 — 인스턴스 파싱·렌더·
명칭 대사는 조립층(라우터·게이트 스크립트) 소관.
"""
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .cellsafe import put as safe_put
from .units import region_scales

SUMMARY_SHEET = "총괄표"
_BOLD = Font(bold=True)
_HDR_FILL = PatternFill("solid", start_color="D9E1F2")
_FALSE_FILL = PatternFill("solid", start_color="FFC7CE")
_WARN_FILL = PatternFill("solid", start_color="FFEB9C")
_TOL = 2                                        # 합계검증 단수차 관례


def _numeric(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _copy_block(dst, src_ws, col0, max_cols=40):
    """시트 값 블록 복사 (col0부터). 반환: (사용 행수, 사용 열수)."""
    used_r = used_c = 0
    for row in src_ws.iter_rows():
        for cell in row:
            if cell.value is None or cell.column > max_cols:
                continue
            safe_put(dst, cell.row, col0 + cell.column - 1,
                     value=cell.value)
            used_r = max(used_r, cell.row)
            used_c = max(used_c, cell.column)
    return used_r, used_c


def _left_values(ctx_wb, sheet, scales):
    """좌 블록 값 후보 집합: (표시값, 원환산값 or None)."""
    ws = ctx_wb[sheet]
    out = []
    for row in ws.iter_rows():
        for cell in row:
            if _numeric(cell.value):
                sc = scales.get(cell.row)
                out.append((float(cell.value),
                            float(cell.value) * sc if sc else None))
    return out


_EXCLUDE_HDR_RE = re.compile(r"전전기|\d+기 전")
_HDR_YEAR_RE = re.compile(r"(20\d{2})[-./]")


def _excluded_cols(rws, max_col, scan_rows=6):
    """우 블록 판정 제외 열 — 원문(2개년)에 없는 기간.

    판별 2중: ① 헤더 명칭('전전기'·'N기 전') ② 헤더 일자 연도 —
    실측(공시화면 렌더)은 기간명이 아니라 '당기말 YYYY-MM-DD' 일자
    표기이므로, 열별 연도를 읽어 **최근 2개 연도만 판정**하고 그 이전
    연도 열은 제외한다. 침묵 제외 금지: 호출부가 제외 건수를 집계해
    총괄표에 사유와 함께 노출 (조서 스코프 정의 — 회계사 확인 대상).
    """
    cols = set()
    col_year = {}
    for r in range(1, scan_rows + 1):
        for c in range(1, max_col + 1):
            v = rws.cell(r, c).value
            if not isinstance(v, str):
                continue
            if _EXCLUDE_HDR_RE.search(v):
                cols.add(c)
            m = _HDR_YEAR_RE.search(v)
            if m and c not in col_year:
                col_year[c] = int(m.group(1))
    if col_year:
        keep_years = sorted(set(col_year.values()), reverse=True)[:2]
        for c, y in col_year.items():
            if y not in keep_years:
                cols.add(c)
    return cols


def _match(rv, left_vals):
    """우 값(원 단위) ↔ 좌 후보 대조. 반환 (일치 여부, 비고)."""
    a = abs(float(rv))
    for lv, lw in left_vals:
        if abs(abs(lv) - a) <= _TOL:
            return True, ""
    for lv, lw in left_vals:
        if lw is None:
            continue
        scale = lw / lv if lv else 0
        tol = max(abs(scale) - 1, _TOL)
        if abs(abs(lw) - a) <= tol:
            return True, f"단위환산 ×{int(abs(scale)):,}"
    return False, f"좌측에서 미발견: {rv:,.0f}"


def build_mapping_workbook(left_xlsx, right_xlsx, pairs, out_path,
                           progress=None):
    """대사 조서 생성. 요약 dict 반환.

    pairs: [{"left": 좌 시트명|None, "right": 우 시트명|None,
             "name": 조서 시트명, "left_title": 감사보고서명,
             "right_def": XBRL명(role 정의), "basis": 짝 근거}]
    """
    lwb = load_workbook(left_xlsx)
    rwb = load_workbook(right_xlsx)

    # 좌 단위 배율 (A-7 리졸버 — 표별 상이·미상 보류)
    class _Ctx:                                 # region_scales 최소 어댑터
        def __init__(self, wb):
            self.wb = wb
            self.rowmaps = {}

    out = Workbook()
    out.remove(out.active)
    ws0 = out.create_sheet(SUMMARY_SHEET)
    rows_summary = []

    for p in pairs:
        name = p["name"][:31]
        ws = out.create_sheet(name)
        left, right = p.get("left"), p.get("right")
        stats = {"checks": 0, "false": 0, "excluded": 0, "notes": []}

        if left and left in lwb.sheetnames:
            lws = lwb[left]
            _lr, lc = _copy_block(ws, lws, 1)
        else:
            lws, lc = None, 0
        col0 = lc + 2 if lc else 1              # 우 블록 = 좌 최대열+2
        if right and right in rwb.sheetnames:
            rws = rwb[right]
            _rr, rc = _copy_block(ws, rws, col0)
        else:
            rws, rc = None, 0

        verdict_col = col0 + rc if rws is not None else None   # 우 끝+1
        if lws is not None and rws is not None:
            excl_cols = _excluded_cols(rws, rc)
            # 좌 값 후보 (표시값 + 단위환산 원값) — 배율은 단위 마커
            scales = {}
            try:
                from .foot import FootingContext
                ctx = FootingContext(left_xlsx)
                if left in ctx.rowmaps:
                    scales = region_scales(ctx, left)
            except Exception:
                scales = {}
            if not scales:
                from .units import detect_sheet_unit
                hit = detect_sheet_unit(lws)
                if hit:
                    scales = {r: hit[0] for r in range(1, lws.max_row + 1)}
            left_vals = _left_values(lwb, left, scales)
            safe_put(ws, 2, verdict_col, value="판정", font=_BOLD,
                     fill=_HDR_FILL)
            for row in rws.iter_rows():
                all_nums = [c for c in row
                            if _numeric(c.value) and c.column <= rc]
                nums = [c for c in all_nums
                        if c.column not in excl_cols]
                stats["excluded"] += len(all_nums) - len(nums)
                if not nums:
                    continue
                ok_all, notes = True, []
                for c in nums:
                    ok, note = _match(c.value, left_vals)
                    if not ok:
                        ok_all = False
                    if note:
                        notes.append(note)
                stats["checks"] += 1
                cell = safe_put(ws, row[0].row, verdict_col, value=ok_all)
                if not ok_all:
                    stats["false"] += 1
                    if cell is not None:
                        cell.fill = _FALSE_FILL
                    stats["notes"].extend(
                        n for n in notes if n.startswith("좌측"))
                elif notes:
                    safe_put(ws, row[0].row, verdict_col + 1,
                             value=" · ".join(sorted(set(notes))[:2]))
        rows_summary.append((name, p.get("left_title", ""),
                             p.get("right_def", ""), p.get("basis", ""),
                             lws is not None, rws is not None, stats))
        if progress:
            progress(f"  [{name}] 대조 {stats['checks']}행")

    # ── 총괄표 (승인 스키마) ────────────────────────────────
    ws0.append(["대사 조서 — 감사보고서 원문 ↔ XBRL 공시화면 "
                "(판정은 기계, 확정은 회계사)"])
    ws0.cell(1, 1).font = Font(bold=True, size=12)
    ws0.append([])
    ws0.append(["시트명", "감사보고서명", "XBRL명", "짝 근거", "대응",
                "수행(행)", "오류 개수",
                "대조 제외(값) — 원문 미표시 기간",
                "오류상세 (비고 초안 — 확정은 "
                "회계사)"])
    hr = ws0.max_row
    for c in ws0[hr]:
        c.font = _BOLD
        c.fill = _HDR_FILL
    n_pairs = n_lonely = 0
    for name, lt, rd, basis, has_l, has_r, stats in rows_summary:
        r = ws0.max_row + 1
        if has_l and has_r:
            pair_txt = "대응"
            n_pairs += 1
        else:
            pair_txt = ("대응 없음 — 감사보고서에만" if has_l
                        else "대응 없음 — 제출파일에만")
            n_lonely += 1
        draft = " / ".join(dict.fromkeys(stats["notes"]))[:120]
        ws0.append([name, lt, rd, basis, pair_txt,
                    stats["checks"] if (has_l and has_r) else "대조 없음",
                    f'=COUNTIF(INDIRECT("\'"&A{r}&"\'!A1:ZZ10000"),'
                    f'"FALSE")',
                    stats["excluded"] or "",
                    draft])
        if not (has_l and has_r):
            ws0.cell(r, 5).fill = _WARN_FILL
        if stats["excluded"]:
            ws0.cell(r, 8).fill = _WARN_FILL
        if stats["false"]:
            ws0.cell(r, 9).fill = _FALSE_FILL
    for col, w in (("A", 14), ("B", 30), ("C", 40), ("D", 22), ("E", 22),
                   ("H", 24), ("I", 60)):
        ws0.column_dimensions[col].width = w

    out.save(out_path)
    total_checks = sum(s["checks"] for *_x, s in rows_summary)
    total_false = sum(s["false"] for *_x, s in rows_summary)
    return {
        "out_path": out_path, "sheets": [r[0] for r in rows_summary],
        "pairs": n_pairs, "lonely": n_lonely,
        "checks": total_checks, "false": total_false,
        "excluded": sum(st["excluded"] for *_x, st in rows_summary),
        "pair_rate": n_pairs / len(rows_summary) if rows_summary else 0,
    }

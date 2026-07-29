"""F-2: 주석 워크시트 — role 배정 + 표→차원 매핑 + 서술문 매핑.

입력 자산 (dart_explorer 산출 JSON — 파일 교환, 네트워크 경계 유지):
- standard_roles.json: 주석(D8) role 정의 목록
- standard_hypercubes.json: role별 축·member
- mapping_corpus.sqlite: role 사용 실증(회사수), element 추천(D-3b)

로직:
1. role 배정: DSD 주석 제목 ↔ 표준 role 정의 제목부 유사도(D-3b 스코어러)
   + 코퍼스 role 사용 회사수 실증 가중 → Top3 제안 (미매칭은 목록화)
2. 표→차원: 열 라벨 분류(기간/증감 element 패턴/기타), 행 라벨은 배정
   role의 표준 하이퍼큐브 member와 매칭 — 매칭 최다 축을 행 축으로 제안,
   member 미매칭 행은 element 추천(상태 4종)
3. 서술문 주석: role 소속 TextBlock/Explanatory element 매핑 + 원문 첨부
상태 4종(Design v2): 표준 | 표준 사용 권장 | 확장 필요 | 수동 확인
"""
import collections
import json
import os
import re
import sqlite3

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .foot import _label, _regions
from .mapping import (DEFAULT_CORPUS, STATE_EXTENSION, STATE_MANUAL,
                      STATE_PREFER_STANDARD, STATE_STANDARD,
                      classify_suggestion, normalize, similarity, suggest)

_BOLD = Font(bold=True)
_HDR_FILL = PatternFill("solid", start_color="D9E1F2")
_STATE_FILL = {
    STATE_STANDARD: None,
    STATE_PREFER_STANDARD: PatternFill("solid", start_color="DDEBF7"),
    STATE_EXTENSION: PatternFill("solid", start_color="FCE4D6"),
    STATE_MANUAL: PatternFill("solid", start_color="FFEB9C"),
}

# 증감표 열 패턴 (기초/취득/처분/상각/대체/손상/기말 → 증감 element 계열)
_ROLLFORWARD = ["기초", "취득", "처분", "감가상각", "상각", "대체", "손상",
                "환율", "기말", "증가", "감소"]
_PERIOD_RE = re.compile(
    r"당\s*기|전\s*기|제\s*\d+\s*(?:\([^)]*\)\s*)?기|기\s*말|기\s*초")

# 이 축들은 열 분화가 아니라 role 필터/보고단위 성격 — 행 축 후보에서 제외
_FILTER_AXES = ("ConsolidatedAndSeparateFinancialStatementsAxis",)


class NoteAssets:
    """표준 role·하이퍼큐브 + 코퍼스 role 실증 (읽기 전용 스냅샷)."""

    def __init__(self, corpus_dir=None, db_path=None):
        corpus_dir = corpus_dir or os.path.dirname(DEFAULT_CORPUS)
        with open(os.path.join(corpus_dir, "standard_roles.json"),
                  encoding="utf-8") as f:
            all_roles = json.load(f)
        self.note_roles = []
        for r in all_roles:
            if not r["code"].startswith("D8"):
                continue
            # 정의: '[D822100] 주석 - 유형자산 - 연결 | Notes ...' → 제목부
            m = re.match(r"^\[(D\d{6})\]\s*(?:주석\s*-\s*)?([^|]+)",
                         r["definition"])
            title = m.group(2) if m else r["definition"]
            title = re.sub(r"-\s*(연결|별도)\s*$", "", title).strip()
            self.note_roles.append({**r, "title": title,
                                    "title_norm": normalize(title)})
        with open(os.path.join(corpus_dir, "standard_hypercubes.json"),
                  encoding="utf-8") as f:
            self.hypercubes = json.load(f)

        # 코퍼스 role 실증: role 코드 → 사용 회사수 + 대표 element 라벨
        # (동의 표현 대응 — 예: 주석 '순확정급여부채' ↔ 종업원급여 role의
        #  소속 element 라벨에서 매칭)
        self.role_corps = collections.Counter()
        self.role_labels = {}
        self.axis_corps = {}                   # 축 → D8 주석 실증 회사수
        path = db_path or DEFAULT_CORPUS
        if os.path.exists(path):
            con = sqlite3.connect(f"file:{path}?mode=ro", uri=True,
                                  timeout=30)
            seen = collections.defaultdict(set)
            label_freq = collections.defaultdict(collections.Counter)
            axis_seen = collections.defaultdict(set)
            for corp, roles, label, dims in con.execute(
                    "SELECT corp_code, roles, label_ko, dims FROM usages "
                    "WHERE roles LIKE '%D8%' AND is_ext = 0"):
                for code in (roles or "").split(","):
                    if code.startswith("D8"):
                        seen[code[:6]].add(corp)
                        if label:
                            label_freq[code[:6]][label] += 1
                for ax in re.split(r"[|;]", dims or ""):
                    if ax.endswith("Axis"):
                        axis_seen[ax].add(corp)
            con.close()
            self.axis_corps = {k: len(v) for k, v in axis_seen.items()}
            self.role_corps = {k: len(v) for k, v in seen.items()}
            # 특이 라벨만 유지: 여러 role에 두루 나오는 라벨(부문표의 온갖
            # 계정 등)은 role 식별 신호가 아님 — ≤2개 role 등장 라벨로 제한
            label_roles = collections.defaultdict(set)
            for code, c in label_freq.items():
                for lb in c:
                    label_roles[lb].add(code)
            self.role_labels = {
                k: [lb for lb, _ in c.most_common(80)
                    if len(label_roles[lb]) <= 2][:50]
                for k, c in label_freq.items()}
        self._max_corps = max(self.role_corps.values(), default=1)

    # ------------------------------------------------------------------
    @staticmethod
    def _title_sim(qn, qtitle, role_title, role_title_norm):
        """제목 유사도 = max(D-3b 유사도, 포함관계 스코어).

        포함관계: '충당부채' ⊂ '충당부채, 우발부채 및 우발자산' —
        자모거리는 길이차에 취약하므로 짧은 쪽이 통째로 포함되면
        0.6 + 0.4×길이비로 보정.
        """
        sim = similarity(qn, qtitle, role_title)
        rn = role_title_norm
        if qn and rn:
            if qn in rn or rn in qn:
                ratio = min(len(qn), len(rn)) / max(len(qn), len(rn))
                sim = max(sim, 0.6 + 0.4 * ratio)
            else:
                # 공통 접두 (예: '기타자본항목' vs '기타자본구성요소')
                k = 0
                for a, b in zip(qn, rn):
                    if a != b:
                        break
                    k += 1
                if k >= 4:
                    sim = max(sim, 0.4 + 0.3 * k / max(len(qn), len(rn)))
        return sim

    def assign_role(self, note_title, prefer_consol="별도", top=3):
        """주석 제목 → role 후보 Top N.

        점수 = 0.7×제목신호 + 0.3×실증빈도.
        제목신호 = max(role 제목 유사도, 0.9×소속 element 라벨 최고 유사도
        — 동의 표현 대응: '순확정급여부채' ↔ 종업원급여 role의 element 라벨).
        """
        import math
        qtitle = re.sub(r"^\d+\.\s*", "", str(note_title)).strip()
        qn = normalize(qtitle)
        scored = []
        for r in self.note_roles:
            sim = self._title_sim(qn, qtitle, r["title"], r["title_norm"])
            label_sim = 0.0
            if sim < 0.9:
                for lb in self.role_labels.get(r["code"][:6], []):
                    s = similarity(qn, qtitle, lb)
                    if s > label_sim:
                        label_sim = s
                        if label_sim >= 0.99:
                            break
            signal = max(sim, 0.9 * label_sim)
            if signal < 0.2:
                continue
            n = self.role_corps.get(r["code"][:6], 0)
            freq = math.log1p(n) / math.log1p(self._max_corps)
            bonus = 0.02 if r["consol"] == prefer_consol else 0.0
            scored.append({"code": r["code"], "definition": r["definition"],
                           "title": r["title"], "sim": round(signal, 3),
                           "title_sim": round(sim, 3),
                           "label_sim": round(label_sim, 3),
                           "n_corps": n,
                           "score": round(0.7 * signal + 0.3 * freq
                                          + bonus, 4)})
        scored.sort(key=lambda x: x["score"], reverse=True)
        return scored[:top]

    def global_axes(self):
        """전 role 하이퍼큐브의 축 풀 (축 id로 member 합집합, 필터 축 제외).

        2차 축 탐색용 — 배정 role에 표의 축이 등재되지 않은 경우
        (실측: 범주별 금융상품 주석의 차입금 축, 재무위험관리의 통화 축).
        """
        if not hasattr(self, "_global_axes"):
            merged = {}
            for axes in self.hypercubes.values():
                for a in axes:
                    if any(f in a["axis"] for f in _FILTER_AXES):
                        continue
                    g = merged.setdefault(
                        a["axis"], {"axis": a["axis"],
                                    "axis_label": a["axis_label"],
                                    "members": {}})
                    for mid, lb in a["members"]:
                        if mid.endswith("Member"):
                            g["members"].setdefault(mid, lb)
            self._global_axes = [
                {**g, "members": sorted(g["members"].items())}
                for g in merged.values()]
        return self._global_axes

    def role_axes(self, code):
        """role의 행 축 후보 (필터 축 제외). 별도(…5) 미등재 시 연결(…0) 폴백."""
        axes = self.hypercubes.get(code)
        if not axes and code.endswith("5"):
            axes = self.hypercubes.get(code[:-1] + "0")
        axes = axes or []
        out = []
        for a in axes:
            if any(f in a["axis"] for f in _FILTER_AXES):
                continue
            members = [(mid, lb) for mid, lb in a["members"]
                       if mid.endswith("Member")]
            if members:
                out.append({"axis": a["axis"], "axis_label": a["axis_label"],
                            "members": members})
        return out


def _classify_column(header_text):
    """표 열 라벨 분류: 기간 / 증감 / 일반."""
    t = str(header_text or "").replace("\n", "")
    if not t.strip():
        return "빈칸"
    hit = next((k for k in _ROLLFORWARD if k in t), None)
    if hit and not _PERIOD_RE.search(t.replace(hit, "")):
        return f"증감({hit})"
    if _PERIOD_RE.search(t):
        return "기간"
    return "항목"


def match_member(assets_axes, row_label, min_sim=0.5):
    """행 라벨 → member 후보. (axis, member_id, member라벨, sim)."""
    qn = normalize(row_label)
    if not qn:
        return None
    best = None
    for ax in assets_axes:
        for mid, lb in ax["members"]:
            sim = similarity(qn, row_label, lb)
            if sim >= min_sim and (best is None or sim > best[3]):
                best = (ax["axis"], mid, lb, sim)
    return best


def find_text_block(corpus, role_code, note_title):
    """서술문 주석 → role 소속 TextBlock/Explanatory element 추천."""
    base = role_code[:6]
    cands = []
    for eid, e in corpus.elements.items():
        if not any(r.startswith(base) for r in e["roles"]):
            continue
        if not ("TextBlock" in eid or "Explanatory" in eid or
                "Disclosure" in eid):
            continue
        label = corpus.std_labels.get(eid) or next(iter(e["labels"]), "")
        sim = similarity(normalize(note_title), note_title, label)
        cands.append({"element_id": eid, "label": label,
                      "n_corps": len(e["corps"]), "sim": round(sim, 3)})
    cands.sort(key=lambda c: (c["sim"], c["n_corps"]), reverse=True)
    return cands[:3]


# ---------------------------------------------------------------------------
# 주석 시트 생성
# ---------------------------------------------------------------------------

# 라우팅 3분류 (확정): 워크시트의 모든 주석 표 행은 반드시 셋 중 하나 —
#   a. element: 항목 행 → D-3b element 매핑 (계정과목과 동일 처리)
#   b. member: 축 배정된 표의 분류 행 → 표준 하이퍼큐브 member 매칭
#   c. manual: 판단 불가 → "수동 확인" 상태
# 침묵 누락 0 — 라벨 있는 모든 행이 rows_detail에 라우팅·상태와 함께 기록됨
ROUTE_MEMBER, ROUTE_ELEMENT, ROUTE_MANUAL = "member", "element", "manual"


def build_note_sheets(wb, ctx, corpus, assets: NoteAssets, induty=None,
                      prefer_consol="별도", succession=None, progress=None):
    """extract 컨텍스트의 주석들 → 워크시트 시트 추가. 결과 dict 반환.

    succession(F-3): 자기 기말 자산 — role은 주석 번호·제목으로 결정적
    승계, 행 element/member도 기말 사용분 우선. D-3b는 신규에만.
    """
    results = {"roles": {}, "unassigned": [], "tables": [],
               "routes": collections.Counter(), "rows_detail": []}
    for sheet in ctx.note_sheets:
        src = ctx.wb[sheet]
        title = str(src.cell(1, 1).value or "")
        inherited_role = succession.find_role(title) if succession else None
        cands = assets.assign_role(title, prefer_consol=prefer_consol)
        if inherited_role is not None:
            assigned = {"code": inherited_role["code"],
                        "definition": inherited_role["definition"],
                        "sim": 1.0, "n_corps": 0, "inherited": True}
        else:
            assigned = cands[0] if cands and cands[0]["sim"] >= 0.4 else None
        results["roles"][sheet] = {"title": title, "assigned": assigned,
                                   "candidates": cands}
        if assigned is None:
            results["unassigned"].append((sheet, title))

        ws = wb.create_sheet(f"주석{sheet}")
        ws.append([title])
        ws.cell(1, 1).font = Font(bold=True, size=12)

        # --- role 배정 블록 ---
        ws.append(["[role 배정]", "", "", "", "", "확정 ☐"])
        ws.cell(ws.max_row, 1).font = _BOLD
        if inherited_role is not None:
            ws.append(["→ 승계 (기말)", inherited_role["code"],
                       inherited_role["definition"][:60],
                       "자기 기말 role 구성 그대로", "", "☐"])
        if cands:
            for i, c in enumerate(cands, 1):
                mark = "→ 배정 제안" if (assigned and i == 1) else f"후보{i}"
                ws.append([mark, c["code"], c["definition"][:60],
                           f"유사도 {c['sim']:.2f}",
                           f"실증 {c['n_corps']}사", "☐" if i == 1 else ""])
        else:
            ws.append(["미매칭 — 회계사 판단 필요"])
            ws.cell(ws.max_row, 1).fill = _STATE_FILL[STATE_MANUAL]
        ws.append([])

        rm = ctx.rowmaps[sheet]
        regions = _regions(rm)
        role_axes = assets.role_axes(assigned["code"]) if assigned else []

        if not regions:
            # --- 서술문 주석: TextBlock 매핑 + 원문 첨부 ---
            ws.append(["[서술문 주석 — 텍스트 블록 element]"])
            ws.cell(ws.max_row, 1).font = _BOLD
            tbs = find_text_block(corpus, assigned["code"],
                                  title) if assigned else []
            if tbs:
                for i, tb in enumerate(tbs, 1):
                    ws.append([f"후보{i}", tb["element_id"],
                               tb["label"][:50],
                               f"실증 {tb['n_corps']}사 · 유사도 {tb['sim']:.2f}",
                               "", "☐" if i == 1 else ""])
            else:
                ws.append(["후보 없음 — 수동 확인"])
                ws.cell(ws.max_row, 1).fill = _STATE_FILL[STATE_MANUAL]
            ws.append([])
            ws.append(["[원문]"])
            ws.cell(ws.max_row, 1).font = _BOLD
            for r in sorted(rm):
                v = src.cell(r, 1).value
                if v and r > 1:
                    ws.append([str(v)[:250]])
        else:
            # --- 표 주석: 표별 차원 매핑 ---
            for ti, region in enumerate(regions, 1):
                header = region[0]
                cols = sorted({c for r in region for c in rm.get(r, [])})
                col_class = {c: _classify_column(
                    src.cell(header, c).value) for c in cols[1:]}

                # 행 member 매칭 → 최다 매칭 축 = 행 축 제안
                axis_votes = collections.Counter()
                row_matches = {}
                for r in region[1:]:
                    label = _label(src, r)
                    if not label:
                        continue
                    hit = match_member(role_axes, label)
                    if hit:
                        axis_votes[hit[0]] += 1
                        row_matches[r] = hit
                best_axis = axis_votes.most_common(1)[0][0] \
                    if axis_votes else None
                axis_source = "role" if best_axis else None
                if best_axis is None:
                    # 2차 전역 축 탐색 (배정 role에 표의 축 미등재 대응) —
                    # 이중 안전장치: 같은 축에 ≥2행 매칭 + 코퍼스 D8
                    # 실증 있는 축만 채택 (우연 단일 매칭 배제)
                    votes2 = collections.Counter()
                    matches2 = {}
                    for r in region[1:]:
                        label = _label(src, r)
                        if not label:
                            continue
                        hit = match_member(assets.global_axes(), label)
                        if hit and assets.axis_corps.get(hit[0]):
                            votes2[hit[0]] += 1
                            matches2[r] = hit
                    top2 = votes2.most_common(1)
                    if top2 and top2[0][1] >= 2:
                        best_axis = top2[0][0]
                        axis_source = "global"
                        row_matches = matches2

                # 축 채택 후 행 재매칭 — 채택 축 우선. 미매칭 행은
                # 고신뢰(≥0.7) 전역 보조 축 허용 (표가 차원 표로 확인된
                # 경우에 한함 — 실측: 표 제2축·자매 축의 member 행)
                if best_axis is not None:
                    pool = role_axes if axis_source == "role" \
                        else assets.global_axes()
                    best_ax_only = [ax for ax in pool
                                    if ax["axis"] == best_axis]
                    row_matches = {}
                    for r in region[1:]:
                        label = _label(src, r)
                        if not label:
                            continue
                        hit = match_member(best_ax_only, label)
                        if hit is None:
                            h2 = match_member(assets.global_axes(), label,
                                              min_sim=0.7)
                            if h2 and assets.axis_corps.get(h2[0]):
                                hit = h2
                        if hit:
                            row_matches[r] = hit

                _src_tag = ""
                if axis_source == "global":
                    _src_tag = (f" (전역 탐색 — 배정 role 외 축, 실증 "
                                f"{assets.axis_corps.get(best_axis, 0)}사)")
                ws.append([f"[표 {ti}] 행 축 제안: "
                           f"{best_axis.split('_')[-1] + _src_tag if best_axis else '(축 매칭 없음 — 행은 element로)'}"])
                ws.cell(ws.max_row, 1).font = _BOLD
                ws.append(["열 분류: " + ", ".join(
                    f"'{str(src.cell(header, c).value or '')[:12]}'"
                    f"→{col_class[c]}" for c in cols[1:6])])

                _hdr = ["행 라벨", "매핑 대상", "ID", "라벨/검색어", "근거",
                        "상태", "확정 ☐"]
                ws.append(_hdr)
                for cc in ws[ws.max_row]:
                    cc.font = _BOLD
                    cc.fill = _HDR_FILL
                n_rows = 0
                for r in region[1:]:
                    label = _label(src, r)
                    if not label:
                        continue
                    n_rows += 1
                    hit = row_matches.get(r)
                    detail = {"sheet": sheet, "table": ti, "row": r,
                              "label": label.strip(), "axis": best_axis,
                              "member_id": None, "cand_ids": [],
                              "inherited": False}

                    # F-3 승계: 기말 사용 member/element 우선 (신규만 D-3b)
                    if succession is not None:
                        own_m = succession.inherit_member(label)
                        own_e = succession.inherit(label)
                        if own_m and (best_axis is not None or own_e is None):
                            route, state = ROUTE_MEMBER, STATE_STANDARD
                            detail["member_id"] = own_m["id"]
                            detail["inherited"] = True
                            warn = ("" if best_axis is not None else
                                    " ⚠ 표 축 미확정 — 배치 확인")
                            ws.append([label.strip(), "member",
                                       own_m["id"].replace("_", ":", 1),
                                       own_m.get("label_ko") or "",
                                       "승계 — 기말 문맥 member" + warn,
                                       STATE_STANDARD, "☐"])
                            detail["route"] = route
                            detail["state"] = state
                            results["routes"][route] += 1
                            results["rows_detail"].append(detail)
                            continue
                        if own_e is not None:
                            route, state = ROUTE_ELEMENT, STATE_STANDARD
                            detail["inherited"] = True
                            detail["cand_ids"] = [own_e["element_id"]]
                            tc = succession.taxcheck_of(
                                own_e["element_id"]) or {}
                            ev = (f"승계 — 기말 사용 "
                                  f"(팩트 {own_e.get('n_facts', 0)})")
                            if tc:
                                ev += f" · 택소노미 점검 {tc['status']}"
                                if tc["status"] == "노랑":
                                    ev += f" ⚠ {tc['detail']}"
                            ws.append([label.strip(), "element",
                                       own_e["element_id"].replace(
                                           "_", ":", 1),
                                       own_e.get("label_ko") or "",
                                       ev, STATE_STANDARD, "☐"])
                            if tc.get("status") == "노랑":
                                for cc in ws[ws.max_row]:
                                    cc.fill = _STATE_FILL[STATE_MANUAL]
                            detail["route"] = route
                            detail["state"] = state
                            results["routes"][route] += 1
                            results["rows_detail"].append(detail)
                            continue

                    if hit:
                        # b. member 행 (축 배정된 표의 분류 행 —
                        #    채택 축 또는 고신뢰 보조 축)
                        route, state = ROUTE_MEMBER, STATE_STANDARD
                        detail["member_id"] = hit[1]
                        aux = "" if hit[0] == best_axis else " (보조 축)"
                        ws.append([label.strip(), "member",
                                   hit[1].replace("_", ":", 1), hit[2],
                                   f"축 {hit[0].split('_')[-1][:28]}{aux} · "
                                   f"유사도 {hit[3]:.2f}",
                                   STATE_STANDARD, "☐"])
                    else:
                        res = suggest(corpus, label, "주석", induty=induty)
                        state, top = classify_suggestion(res)
                        exts = res.get("similar_extensions") or []
                        detail["cand_ids"] = [
                            c["element_id"]
                            for c in (res.get("candidates") or [])[:4]]
                        # a. 항목 행 (D-3b 신호 있음) / c. 판단 불가
                        route = ROUTE_MANUAL if state == STATE_MANUAL \
                            else ROUTE_ELEMENT
                        # 안전장치: 축 배정 실패 표의 행은 member일 가능성이
                        # 남아 있음 — 잘못된 확신보다 명시된 불확실이 낫다
                        warn = ("" if best_axis is not None else
                                " ⚠ 표 구조 미확정 — member 가능성 확인")
                        if top is not None:
                            evidence = top["evidence"][:46]
                            if state == STATE_PREFER_STANDARD and exts:
                                evidence += (f" | 확장사례 '{exts[0]['label']}'"
                                             f" {exts[0]['n_companies']}사")
                            ws.append([label.strip(), "element",
                                       top["element_id"].replace("_", ":", 1),
                                       top["std_label"], evidence + warn,
                                       state, "☐"])
                        else:
                            note = ("; ".join(
                                f"확장 '{e['label'][:20]}' {e['n_companies']}사"
                                for e in exts[:2]) or "신호 없음") + warn
                            ws.append([label.strip(),
                                       "수동 확인" if route == ROUTE_MANUAL
                                       else "element",
                                       "[확장 필요]"
                                       if state == STATE_EXTENSION else "",
                                       "", note, state, "☐"])
                        fill = _STATE_FILL.get(state)
                        if fill:
                            for cc in ws[ws.max_row]:
                                cc.fill = fill
                    detail["route"] = route
                    detail["state"] = state
                    results["routes"][route] += 1
                    results["rows_detail"].append(detail)
                ws.append([])
                results["tables"].append({
                    "sheet": sheet, "table": ti, "axis": best_axis,
                    "axis_source": axis_source, "rows": n_rows,
                    "member_matched": len(row_matches)
                    if best_axis else 0})
        for col, w in (("A", 34), ("B", 10), ("C", 44), ("D", 30),
                       ("E", 44), ("F", 14), ("G", 8)):
            ws.column_dimensions[col].width = w
        if progress:
            progress(f"  [주석{sheet}] {title[:24]} → "
                     f"{assigned['code'] if assigned else '미매칭'}")
    return results

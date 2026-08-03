"""V-2 게이트: 제출파일 속성 검증 — 3사 왕복 + 변조 3종.

G1 신원·무벡스·삼성(B-4 첨부) 왕복 — 기간·단위·주석 명칭 리포트
G2 변조 3종: periodType 위반·decimals 불일치·주석명 변조 각 1건 검출
(G3 전 테스트 회귀는 스위트 전체 실행으로 별도)
"""
import copy
import os
import sys
import zipfile

import pytest

_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__),
                                     "..", "..", ".."))
_XB = os.path.join(_ROOT, "dart_explorer", "cache", "xbrl")
_DOC = os.path.join(_ROOT, "dart_explorer", "cache", "document")

_CASES = {
    "신원": {
        "pkg": os.path.join(_XB, "00136925", "20260323000596_11011"),
        "dsd": os.path.join(_DOC, "00136925", "20260323000596.dsd"),
    },
    "무벡스": {
        "pkg": os.path.join(_XB, "01358463", "20260318001359_11011"),
        "dsd": os.path.join(_DOC, "01358463",
                            "20260318001359_별도감사보고서.dsd"),
    },
    "삼성": {
        "pkg": os.path.join(_XB, "00126380", "20260310002820_11011"),
        "zip": os.path.join(_DOC, "00126380", "20260310002820.zip"),
        "entry": "20260310002820_00761.xml",   # B-4 첨부(연결감사보고서)
    },
}
_ready = all(os.path.isdir(c["pkg"]) for c in _CASES.values())
pytestmark = pytest.mark.skipif(not _ready, reason="3사 패키지 캐시 없음")


def _load_case(name, tmp):
    if _ROOT not in sys.path:
        sys.path.insert(0, _ROOT)
    from dart_explorer.xbrl.dimension_table import XbrlInstance
    from dart_explorer.xbrl.taxonomy import TaxonomyPackage
    from dsd_tool.excel_out import extract
    from dsd_tool.foot import FootingContext
    from dsd_tool.taxonomy_labels import get_resolver, normalize_concept

    case = _CASES[name]
    inst = XbrlInstance(case["pkg"])
    facts = {}
    for eid, fl in inst.facts.items():
        facts[eid] = [{
            "value": f["value"], "decimals": f.get("decimals"),
            "unit": f.get("unit", ""), "type": f["ctx"]["type"],
            "start": f["ctx"].get("start"), "end": f["ctx"].get("end"),
            "dims": dict(f["ctx"].get("dims") or {}),
        } for f in fl]
    tp = TaxonomyPackage(case["pkg"])
    resolver = get_resolver()

    def attrs_of(eid):
        a = resolver.attrs.get(normalize_concept(eid))
        if a:
            return a
        return tp.ext_attrs.get(normalize_concept(eid))

    def std_label(eid):
        r = resolver.resolve(eid)
        return r["ko"] if r["standard"] else "(확장)"

    # DSD 주석 제목 (B-4 첨부는 zip에서 래핑)
    dsd = case.get("dsd")
    if not dsd:
        from dart_explorer.converters.document_wrap import wrap_as_contents
        z = zipfile.ZipFile(case["zip"])
        dsd = str(tmp / f"{name}.dsd")
        wrap_as_contents(z.read(case["entry"]), dsd)
    xlsx = str(tmp / f"{name}.xlsx")
    extract(dsd, xlsx)
    ctx = FootingContext(xlsx)
    from dsd_tool.attr_check import collect_note_titles
    titles = collect_note_titles(ctx)   # D-1 ②: 시트+원문 병행
    return {"facts": facts, "attrs_of": attrs_of, "std_label": std_label,
            "role_defs": tp.role_defs, "titles": titles,
            "doc_end": str(inst.doc_period_end)}


@pytest.fixture(scope="module")
def cases(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("v2")
    return {n: _load_case(n, tmp) for n in _CASES}, tmp


def test_g1_three_companies(cases):
    from dsd_tool.attr_check import attr_check
    data, tmp = cases
    for name, c in data.items():
        res = attr_check(
            c["facts"], c["attrs_of"], c["role_defs"], c["titles"],
            c["doc_end"], std_label=c["std_label"],
            out_path=str(tmp / f"속성검증_{name}.xlsx"))
        s = res["summary"]
        print(f"\n[V-2 G1 {name}] 기간 {s['period']['true']}/"
              f"{s['period']['total']} (위반 {s['period']['false']}·"
              f"불가 {s['period']['na']}) | 단위 {s['unit']['true']}/"
              f"{s['unit']['total']} (위반 {s['unit']['false']}) | "
              f"주석 명칭 매칭 {s['name']['true']}/{s['name']['total']} | "
              f"쌍 한쪽만 {s['pair_only']} · role만 {s['only_roles']} · "
              f"관례 decimals {s['doc_dec']}")
        assert s["period"]["total"] >= 100
        # 실공시에서도 위반이 나올 수 있다(V-2의 표면화 대상 — 회계사
        # 확인 영역). 붕괴 방지 상한만 검사하고 위반 목록은 출력.
        assert s["period"]["false"] <= 10
        for r in [x for x in res["period_rows"] if x["true"] is False][:6]:
            print(f"   위반: {r['element']} — {r['note']}")
        assert os.path.exists(res["out_path"])
        from openpyxl import load_workbook
        wb = load_workbook(res["out_path"])
        assert wb.sheetnames[0] == "요약"
        assert {"기간속성", "단위속성", "주석명칭"} <= set(wb.sheetnames)


def test_g2_tamper_three_kinds(cases):
    """변조 3종 — 각 1건 심어 해당 시트에서 검출."""
    from dsd_tool.attr_check import attr_check
    data, _tmp = cases
    c = data["신원"]

    # ① periodType 위반: instant 정의 element의 팩트를 duration으로
    f1 = copy.deepcopy(c["facts"])
    base0 = attr_check(c["facts"], c["attrs_of"], c["role_defs"],
                       c["titles"], c["doc_end"])
    ok_elems = {r["element"] for r in base0["period_rows"]
                if r["true"] is True}
    eid1 = next(e for e in f1
                if e in ok_elems
                and (c["attrs_of"](e) or {}).get("periodType") == "instant"
                and f1[e])
    base = attr_check(c["facts"], c["attrs_of"], c["role_defs"],
                      c["titles"], c["doc_end"])
    base_bad = {r["element"] for r in base["period_rows"]
                if r["true"] is False}
    f1[eid1][0]["type"] = "duration"
    r1 = attr_check(f1, c["attrs_of"], c["role_defs"], c["titles"],
                    c["doc_end"])
    new_bad = {r["element"] for r in r1["period_rows"]
               if r["true"] is False} - base_bad
    assert new_bad == {eid1}, new_bad          # 신규 전환은 변조 1건만

    # ② decimals 불일치: 화폐 팩트에 비관례 값
    f2 = copy.deepcopy(c["facts"])
    eid2 = next(e for e in f2
                if "monetary" in ((c["attrs_of"](e) or {})
                                  .get("type") or "").lower() and f2[e])
    base_u = {r["element"] for r in base0["unit_rows"]
              if r["true"] is False}
    f2[eid2][0]["decimals"] = "-4"
    r2 = attr_check(f2, c["attrs_of"], c["role_defs"], c["titles"],
                    c["doc_end"])
    bad2 = [r for r in r2["unit_rows"]
            if r["true"] is False and r["element"] not in base_u]
    assert len(bad2) == 1 and bad2[0]["element"] == eid2
    assert "비관례 decimals" in bad2[0]["note"]

    # ③ 주석명 변조: 매칭되던 제목을 임의 문자열로
    r0 = attr_check(c["facts"], c["attrs_of"], c["role_defs"],
                    c["titles"], c["doc_end"])
    matched_idx = next(i for i, r in enumerate(r0["name_rows"])
                       if r["true"])
    titles2 = list(c["titles"])
    titles2[matched_idx] = "99. 게이트변조주석명칭"
    r3 = attr_check(c["facts"], c["attrs_of"], c["role_defs"], titles2,
                    c["doc_end"])
    tampered = r3["name_rows"][matched_idx]
    assert tampered["true"] is False and "게이트변조" in tampered["title"]
    print(f"\n[V-2 G2] 변조 3종 검출 — 기간({eid1}) · decimals({eid2}) "
          f"· 주석명(행 {matched_idx + 1}) 각 1건만 FALSE 전환")

def _guide_inputs(pkg):
    """F-4b-lite+ 조립층 — 제출파일 단독 3종 입력 (HypercubeDef 재사용)."""
    from dart_explorer.xbrl.dimension_table import HypercubeDef, XbrlInstance
    cube = HypercubeDef(pkg)
    axis_members, domain_in_members = {}, {}
    for cubes in cube.by_base.values():
        for c in cubes:
            for ax in c["axes"]:
                key = f"{c['code']}:{ax['axis']}"
                axis_members[key] = list(ax["members"])
                domain_in_members[key] = any(
                    d in ax["members"] for d in ax["domains"])
    inst = XbrlInstance(pkg)
    axes_used = sorted({(a, eid) for eid, fl in inst.facts.items()
                        for f in fl
                        for a in (f["ctx"].get("dims") or {})})
    return {"axis_members": axis_members,
            "domain_in_members": domain_in_members,
            "axes_used": axes_used}


def test_g3_guide_sheet_in_v2_report(cases, tmp_path):
    """F-4b-lite+ — 삼성 첨부 인스턴스로 가이드검증형 시트 + 변조 검출."""
    from dsd_tool.attr_check import attr_check
    data, _tmp = cases
    c = data["삼성"]
    gi = _guide_inputs(_CASES["삼성"]["pkg"])
    out = str(tmp_path / "v2_guide.xlsx")
    res = attr_check(c["facts"], c["attrs_of"], c["role_defs"],
                     c["titles"], c["doc_end"], guide_inputs=gi,
                     out_path=out)
    g = res["summary"]["guide"]
    rules = {r["rule"] for r in res["guide_rows"]}
    assert rules == {"5.Ⅱ.4(1)나", "5.Ⅱ.4(1)라", "5.Ⅱ.3(1)아"}
    from openpyxl import load_workbook
    assert "가이드검증형" in load_workbook(out).sheetnames
    base_false = g["false"]
    # 변조 1건: 임의 축에 member 중복을 심는다
    import copy
    gi2 = copy.deepcopy(gi)
    key = next(k for k, v in gi2["axis_members"].items() if v)
    gi2["axis_members"][key] = gi2["axis_members"][key] +         [gi2["axis_members"][key][0]]
    res2 = attr_check(c["facts"], c["attrs_of"], c["role_defs"],
                      c["titles"], c["doc_end"], guide_inputs=gi2)
    assert res2["summary"]["guide"]["false"] == base_false + 1
    tampered = [r for r in res2["guide_rows"] if r["true"] is False and
                r["spot"] == key]
    assert tampered and "member 중복" in tampered[0]["note"]
    print(f"\n[F-4b-lite+ G] 삼성 가이드검증형 {g['total']}행 (위반 "
          f"{g['false']}) · 변조 1건 → +1 검출")

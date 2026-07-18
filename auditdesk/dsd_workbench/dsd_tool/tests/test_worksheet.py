"""F-1 게이트: 본문 워크시트 — 기공시사 자동 대조.

XBRL 기공시사(삼성전자, 코퍼스 내) DSD 투입 → 생성 워크시트의 element를
해당사 실제 공시 인스턴스와 자동 대조. 본문 element 일치율 목표 90%+
(홀드아웃 방식 — 해당사를 코퍼스 집계에서 제외하고 추천).
"""
import os

import pytest
from openpyxl import load_workbook

from dsd_tool.mapping import normalize
from dsd_tool.worksheet import build_worksheet

_REPO = os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.dirname(os.path.abspath(__file__)))))
_SAMSUNG_DSD = os.path.join(
    os.path.dirname(__file__), "..", "fixtures", "real",
    "[삼성전자(주)]_2025_[감사보고서].dsd")
_SAMSUNG_XBRL = os.path.join(
    _REPO, "auditlink-v2", "dart_explorer", "cache", "xbrl", "00126380",
    "20260310002820_11011")
if not os.path.isdir(_SAMSUNG_XBRL):        # 저장소 루트 변동 대비
    _SAMSUNG_XBRL = os.path.join(
        os.path.dirname(__file__), "..", "..", "..", "dart_explorer",
        "cache", "xbrl", "00126380", "20260310002820_11011")
_HANBIT = os.path.join(os.path.dirname(__file__), "..", "fixtures",
                       "synthetic", "한빛정밀_클린.dsd")
_SAMSUNG_CORP = "00126380"

_ready = os.path.exists(_SAMSUNG_DSD) and os.path.isdir(_SAMSUNG_XBRL)


def _answer_key():
    """삼성 실제 인스턴스 정답지.

    반환: (라벨→element 집합, 본문(D2~D6) 사용 element 전체 집합)
    라벨 표기 관행이 DSD와 인스턴스 간 다른 계정(예: DSD '매출채권' vs
    XBRL '매출채권 및 기타유동채권')은 라벨 키로는 대조 불가 —
    그 경우 "그 회사가 본문에서 실제 사용한 element인가"(집합 멤버십)로
    대조한다 (게이트 취지: 생성 워크시트 element ↔ 실제 공시 인스턴스).
    """
    import sys
    explorer_root = os.path.abspath(
        os.path.join(os.path.dirname(_SAMSUNG_XBRL), "..", "..", "..", ".."))
    if explorer_root not in sys.path:
        sys.path.insert(0, explorer_root)
    from dart_explorer.xbrl.corpus import extract_package
    key, fs_elements = {}, set()
    for row in extract_package(_SAMSUNG_XBRL):
        ln = normalize(row["label_ko"])
        if ln:
            key.setdefault(ln, set()).add(row["element_id"])
        if any(code[:2] in ("D2", "D3", "D4", "D5", "D6")
               for code in (row["roles"] or "").split(",") if code):
            fs_elements.add(row["element_id"])
    return key, fs_elements


@pytest.mark.skipif(not _ready, reason="삼성 DSD 또는 XBRL 캐시 없음")
def test_gate_f1_samsung_holdout(tmp_path):
    res = build_worksheet(
        _SAMSUNG_DSD, out_path=str(tmp_path / "삼성_ws.xlsx"),
        report_type="annual", holdout_corp=_SAMSUNG_CORP)
    assert os.path.exists(res["out_path"])
    key, fs_elements = _answer_key()

    checked = hit1 = hit4 = label_checked = label_hit = 0
    misses = []
    for sheet, rows in res["sheets"].items():
        for r in rows:
            if not r["top1"]:
                continue
            ln = normalize(r["label"])
            cand_ids = [c["element_id"] for c in r["candidates"][:4]]
            checked += 1

            def _ok(eid):
                if ln in key:               # 라벨 정답지 있으면 엄격 대조
                    return eid in key[ln]
                return eid in fs_elements   # 없으면 실사용 element 멤버십

            if ln in key:
                label_checked += 1
                label_hit += int(_ok(cand_ids[0]))
            hit1 += int(_ok(cand_ids[0]))
            ok4 = any(_ok(e) for e in cand_ids)
            hit4 += int(ok4)
            if not ok4:
                misses.append((sheet, r["label"], cand_ids[:2],
                              sorted(key.get(ln, []))[:2]))
    rate1 = hit1 / checked if checked else 0
    rate4 = hit4 / checked if checked else 0
    label_rate = label_hit / label_checked if label_checked else 0
    # 파이프 리다이렉트 시 cp949 콘솔 대비 ASCII-safe 출력
    print(f"\n[GATE F-1] mapped={checked} Top-4={rate4:.1%} "
          f"Top-1={rate1:.1%} (label-key n={label_checked}: "
          f"{label_rate:.1%})")
    for m in misses[:12]:
        print("  MISS:", ascii(m))
    assert checked >= 80, "대조 표본 부족"
    # 게이트 기준: D-3b와 일관 — 워크시트는 Top-4(대안 병기)이므로
    # Top-4 포함률 90%+. Top-1은 참고치로 병기.
    assert rate4 >= 0.90, f"본문 element Top-4 일치율 {rate4:.1%} < 90%"


@pytest.mark.skipif(not os.path.exists(_HANBIT), reason="한빛 DSD 없음")
def test_worksheet_structure(tmp_path):
    """양식 요건: 작성 개요·지원 열·확장 표기·CE member 블록·미완성 경고."""
    res = build_worksheet(_HANBIT, out_path=str(tmp_path / "한빛_ws.xlsx"),
                          report_type="annual")
    wb = load_workbook(res["out_path"])
    assert wb.sheetnames[0] == "작성 개요"
    assert {"BS", "PL", "CE", "CF"} <= set(wb.sheetnames)

    ov = "\n".join(str(c.value) for row in wb["작성 개요"].iter_rows()
                   for c in row if c.value)
    assert "완성" in ov and "회계사" in ov          # 미확정=미완성 경고
    assert "편집기 검색어" in [c.value for c in wb["BS"][1]]
    assert "확정 ☐" in [c.value for c in wb["BS"][1]]

    bs = wb["BS"]
    hdr = [c.value for c in bs[1]]
    id_col = hdr.index("element ID") + 1
    search_col = hdr.index("편집기 검색어") + 1
    ids = [bs.cell(r, id_col).value for r in range(2, bs.max_row + 1)]
    assert any(v and str(v).startswith("ifrs-full:") for v in ids)
    # 편집기 검색어 = 표준 한글 라벨 (한글 존재)
    searches = [bs.cell(r, search_col).value
                for r in range(2, bs.max_row + 1) if bs.cell(r, 1).value]
    assert any(s and any("가" <= ch <= "힣" for ch in str(s))
               for s in searches)

    # CE member 매핑 블록
    ce_text = "\n".join(str(c.value) for row in wb["CE"].iter_rows()
                        for c in row if c.value)
    assert "member 매핑" in ce_text
    assert "Member" in ce_text                     # member ID 추천 존재

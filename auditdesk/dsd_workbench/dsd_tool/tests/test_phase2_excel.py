"""Phase 2 workbook regressions; actual extract and writer, small corpus stub."""
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from dsd_tool.excel_out import extract
from dsd_tool.foot import foot
from dsd_tool.foot_excel import write_ai_footing, DETAIL_SHEET
from dsd_tool.tests.fixture import build_dsd
from dsd_tool.tests.test_h2 import _MERGED


@pytest.fixture
def extracted(tmp_path):
    path = str(tmp_path / 'source.xlsx')
    extract(build_dsd(str(tmp_path / 'source.dsd'), contents=_MERGED), path)
    return path


@pytest.mark.parametrize('value', ['Original title', 123456])
def test_m09_merged_return_link_keeps_source(extracted, tmp_path, value):
    result = foot(extracted, save=False)
    wb = load_workbook(extracted)
    wb['BS'].merge_cells('A2:M2')
    wb['BS']['A2'] = value
    wb['3'].merge_cells('C8:N8')
    wb['3']['C8'] = value
    wb.save(extracted)
    output = write_ai_footing(extracted, result, out_path=str(tmp_path/'out.xlsx'))
    out = load_workbook(output['out_path'])
    assert out['BS']['A2'].value == value
    assert out['3']['C8'].value == value
    assert any('HYPERLINK' in str(c.value) for row in out['BS'] for c in row)


@pytest.mark.parametrize('limit', [0, 2, 5])
def test_m10_core_result_retains_limit(extracted, limit):
    assert foot(extracted, limit=limit, save=False)['limit'] == limit


def test_m10_excel_preserves_fuzzy_semantics_and_limit(extracted, tmp_path):
    result = foot(extracted, limit=5, save=False)
    assert result['fuzzy'] > 0
    output = write_ai_footing(extracted, result, out_path=str(tmp_path/'out.xlsx'))
    wb = load_workbook(output['out_path'])
    ws = wb[DETAIL_SHEET]
    assert '±5' in ws['A1'].value
    fuzzy_rows = [r for r in range(6, ws.max_row+1) if ws.cell(r,1).value=='푸팅' and '단수차' in str(ws.cell(r,9).value)]
    assert fuzzy_rows, 'FUZZY must remain explicit, not TRUE while counted as an error'
    for r in fuzzy_rows:
        assert ws.cell(r,9).fill.fgColor.rgb == 'FFFFFF00'
        assert f'ABS(E{r}-H{r})<=5' in ws.cell(r,9).value


def test_m12_extracted_narrow_ce_builds(tmp_path, monkeypatch):
    from dsd_tool import worksheet
    monkeypatch.setattr(worksheet, 'MappingCorpus', lambda *a, **k: SimpleNamespace(std_labels={}))
    monkeypatch.setattr(worksheet, 'suggest', lambda *a, **k: {'verdict':'없음','similar_extensions':[]})
    result = worksheet.build_worksheet(build_dsd(str(tmp_path/'source.dsd')),
             out_path=str(tmp_path/'worksheet.xlsx'), include_notes=False)
    wb = load_workbook(result['out_path'])
    assert 'CE' in wb.sheetnames
    assert any('자본금' in str(c.value) for row in wb['CE'] for c in row)
    assert any(c.value == 100000 for row in wb['CE'] for c in row)


def test_m12_unsupported_ce_has_table_specific_error(tmp_path, monkeypatch):
    from dsd_tool import worksheet
    from dsd_tool.tests.fixture import CONTENTS_XML
    monkeypatch.setattr(worksheet, 'MappingCorpus', lambda *a, **k: SimpleNamespace(std_labels={}))
    monkeypatch.setattr(worksheet, 'suggest', lambda *a, **k: {'verdict':'없음','similar_extensions':[]})
    contents = CONTENTS_XML.replace('<TH>자본금</TH><TH>이익잉여금</TH>', '<TH>주석</TH>')
    contents = contents.replace('<TD ALIGN="RIGHT">100,000</TD><TD ALIGN="RIGHT">200,000</TD>', '<TD>1</TD>')
    with pytest.raises(ValueError, match='CE.*자본.*금액'):
        worksheet.build_worksheet(build_dsd(str(tmp_path/'unsupported.dsd'), contents=contents),
                                  out_path=str(tmp_path/'out.xlsx'), include_notes=False)


@pytest.mark.parametrize('limit,delta,verdict', [
    (0, 0, '일치'), (0, 1, '불일치'), (5, 1, '단수차'),
    (5, 5, '단수차'), (5, 6, '불일치'),
])
def test_m10_limit_boundary_core_excel_and_session_agree(tmp_path, monkeypatch, limit, delta, verdict):
    import json
    from openpyxl import Workbook
    from auditdesk.routers import workbench
    from dsd_tool.foot import FOOT_SHEET
    from dsd_tool.excel_out import MAP_SHEET
    source = str(tmp_path/'manual.xlsx')
    wb = Workbook(); ws = wb.active; ws.title = 'BS'
    ws.append(['계정', '주석', '당기', '전기'])
    for label, value in [('child1',100),('child2',200),('total',300+delta)]:
        ws.append([label, None, value, 300 if label == 'total' else value])
    mapping = wb.create_sheet(MAP_SHEET); mapping.append(['sheet','row','col'])
    for r in range(1,5):
        for c in range(1,5): mapping.append(['BS',r,c])
    levels = wb.create_sheet(FOOT_SHEET); levels.append(['[레벨]'])
    for r, level in [(2,1),(3,1),(4,2)]: levels.append(['BS',r,'label',None,level])
    wb.save(source)
    core = foot(source, limit=limit, save=False)
    assert len(core['foot']) == 2
    assert {row['scope']:row['verdict'] for row in core['foot']} == {'당기':verdict, '전기':'일치'}
    captured = {}
    monkeypatch.setattr(workbench, '_update', lambda sid, **values: captured.update(values))
    result = workbench._run_foot('s', source, True, limit, lambda *a: None)
    persisted = json.loads(captured['foot'])
    assert result['limit'] == persisted['limit'] == limit
    expected = {'match':2 if verdict=='일치' else 1,
                'rounding':1 if verdict=='단수차' else 0,
                'mismatch':1 if verdict=='불일치' else 0}
    assert {k:result['summary'][k] for k in expected} == expected
    assert result['summary'] == persisted['summary']
    out = load_workbook(result['ai_excel_path']); detail = out[DETAIL_SHEET]
    assert f'±{limit}' in detail['A1'].value
    for r in [6,7]:
        assert detail.cell(r,5).value.startswith('=')
        assert detail.cell(r,8).value.startswith('=')
        assert detail.cell(r,9).value == f'=IF(E{r}=H{r},TRUE,IF(ABS(E{r}-H{r})<={limit},"단수차",FALSE))'
        assert (detail.cell(r,9).fill.fgColor.rgb == 'FFFFFF00') == (r == 6 and verdict != '일치')
    summary = out['총괄표']
    row = next(r for r in range(5,summary.max_row+1) if summary.cell(r,7).value=='BS')
    assert summary.cell(row,9).value == (0 if verdict=='일치' else 1)


def test_m09_occupied_unmerged_cell_and_full_sheet_fallback():
    from openpyxl import Workbook
    from dsd_tool.cellsafe import put
    ws = Workbook().active; ws['M2'] = 0
    link = '=HYPERLINK("#summary!A1","return")'
    assert put(ws,2,13,value=link,preserve_value=True).coordinate == 'N2'
    assert ws['M2'].value == 0
    ws.cell(2,16384,'original'); failures=[]
    assert put(ws,2,13,value=link,preserve_value=True,failures=failures) is None
    assert ws['M2'].value == 0 and ws.cell(2,16384).value == 'original'
    assert failures and '원문 보존' in failures[0]['reason']


@pytest.mark.parametrize('note_header', ['주석', '주 석', '주\n석'])
def test_m12_ce_note_references_are_not_capital(tmp_path, monkeypatch, note_header):
    from dsd_tool import worksheet
    from dsd_tool.tests.fixture import CONTENTS_XML
    monkeypatch.setattr(worksheet, 'MappingCorpus', lambda *a, **k: SimpleNamespace(std_labels={}))
    monkeypatch.setattr(worksheet, 'suggest', lambda *a, **k: {'verdict':'없음','similar_extensions':[]})
    content = CONTENTS_XML.replace('<TH>과목</TH><TH>자본금</TH>', f'<TH>과목</TH><TH>{note_header}</TH><TH>자본금</TH>')
    content = content.replace('<TD>기초잔액</TD>', '<TD>기초잔액</TD><TD>9</TD>')
    result = worksheet.build_worksheet(build_dsd(str(tmp_path/'wide.dsd'), contents=content),
                                      out_path=str(tmp_path/'out.xlsx'), include_notes=False)
    ws = load_workbook(result['out_path'])['CE']
    values = [c.value for row in ws for c in row]
    assert 100000 in values and 200000 in values
    assert 9 not in values, 'note references must not become capital values'
    assert [c.value for c in ws[1] if str(c.value).startswith('값:')] == ['값: 자본금','값: 이익잉여금']

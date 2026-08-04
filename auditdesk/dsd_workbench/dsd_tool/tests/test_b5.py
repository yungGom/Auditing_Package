"""B-5 격상판 게이트: 산출물 완전성·정직성 — 합성 회귀 고정.

① 소스 식별 — 표지·사용안내에 소스 문서 파일명·수신 시각 명기
② 커버리지 — 원문 대비 시트화 %(가시 문자량) + 미시트화 블록 목록
③ 흡수 감지 — 주석 인식 실패 시 FS(CF) 시트로 흡수된 이질 콘텐츠를
   배너(수신물)+사용안내 목록으로 고지 (침묵 흡수 0)
(실파일 게이트는 조선내화 20250814002901 — 연결CF·CF 흡수 2면 +
미시트화 25블록 = 27개 지점 전부 고지, 커버리지 76.1% / 삼성 별도
_00760 — 감사의견·내부회계 본문 2블록 고지·오탐 0, 92.5% — 로컬 실측)
"""
import pytest
from openpyxl import load_workbook

from dsd_tool.excel_out import extract
from dsd_tool.tests.fixture import CONTENTS_XML, build_dsd

# 미시트화 블록: 첨부 재무제표 뒤·외부감사 앞의 큰 서술문 (섹션 밖 P)
_LOOSE = ("<P>" + "사업의 개요 서술문입니다. " * 30 + "</P>\r\n")
_WITH_LOOSE = CONTENTS_XML.replace(
    "<INSERTION>◆부속명세서 삽입◆</INSERTION>",
    _LOOSE + "<INSERTION>◆부속명세서 삽입◆</INSERTION>")

# 흡수형: 주석 TITLE이 '주석'이 아니고(폴백 미발동) 헤더가 문단 중간
# 텍스트 직후 번호(&cr; 경계도 없음) — 어떤 경계 규칙에도 걸리지 않아
# 주석 전체가 CF 블록으로 흡수되는 문서 (구 흡수 재현형)
_lo = CONTENTS_XML.find("<SECTION-2><TITLE>주석</TITLE>")
_hi = CONTENTS_XML.find("</SECTION-2>") + len("</SECTION-2>\r")
_ABSORBED = CONTENTS_XML[:_lo] + (
    "<SECTION-2><TITLE>재무제표 부속내역</TITLE>\r\n"
    "<P>부속내역은 다음과 같습니다. 2. 재무제표 작성기준 회사는 K-IFRS를"
    " 적용하고 있습니다. 3. 현금및현금성자산 내역은 생략합니다.</P>\r\n"
    '<TABLE WIDTH="600"><TR><TH>구분</TH><TH>금액</TH></TR>'
    '<TR><TD>보통예금</TD><TD ALIGN="RIGHT">500,000</TD></TR></TABLE>\r\n'
    "</SECTION-2>\r") + CONTENTS_XML[_hi:]


def _extract(tmp_path, contents, name, no_meta=True):
    dsd = build_dsd(str(tmp_path / f"{name}.dsd"), contents=contents,
                    no_meta=no_meta)
    xlsx = str(tmp_path / f"{name}.xlsx")
    return extract(dsd, xlsx), xlsx


def _guide_text(xlsx):
    wb = load_workbook(xlsx, read_only=True)
    t = "\n".join(str(c.value) for row in wb["사용안내"].iter_rows()
                  for c in row if c.value)
    wb.close()
    return t


def test_source_identification(tmp_path):
    """소스 문서 파일명·수신 시각 — 사용안내·표지 모두 명기."""
    info, xlsx = _extract(tmp_path, CONTENTS_XML, "src")
    g = _guide_text(xlsx)
    assert "소스 문서: src.dsd" in g and "수신 시각" in g
    wb = load_workbook(xlsx, read_only=True)
    cover = "\n".join(str(c.value) for row in wb["표지"].iter_rows()
                      for c in row if c.value)
    wb.close()
    assert "소스 문서: src.dsd" in cover
    assert "시트화 커버리지" in g               # 커버리지 상시 표기


def test_coverage_uncovered_blocks(tmp_path):
    """시트 밖 큰 서술문 → 미시트화 블록 목록 + 커버리지 하락."""
    base, _ = _extract(tmp_path, CONTENTS_XML, "base")
    loose, xlsx = _extract(tmp_path, _WITH_LOOSE, "loose")
    assert loose["coverage"] < base["coverage"]
    assert len(loose["uncovered_blocks"]) == \
        len(base["uncovered_blocks"]) + 1
    blk = loose["uncovered_blocks"][-1]
    assert blk["chars"] >= 200 and "사업의개요" in blk["preview"]
    assert "미시트화 블록" in _guide_text(xlsx)


def test_absorbed_banner_and_list(tmp_path):
    """주석 인식 실패 문서 — CF 흡수 배너(수신물)+사용안내 목록."""
    info, xlsx = _extract(tmp_path, _ABSORBED, "abs")
    assert info["note_count"] == 0              # 흡수 상황 재현
    assert "CF" in info["absorbed"]
    marks = " ".join(info["absorbed"]["CF"])
    assert "재무제표 부속내역" in marks or "주석형 헤더" in marks
    g = _guide_text(xlsx)
    assert "흡수 의심" in g
    wb = load_workbook(xlsx)
    assert "흡수" in str(wb["CF"].cell(1, 9).value or "")   # 배너 (행 불변)
    assert str(wb["CF"].cell(1, 9).fill.start_color.rgb).endswith("FFEB9C")


def test_absorbed_banner_wrapped_only(tmp_path):
    """편집기 계열(meta 있음)은 시트 배너 없음 — 사용안내 목록만."""
    info, xlsx = _extract(tmp_path, _ABSORBED, "abs_edit", no_meta=False)
    assert "CF" in info["absorbed"]             # 통계는 동일 노출
    wb = load_workbook(xlsx)
    assert wb["CF"].cell(1, 9).value in (None, "")

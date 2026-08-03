"""F-4b-lite 게이트: [가이드 체크] 시트.

G1 검증형 4종 단위 — 위반 심은 합성 입력으로 검출 확인
G2 실전 스캐폴드 재생성(신원) — 체크 시트 부착·충전율 리포트
(G3 전 테스트 회귀는 스위트 전체 실행으로 별도)
"""
import os
import sys

import pytest

from dsd_tool.guide_check import (MACHINE_RULES, attach_guide_check,
                                  load_active_rules, machine_check)

_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__),
                                     "..", "..", ".."))
_ASSET = os.path.join(_ROOT, "assets", "guide", "guide_rules_2026.json")
pytestmark = pytest.mark.skipif(not os.path.exists(_ASSET),
                                reason="guide_rules 자산 없음")


def test_active_rules_loaded():
    rules = load_active_rules()
    assert len(rules) == 133                    # 검수 반영본 전원 승인
    assert all(r["approved"] for r in rules)


def test_machine_four_kinds_detect_violations():
    """위반 심은 합성 입력 — 4종 전부 검출 + 정상 입력 통과."""
    bad = {
        "axis_members": {"축A": ["M1", "M2", "M1"]},          # 중복
        "table_headers": {"표1": ["구분", "당기", "합계"]},    # 합계열
        "axes_used": [("ifrs-full_CurrentAndNoncurrentAxis",
                       "ifrs-full_Inventories")],             # 비변동 사용
        "open_close": [("표2", "eid_open", "eid_close")],     # 기초≠기말
    }
    good = {
        "axis_members": {"축A": ["M1", "M2"]},
        "table_headers": {"표1": ["구분", "당기", "전기"]},
        "axes_used": [("ifrs-full_CurrentAndNoncurrentAxis",
                       "ChangesInInventories")],              # 변동 예외
        "open_close": [("표2", "eid_same", "eid_same")],
    }
    for rid in MACHINE_RULES:
        bad_out = machine_check(rid, bad)
        assert any(v == "위반" for _s, v, _n in bad_out), rid
        good_out = machine_check(rid, good)
        assert all(v != "위반" for _s, v, _n in good_out), rid
        # 입력 없음 → 판단필요 (침묵 통과 금지)
        none_out = machine_check(rid, {})
        assert none_out[0][1] == "판단필요"

    # 조건부 승인(2026-08-03) 3갈래 검증 — 5.Ⅱ.3(1)아:
    from dsd_tool.guide_check import CNC_CHANGE_AXIS_WHITELIST
    wl_axis = next(iter(CNC_CHANGE_AXIS_WHITELIST))
    tri = machine_check("5.Ⅱ.3(1)아", {"axes_used": [
        (wl_axis, "ifrs-full_FinancialAssets"),                # ① 화이트리스트
        ("entity99_CurrentAndNoncurrentChangesAxis",
         "ifrs-full_Inventories"),                             # ② 확장 변동성
        ("ifrs-full_CurrentAndNoncurrentAxis",
         "ifrs-full_Inventories"),                             # ③ 위반 유지
    ]})
    verdicts = {spot.split(" @ ")[0]: v for spot, v, _ in tri}
    assert verdicts[wl_axis] == "통과"
    assert verdicts["entity99_CurrentAndNoncurrentChangesAxis"] ==         "판단필요"                                       # 오탐 후보 정지
    assert verdicts["ifrs-full_CurrentAndNoncurrentAxis"] == "위반"
    print("\n[F-4b-lite G1] 검증형 4종 — 위반 검출·정상 통과·입력 없음"
          " 판단필요 전부 확인")


@pytest.fixture(scope="module")
def scaffold(tmp_path_factory):
    """신원 실전 스캐폴드 재생성 (F-3b 게이트 조립 재사용)."""
    if _ROOT not in sys.path:
        sys.path.insert(0, _ROOT)
    pkg = os.path.join(_ROOT, "dart_explorer", "cache", "xbrl",
                       "00136925", "20260323000596_11011")
    doc = os.path.join(_ROOT, "dart_explorer", "cache", "document",
                       "00136925")
    if not os.path.isdir(pkg):
        pytest.skip("신원 캐시 없음")
    from dart_explorer.xbrl.dimension_table import HypercubeDef, XbrlInstance
    from dsd_tool.excel_out import extract
    from dsd_tool.rollforward import rollforward
    from dsd_tool.succession import load_assets
    tmp = tmp_path_factory.mktemp("f4b")
    half = str(tmp / "half.xlsx")
    ye = str(tmp / "ye.xlsx")
    extract(os.path.join(doc, "20250814001572.dsd"), half)
    extract(os.path.join(doc, "20260323000596.dsd"), ye)
    inst = XbrlInstance(pkg)
    facts = {}
    for eid, fl in inst.facts.items():
        facts[eid] = [{
            "value": f["value"], "decimals": f.get("decimals"),
            "type": f["ctx"]["type"], "start": f["ctx"].get("start"),
            "end": f["ctx"].get("end"),
            "dims": dict(f["ctx"].get("dims") or {}),
        } for f in fl]
    out = str(tmp / "스캐폴드.xlsx")
    res = rollforward(
        half, facts, inst.doc_period_end,
        load_assets(os.path.join(pkg, "succession_assets.json")),
        ye_xlsx=ye, out_path=out)

    # 검증형 입력 조립 (기존 파서 재사용 — HypercubeDef·인스턴스 dims)
    cube = HypercubeDef(pkg)
    # 중복 검사는 '동일 표(role)의 동일 축' 단위 — role 간 합산은
    # 정상적 재사용(5.Ⅱ.3(1)다)을 가짜 중복으로 만든다
    axis_members = {}
    for cubes in cube.by_base.values():
        for c in cubes:
            for ax in c["axes"]:
                axis_members[f"{c['code']}:{ax['axis']}"] =                     list(ax["members"])
    axes_used = sorted({(a, eid) for eid, fl in facts.items()
                        for f in fl for a in (f["dims"] or {})})
    inputs = {"axis_members": axis_members, "axes_used": axes_used}
    from dsd_tool.foot import FootingContext
    ctx = FootingContext(half)
    scope = {"note_titles": [str(ctx.wb[s].cell(1, 1).value or "")
                             for s in ctx.note_sheets]}
    return out, inputs, scope


def test_scaffold_guide_check_attached(scaffold):
    out, inputs, scope = scaffold
    stats = attach_guide_check(out, inputs=inputs, scope=scope)
    print(f"\n[F-4b-lite G2] 신원 스캐폴드 체크 시트 — 행 {stats['rows']}"
          f" · 기계 판정 {stats['machine']} · 위반 {stats['violations']}"
          f" · 판단필요 {stats['need_judge']} · 해당없음 {stats['na']}"
          f" · 충전율 {stats['fill_rate']:.1%}")
    assert stats["rows"] >= 133                 # 규칙 × 적용 지점 전개
    assert stats["machine"] >= 4                # 검증형 실제 판정 발생
    # 충전율 = 기계 판정+판단필요 노출 / 전 행 (해당없음은 Ⅲ절 스코프
    # 밖 규칙 — 정상 분류). 전 행이 상태를 부여받았는지도 검사.
    assert stats["fill_rate"] is not None and stats["fill_rate"] >= 0.35
    assert stats["machine"] + stats["need_judge"] + stats["na"] ==         stats["rows"]                           # 침묵 공란 0
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert "가이드 체크" in wb.sheetnames

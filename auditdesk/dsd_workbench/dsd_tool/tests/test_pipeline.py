"""검증 게이트 (스펙 5장) — 합성 DSD 기반 자동 테스트.

G1 추출 / G2 무변경 / G3 단일변경 / G4 대량변경 / G5 빈셀삽입
(DART 편집기 열림 확인은 실제 파일로 수동 검증)
"""
import os
import zipfile

import pytest
from openpyxl import load_workbook

from dsd_tool.excel_out import MAP_SHEET, extract
from dsd_tool.repack import diff, repack
from dsd_tool.scanner import scan
from dsd_tool.textutil import (clean_text, escape_for_dsd, format_number_like,
                               match_fs_title, try_number)
from dsd_tool.zipsplice import read_contents, replace_contents

from .fixture import CONTENTS_XML, build_dsd


# ---------------------------------------------------------------------------
# 단위: textutil
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("value,expected", [
    ("1,234,567", 1234567.0),
    ("(234,567)", -234567.0),
    ("-1,234", -1234.0),
    ("1,234.56", 1234.56),
    ("0", 0.0),
    ("500000", 500000.0),
    ("3,4,5,6", None),        # 주석 참조 → 텍스트 유지
    ("12,34", None),          # 천단위 위반
    ("1,2345", None),
    ("-", None),              # 대시 placeholder
    ("", None),
    ("012", None),            # 선행 0 → 텍스트
    ("2026", 2026.0),
    ("5.5%", None),
    ("2026-12-31", None),
    ("2026.12.31", None),
])
def test_try_number(value, expected):
    got = try_number(value)
    if expected is None:
        assert got is None
    else:
        assert got == pytest.approx(expected)


def test_clean_text_and_escape_roundtrip():
    raw = '<SPAN USERMARK=" B">1. 일반</SPAN>&amp;cr;&amp;cr;본문 A&amp;B'
    assert clean_text(raw) == "1. 일반\n\n본문 A&B"
    # 역변환 이스케이프 순서: \n → &cr; 먼저, 그다음 & → &amp;
    assert escape_for_dsd("본문 A&B\n다음줄") == "본문 A&amp;B&amp;cr;다음줄"


@pytest.mark.parametrize("orig,value,expected", [
    ("1,234,567", 9876543, "9,876,543"),
    ("(234,567)", -300000, "(300,000)"),
    ("500000", 7, "7"),
    ("", 12345, "12,345"),       # 빈 셀 삽입 기본: 콤마
    ("1,234", -5, "-5"),         # 원본이 괄호 아님 → '-'
])
def test_format_number_like(orig, value, expected):
    assert format_number_like(orig, value) == expected


@pytest.mark.parametrize("title,expected", [
    ("재 무 상 태 표", ("", "", "재무상태표")),
    ("재무상태표", ("", "", "재무상태표")),
    ("연 결 재 무 상 태 표", ("", "연결", "재무상태표")),
    ("반기재무상태표", ("반기", "", "재무상태표")),
    ("포괄손익계산서", ("", "", "포괄손익계산서")),
    ("이익잉여금처분계산서", ("", "", "이익잉여금처분계산서")),
    ("주석", None),
])
def test_match_fs_title(title, expected):
    assert match_fs_title(title) == expected


# ---------------------------------------------------------------------------
# 스캐너
# ---------------------------------------------------------------------------

def test_scan_structure():
    doc = scan(CONTENTS_XML)
    # FS 4종 동적 감지, 포괄손익 단독 → PL
    assert [b.sheet_name for b in doc.fs_blocks] == ["BS", "PL", "CE", "CF"]
    # 주석 3개, USERMARK 방식
    assert doc.note_mode == "usermark"
    assert [n.number for n in doc.notes] == [1, 2, 3]
    assert doc.notes[0].title == "일반적 사항"
    # 외부감사 TE 테이블
    assert len(doc.te_tables) == 1
    te_cells = [c for c in doc.te_tables[0].all_cells() if c.tag == "TE"]
    assert len(te_cells) == 4
    assert te_cells[0].attrs.get("ACODE") == "A001"
    # 표지
    assert any("제 57 기" in p.text for p in doc.cover_paragraphs)


def test_scan_consolidated_and_five_fs():
    """연결 접두사 + 손익/포괄손익 분리(5종) → 연결PL / 연결PL1."""
    xml = CONTENTS_XML.replace("재 무 상 태 표", "연 결 재 무 상 태 표") \
        .replace("포괄손익계산서</TD>", "연결손익계산서</TD>") \
        .replace(">자 본 변 동 표<", ">연결포괄손익계산서<") \
        .replace(">현 금 흐 름 표<", ">연결자본변동표<")
    doc = scan(xml)
    assert [b.sheet_name for b in doc.fs_blocks] == \
        ["연결BS", "연결PL", "연결PL1", "연결CE"]


def test_note_fallback_span_id():
    xml = CONTENTS_XML.replace('USERMARK=" B"', 'ID="note"')
    doc = scan(xml)
    assert doc.note_mode == "span-id"
    assert [n.number for n in doc.notes] == [1, 2, 3]


# ---------------------------------------------------------------------------
# G1: 추출
# ---------------------------------------------------------------------------

def _make(tmp_path):
    dsd = build_dsd(str(tmp_path / "테스트.dsd"))
    info = extract(dsd, str(tmp_path / "테스트.xlsx"))
    return dsd, info


def test_g1_extract(tmp_path):
    dsd, info = _make(tmp_path)
    wb = load_workbook(info["out_path"])
    assert wb.sheetnames == ["사용안내", "원문", "표지", "BS", "PL", "CE",
                             "CF", "1", "2", "3", "외부감사", "_MAP",
                             "_META"]
    # 원문 통합 시트 (스펙 7.5 P3) — 참조용, _MAP 미포함이라 편집 불가
    verbatim = wb["원문"]
    vtext = "\n".join(str(c.value) for row in verbatim.iter_rows()
                      for c in row if c.value)
    assert "[[ BS ]]" in vtext and "[[ 3 ]]" in vtext
    assert 1234567 in [c.value for row in verbatim.iter_rows() for c in row]
    assert wb[MAP_SHEET].sheet_state == "hidden"
    assert info["mapped_cells"] > 30
    # 숫자 셀은 진짜 숫자로 저장 (ACCIO 한계 극복)
    bs = wb["BS"]
    values = [c.value for row in bs.iter_rows() for c in row]
    assert 1234567 in values
    assert -234567 in values          # 괄호 음수 → 음수
    assert "3,4" in values            # 주석 참조는 텍스트 유지


# ---------------------------------------------------------------------------
# G2: 무변경 역변환 → 바이트 동일
# ---------------------------------------------------------------------------

def test_g2_no_change_roundtrip(tmp_path):
    """G2 바이트 동일은 원문 보존 모드(--keep-cr)로 검증."""
    dsd, info = _make(tmp_path)
    out = str(tmp_path / "out.dsd")
    result = repack(info["out_path"], dsd, out, clean_cr=False)
    assert result["changes"] == []
    assert result["stats"]["cr_only_total"] == 2   # 보존 경고용 카운트
    assert open(out, "rb").read() == open(dsd, "rb").read()


def test_g2_replace_contents_identity(tmp_path):
    """스플라이스 경로 자체 검증: 동일 내용 교체 → contents.xml 바이트 동일."""
    dsd = build_dsd(str(tmp_path / "t.dsd"))
    data = open(dsd, "rb").read()
    new_zip = replace_contents(data, CONTENTS_XML.encode("utf-8"))
    assert read_contents(new_zip) == CONTENTS_XML.encode("utf-8")
    with zipfile.ZipFile(str(tmp_path / "t.dsd")) as zin:
        orig_others = {n: zin.read(n) for n in zin.namelist()
                       if n != "contents.xml"}
    tmp2 = tmp_path / "t2.dsd"
    tmp2.write_bytes(new_zip)
    with zipfile.ZipFile(str(tmp2)) as z:
        assert z.testzip() is None
        for n, blob in orig_others.items():
            assert z.read(n) == blob


# ---------------------------------------------------------------------------
# G3: 단일 변경 → 정확히 1곳만 변경
# ---------------------------------------------------------------------------

def _set_cell(xlsx, sheet, find_value, new_value):
    wb = load_workbook(xlsx)
    ws = wb[sheet]
    for row in ws.iter_rows():
        for c in row:
            if c.value == find_value:
                c.value = new_value
                wb.save(xlsx)
                return c.row, c.column
    raise AssertionError(f"{find_value!r} not found in {sheet}")


def test_g3_single_change(tmp_path):
    dsd, info = _make(tmp_path)
    _set_cell(info["out_path"], "BS", 1234567, 9876543)
    out = str(tmp_path / "out.dsd")
    result = repack(info["out_path"], dsd, out, clean_cr=False)
    assert len(result["changes"]) == 1
    new_text = read_contents(open(out, "rb").read()).decode("utf-8")
    assert new_text.count("9,876,543") == 1
    assert "1,234,567" not in new_text
    # 변경 구간 외 나머지는 완전 동일
    ch = result["changes"][0]
    assert new_text[:ch["xml_start"]] == CONTENTS_XML[:ch["xml_start"]]
    assert new_text[ch["xml_start"] + len(ch["new_raw"]):] == \
        CONTENTS_XML[ch["xml_end"]:]


def test_g3_false_change_guard(tmp_path):
    """숫자를 콤마 문자열로 다시 타이핑해도 거짓변경 없음 (스펙 3.4)."""
    dsd, info = _make(tmp_path)
    _set_cell(info["out_path"], "BS", 1234567, "1,234,567")
    changes, _, _, _ = diff(info["out_path"], dsd, clean_cr=False)
    assert changes == []


# ---------------------------------------------------------------------------
# G4: 대량 변경
# ---------------------------------------------------------------------------

def test_g4_multi_change(tmp_path):
    dsd, info = _make(tmp_path)
    xlsx = info["out_path"]
    _set_cell(xlsx, "BS", 1234567, 9999999)
    _set_cell(xlsx, "BS", -234567, -300000)     # 괄호 음수 스타일 유지
    _set_cell(xlsx, "BS", "3,4", "3,4,5")       # 텍스트 셀
    _set_cell(xlsx, "PL", 999999, 1000001)
    _set_cell(xlsx, "2", 12581632, 12581700)    # 주석 테이블
    out = str(tmp_path / "out.dsd")
    result = repack(xlsx, dsd, out, clean_cr=False)
    assert len(result["changes"]) == 5
    new_text = read_contents(open(out, "rb").read()).decode("utf-8")
    for token in ["9,999,999", "(300,000)", ">3,4,5<", "1,000,001", "12,581,700"]:
        assert token in new_text, token
    # 원본 스타일 보존: 음수는 괄호로
    assert "-300,000" not in new_text
    with zipfile.ZipFile(out) as z:
        assert z.testzip() is None


# ---------------------------------------------------------------------------
# G5: 빈 셀 삽입
# ---------------------------------------------------------------------------

def _find_empty_mapped(xlsx, sheet):
    wb = load_workbook(xlsx)
    for row in wb[MAP_SHEET].iter_rows(min_row=2, values_only=True):
        if row[0] == sheet and (row[5] is None or row[5] == ""):
            return int(row[1]), int(row[2])
    raise AssertionError("빈 매핑 셀 없음")


def test_g5_empty_cell_insert(tmp_path):
    dsd, info = _make(tmp_path)
    xlsx = info["out_path"]
    r, c = _find_empty_mapped(xlsx, "BS")
    wb = load_workbook(xlsx)
    wb["BS"].cell(row=r, column=c).value = 12345
    wb.save(xlsx)
    out = str(tmp_path / "out.dsd")
    result = repack(xlsx, dsd, out, clean_cr=False)
    assert len(result["changes"]) == 1
    assert result["changes"][0]["new_raw"] == "12,345"
    # 재추출로 왕복 확인
    info2 = extract(out, str(tmp_path / "재추출.xlsx"))
    wb2 = load_workbook(info2["out_path"])
    assert wb2["BS"].cell(row=r, column=c).value == 12345


def test_te_cell_edit(tmp_path):
    """외부감사 TE 셀 수정 (빈 TE 포함)."""
    dsd, info = _make(tmp_path)
    xlsx = info["out_path"]
    _set_cell(xlsx, "외부감사", 10, 55)
    out = str(tmp_path / "out.dsd")
    result = repack(xlsx, dsd, out, clean_cr=False)
    assert len(result["changes"]) == 1
    new_text = read_contents(open(out, "rb").read()).decode("utf-8")
    assert '<TE ACODE="A003">55</TE>' in new_text


def test_escape_and_newline_edit(tmp_path):
    dsd, info = _make(tmp_path)
    xlsx = info["out_path"]
    _set_cell(xlsx, "BS", "자산총계", "자산&부채\n총계")
    out = str(tmp_path / "out.dsd")
    repack(xlsx, dsd, out)
    new_text = read_contents(open(out, "rb").read()).decode("utf-8")
    assert "자산&amp;부채&amp;cr;총계" in new_text


# 실제 공시 DSD (리포지토리에 있을 때만 실행)
_SAMSUNG = os.path.join(os.path.dirname(__file__), "..", "fixtures", "real",
                        "[삼성전자(주)]_2025_[감사보고서].dsd")


# ---------------------------------------------------------------------------
# 주석 번호 중복 정리 ("1. 1. 제목" → "1. 제목")
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("text,expected", [
    ("1. 1. 일반적 사항", "1. 일반적 사항"),
    ("12. 12. 우발부채", "12. 우발부채"),
    ("1. 2. 어쩌구", "1. 2. 어쩌구"),      # 비중복(다른 번호)은 그대로
    ("1. 일반적 사항", "1. 일반적 사항"),   # 이미 정상
    ("1. 12. 항목", "1. 12. 항목"),        # \1 부분일치 방지
])
def test_dedup_note_number(text, expected):
    from dsd_tool.textutil import dedup_note_number
    assert dedup_note_number(text) == expected


def _dup_fixture_xml():
    """주석 1 헤더를 dartdb식 중복 번호로 오염시킨 변형."""
    return CONTENTS_XML.replace(
        '<SPAN USERMARK=" B">1. 일반적 사항</SPAN>',
        '<SPAN USERMARK=" B">1. 1. 일반적 사항</SPAN>')


def test_note_dedup_extract_and_repack(tmp_path):
    dsd = build_dsd(str(tmp_path / "dup.dsd"), _dup_fixture_xml())

    # 기본: 정리된 값이 엑셀에 기록
    info = extract(dsd, str(tmp_path / "dup.xlsx"))
    assert info["deduped_notes"] == 1
    wb = load_workbook(info["out_path"])
    assert wb["1"].cell(1, 1).value == "1. 일반적 사항"
    assert wb["2"].cell(1, 1).value == "2. 재무제표 작성기준"  # 비오염은 그대로

    # 사용자 수정 없이 repack → 헤더 차이만 자동 반영
    out = str(tmp_path / "dup_수정.dsd")
    result = repack(info["out_path"], dsd, out)
    dedups = [c for c in result["changes"] if c["reason"] == "note-dedup"]
    cleans = [c for c in result["changes"] if c["reason"] == "clean-cr"]
    assert len(dedups) == 1 and len(cleans) == 2   # 기본: &cr; 정리 동반
    assert len(result["changes"]) == 3
    new_text = read_contents(open(out, "rb").read()).decode("utf-8")
    assert '<SPAN USERMARK=" B">1. 일반적 사항</SPAN>' in new_text
    assert "1. 1." not in new_text

    # 멱등성: 정리된 DSD 재추출 → 추가 변경 없음
    info2 = extract(out, str(tmp_path / "re.xlsx"))
    assert info2["deduped_notes"] == 0
    assert info2["note_count"] == 3      # SPAN 보존 → 주석 감지 유지
    changes2, _, _, stats2 = diff(info2["out_path"], out)
    assert changes2 == [] and stats2["cr_only_total"] == 0


def test_note_dedup_keep_option(tmp_path):
    dsd = build_dsd(str(tmp_path / "dup.dsd"), _dup_fixture_xml())
    info = extract(dsd, str(tmp_path / "dup.xlsx"), keep_note_numbers=True)
    assert info["deduped_notes"] == 0
    wb = load_workbook(info["out_path"])
    assert wb["1"].cell(1, 1).value == "1. 1. 일반적 사항"
    # 원본 유지 모드에서는 무변경 repack이 바이트 동일 (G2)
    out = str(tmp_path / "out.dsd")
    result = repack(info["out_path"], dsd, out, clean_cr=False)
    assert result["changes"] == []
    assert open(out, "rb").read() == open(dsd, "rb").read()


@pytest.mark.skipif(not os.path.exists(_SAMSUNG), reason="실제 DSD 없음")
def test_note_dedup_real_samsung(tmp_path):
    """실제 dartdb DSD: 주석 32개 전부 "N. N." 오염 → 전부 정리."""
    xlsx = str(tmp_path / "삼성.xlsx")
    info = extract(_SAMSUNG, xlsx)
    assert info["deduped_notes"] == 32
    out = str(tmp_path / "삼성_수정.dsd")
    result = repack(xlsx, _SAMSUNG, out)
    dedups = [c for c in result["changes"] if c["reason"] == "note-dedup"]
    cleans = [c for c in result["changes"] if c["reason"] == "clean-cr"]
    assert len(dedups) == 32 and len(cleans) == 319
    new_text = read_contents(open(out, "rb").read()).decode("utf-8")
    assert '<SPAN USERMARK=" B">5. 금융자산의 양도</SPAN>' in new_text
    info2 = extract(out, str(tmp_path / "재추출.xlsx"))
    assert info2["note_count"] == 32 and info2["deduped_notes"] == 0


# ---------------------------------------------------------------------------
# --clean-cr: &cr;-only 셀 정리
# ---------------------------------------------------------------------------

def test_extract_reports_cr_only(tmp_path):
    dsd, info = _make(tmp_path)
    # 픽스처: <TD>&cr;</TD> + <TD>&cr;&cr;</TD> 2개 (P 문단의 &cr;은 제외)
    assert info["cr_only_cells"] == 2
    assert info["cr_only_by_sheet"] == {"BS": 2}


def test_clean_cr(tmp_path):
    dsd, info = _make(tmp_path)
    out = str(tmp_path / "clean.dsd")
    result = repack(info["out_path"], dsd, out, clean_cr=True)
    cleans = [c for c in result["changes"] if c["reason"] == "clean-cr"]
    assert len(cleans) == 2
    assert all(c["new_raw"] == "" for c in cleans)

    new_text = read_contents(open(out, "rb").read()).decode("utf-8")
    assert "<TD>&amp;cr;" not in new_text        # cr-only 셀 정리됨
    assert "주석&amp;cr;참조" in new_text         # 값+&cr; 혼합 셀 보존
    assert "<P>&amp;cr;</P>" in new_text          # P 문단의 의도적 빈 줄 보존

    # 클린업 후 재추출: 해당 셀은 빈 셀, 혼합 셀 텍스트는 그대로
    info2 = extract(out, str(tmp_path / "재추출.xlsx"))
    assert info2["cr_only_cells"] == 0
    wb = load_workbook(info2["out_path"])
    for ch in cleans:
        assert wb[ch["sheet"]].cell(row=ch["row"], column=ch["col"]).value \
            in (None, "")
    vals = [c.value for row in wb["BS"].iter_rows() for c in row]
    assert "주석\n참조" in vals


def test_clean_cr_user_edit_wins(tmp_path):
    """cr-only 셀에 사용자가 값을 입력하면 clean-cr이 아니라 edit으로 처리."""
    dsd, info = _make(tmp_path)
    wb = load_workbook(info["out_path"])
    target = None
    for row in wb["_MAP"].iter_rows(min_row=2, values_only=True):
        if row[5] == "&amp;cr;":
            target = (row[0], int(row[1]), int(row[2]))
            break
    assert target
    wb[target[0]].cell(row=target[1], column=target[2]).value = 42
    wb.save(info["out_path"])
    changes, _, _, _ = diff(info["out_path"], dsd, clean_cr=True)
    hit = [c for c in changes
           if (c["sheet"], c["row"], c["col"]) == target]
    assert hit[0]["reason"] == "edit" and hit[0]["new_raw"] == "42"


def test_keep_cr_disables_clean(tmp_path):
    """clean_cr=False(--keep-cr)면 &cr;-only 셀 보존, stats로 경고 가능."""
    dsd, info = _make(tmp_path)
    changes, _, _, stats = diff(info["out_path"], dsd, clean_cr=False)
    assert changes == []
    assert stats["cr_only_total"] == 2


# ---------------------------------------------------------------------------
# 실제 공시 DSD (있을 때만): 삼성전자 별도 clean-cr
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not os.path.exists(_SAMSUNG), reason="실제 DSD 없음")
def test_clean_cr_real_samsung(tmp_path):
    from dsd_tool.textutil import CLEANABLE_TAGS, CR_ONLY_RE, enclosing_tag

    xlsx = str(tmp_path / "삼성.xlsx")
    # 주석 번호 정리 변경이 섞이지 않도록 원본 유지 모드로 추출
    info = extract(_SAMSUNG, xlsx, keep_note_numbers=True)
    text = read_contents(open(_SAMSUNG, "rb").read()).decode("utf-8")

    # 매핑 셀 중 혼합 셀(값+&cr;) 실측: 91개 (2026-07-07 기준 회귀값)
    wb = load_workbook(xlsx)
    mixed_spans = set()
    for row in wb["_MAP"].iter_rows(min_row=2, values_only=True):
        s, e = int(row[3]), int(row[4])
        raw = text[s:e]
        if enclosing_tag(text, s) in CLEANABLE_TAGS and \
                "&amp;cr;" in raw and not CR_ONLY_RE.match(raw):
            mixed_spans.add((s, e))
    assert len(mixed_spans) == 91
    assert info["cr_only_cells"] == 319

    out = str(tmp_path / "삼성_clean.dsd")
    result = repack(xlsx, _SAMSUNG, out, clean_cr=True)
    cleans = result["changes"]
    assert len(cleans) == 319
    assert all(c["reason"] == "clean-cr" for c in cleans)
    # 혼합 셀 구간은 전혀 건드리지 않음
    clean_spans = {(c["xml_start"], c["xml_end"]) for c in cleans}
    assert not (clean_spans & mixed_spans)

    # 재추출: cr-only 0, 혼합 셀 수는 그대로 91
    xlsx2 = str(tmp_path / "삼성2.xlsx")
    info2 = extract(out, xlsx2)
    assert info2["cr_only_cells"] == 0
    new_text = read_contents(open(out, "rb").read()).decode("utf-8")
    wb2 = load_workbook(xlsx2)
    mixed2 = 0
    for row in wb2["_MAP"].iter_rows(min_row=2, values_only=True):
        s, e = int(row[3]), int(row[4])
        raw = new_text[s:e]
        if enclosing_tag(new_text, s) in CLEANABLE_TAGS and \
                "&amp;cr;" in raw and not CR_ONLY_RE.match(raw):
            mixed2 += 1
    assert mixed2 == 91


# ---------------------------------------------------------------------------
# 엑셀 테두리: 테이블 영역만 thin, 서술문/제목은 제외
# ---------------------------------------------------------------------------

def _find_cell(ws, value):
    for row in ws.iter_rows():
        for c in row:
            if c.value == value:
                return c
    raise AssertionError(f"{value!r} not found in {ws.title}")


def test_borders(tmp_path):
    dsd, info = _make(tmp_path)
    wb = load_workbook(info["out_path"])

    # FS 시트: 데이터 테이블 셀 테두리 있음 (라벨/숫자/빈 병합 영역 포함)
    bs = wb["BS"]
    assert _find_cell(bs, "자산총계").border.left.style == "thin"
    assert _find_cell(bs, 1234567).border.left.style == "thin"
    assert _find_cell(bs, "당기").border.left.style == "thin"   # 헤더 TH
    # 제목 행(1행)은 테두리 없음
    assert bs.cell(1, 1).border.left.style is None

    # 주석 시트: 테이블 셀은 테두리, 서술문(A열 문단)·헤더행은 없음
    ws2 = wb["2"]
    assert _find_cell(ws2, "금융자산").border.left.style == "thin"
    assert _find_cell(ws2, 12581632).border.left.style == "thin"
    assert ws2.cell(1, 1).border.left.style is None              # 헤더 행
    ws1 = wb["1"]
    para = _find_cell(ws1, "회사의 본사는 서울특별시에 위치하고 있습니다.")
    assert para.border.left.style is None                        # 서술문

    # 단위 행(TU, 테이블 밖)과 외부감사·사용안내는 테두리 없음
    unit = _find_cell(ws2, "(단위 : 백만원)")
    assert unit.border.left.style is None
    assert wb["사용안내"].cell(1, 1).border.left.style is None


def test_repack_refuses_overwrite(tmp_path):
    dsd, info = _make(tmp_path)
    from dsd_tool.repack import RepackError
    with pytest.raises(RepackError):
        repack(info["out_path"], dsd, dsd)


def test_repack_wrong_original(tmp_path):
    dsd, info = _make(tmp_path)
    other = build_dsd(str(tmp_path / "다른.dsd"),
                      CONTENTS_XML.replace("테스트주식회사", "다른주식회사"))
    from dsd_tool.repack import RepackError
    with pytest.raises(RepackError):
        repack(info["out_path"], other, str(tmp_path / "x.dsd"))

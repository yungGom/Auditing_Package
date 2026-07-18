"""Footing 검증 게이트 (패치 A-4): G-F1 ~ G-F4 + 레벨 오버라이드.

- G-F1: 한빛정밀(대사 완결 설계)로 전 항목 일치
- G-F2: 셀 1개 +1 변조 → 해당 트리만 단수차, 나머지 무영향
- G-F3: 삼성전자 실파일 BS/PL 통과율 (불일치 0)
- G-F4: BS 주석5 현금 ↔ 주석 5 시트 값 매칭
"""
import os

import pytest
from openpyxl import load_workbook

from dsd_tool.excel_out import extract
from dsd_tool.foot import (FOOT_SHEET, infer_relations, foot,
                           relations_from_levels)

_SYN = os.path.join(os.path.dirname(__file__), "..", "fixtures", "synthetic",
                    "한빛정밀_클린.dsd")
_SAMSUNG = os.path.join(os.path.dirname(__file__), "..", "fixtures", "real",
                        "[삼성전자(주)]_2025_[감사보고서].dsd")


# ---------------------------------------------------------------------------
# 알고리즘 단위 테스트
# ---------------------------------------------------------------------------

def test_infer_bottom_sum():
    """하단 합계: 9,830 + 5,400 → 15,230."""
    rels = infer_relations([(1, 9830), (2, 5400), (3, 15230)])
    assert len(rels) == 1
    assert rels[0]["parent"]["key"] == 3
    assert [c["key"] for c in rels[0]["children"]] == [1, 2]
    assert rels[0]["diff"] == 0


def test_infer_top_subtotal_and_multilevel():
    """상단 소계 + 다단: 총계(하단) ← {소계1(상단), 소계2(상단)}."""
    seq = [(1, 100), (2, 60), (3, 40),      # 소계 100 = 60+40 (상단)
           (4, 50), (5, 30), (6, 20),       # 소계 50 = 30+20 (상단)
           (7, 150)]                        # 총계 150 = 100+50 (하단)
    rels = infer_relations(seq)
    parents = {r["parent"]["key"]: [c["key"] for c in r["children"]]
               for r in rels}
    assert parents[1] == [2, 3]
    assert parents[4] == [5, 6]
    assert parents[7] == [1, 4]


def test_infer_fuzzy_within_limit():
    rels = infer_relations([(1, 100), (2, 50), (3, 151)])   # 단수차 1
    assert len(rels) == 1 and rels[0]["diff"] == 1
    assert infer_relations([(1, 100), (2, 50), (3, 160)]) == []  # 한도 초과


def test_relations_from_levels_mismatch():
    """수동 레벨 트리는 한도 초과 불일치도 그대로 노출한다."""
    nodes = [{"key": 1, "val": 100, "level": 1},
             {"key": 2, "val": 50, "level": 1},
             {"key": 3, "val": 999, "level": 2}]
    rels = relations_from_levels(nodes)
    assert len(rels) == 1 and rels[0]["diff"] == 849


# ---------------------------------------------------------------------------
# G-F1 / G-F2 / G-F4 — 한빛정밀
# ---------------------------------------------------------------------------

pytestmark = pytest.mark.skipif(
    not os.path.exists(_SYN), reason="한빛정밀 합성 DSD 없음")


@pytest.fixture()
def hanbit_xlsx(tmp_path):
    xlsx = str(tmp_path / "한빛.xlsx")
    extract(_SYN, xlsx)
    return xlsx


def test_gf1_all_match(hanbit_xlsx):
    res = foot(hanbit_xlsx)
    assert res["mismatch"] == 0
    assert res["fuzzy"] == 0
    assert res["match"] >= 40           # BS 트리 양기간 + PL/CE/CF + 주석 표
    # BS 완전 트리 확인: 자산총계 = 유동 + 비유동 (당기)
    bs = [r for r in res["foot"] if r["sheet"] == "BS" and r["scope"] == "당기"]
    labels = {r["label"]: r for r in bs}
    assert "자산총계" in labels and labels["자산총계"]["n_children"] == 2
    assert labels["자산총계"]["verdict"] == "일치"
    assert "부채와자본총계" in labels
    # _FOOT 시트 생성 확인
    wb = load_workbook(hanbit_xlsx)
    assert FOOT_SHEET in wb.sheetnames


def test_gf2_single_corruption(hanbit_xlsx):
    base = foot(hanbit_xlsx)
    base_keys = {(r["sheet"], r["scope"], r["loc"]): r["verdict"]
                 for r in base["foot"]}

    # BS 현금및현금성자산 당기 15,230 → +1 변조
    wb = load_workbook(hanbit_xlsx)
    target_row = None
    for row in wb["BS"].iter_rows():
        for c in row:
            if c.value == 15230:
                c.value = 15231
                target_row = c.row
                break
        if target_row:
            break
    assert target_row
    wb.save(hanbit_xlsx)

    res = foot(hanbit_xlsx)
    assert res["mismatch"] == 0
    fuzzies = [r for r in res["foot"] if r["verdict"] == "단수차"]
    assert fuzzies, "단수차 미검출"
    # 변조 셀이 속한 트리(유동자산, 당기)만 단수차
    for r in fuzzies:
        assert r["sheet"] == "BS" and r["scope"] == "당기"
        assert target_row in r["child_keys"] or r["parent_key"] == target_row
    # 나머지 관계는 무영향 (동일 판정 유지)
    for r in res["foot"]:
        key = (r["sheet"], r["scope"], r["loc"])
        if r["verdict"] == "일치" and key in base_keys:
            assert base_keys[key] == "일치"
    assert res["match"] == base["match"] - len(fuzzies)


def test_gf4_note_matching(hanbit_xlsx):
    res = foot(hanbit_xlsx)
    cash = [r for r in res["notes"]
            if r["sheet"] == "BS" and "현금및현금성자산" in r["label"]
            and "5" in r["refs"]]
    assert cash, "BS 현금및현금성자산 주석참조 행 없음"
    assert all(r["found"] for r in cash)
    assert all(r["where"].startswith("주석5") for r in cash)


def test_level_override_detects_mismatch(hanbit_xlsx):
    """수동 레벨로 틀린 트리 지정 → 불일치(빨강) 검출."""
    foot(hanbit_xlsx)                      # _FOOT [레벨] 섹션 생성
    wb = load_workbook(hanbit_xlsx)
    ws = wb[FOOT_SHEET]
    # [레벨] 섹션에서 BS 현금/매출채권/재고 행을 찾아 잘못된 수동 레벨 부여
    in_lv = False
    hits = 0
    for row in ws.iter_rows():
        if row[0].value == "[레벨]":
            in_lv = True
            continue
        if in_lv and isinstance(row[0].value, str) and \
                row[0].value.startswith("["):
            break
        if in_lv and row[0].value == "BS" and hits < 3:
            label = str(row[2].value or "")
            if "현금및현금성자산" in label or "매출채권" in label:
                row[4].value = 1
                hits += 1
            elif "재고자산" in label:
                row[4].value = 2          # 재고를 가짜 부모로
                hits += 1
    assert hits == 3
    wb.save(hanbit_xlsx)

    res = foot(hanbit_xlsx)
    assert res["manual_overrides"] == 3
    manual = [r for r in res["foot"] if r["direction"] == "수동"]
    assert manual
    mism = [r for r in manual if r["verdict"] == "불일치"]
    assert mism, "수동 트리 불일치 미검출"
    assert any("재고자산" in r["label"] for r in mism)


def test_foot_cli(hanbit_xlsx, capsys):
    from dsd_tool.cli import main
    assert main(["foot", hanbit_xlsx]) == 0
    out = capsys.readouterr().out
    assert "합계검증: 일치" in out and "주석대사" in out


# ---------------------------------------------------------------------------
# G-F3 — 삼성전자 실파일 BS/PL 통과율
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not os.path.exists(_SAMSUNG), reason="실제 DSD 없음")
def test_gf3_samsung_real(tmp_path, capsys):
    xlsx = str(tmp_path / "삼성.xlsx")
    extract(_SAMSUNG, xlsx, keep_note_numbers=True)
    res = foot(xlsx)
    bs_pl = [r for r in res["foot"] if r["sheet"] in ("BS", "PL")]
    assert bs_pl, "BS/PL 합계관계 미검출"
    # 실파일은 백만원 반올림 단수차 허용, 불일치는 0이어야 함
    assert all(r["verdict"] in ("일치", "단수차") for r in bs_pl)
    assert sum(1 for r in bs_pl if r["sheet"] == "BS") >= 5

    # 통과율 리포트 (원인 분류)
    from collections import Counter
    for sheet in ("BS", "PL"):
        rels = [r for r in res["foot"] if r["sheet"] == sheet]
        c = Counter(r["verdict"] for r in rels)
        covered = {k for r in rels for k in r["child_keys"]} | \
                  {r["parent_key"] for r in rels}
        print(f"[G-F3] {sheet}: 관계 {len(rels)}개 "
              f"(일치 {c['일치']}, 단수차 {c['단수차']}, 불일치 {c['불일치']}) "
              f"/ 참여 행 {len(covered)}개")
    print("[G-F3] 미검출 원인 후보: 차감 표시 관행(양방향 스캔으로 대부분 해소), "
          "단일 항목 소계, 반올림 누적이 한도(±2) 초과인 경우")
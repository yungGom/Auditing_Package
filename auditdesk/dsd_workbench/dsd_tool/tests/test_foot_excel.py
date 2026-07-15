"""A-5a 게이트: Footing 결과 엑셀 (AI_Footing 형식 재현).

1. 삼성전자 FY2025로 생성 → 예시 AI_Footing과 시트 구성·총괄표 열 구조 동일
   (예시 AI_Footing_삼성전자는 미확보 — 확보된 예시 페어[두원냉기]의 형식을
   기준으로 열 구조를 대조. 예시 없으면 실측 상수로 검증)
2. 총괄표 오류 카운트가 A-4 _FOOT 결과와 정합
3. 하이퍼링크 왕복(총괄표↔시트) 동작
"""
import os
import re

import pytest
from openpyxl import load_workbook

from dsd_tool.excel_out import extract
from dsd_tool.foot import FUZZY, MISMATCH, foot
from dsd_tool.foot_excel import (_SUMMARY_LEFT_HDR, _SUMMARY_RIGHT_HDR,
                                 write_ai_footing)

_SAMSUNG = os.path.join(
    os.path.dirname(__file__), "..", "fixtures", "real",
    "[삼성전자(주)]_2025_[감사보고서].dsd")
_SAMPLE = os.path.join(
    os.path.dirname(__file__), "..", "fixtures", "ai_footing",
    "AI Footing_[두원냉기(연결)-감사보고서(2025.12.31)].xlsx")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("a5a")
    xlsx = str(tmp / "삼성.xlsx")
    extract(_SAMSUNG, xlsx)
    res = foot(xlsx)
    out = write_ai_footing(xlsx, res, out_path=str(tmp / "AI_Footing.xlsx"))
    return res, out


pytestmark = pytest.mark.skipif(not os.path.exists(_SAMSUNG),
                                reason="삼성 DSD 없음")


# ---------------------------------------------------------------------------
# 게이트 1: 시트 구성·총괄표 열 구조가 예시와 동일
# ---------------------------------------------------------------------------

def test_gate1_sheet_and_summary_structure(built):
    res, out = built
    wb = load_workbook(out["out_path"])
    names = wb.sheetnames
    assert names[0] == "총괄표" and names[1] == "원문"
    assert {"BS", "PL", "CE", "CF"} <= set(names)
    assert any(n.isdigit() for n in names)          # 주석 1..N

    ws = wb["총괄표"]
    left = [ws.cell(4, 2 + j).value for j in range(len(_SUMMARY_LEFT_HDR))]
    right = [ws.cell(4, 7 + j).value for j in range(len(_SUMMARY_RIGHT_HDR))]
    assert left == _SUMMARY_LEFT_HDR
    assert right == _SUMMARY_RIGHT_HDR
    # 주석 행에 제목 병기 ("22. 판매비와관리비" → 제목부)
    titles = [ws.cell(r, 3).value for r in range(5, ws.max_row + 1)]
    assert any(t and "판매비와관리비" in str(t) for t in titles)

    if os.path.exists(_SAMPLE):                     # 예시 실물과 직접 대조
        swb = load_workbook(_SAMPLE)
        assert swb.sheetnames[0] == "총괄표" and swb.sheetnames[1] == "원문"
        sws = swb["총괄표"]
        s_left = [sws.cell(4, 2 + j).value
                  for j in range(len(_SUMMARY_LEFT_HDR))]
        s_right = [sws.cell(4, 7 + j).value
                   for j in range(len(_SUMMARY_RIGHT_HDR))]
        assert left == s_left and right == s_right   # 열 구조 동일
        assert sws.cell(2, 2).value == ws.cell(2, 2).value == "시트정보"


# ---------------------------------------------------------------------------
# 게이트 2: 총괄표 카운트 = A-4 _FOOT 결과 정합
# ---------------------------------------------------------------------------

def test_gate2_counts_match_foot(built):
    res, out = built
    exp_foot = sum(1 for r in res["foot"]
                   if r["verdict"] in (FUZZY, MISMATCH))
    exp_cross = sum(1 for r in res["notes"] if not r["found"])
    assert out["foot_errors"] == exp_foot
    assert out["cross_errors"] == exp_cross

    wb = load_workbook(out["out_path"])
    ws = wb["총괄표"]
    got_foot = got_cross = 0
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 7).value:
            got_foot += ws.cell(r, 9).value or 0
            got_cross += ws.cell(r, 10).value or 0
        v = ws.cell(r, 11).value
        if ws.cell(r, 7).value:
            assert v == f"=총괄표!I{r}+총괄표!J{r}"   # 총합 수식 (예시 그대로)
    assert int(got_foot) == exp_foot
    assert int(got_cross) == exp_cross


# ---------------------------------------------------------------------------
# 게이트 3: 하이퍼링크 왕복 + 오류 셀 표시
# ---------------------------------------------------------------------------

def test_gate3_hyperlink_roundtrip_and_marks(built):
    res, out = built
    wb = load_workbook(out["out_path"])
    ws = wb["총괄표"]
    for r in range(5, ws.max_row + 1):
        for col in (2, 8):
            v = ws.cell(r, col).value
            if not v:
                continue
            m = re.match(r'=HYPERLINK\("#\'?([^\'!]+)\'?!', str(v))
            assert m, v
            assert m.group(1) in wb.sheetnames       # 링크 대상 시트 실존
    # 각 데이터 시트 M2 복귀 링크
    for name in ("BS", "PL", "CE", "CF"):
        v = wb[name].cell(2, 13).value
        assert v and "총괄표" in str(v) and v.startswith("=HYPERLINK")

    # 오류 셀 노랑 + 메모 (A-4 좌표)
    err = next(r for r in res["foot"] if r["verdict"] in (FUZZY, MISMATCH))
    row = int(re.match(r"R(\d+)", err["loc"]).group(1))
    cell = wb[err["sheet"]].cell(row=row, column=1)
    assert cell.fill.start_color.rgb == "FFFFFF00"
    assert cell.comment and "푸팅" in cell.comment.text

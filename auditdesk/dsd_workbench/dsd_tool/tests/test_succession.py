"""F-3 게이트: 승계 모드 — 신원종합개발 기말↔반기 왕복 대조.

스펙 게이트: 코퍼스에서 기말+반기 모두 공시한 회사(신원종합개발) 선정 →
"기말 인스턴스로 반기 생성 vs 실제 공시 반기 인스턴스" 대조.
리포트 지표 (수치 자체가 벤치마크 — 고정 임계는 보고 후 확정):
1. 승계율: 반기 실제 사용 element 중 기말 자산으로 커버되는 비율
2. 승계 정확도: 라벨→element 승계 예측이 반기 실제와 일치(Top-1)
3. 신규 계정 검출: 반기 전용 element가 '신규(D-3b 경로)'로 라우팅되는 비율
4. 기간 변환 정확성: 반기 인스턴스의 실제 기간 문맥이 F-1 반기 기간
   체계(당기/전기 × 누적/3개월 + 시점)로 전수 분류되는지
"""
import glob
import os
import re
import sys
import xml.etree.ElementTree as ET

import pytest

_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__),
                                     "..", "..", ".."))
_ANNUAL = os.path.join(_ROOT, "dart_explorer", "cache", "xbrl", "00136925",
                       "20260323000596_11011")
_HALF = os.path.join(_ROOT, "dart_explorer", "cache", "xbrl", "00136925",
                     "20250814001572_11012")
_ASSETS_JSON = os.path.join(_ANNUAL, "succession_assets.json")
_SAMSUNG_DSD = os.path.join(
    os.path.dirname(__file__), "..", "fixtures", "real",
    "[삼성전자(주)]_2025_[감사보고서].dsd")
_SAMSUNG_ANNUAL = os.path.join(_ROOT, "dart_explorer", "cache", "xbrl",
                               "00126380", "20260310002820_11011")

_ready = (os.path.isdir(_ANNUAL) and os.path.isdir(_HALF) and
          os.path.exists(_ASSETS_JSON))


def _half_pairs():
    """반기 인스턴스 정답지: [(라벨, element_id, is_ext)] — dart_explorer
    파서 재사용 (테스트 한정 크로스 임포트, 정답지 구축 목적)."""
    if _ROOT not in sys.path:
        sys.path.insert(0, _ROOT)
    from dart_explorer.xbrl.corpus import extract_package
    rows = extract_package(_HALF)
    return [(r["label_ko"] or "", r["element_id"], r["is_ext"])
            for r in rows if not r["element_id"].startswith("dart-gcd")]


# ---------------------------------------------------------------------------
# 게이트 1: 승계율·정확도·신규 검출 (기말→반기 왕복)
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not _ready, reason="신원종합개발 기말/반기 캐시 없음")
def test_gate1_succession_roundtrip():
    from dsd_tool.mapping import normalize
    from dsd_tool.succession import load_assets

    assets = load_assets(_ASSETS_JSON)
    pairs = _half_pairs()
    assert len(pairs) >= 80          # 신원 반기 실측 101 (본문 위주 공시)

    half_eids = {e for _l, e, _x in pairs}
    covered = half_eids & assets.element_ids
    coverage = len(covered) / len(half_eids)

    # 승계 정확도: 라벨이 기말에 존재하는 행 → 승계 예측 == 반기 실제
    hit = n_pred = 0
    misses = []
    # 신규 검출: 반기 전용 element의 행 → 승계 미적용(신규 라우팅)이 정답
    new_total = new_detected = 0
    for label, eid, _is_ext in pairs:
        if len(normalize(label)) < 2:
            continue
        inh = assets.inherit(label)
        if eid in assets.element_ids:
            if inh is not None:
                n_pred += 1
                if inh["element_id"] == eid:
                    hit += 1
                elif len(misses) < 10:
                    misses.append((label, eid, inh["element_id"]))
        else:
            new_total += 1
            if inh is None:
                new_detected += 1
            elif len(misses) < 10:
                misses.append((label, eid, f"오승계:{inh['element_id']}"))

    acc = hit / n_pred if n_pred else None
    det = new_detected / new_total if new_total else None
    print(f"\n[F-3 게이트] 반기 element {len(half_eids)} | "
          f"승계율(구조 커버) {coverage:.1%} | "
          f"승계 정확도 {acc:.1%} (n={n_pred}) | "
          f"신규 검출 {new_detected}/{new_total}"
          + (f" ({det:.1%})" if det is not None else ""))
    for m in misses:
        print("  miss:", m)

    # 리포트 게이트 (임계는 보고 후 확정 — 붕괴 방지용 하한만)
    # 주: '오승계'로 표시되는 반기 전용 element 다수는 회사가 연도 간
    # element 선택을 바꾼 동개념 이형(단기차입금의 상환 ↔ 유동차입금의
    # 상환 등) — 기말 element 승계 제안이 오히려 일관성 있는 방향이며
    # 확정 ☐에서 회계사가 판단한다. 진짜 신규(기말에 유사 계정 없음)만
    # 신규 검출로 잡히는 것이 의도된 동작.
    assert coverage >= 0.5
    assert acc is not None and acc >= 0.8
    assert n_pred >= 60              # 신원 반기 실측 81


# ---------------------------------------------------------------------------
# 게이트 2: 기간 변환 — 반기 문맥이 F-1 반기 기간 체계로 전수 분류
# ---------------------------------------------------------------------------

def _classify_period(start, end, instant, fy_year):
    """반기 문맥 분류 (F-1 PERIOD_LABELS['half'] 어휘)."""
    if instant:
        y, md = int(instant[:4]), instant[5:]
        if y == fy_year:
            return "당기말" if md != "01-01" else "당기초(전기말)"
        return "전기말" if md in ("12-31",) else "전기 동기말(비교 시점)"
    sy, ey = int(start[:4]), int(end[:4])
    sm, em = start[5:7], end[5:7]
    if sy != ey:
        return "연간(경과 비교)" if em == "12" else None
    tag = "당기" if sy == fy_year else "전기"
    if sm == "01":
        return f"{tag} 누적" if em != "12" else f"{tag} 연간"
    if sm == "04":
        return f"{tag} 3개월"
    return None


@pytest.mark.skipif(not _ready, reason="신원종합개발 기말/반기 캐시 없음")
def test_gate2_period_conversion():
    xbrli = "http://www.xbrl.org/2003/instance"
    inst = sorted(glob.glob(os.path.join(_HALF, "*.xbrl")))[0]
    root = ET.parse(inst).getroot()
    periods = set()
    for ctx in root.iter(f"{{{xbrli}}}context"):
        p = ctx.find(f"{{{xbrli}}}period")
        i = p.find(f"{{{xbrli}}}instant")
        if i is not None:
            periods.add((None, None, i.text.strip()))
        else:
            periods.add((p.find(f"{{{xbrli}}}startDate").text.strip(),
                         p.find(f"{{{xbrli}}}endDate").text.strip(), None))
    assert periods
    fy = max(int((s or i)[:4]) for s, e, i in periods
             for i in [i] if (s or i))

    unclassified = []
    dist = {}
    for s, e, i in sorted(periods, key=str):
        cls = _classify_period(s, e, i, fy)
        if cls is None:
            unclassified.append((s, e, i))
        else:
            dist[cls] = dist.get(cls, 0) + 1
    print(f"\n[F-3 기간 변환] 문맥 기간 유형 {len(periods)}종 분류: {dist}")
    assert not unclassified, f"미분류 기간: {unclassified}"
    # F-1 반기 기간 블록 어휘가 실제 공시 문맥을 포괄
    assert "당기 누적" in dist and "전기 누적" in dist
    assert "당기말" in dist and "전기말" in dist


# ---------------------------------------------------------------------------
# 게이트 3: 승계 워크시트 생성 (본문) — 삼성 자기 기말 자산으로 half 생성
# ---------------------------------------------------------------------------

@pytest.mark.skipif(
    not (os.path.exists(_SAMSUNG_DSD) and os.path.isdir(_SAMSUNG_ANNUAL)),
    reason="삼성 DSD/기말 캐시 없음")
def test_gate3_succession_worksheet(tmp_path):
    from openpyxl import load_workbook

    from dsd_tool.succession import load_assets
    from dsd_tool.worksheet import build_worksheet

    if _ROOT not in sys.path:
        sys.path.insert(0, _ROOT)
    sj = os.path.join(_SAMSUNG_ANNUAL, "succession_assets.json")
    if not os.path.exists(sj):
        from dart_explorer.xbrl.corpus import export_succession_assets
        export_succession_assets(_SAMSUNG_ANNUAL)
    assets = load_assets(sj)

    out = str(tmp_path / "승계.xlsx")
    res = build_worksheet(_SAMSUNG_DSD, out_path=out, report_type="half",
                          succession=assets, include_notes=False,
                          holdout_corp="00126380")
    st = res["stats"]
    print(f"\n[F-3 워크시트] 본문 {st['rows']}행 — 승계 {st['inherited']} / "
          f"신규 {st['new']} / D-4c 폐지 경고 {st['deprecated']}")
    assert st["inherited"] >= st["rows"] * 0.5      # 자기 기말 → 승계가 다수
    assert st["inherited"] + st["new"] + st["skipped"] == st["rows"]

    wb = load_workbook(out)
    text = "\n".join(str(c.value) for ws in wb for row in ws.iter_rows()
                     for c in row if c.value)
    assert "승계 모드(기말 태깅 이어받기)" in text
    assert "승계 — 기말 사용" in text
    # 기간 체계가 분반기로 변환 (F-1 규칙 재사용)
    assert re.search(r"값: (당기 누적|당기 3개월)", text)

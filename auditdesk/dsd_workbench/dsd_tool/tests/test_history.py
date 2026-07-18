"""수정이력 관리 게이트 (패치 A-2).

게이트: repack 2회 실행 후 이력 2건 정확 조회, --dry-run은 기록하지 않음.
"""
import pytest
from openpyxl import load_workbook

from dsd_tool.excel_out import extract
from dsd_tool.history import query_changes, query_runs
from dsd_tool.repack import diff, repack

from .fixture import build_dsd


@pytest.fixture()
def hist_env(tmp_path, monkeypatch):
    db = str(tmp_path / "history.sqlite")
    monkeypatch.setenv("DSD_TOOL_HISTORY", db)
    dsd = build_dsd(str(tmp_path / "이력테스트.dsd"))
    info = extract(dsd, str(tmp_path / "이력테스트.xlsx"),
                   keep_note_numbers=True)
    return dsd, info["out_path"], tmp_path


def test_history_gate(hist_env):
    dsd, xlsx, tmp = hist_env

    # 실행 1: 무변경 (keep-cr) — 실행 자체가 기록됨
    repack(xlsx, dsd, str(tmp / "out1.dsd"), clean_cr=False)
    # 실행 2: 셀 1개 수정
    wb = load_workbook(xlsx)
    for row in wb["BS"].iter_rows():
        for c in row:
            if c.value == 1234567:
                c.value = 7654321
    wb.save(xlsx)
    repack(xlsx, dsd, str(tmp / "out2.dsd"), clean_cr=False)

    # dry-run(diff)은 기록되지 않음
    diff(xlsx, dsd, clean_cr=False)

    runs = query_runs()
    assert len(runs) == 2
    runs_for_file = query_runs(dsd)
    assert len(runs_for_file) == 2

    latest = runs[0]           # 최신순
    assert latest["edits"] == 1
    assert latest["cleans"] == 0 and latest["dedups"] == 0
    assert latest["clean_cr"] == 0
    assert latest["out_path"].endswith("out2.dsd")
    assert len(latest["dsd_sha1"]) == 40

    changes = query_changes(latest["id"])
    assert len(changes) == 1
    ch = changes[0]
    assert ch["sheet"] == "BS" and ch["reason"] == "edit"
    assert ch["old"] == "1,234,567" and ch["new"] == "7,654,321"

    first = runs[1]
    assert first["edits"] == 0
    assert query_changes(first["id"]) == []


def test_history_reason_counts(hist_env):
    """기본 모드 repack: clean-cr/note-dedup 건수가 이력에 분리 집계."""
    dsd, xlsx, tmp = hist_env
    # 기본 모드용으로 기본 추출본 사용 (dedup 표시 포함)
    xlsx2 = str(tmp / "기본추출.xlsx")
    extract(dsd, xlsx2)
    repack(xlsx2, dsd, str(tmp / "out3.dsd"))
    latest = query_runs()[0]
    assert latest["clean_cr"] == 1
    assert latest["cleans"] == 2       # 픽스처의 &cr;-only 셀 2개
    assert latest["edits"] == 0


def test_history_cli(hist_env, capsys):
    dsd, xlsx, tmp = hist_env
    repack(xlsx, dsd, str(tmp / "out.dsd"), clean_cr=False)
    from dsd_tool.cli import main
    assert main(["history", dsd, "--changes"]) == 0
    out = capsys.readouterr().out
    assert "수정이력 1건" in out
    assert "SHA1=" in out


def test_history_opt_out(hist_env):
    dsd, xlsx, tmp = hist_env
    repack(xlsx, dsd, str(tmp / "out.dsd"), clean_cr=False,
           record_history=False)
    assert query_runs() == []

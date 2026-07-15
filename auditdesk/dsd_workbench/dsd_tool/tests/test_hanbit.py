"""한빛정밀 합성 DSD (실제 뼈대 + 값 교체) — §9 시나리오 자동 테스트.

hanbit_skeleton.build()가 대량 편집(5,000+셀) 통합 테스트를 겸한다.
시나리오 5(DART 편집기 열림)는 수동 검증.
"""
import os

import pytest
from openpyxl import load_workbook

from dsd_tool.excel_out import extract
from dsd_tool.repack import repack
from dsd_tool.zipsplice import read_contents

from .hanbit_skeleton import SKELETON, build

pytestmark = pytest.mark.skipif(
    not os.path.exists(SKELETON), reason="뼈대 DSD(삼성전자 별도) 없음")

V6_MIXED = [
    "주식발행&amp;cr;초과금",
    "연이자율&amp;cr;(%)",
    "상각후원가&amp;cr;측정 금융자산",
]


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    work = tmp_path_factory.mktemp("hanbit_build")
    dirty, clean, report = build(str(work))
    return {"dirty": dirty, "clean": clean, "report": report, "work": work}


def _xml(path):
    return read_contents(open(path, "rb").read()).decode("utf-8")


# ---------------------------------------------------------------------------
# 빌드 자체 검증 (대량 편집 + V 포인트)
# ---------------------------------------------------------------------------

def test_build_report(built):
    rep = built["report"]
    assert rep["edited_cells"] > 3000          # 대량 편집 통합 검증
    assert rep["residual_samsung"] == 0        # 뼈대 텍스트 완전 교체
    assert rep.get("paras_dropped", 0) == 0


def test_v_points(built):
    xml = _xml(built["dirty"])
    # V1 같은 숫자 다중 출현 (BS·CF기말·주석5합계·주석23)
    assert xml.count(">15,230<") >= 3
    assert xml.count(">12,480<") >= 3
    # V2 괄호 음수 / V3 "-" 셀 / V7 복수 주석참조 / V9 단위 혼재
    assert ">(250)<" in xml and ">(2,500)<" in xml and ">(3,640)<" in xml
    assert ">-<" in xml
    for ref in (">6, 23<", ">8, 10<", ">21, 23<"):
        assert ref in xml, ref
    assert "(단위 : 원)" in xml
    # V6 혼합 셀 (뼈대 CE 헤더 보존 + 주석 11/23에 심음)
    for token in V6_MIXED:
        assert token in xml, token
    # 회사 정보 교체
    assert "주식회사 한빛정밀" in xml


# ---------------------------------------------------------------------------
# 시나리오 1: 오염본 extract → 헤더 정리 표시 + &cr; 리포트
# ---------------------------------------------------------------------------

def test_s1_extract_dirty(built):
    xlsx = str(built["work"] / "s1.xlsx")
    info = extract(built["dirty"], xlsx)
    assert info["fs_sheets"] == ["BS", "PL", "PL1", "CE", "CF"]
    assert info["note_count"] == 32            # 뼈대 구조 유지 (25 + 여백 7)
    assert info["deduped_notes"] == 32         # V8: "N. N." 오염 전체
    assert info["cr_only_cells"] > 200         # V5: 뼈대 &cr; 셀 활용
    wb = load_workbook(xlsx)
    assert wb["1"].cell(1, 1).value == "1. 일반사항"
    assert wb["26"].cell(1, 1).value == "26. (여백)"
    ce_vals = [c.value for row in wb["CE"].iter_rows() for c in row]
    assert "-" in ce_vals                      # V3


# ---------------------------------------------------------------------------
# 시나리오 2: V1 다중 출현 셀 중 1곳만 수정 → 다른 곳 불변
# ---------------------------------------------------------------------------

def test_s2_v1_positional_replace(built):
    xlsx = str(built["work"] / "s2.xlsx")
    extract(built["dirty"], xlsx, keep_note_numbers=True)
    before = _xml(built["dirty"]).count(">15,230<")
    wb = load_workbook(xlsx)
    hit = None
    for row in wb["BS"].iter_rows():
        for c in row:
            if c.value == 15230:
                c.value = 99999
                hit = True
                break
        if hit:
            break
    assert hit
    wb.save(xlsx)
    out = str(built["work"] / "s2.dsd")
    result = repack(xlsx, built["dirty"], out, clean_cr=False)
    assert len(result["changes"]) == 1
    new_xml = _xml(out)
    assert new_xml.count(">99,999<") == 1
    assert new_xml.count(">15,230<") == before - 1   # 나머지 위치 불변


# ---------------------------------------------------------------------------
# 시나리오 3: 클린본 = 기본 파이프라인 통과분 (정리 완료 상태)
# ---------------------------------------------------------------------------

def test_s3_clean_output(built):
    xlsx = str(built["work"] / "s3.xlsx")
    info = extract(built["clean"], xlsx)
    assert info["cr_only_cells"] == 0          # &cr;-only 전부 정리
    assert info["deduped_notes"] == 0          # 주석 번호 정리 완료
    assert info["note_count"] == 32
    xml = _xml(built["clean"])
    assert '>1. 일반사항</SPAN>' in xml
    assert "1. 1." not in xml
    for token in V6_MIXED:                     # 혼합 셀(의도적 줄바꿈)은 보존
        assert token in xml, token
    # 클린본 무변경 재통과 → 멱등 (기본 모드 그대로)
    out = str(built["work"] / "s3_re.dsd")
    result = repack(xlsx, built["clean"], out)
    assert result["changes"] == []
    assert open(out, "rb").read() == open(built["clean"], "rb").read()


# ---------------------------------------------------------------------------
# 시나리오 4: --keep-note-numbers --keep-cr → 무변경 시 바이트 동일 (G2)
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("key", ["dirty", "clean"])
def test_s4_keep_modes_byte_identity(built, key):
    xlsx = str(built["work"] / f"s4_{key}.xlsx")
    extract(built[key], xlsx, keep_note_numbers=True)
    out = str(built["work"] / f"s4_{key}.dsd")
    result = repack(xlsx, built[key], out, clean_cr=False)
    assert result["changes"] == []
    assert open(out, "rb").read() == open(built[key], "rb").read()

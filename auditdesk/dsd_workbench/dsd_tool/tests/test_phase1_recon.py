"""M-02/M-03: core verdicts must survive traceable Excel rendering."""
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from dsd_tool.recon import recon_ce, recon_statement, _write_recon_excel, count_rendered_verdicts


def context(kind, prior=False, missing=False, delta=0):
    wb = Workbook()
    ws = wb.active
    ws.title = kind
    if kind == "CE":
        ws.append(["title"])
        ws.append(["항목", "주석", "Capital", "Other" if missing else "Retained"])
        block = "당기" if prior else "전기"
        ws.append([f"2025.1.1({block}초)", "", 100 + delta, 200])
        ws.append([f"2025.12.31({block}말)", "", 100 + delta, 200])
    else:
        ws.append(["title"])
        ws.append(["항목", "주석", "당기", "전기"])
        ws.append(["현금", "", 100 + delta, 100])
        ws.append(["부채", "", -100 - delta, -100])
    ctx = SimpleNamespace(wb=wb, rowmaps={kind: {r: [1, 2, 3, 4] for r in (2, 3, 4)}},
                           fs_sheets=[kind], note_sheets=[],
                           fs_sequences=lambda s: ([("당기", [(3, 100 + delta), (4, -100 - delta)]), ("전기", [(3, 100), (4, -100)])], [3, 4]))
    from dsd_tool.foot import FootingContext
    ctx.fs_value_col = lambda s, r, p: FootingContext.fs_value_col(ctx, s, r, p)
    return ctx


@pytest.mark.parametrize("tolerance", [0, 2])
def test_m02_missing_ce_column_never_passes(tmp_path, tolerance):
    cur, pri = context("CE"), context("CE", prior=True, missing=True)
    rows = recon_ce(cur, pri, "CE", tolerance)
    assert len(rows) == 2 and all(not r["true"] for r in rows)
    out = _write_recon_excel({"CE": rows}, [], {}, {}, str(tmp_path / "out.xlsx"), "c", "p", tolerance, cur_ctx=cur, pri_ctx=pri)
    assert out["stmt"] == {"n": 2, "true": 0, "false": 2}


@pytest.mark.parametrize("kind", ["BS", "CE"])
@pytest.mark.parametrize("tolerance,delta,expected", [(0, 0, True), (0, 1, False), (2, 1, True), (2, 2, True), (2, 3, False)])
def test_m03_core_excel_summary_agree(tmp_path, kind, tolerance, delta, expected):
    cur, pri = context(kind), context(kind, prior=True, delta=delta)
    fn = recon_ce if kind == "CE" else recon_statement
    rows = fn(cur, pri, kind, tolerance)
    assert len(rows) == 2 and all(r["true"] == expected for r in rows)
    path = str(tmp_path / "out.xlsx")
    out = _write_recon_excel({kind: rows}, [], {}, {}, path, "c", "p", tolerance, cur_ctx=cur, pri_ctx=pri)
    assert out["stmt"] == {"n": 2, "true": 2 if expected else 0, "false": 0 if expected else 2}
    wb = load_workbook(path)
    assert count_rendered_verdicts(wb[kind]) == (2, 2 if expected else 0, 0 if expected else 2)
    assert wb[kind]["E3"].value.startswith("=")  # cell references remain traceable
    if tolerance > 0:
        assert wb[kind]["E3"].value == (
            f"=ABS(D3-I3)<={tolerance}" if kind == "BS" else
            f"=AND(ABS(C3-I3)<={tolerance},ABS(D3-J3)<={tolerance})")
    wb.close()


@pytest.mark.parametrize("tolerance,delta", [(0, 0), (0, 1), (2, 1), (2, 2), (2, 3)])
def test_m03_note_formula_matches_core(tmp_path, tolerance, delta):
    from dsd_tool.recon import recon_note_tables
    cur, pri = context("1"), context("1", prior=True, delta=delta)
    for ctx in (cur, pri):
        ctx.fs_sheets = []
        ctx.note_sheets = ["1"]
        ctx.wb["1"].insert_rows(3)
        ctx.wb["1"].cell(3, 3, "금액")
        ctx.wb["1"].cell(3, 4, "금액")
        ctx.rowmaps["1"] = {r: [1, 2, 3, 4] for r in (2, 3, 4, 5)}
    rows, skipped, unpaired = recon_note_tables(cur, pri, "1", "1", tolerance)
    assert rows and all(r["true"] == (delta <= tolerance) for r in rows)
    path = str(tmp_path / "notes.xlsx")
    result = _write_recon_excel({}, [("1", "1", "note")], {"1": rows}, {"1": skipped}, path, "cur", "pri", tolerance, cur_ctx=cur, pri_ctx=pri)
    assert result["notes"]["n"] > 0
    assert result["notes"]["false"] == (0 if delta <= tolerance else result["notes"]["n"])


@pytest.mark.parametrize("tolerance,delta", [(0, 0), (2, 1)])
def test_m03_four_amount_columns_use_core_selected_cells(tmp_path, tolerance, delta):
    from dsd_tool.foot import FootingContext
    cur, pri = context("BS"), context("BS", prior=True)
    for ctx in (cur, pri):
        ctx.rowmaps["BS"] = {r: [1, 2, 3, 4, 5, 6] for r in (2, 3, 4)}
        ctx.fs_sequences = lambda s, ctx=ctx: FootingContext.fs_sequences(ctx, s)
        ctx.fs_value_col = lambda s, r, p, ctx=ctx: FootingContext.fs_value_col(ctx, s, r, p)
    for r in (3, 4):
        cur.wb["BS"].cell(r, 5, 100)
        cur.wb["BS"].cell(r, 6).value = None
        pri.wb["BS"].cell(r, 3).value = None
        pri.wb["BS"].cell(r, 4).value = 100 + delta
    rows = recon_statement(cur, pri, "BS", tolerance)
    assert len(rows) == 2 and all(r["true"] for r in rows)
    out = _write_recon_excel({"BS": rows}, [], {}, {}, str(tmp_path / "out.xlsx"), "c", "p", tolerance, cur_ctx=cur, pri_ctx=pri)
    assert out["stmt"] == {"n": 2, "true": 2, "false": 0}

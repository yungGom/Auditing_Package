"""A-6 게이트: 합계검증 [검증내역] 시트 — 전수 목록·수식 판정.

G1 검증내역 전수 목록 — 수행된 모든 검증이 행 단위, 판정·차이는 엑셀
   수식(셀 참조 — 근거 추적), 미발견 크로스는 사유 분류 필수
G2 총괄표 수행 건수 병기 — "0/0 검증불능"과 "0/N 전부 일치" 구분,
   검증 대상 모수 0 시트는 경고색
G3 변조 1건 → 해당 행 FALSE·총괄표 크로스 오류 +1
(법인 원형 형식 정합은 test_foot_excel 삼성 게이트가 회귀로 보증)
"""
import pytest
from openpyxl import load_workbook

from dsd_tool.excel_out import extract
from dsd_tool.foot import foot
from dsd_tool.foot_excel import DETAIL_SHEET, write_ai_footing
from dsd_tool.tests.fixture import CONTENTS_XML, build_dsd

# 현금및현금성자산의 주석 참조를 실존 주석(3)으로 — refs 기반 발견 1건
_BASE = CONTENTS_XML.replace("<TD>4, 28</TD>", "<TD>3</TD>")
# 변조: 주석 3의 값을 바꿔 발견→미발견 전환
_TAMPERED = _BASE.replace(
    '<TR><TD>보통예금</TD><TD ALIGN="RIGHT">500,000</TD></TR>',
    '<TR><TD>보통예금</TD><TD ALIGN="RIGHT">999</TD></TR>')


def _build(tmp_path, contents, name):
    dsd = build_dsd(str(tmp_path / f"{name}.dsd"), contents=contents)
    xlsx = str(tmp_path / f"{name}.xlsx")
    extract(dsd, xlsx)
    res = foot(xlsx, save=False)
    out = write_ai_footing(xlsx, res,
                           out_path=str(tmp_path / f"AI_{name}.xlsx"))
    return res, out


@pytest.fixture(scope="module")
def base(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("a6")
    return _build(tmp, _BASE, "base")


def test_g1_detail_sheet_formulas_and_reasons(base):
    res, out = base
    wb = load_workbook(out["out_path"])
    assert wb.sheetnames[:3] == ["총괄표", "원문", DETAIL_SHEET]
    ws = wb[DETAIL_SHEET]
    # 요약: 수행 모수 명시 (found/missing 사유별 집계)
    assert "수행" in str(ws.cell(2, 1).value)
    s3 = str(ws.cell(3, 1).value)
    assert "수행" in s3 and "사유별" in s3
    # 전수: 검증내역 행 수 = 푸팅 전체 + 크로스 전체
    n_rows = ws.max_row - 5                     # 헤더(5행) 이후
    assert n_rows == out["checks"]["foot_total"] + out["checks"]["cross_total"]
    assert out["checks"]["cross_total"] >= 4    # 자산총계×2·현금×2
    assert out["checks"]["cross_found"] >= 1    # refs 3 → 보통예금 발견
    found_row = missing_row = None
    for r in range(6, ws.max_row + 1):
        typ, verdict = ws.cell(r, 1).value, ws.cell(r, 9).value
        if typ and str(typ).startswith("크로스"):
            if verdict is False:
                missing_row = r
                # 사유 없는 미발견 금지
                assert str(ws.cell(r, 11).value or "").startswith(
                    ("①", "②", "③")), r
            elif str(verdict).startswith("="):
                found_row = r
    assert found_row and missing_row
    # 판정·차이·좌변은 수식 (근거 추적 가능)
    assert str(ws.cell(found_row, 5).value).startswith("=")
    assert "ABS" in str(ws.cell(found_row, 9).value)
    assert str(ws.cell(found_row, 12).value).startswith("=HYPERLINK")
    reasons = out["checks"]["missing_reasons"]
    assert sum(reasons.values()) == \
        out["checks"]["cross_total"] - out["checks"]["cross_found"]


def test_g2_summary_counts_and_warn(base):
    res, out = base
    wb = load_workbook(out["out_path"])
    ws0 = wb["총괄표"]
    assert ws0.cell(4, 12).value == "푸팅 수행"
    assert ws0.cell(4, 13).value == "크로스 수행"
    rows = {ws0.cell(r, 7).value: r for r in range(5, ws0.max_row + 1)
            if ws0.cell(r, 7).value}
    # BS: 크로스 수행 > 0 (0/N 구분 가능)
    assert ws0.cell(rows["BS"], 13).value >= 4
    # 검증 대상인데 푸팅 모수 0인 시트 존재 시 경고색 (합성 CE 등)
    warn = [n for n, r in rows.items()
            if ws0.cell(r, 12).fill.start_color.rgb == "FFFFC000"
            or ws0.cell(r, 13).fill.start_color.rgb == "FFFFC000"]
    assert warn                                  # 침묵 무결 금지 — 표시됨
    # 비검증 시트(표지·사용안내)는 경고 없음
    assert "표지" not in warn and "사용안내" not in warn


def test_g3_tamper_flips_row_and_summary(base, tmp_path):
    res0, out0 = base
    res, out = _build(tmp_path, _TAMPERED, "tampered")
    assert out["checks"]["cross_found"] == out0["checks"]["cross_found"] - 1
    assert out["cross_errors"] == out0["cross_errors"] + 1
    wb = load_workbook(out["out_path"])
    ws = wb[DETAIL_SHEET]
    flipped = [r for r in range(6, ws.max_row + 1)
               if str(ws.cell(r, 1).value or "").startswith("크로스")
               and ws.cell(r, 9).value is False
               and "현금및현금성자산" in str(ws.cell(r, 4).value)]
    assert flipped                               # 해당 행 FALSE + 사유
    assert str(ws.cell(flipped[0], 11).value or "").startswith(
        ("①", "②", "③"))

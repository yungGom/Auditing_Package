"""H-2 게이트: 병합 셀 안전 기입 — MergedCell 오류 회귀 고정.

재현: 주석 표 첫 열 ROWSPAN 병합 + 하단합계 단수차 → 오류 메모가
병합 내부 좌표(R8C1)에 기입되며 AttributeError로 전체 생성 중단
(B-6c 분할로 뭉텅이 문서 주석 표가 처음 오류 표시 대상에 진입).
G1 병합 픽스처 — 생성 성공·좌상단 앵커에 메모·채우기
G2 메모 이어붙임 — 같은 앵커에 2건 기입 시 덮어쓰기 금지
G3 한 셀 실패 ≠ 전체 중단 — failures 축적·산출물 노출
"""
import pytest
from openpyxl import load_workbook

from dsd_tool.cellsafe import anchor, put
from dsd_tool.excel_out import extract
from dsd_tool.foot import foot
from dsd_tool.foot_excel import DETAIL_SHEET, write_ai_footing
from dsd_tool.tests.fixture import CONTENTS_XML, build_dsd

# 주석 3 표 교체: 첫 열 ROWSPAN 병합 + 100+200 vs 301(단수차 유발)
_lo = CONTENTS_XML.find('<P><SPAN USERMARK=" B">3. 현금및현금성자산')
_hi = CONTENTS_XML.find("</SECTION-2>")
_MERGED = CONTENTS_XML[:_lo] + (
    '<P><SPAN USERMARK=" B">3. 현금및현금성자산</SPAN>&amp;cr;&amp;cr;'
    "내역은 다음과 같습니다.</P>\r\n"
    '<TABLE WIDTH="600">\r\n'
    "<TR><TH>구분</TH><TH>금액</TH></TR>\r\n"
    '<TR><TD ROWSPAN="3">현금성</TD><TD ALIGN="RIGHT">100</TD></TR>\r\n'
    '<TR><TD ALIGN="RIGHT">200</TD></TR>\r\n'
    '<TR><TD ALIGN="RIGHT">301</TD></TR>\r\n'
    "</TABLE>\r\n") + CONTENTS_XML[_hi:]


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("h2")
    dsd = build_dsd(str(tmp / "m.dsd"), contents=_MERGED)
    xlsx = str(tmp / "m.xlsx")
    extract(dsd, xlsx)
    res = foot(xlsx, save=False)
    out = write_ai_footing(xlsx, res, out_path=str(tmp / "AI.xlsx"))
    return res, out


def test_g1_merged_interior_comment_to_anchor(built):
    """병합 내부 좌표 오류 표시 — 생성 성공 + 좌상단 앵커 기입."""
    res, out = built
    err = next(r for r in res["foot"] if r["verdict"] != "일치")
    assert err["sheet"] == "3" and err["loc"] == "R8"   # 병합 내부
    wb = load_workbook(out["out_path"])
    ws = wb["3"]
    a = anchor(ws, 8, 1)
    assert a.coordinate == "A6"                 # 좌상단 앵커 리다이렉트
    assert a.comment is not None and "푸팅" in a.comment.text
    assert out["write_failures"] == []          # 실패 0 — 전부 기입 성공


def test_g2_comment_append_no_overwrite(built):
    """같은 앵커에 재기입 — 기존 메모에 이어붙임 (덮어쓰기 금지)."""
    _res, out = built
    wb = load_workbook(out["out_path"])
    ws = wb["3"]
    put(ws, 8, 1, comment=("두 번째 메모", "t"))
    text = anchor(ws, 8, 1).comment.text
    assert "푸팅" in text and "두 번째 메모" in text


def test_g3_failure_accumulates_not_abort():
    """기입 자체가 불가한 경우 — 실패 축적, 예외 전파 없음."""
    class _Boom:
        title = "X"

        class merged_cells:
            ranges = []

        def cell(self, row=None, column=None):
            raise RuntimeError("기입 불가 모의")

    fails = []
    r = put(_Boom(), 1, 1, value=1, failures=fails, what="모의")
    assert r is None and len(fails) == 1
    assert fails[0]["what"] == "모의" and "RuntimeError" in fails[0]["reason"]


def test_g3_failures_exposed_in_output(built, tmp_path):
    """(형식 확인) 검증내역에 '기입 불가' 노출 경로 — 정상 파일은 0건."""
    _res, out = built
    wb = load_workbook(out["out_path"])
    text = "\n".join(str(c.value) for row in wb[DETAIL_SHEET].iter_rows()
                     for c in row if c.value)
    assert "기입 불가" not in text              # 실패 0이면 침묵도 없음

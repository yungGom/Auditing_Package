"""B-4 게이트: 산출물 정직성 3종 — 열람용 안내 분기·뭉침 배너·탈락 노출.

합성 픽스처로 회귀 고정: 수신 래핑본(meta 없음)에서만 배너·안내가
붙고, 편집기 계열(meta 있음)은 불변이어야 한다.
"""
import os

import pytest

from dsd_tool.excel_out import extract
from dsd_tool.tests.fixture import CONTENTS_XML, build_dsd

# 뭉침 잔여: 2차 분할(B-6-mini)이 가를 수 없는 형태 — 표 셀 안에
# 주석 헤더 모양 텍스트가 남은 수신물 모의 (분할 실패 잔여는 뭉침
# 배너로 노출 — B-6-mini 이후에도 유지되는 정직성 경로 검증).
# 종전 평문 P 뭉침 케이스는 B-6-mini가 실제로 분할해 배너가
# 불필요해졌다 (test_b6_mini에서 분할 자체를 검증).
_MERGED = CONTENTS_XML.replace(
    '<P><SPAN USERMARK=" B">2. 재무제표 작성기준</SPAN>'
    "&amp;cr;&amp;cr;회사는 K-IFRS를 적용하고 있습니다.</P>",
    '<TABLE WIDTH="600"><TR><TD><P>2. 재무제표 작성기준 내용'
    "</P></TD><TD>값</TD></TR></TABLE>")
# 탈락: 주석 시작 이후에 유효 FS 제목(별도 섹션 모의) 삽입
_DROPPED = CONTENTS_XML.replace(
    "</SECTION-2>\r",
    '<TABLE WIDTH="600"><TR><TD ALIGN="CENTER">연결 재무상태표'
    "</TD></TR></TABLE>\r\n"
    '<TABLE WIDTH="600"><TR><TD>자산총계</TD>'
    '<TD ALIGN="RIGHT">9,999</TD></TR></TABLE>\r\n'
    "</SECTION-2>\r", 1)


def _run(tmp_path, contents, no_meta):
    dsd = build_dsd(str(tmp_path / f"b4_{no_meta}.dsd"),
                    contents=contents, no_meta=no_meta)
    xlsx = str(tmp_path / f"b4_{no_meta}.xlsx")
    meta = extract(dsd, xlsx)
    from openpyxl import load_workbook
    return meta, load_workbook(xlsx)


def test_viewonly_guide_branch(tmp_path):
    """수신물 → 열람용 안내 / 편집기 계열 → 기존 안내 (분기)."""
    m1, wb1 = _run(tmp_path, CONTENTS_XML, no_meta=True)
    assert m1["viewonly"] is True
    assert "열람용" in str(wb1["사용안내"].cell(1, 1).value)
    text = "\n".join(str(c.value) for row in wb1["사용안내"].iter_rows()
                     for c in row if c.value)
    assert "회사 보유 DSD" in text
    assert "repack" not in text                 # 수정·역변환 안내 제거

    m2, wb2 = _run(tmp_path, CONTENTS_XML, no_meta=False)
    assert m2["viewonly"] is False
    assert "repack" in "\n".join(
        str(c.value) for row in wb2["사용안내"].iter_rows()
        for c in row if c.value)                # 기존 안내 불변


def test_merged_banner_wrapped_only(tmp_path):
    """뭉침 감지 — 수신물만 배너, 편집기 계열은 배너 없음."""
    m1, wb1 = _run(tmp_path, _MERGED, no_meta=True)
    assert "1" in m1["merged_notes"], m1["merged_notes"]
    banner = str(wb1["1"].cell(2, 1).value or "")
    assert "미분할 포함" in banner

    m2, wb2 = _run(tmp_path, _MERGED, no_meta=False)
    assert m2["merged_notes"] == {}             # 편집기 계열 meta 비노출
    assert wb2["1"].cell(2, 1).value in (None, "")


def test_dropped_fs_exposed(tmp_path):
    """탈락 감지 — 주석 이후 FS 제목은 '미분할 원문' 시트로 노출."""
    m, wb = _run(tmp_path, _DROPPED, no_meta=True)
    assert m["fs_dropped"] == ["연결 재무상태표"], m["fs_dropped"]
    assert "미분할 원문" in wb.sheetnames
    text = "\n".join(str(c.value) for row in wb["미분할 원문"].iter_rows()
                     for c in row if c.value)
    assert "연결 재무상태표" in text and "9,999" in text
    # 본문 FS 시트(BS 등)는 그대로 — 기존 인식 불변
    assert "BS" in wb.sheetnames

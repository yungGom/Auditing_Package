"""M-1 게이트: 대사 조서 자동 조립 — 합성 회귀 고정.

G1 병렬 배치 — 우 블록 시작열 = 좌 최대열 + 2 (승인 규칙)
G2 판정 — 동일 값 TRUE·단위환산 TRUE(비고)·미발견 FALSE
G3 기간 제외 — 최근 2개 연도만 판정, 이전 연도 열은 명시 집계
G4 총괄표 — COUNTIF(INDIRECT) 수식·대응 없음 자동 기재
G5 변조 — 우측 값 1건 변조 → FALSE +1
(실파일: 바이오 기말 공개 DSD+인스턴스 — 대조 91행·FALSE 4(CE
기초잔액류 표면화)·전전기 73값 명시 제외·변조 +1 — 로컬 실측)
"""
import pytest
from openpyxl import Workbook, load_workbook

from dsd_tool.mapping_sheet import SUMMARY_SHEET, build_mapping_workbook


def _left(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "BS"
    ws["A1"] = "재 무 상 태 표"
    ws["A2"] = "(단위 : 천원)"
    ws.append([])
    ws.append(["과목", "당기", "전기"])
    ws.append(["현금및현금성자산", 500_000, 450_000])
    ws.append(["매출채권", 120_000, 110_000])
    ws2 = wb.create_sheet("1")
    ws2["A1"] = "1. 일반사항"
    wb.save(path)
    return path


def _right(path, tamper=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "R_BS"
    ws["A1"] = "[D210005] 재무상태표 - 별도"
    ws.append([])
    ws.append(["계정과목", "당기말\n2025-12-31", "당기말\n2024-12-31",
               "당기말\n2023-12-31"])
    # 원 단위 인스턴스 값 — 좌(천원)과 단위환산 일치해야 함
    ws.append(["현금및현금성자산", 500_000_000 + (7_777_777 if tamper
                                                  else 0),
               450_000_000, 999_999_999])
    ws.append(["매출채권", 119_999_600, 110_000_400, 888_888_888])
    wb.create_sheet("R_ONLY")["A1"] = "[D888888] 대응없는 공시"
    wb["R_ONLY"]["B3"] = 123
    wb.save(path)
    return path


_PAIRS = [
    {"left": "BS", "right": "R_BS", "name": "BS",
     "left_title": "재무상태표", "right_def": "[D210005] 재무상태표",
     "basis": "재무제표 제목 분류"},
    {"left": "1", "right": None, "name": "주석1",
     "left_title": "1. 일반사항", "right_def": "", "basis": ""},
    {"left": None, "right": "R_ONLY", "name": "X_R",
     "left_title": "", "right_def": "[D888888] 대응없는 공시",
     "basis": ""},
]


@pytest.fixture()
def built(tmp_path):
    left = _left(str(tmp_path / "L.xlsx"))
    right = _right(str(tmp_path / "R.xlsx"))
    out = str(tmp_path / "조서.xlsx")
    res = build_mapping_workbook(left, right, _PAIRS, out)
    return res, out


def test_g1_side_by_side_offset(built):
    res, out = built
    wb = load_workbook(out)
    ws = wb["BS"]
    # 좌 최대열 = 3(과목·당기·전기) → 우 블록 시작 = 5열 (3+2)
    assert ws.cell(4, 1).value == "과목"
    assert ws.cell(3, 5).value == "계정과목"        # 우 블록 헤더


def test_g2_verdicts_and_conversion(built):
    res, out = built
    wb = load_workbook(out)
    ws = wb["BS"]
    # 판정 열 = 우 블록 끝 + 1 (4열 우 블록 → 5+4-1=8, 판정 9열)
    vcol = 9
    assert ws.cell(2, vcol).value == "판정"
    assert ws.cell(4, vcol).value is True           # 현금 행 (환산 일치)
    assert ws.cell(5, vcol).value is True           # 매출채권 (±배율 미만)
    assert "단위환산" in str(ws.cell(4, vcol + 1).value or "") or \
        "단위환산" in str(ws.cell(5, vcol + 1).value or "")
    assert res["false"] == 0
    assert res["excluded"] == 2                     # 2023 열 2값 명시 제외


def test_g4_summary_formula_and_lonely(built):
    res, out = built
    wb = load_workbook(out)
    w0 = wb[SUMMARY_SHEET]
    rows = {w0.cell(r, 1).value: r for r in range(4, w0.max_row + 1)}
    f = w0.cell(rows["BS"], 7).value
    assert str(f).startswith('=COUNTIF(INDIRECT(') and '"FALSE"' in str(f)
    assert "대응 없음" in str(w0.cell(rows["주석1"], 5).value)
    assert "대응 없음" in str(w0.cell(rows["X_R"], 5).value)
    assert w0.cell(rows["BS"], 8).value == 2        # 제외 집계 노출
    assert res["pairs"] == 1 and res["lonely"] == 2


def test_g5_tamper_flips_false(built, tmp_path):
    res0, _ = built
    left = _left(str(tmp_path / "L2.xlsx"))
    right = _right(str(tmp_path / "R2.xlsx"), tamper=True)
    res = build_mapping_workbook(left, right, _PAIRS,
                                 str(tmp_path / "조서2.xlsx"))
    assert res["false"] == res0["false"] + 1

"""Phase 1: isolated, offline reproductions of M-01/M-04/M-06/M-07."""
import importlib
import json
import sys
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from auditdesk import jobs
from auditdesk.routers import explorer, studio, workbench as w
from dsd_tool.tests.fixture import build_dsd


@pytest.fixture
def session(tmp_path, monkeypatch):
    monkeypatch.setattr(jobs, "_DB", str(tmp_path / "app.sqlite"))
    monkeypatch.setattr(w, "_WORKDIR", str(tmp_path / "sessions"))
    monkeypatch.setattr(studio, "_WORKDIR", str(tmp_path / "studio"))
    monkeypatch.setattr(jobs, "submit", lambda kind, fn: fn(lambda *a: None))
    dsd = build_dsd(str(tmp_path / "input.dsd"))
    sid = w.create_session({"dsd_path": dsd})["session_id"]
    w.extract_session(sid)
    return sid, tmp_path


def edit(sid, amount):
    path = w.get_session(sid)["xlsx_path"]
    wb = load_workbook(path)
    for row in wb["BS"]:
        for cell in row:
            if isinstance(cell.value, (float, int)) and cell.value > 1000000:
                cell.value = amount
                wb.save(path)
                wb.close()
                return
    raise AssertionError("fixture edit target missing")


def test_m01_changed_value_same_ids_requires_review(session):
    sid, _ = session
    edit(sid, 1234577)
    review = w.diff_session(sid)
    edit(sid, 1234587)
    with pytest.raises(HTTPException) as err:
        w.repack_session(sid, {"approved_change_ids": [c["id"] for c in review["changes"]]})
    assert err.value.status_code == 409
    assert "다시" in err.value.detail


def test_m01_queued_repack_rechecks_input(session, monkeypatch):
    sid, _ = session
    w.diff_session(sid)
    pending = []
    monkeypatch.setattr(jobs, "submit", lambda kind, fn: pending.append(fn) or "queued")
    w.repack_session(sid)
    edit(sid, 1234587)
    with pytest.raises((HTTPException, ValueError)):
        pending[0](lambda *a: None)
    assert w.get_session(sid)["repack"] is None


def test_m01_unchanged_input_repack_works(session):
    sid, _ = session
    review = w.diff_session(sid, {"options": {"clean_cr": False}})
    result = w.repack_session(sid)["job_id"]
    assert Path(result["output_path"]).read_bytes() == Path(w.get_session(sid)["dsd_path"]).read_bytes()
    assert review["counts"]["total"] == 0


def test_m06_reextract_preserves_edited_workbook_and_invalidates(session):
    sid, _ = session
    edit(sid, 7654321)
    w.diff_session(sid)
    w._update(sid, foot='{"old":true}', recon='{"old":true}', repack='{"old":true}')
    old = Path(w.get_session(sid)["xlsx_path"])
    before = old.read_bytes()
    w.extract_session(sid)
    current = w.get_session(sid)
    assert old.read_bytes() == before
    assert Path(current["xlsx_path"]) != old
    assert all(current[k] is None for k in ("diff", "diff_options", "repack", "foot", "recon"))


def test_m04_result_is_persisted_only_to_own_session(session, monkeypatch):
    sid, tmp = session
    other = w.create_session({"dsd_path": w.get_session(sid)["dsd_path"]})["session_id"]
    package = tmp / "package"
    package.mkdir()
    (package / "succession_assets.json").write_text("{}")
    from types import SimpleNamespace
    monkeypatch.setattr(importlib.import_module("dart_explorer.xbrl.dimension_table"), "XbrlInstance", lambda p: SimpleNamespace(facts={}, doc_period_end="2025-12-31"))
    monkeypatch.setattr(importlib.import_module("dart_explorer.xbrl.corpus"), "_element_roles", lambda p: {})
    monkeypatch.setattr(importlib.import_module("dsd_tool.succession"), "load_assets", lambda p: {})
    monkeypatch.setattr(studio, "_decided_map", lambda: {})
    monkeypatch.setattr(importlib.import_module("dsd_tool.xbrl_recon"), "xbrl_recon", lambda *a, **kw: {"rows": {}, "false": 0})
    w._update(sid, recon='{"prior":true}')
    result = studio.xbrl_recon_route({"session_id": sid, "package_dir": str(package)})["job_id"]
    assert w.get_session(sid).get("xbrl_recon") == result
    assert result["session_id"] == sid
    assert w.get_session(other).get("xbrl_recon") is None
    assert w.get_session(sid)["recon"] == {"prior": True}


def test_m07_cache_clear_preserves_work_products_and_session_inputs(session, monkeypatch):
    sid, tmp = session
    root = tmp / "cache"
    doc = root / "document" / "corp"
    doc.mkdir(parents=True)
    inputs = [doc / "edited.xlsx", doc / "source.dsd", doc / "source.zip"]
    for p in inputs:
        p.write_bytes(b"user/session data")
    w._update(sid, dsd_path=str(inputs[1]), xlsx_path=str(inputs[0]))
    transient = root / "search" / "query.json"
    transient.parent.mkdir()
    transient.write_text("{}")
    monkeypatch.setattr(explorer, "_cache_root", lambda: str(root))
    explorer.clear_cache("DELETE")
    assert all(p.exists() for p in inputs)
    assert not transient.exists()


def test_m01_repack_consumes_reviewed_snapshot_during_concurrent_edit(session, monkeypatch):
    sid, _ = session
    edit(sid, 1234577)
    w.diff_session(sid)
    module = importlib.import_module("dsd_tool.repack")
    actual = module.repack
    changes = []
    def concurrent(xlsx, *args, **kw):
        edit(sid, 1234587)
        result = actual(xlsx, *args, **kw)
        changes.extend(result["changes"])
        return result
    monkeypatch.setattr(module, "repack", concurrent)
    w.repack_session(sid)
    assert any("1,234,577" in str(c["new"]) for c in changes)
    assert all("1,234,587" not in str(c["new"]) for c in changes)


def test_m06_old_foot_completion_cannot_restore_invalid_result(session, monkeypatch):
    sid, _ = session
    xlsx = w.get_session(sid)["xlsx_path"]
    def foot(*a, **kw):
        w.extract_session(sid)
        return dict(match=0, fuzzy=0, mismatch=0, note_missing=0, note_found=0,
                    manual_overrides=0, foot=[], levels=[], notes=[])
    monkeypatch.setattr(importlib.import_module("dsd_tool.foot"), "foot", foot)
    with pytest.raises(HTTPException):
        w._run_foot(sid, xlsx, False, 2, lambda *a: None)
    assert w.get_session(sid)["foot"] is None


def test_m06_failed_extract_keeps_previous_work_and_state(session, monkeypatch):
    sid, _ = session
    w.diff_session(sid)
    before = w.get_session(sid)
    def fail(*a, **kw):
        raise PermissionError("locked")
    monkeypatch.setattr(importlib.import_module("dsd_tool.excel_out"), "extract", fail)
    with pytest.raises(PermissionError):
        w.extract_session(sid)
    assert w.get_session(sid) == before


def test_m07_deletes_only_unreferenced_downloads(session, monkeypatch):
    _, tmp = session
    root = tmp / "cache"
    corp = root / "xbrl" / "corp"
    corp.mkdir(parents=True)
    archive = corp / "20260914000001_11011.zip"
    archive.write_bytes(b"download")
    package = corp / "receipt"
    package.mkdir()
    work = package / "custom.xml"
    work.write_text("user modifications")
    unknown = root / "notes.txt"
    unknown.write_text("user notes")
    monkeypatch.setattr(explorer, "_cache_root", lambda: str(root))
    result = explorer.clear_cache("DELETE")
    assert not archive.exists() and work.exists() and unknown.exists()
    assert result["cleared"] == 1


def test_m07_active_job_blocks_cache_clear(session, monkeypatch):
    _, tmp = session
    root = tmp / "cache"
    root.mkdir()
    monkeypatch.setattr(explorer, "_cache_root", lambda: str(root))
    with jobs.connect() as con:
        con.execute("INSERT INTO jobs(id, state) VALUES('busy', 'running')")
    with pytest.raises(HTTPException) as err:
        explorer.clear_cache("DELETE")
    assert err.value.status_code == 409


def test_m01_history_keeps_source_workbook_path(session):
    sid, _ = session
    w.diff_session(sid)
    w.repack_session(sid)
    from dsd_tool.history import query_runs
    s = w.get_session(sid)
    assert query_runs(s["dsd_path"])[0]["xlsx_path"] == s["xlsx_path"]


def test_m01_background_conflict_explains_rereview():
    message = "수정 확인 이후 엑셀이 변경되었습니다 — 다시 수정 확인하세요"
    assert jobs._humanize_error(HTTPException(409, message)) == message


def test_m06_preextract_job_cannot_overwrite_new_results(session):
    sid, _ = session
    with pytest.raises(HTTPException):
        w._update(sid, expected_xlsx=None, recon='{"stale":true}')


def test_m01_old_review_token_rejected_after_another_review(session):
    sid, _ = session
    edit(sid, 1234577)
    old = w.diff_session(sid)
    edit(sid, 1234587)
    w.diff_session(sid)
    with pytest.raises(HTTPException) as err:
        w.repack_session(sid, {"reviewed_xlsx_sha256": old["xlsx_sha256"]})
    assert err.value.status_code == 409


def test_m01_non_diff_workbook_change_also_requires_review(session):
    sid, _ = session
    w.diff_session(sid)
    path = w.get_session(sid)["xlsx_path"]
    wb = load_workbook(path)
    wb.create_sheet("reviewer_notes")["A1"] = "post-review change"
    wb.save(path)
    wb.close()
    with pytest.raises(HTTPException) as err:
        w.repack_session(sid)
    assert err.value.status_code == 409


def test_m06_reextract_clears_xbrl_result(session):
    sid, _ = session
    w._update(sid, xbrl_recon=json.dumps({"session_id": sid, "false": 0}))
    w.extract_session(sid)
    assert w.get_session(sid)["xbrl_recon"] is None


def test_m06_diff_finishing_after_reextract_is_rejected(session, monkeypatch):
    sid, _ = session
    module = importlib.import_module("dsd_tool.repack")
    original = module.diff
    def interleave(*a, **kw):
        result = original(*a, **kw)
        w.extract_session(sid)
        return result
    monkeypatch.setattr(module, "diff", interleave)
    with pytest.raises(HTTPException):
        w.diff_session(sid)
    assert w.get_session(sid)["diff"] is None


def test_m07_unknown_archive_is_not_a_download_cache(session, monkeypatch):
    _, tmp = session
    root = tmp / "cache"
    folder = root / "xbrl" / "corp"
    folder.mkdir(parents=True)
    work = folder / "reviewer_backup.zip"
    work.write_bytes(b"user work")
    monkeypatch.setattr(explorer, "_cache_root", lambda: str(root))
    explorer.clear_cache("DELETE")
    assert work.exists()

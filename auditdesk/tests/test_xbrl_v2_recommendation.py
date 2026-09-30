"""Issue #9 behavior contracts; only runtime-generated public-shape test inputs."""
from dataclasses import replace, asdict
from hashlib import sha256
from io import BytesIO
import json
from zipfile import ZipFile
import pytest
from openpyxl import Workbook
from auditdesk.xbrl_v2.dsd import parse_dsd, DsdReportContext
from auditdesk.xbrl_v2.taxonomy import parse_taxonomy_workbook, TaxonomyMetadata, NamespaceBinding
from auditdesk.xbrl_v2.recommendation import (
    recommend, RecommendationConstraint, ReferenceEvidence, assert_fresh,
)


def inputs():
    xml = '<DOCUMENT><COMPANY-NAME AREGCIK="sample">Synthetic</COMPANY-NAME><DOCUMENT-NAME>연결감사보고서</DOCUMENT-NAME><PERIODEND>2026-12-31</PERIODEND><TITLE>재무상태표</TITLE><TABLE><TR><TH>항목</TH><TH>당기</TH><TH>전기</TH></TR><TR><TD>현금</TD><TD>100</TD><TD>90</TD></TR><TR><TD>이름없는항목</TD><TD></TD><TD NIL="true"/></TR></TABLE></DOCUMENT>'
    z=BytesIO()
    with ZipFile(z,'w') as archive: archive.writestr('contents.xml',xml)
    report=DsdReportContext('DART:AREGCIK','sample','CONNECTED','2026-01-01','2026-12-31')
    doc=parse_dsd(z.getvalue(),report=report)
    wb=Workbook(); c=wb.active; c.title='Concepts'
    c.append(['#','prefix','name','id','type','periodType','abstract'])
    c.append([1,'std','Root','root','xbrli:stringItemType','duration','true'])
    c.append([2,'std','Cash','cash','xbrli:monetaryItemType','instant','false'])
    c.append([3,'std','WrongPeriod','wrong','xbrli:monetaryItemType','duration','false'])
    r=wb.create_sheet('RoleTypes'); r.append(['id','roleURI','usedOn','definition'])
    r.append(['bs','urn:role:bs','presentationLink','재무상태표 - 연결'])
    r.append(['note','urn:role:note','presentationLink','기타 주석'])
    p=wb.create_sheet('Presentation Link')
    for role in ('urn:role:bs','urn:role:note'):
        p.append([None,'LinkRole',None,role]);p.append([None,'Definition',None,'재무상태표' if role.endswith('bs') else '기타 주석'])
        p.append(['#','prefix',None,'name','label','depth','order','priority','parent','arcrole','preferredLabel'])
        p.append([1,'std',None,'Root','Root',0,1,0,None,'parent-child','label'])
        p.append([2,'std',None,'Cash','현금',1,1,0,'std:Root','parent-child','label'])
        p.append([3,'std',None,'WrongPeriod','현금',1,2,0,'std:Root','parent-child','label'])
    l=wb.create_sheet('Label Link');l.append(['LinkRole','urn:label']);l.append([]);l.append([None,None,None,'ko']);l.append(['#','prefix','name','label'])
    l.append([1,'std','Root','재무상태표']);l.append([2,'std','Cash','현금']);l.append([3,'std','WrongPeriod','현금'])
    rf=wb.create_sheet('Reference Link');rf.append(['LinkRole','urn:ref']);rf.append([]);rf.append([]);rf.append(['#','prefix','name','label','verbose','role','Name','Number','Paragraph'])
    rf.append([1,'std','Cash',None,None,'disclosureRef','Test Standard','1','2'])
    b=BytesIO();wb.save(b);data=b.getvalue()
    metadata=TaxonomyMetadata('release-test','CURRENT_TAXONOMY','2026-12-31','2026-01-01','2026-12-31','test applicability declaration',sha256(data).hexdigest())
    tax=parse_taxonomy_workbook(data,logical_uri='urn:test:current',metadata=metadata,namespaces=(NamespaceBinding('std','urn:std:2026','test namespace declaration'),))
    return doc,tax,z.getvalue(),data


def amount(doc): return next(s for s in doc.subjects if s.text=='100')
def choice(doc, **kw):
    return RecommendationConstraint(subject_id=amount(doc).subject_id,source_snapshot_id=doc.snapshot_id,role_uri='urn:role:bs',period_type='instant',**kw)
def result_for(result,doc): return next(r for r in result.subjects if r.subject_id==amount(doc).subject_id)
def prior(doc,tax,**changes):
    o=next(o for o in tax.occurrences if o.qname and o.qname.local_name=='Cash' and o.role_uri=='urn:role:bs')
    base=ReferenceEvidence('prior-1','PRIOR_COMPANY_XBRL','a'*64,'DART:AREGCIK','sample','CONNECTED','2025-12-31',o.qname.key,o.role_uri,'OLD LABEL','xbrli:monetaryItemType','instant',o.path)
    return replace(base,**changes)


def test_current_only_no_instance_and_no_automatic_confirmation():
    doc,tax,_,_=inputs();r=recommend(doc,tax,constraints=(choice(doc),))
    s=result_for(r,doc)
    assert s.state=='SUGGESTED' and len(s.candidates)==1
    assert s.candidates[0].qname=='{urn:std:2026}Cash'
    assert s.candidates[0].role_uri=='urn:role:bs'
    assert s.candidates[0].current_validation=='PARTIAL'
    assert s.candidates[0].reference_resource_ids
    assert r.coverage.total==len(doc.subjects)
    assert all(s.state not in ('USER_CONFIRMED','DERIVED_FROM_CONFIRMED') for s in r.subjects)


def test_structure_before_label_and_unknown_period_is_not_mismatch():
    doc,tax,_,_=inputs();s=result_for(recommend(doc,tax,constraints=(choice(doc),)),doc)
    assert any('PERIOD_TYPE_MISMATCH' in x.reasons for x in s.rejected)
    unknown=result_for(recommend(doc,tax,constraints=(replace(choice(doc),period_type=None),)),doc)
    assert len(unknown.candidates)==2 and unknown.ambiguous
    assert any('PERIOD_TYPE_UNKNOWN' in c.unverified for c in unknown.candidates)


@pytest.mark.parametrize('changes',[{'qname':'{urn:old}Cash'},{'role_uri':'urn:old-role'},{'entity_identifier':'other'},{'scope':'SEPARATE'},{'period_end':'2026-12-31'},{'type_name':'xbrli:stringItemType'},{'period_type':'duration'},{'path':('{urn:std:2026}Cash',)}])
def test_prior_cannot_override_current_or_seed_candidates(changes):
    doc,tax,_,_=inputs();p=prior(doc,tax,**changes)
    s=result_for(recommend(doc,tax,constraints=(choice(doc),),references=(p,)),doc)
    assert [c.qname for c in s.candidates]==['{urn:std:2026}Cash']
    assert s.references[0].state!='REVALIDATED'
    assert s.references[0].reasons


def test_prior_label_retained_as_reference_not_current_label():
    doc,tax,_,_=inputs();s=result_for(recommend(doc,tax,constraints=(choice(doc),),references=(prior(doc,tax),)),doc)
    assert s.references[0].state=='REVALIDATED'
    assert 'LABEL_CHANGED' in s.references[0].reasons
    assert s.candidates[0].label=='현금' and s.state=='SUGGESTED'


def test_missing_current_never_uses_prior():
    doc,tax,_,_=inputs();r=recommend(doc,None,references=(prior(doc,tax),))
    assert all(not s.candidates for s in r.subjects)
    assert 'CURRENT_TAXONOMY_MISSING' in r.diagnostics


def test_unverified_applicability_is_unresolved():
    doc,tax,_,_=inputs();r=recommend(doc,replace(tax,current_applicable=False))
    assert all(not s.candidates for s in r.subjects)


def test_source_conflict_blocks_and_hash_change_is_stale():
    doc,tax,_,_=inputs();r=recommend(doc,tax)
    assert_fresh(r,doc,tax)
    changed=replace(doc,source_sha256='b'*64)
    with pytest.raises(ValueError,match='STALE'): assert_fresh(r,changed,tax)
    bad=replace(doc,diagnostics=(*doc.diagnostics,'CONFLICT:scope'))
    assert all(s.state=='CONFLICT' for s in recommend(bad,tax).subjects)


def test_replay_scope_and_constraint_identity():
    doc,tax,_,_=inputs();r=recommend(doc,tax,constraints=(choice(doc),))
    assert r==recommend(doc,tax,constraints=(choice(doc),))
    with pytest.raises(ValueError): recommend(doc,tax,constraints=(replace(choice(doc),source_snapshot_id='old'),))
    with pytest.raises(ValueError): recommend(doc,tax,constraints=(choice(doc),choice(doc)))
    with pytest.raises(ValueError,match='STALE'): assert_fresh(r,doc,tax,rule_version='new')


def test_no_prior_extension_or_golden_source_promoted():
    doc,tax,_,_=inputs()
    p=prior(doc,tax,qname='{urn:prior-company}Custom')
    s=result_for(recommend(doc,tax,references=(p,),constraints=(choice(doc),)),doc)
    assert not any(c.qname==p.qname for c in s.candidates)
    with pytest.raises(ValueError): recommend(doc,tax,references=(replace(p,source_kind='GOLDEN'),))


def test_unknown_dimensions_and_no_value_rewrite():
    doc,tax,_,_=inputs();r=recommend(doc,tax,constraints=(choice(doc,dimensions=(('{urn:std}Axis','{urn:std}Member'),)),))
    s=result_for(r,doc)
    assert s.candidates and all('DIMENSIONS_UNVERIFIED' in c.unverified for c in s.candidates)
    assert s.source_text=='100'
    assert any(x.source_value_state=='BLANK' for x in r.subjects)
    assert any(x.source_value_state=='NIL' for x in r.subjects)
    assert json.loads(json.dumps(asdict(r),ensure_ascii=False))['coverage']['total']==len(doc.subjects)


def test_freshness_requires_latest_nonempty_reference_and_constraint_inputs():
    doc,tax,_,_=inputs();c=choice(doc);p=prior(doc,tax)
    r=recommend(doc,tax,constraints=(c,),references=(p,))
    assert_fresh(r,doc,tax,constraints=(c,),references=(p,))
    with pytest.raises(ValueError,match='STALE'): assert_fresh(r,doc,tax,constraints=(c,))
    with pytest.raises(ValueError,match='STALE'): assert_fresh(r,doc,tax,references=(p,))
    with pytest.raises(ValueError,match='STALE'):
        assert_fresh(r,doc,tax,constraints=(c,),references=(replace(p,source_sha256='c'*64),))
    with pytest.raises(ValueError,match='STALE'):
        assert_fresh(r,doc,replace(tax,roles=()),constraints=(c,),references=(p,))


def test_unknown_structure_has_no_global_label_fallback():
    doc,tax,_,_=inputs()
    changed=replace(doc,subjects=tuple(replace(s,table_title='unrelated') if s.subject_id==amount(doc).subject_id else s for s in doc.subjects))
    s=result_for(recommend(changed,tax),doc)
    assert s.state=='UNRESOLVED' and not s.candidates
    assert 'NO_ROLE_CANDIDATE' in s.diagnostics


def test_applicability_flag_cannot_mask_contradictory_metadata():
    doc,tax,_,_=inputs()
    altered=replace(tax,metadata=replace(tax.metadata,content_sha256='c'*64),current_applicable=True)
    assert all(not s.candidates for s in recommend(doc,altered).subjects)
    altered=replace(tax,metadata=replace(tax.metadata,report_period='2025-12-31'),current_applicable=True)
    assert all(s.state=='CONFLICT' for s in recommend(doc,altered).subjects)


@pytest.mark.parametrize('changes',[{'type_name':''},{'period_type':''}])
def test_unknown_reference_compatibility_not_revalidated(changes):
    doc,tax,_,_=inputs()
    s=result_for(recommend(doc,tax,constraints=(choice(doc),),references=(prior(doc,tax,**changes),)),doc)
    assert s.references[0].state=='UNVERIFIED'


def test_empty_report_identity_is_source_gap():
    doc,tax,_,_=inputs()
    changed=replace(doc,metadata=replace(doc.metadata,requested=replace(doc.metadata.requested,entity_scheme=None)))
    assert all(not s.candidates for s in recommend(changed,tax).subjects)


def test_current_interim_reference_is_not_prior_year():
    doc,tax,_,_=inputs()
    s=result_for(recommend(doc,tax,constraints=(choice(doc),),references=(prior(doc,tax,period_end='2026-06-30'),)),doc)
    assert s.references[0].state=='UNVERIFIED'


def test_provenance_and_role_proposals_are_self_contained():
    doc,tax,_,_=inputs();s=result_for(recommend(doc,tax,constraints=(choice(doc),)),doc)
    assert s.source_span==amount(doc).source_span and s.table_id==amount(doc).table_id
    assert s.role_candidates and s.role_candidates[0].source.sheet=='RoleTypes'
    c=s.candidates[0]
    assert c.occurrence_source.sheet=='Presentation Link'
    assert c.label_source.sheet=='Label Link'
    assert c.reference_resources[0].parts
    assert 'KNOWN_PERIOD_TYPE_MATCH' in c.evidence_reasons


def test_irrelevant_roles_do_not_expand_per_subject_payload():
    doc,tax,_,_=inputs()
    irrelevant=tuple(replace(o,role_uri='urn:irrelevant:'+str(i),occurrence_id=o.occurrence_id+str(i)) for i in range(2000) for o in tax.occurrences[:1])
    expanded=replace(tax,occurrences=tax.occurrences+irrelevant)
    before=result_for(recommend(doc,tax,constraints=(choice(doc),)),doc)
    after=result_for(recommend(doc,expanded,constraints=(choice(doc),)),doc)
    assert after.candidates==before.candidates
    assert after.evaluated_occurrences==before.evaluated_occurrences
    assert after.excluded_role_count==before.excluded_role_count+2000
    assert len(after.rejected)==len(before.rejected)


def test_document_unclassified_coverage_remains_visible():
    doc,tax,_,_=inputs()
    changed=replace(doc,coverage=replace(doc.coverage,unclassified_text_nodes=4),diagnostics=(*doc.diagnostics,'UNCLASSIFIED_TEXT:123'))
    result=recommend(changed,tax)
    assert result.source_coverage.unclassified_text_nodes==4
    assert 'UNCLASSIFIED_TEXT:123' in result.diagnostics


def test_known_role_scope_mismatch_is_not_label_candidate():
    doc,tax,_,_=inputs()
    changed=replace(tax,roles=tuple(replace(r,definition='재무상태표 - 별도') if r.role_uri=='urn:role:bs' else r for r in tax.roles))
    result=result_for(recommend(doc,changed,constraints=(choice(doc),)),doc)
    assert not result.candidates
    assert 'ROLE_SCOPE_MISMATCH' in result.diagnostics


def test_output_truncation_is_visible_and_keeps_subject_coverage():
    doc,tax,_,_=inputs()
    occurrence=next(o for o in tax.occurrences if o.qname.local_name=='Cash' and o.role_uri=='urn:role:bs')
    changed=replace(tax,occurrences=tax.occurrences+tuple(replace(occurrence,occurrence_id='extra'+str(i),path=(*occurrence.path,str(i))) for i in range(65)))
    result=recommend(doc,changed,constraints=(choice(doc),))
    subject=result_for(result,doc)
    assert subject.candidate_total==66 and len(subject.candidates)==50
    assert 'CANDIDATES_TRUNCATED:66' in subject.diagnostics
    assert subject.ambiguous and result.coverage.total==len(doc.subjects)

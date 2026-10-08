from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import pytest
from openpyxl import Workbook
from auditdesk.xbrl_v2.taxonomy import TaxonomyMetadata, NamespaceBinding, parse_taxonomy_workbook


def workbook(duplicate=False, malformed=False):
    w=Workbook(); c=w.active; c.title='Concepts'
    c.append(['#','prefix','name','id','type','periodType','abstract'])
    c.append([1,'ifrs','Cash','cash','xbrli:monetaryItemType','instant','false'])
    c.append([2,'ifrs','Root','root','xbrli:stringItemType','duration','true'])
    if duplicate: c.append([3,'ifrs','Cash','cash2','xbrli:stringItemType','duration','false'])
    if malformed: c.cell(1,7,'name')
    r=w.create_sheet('RoleTypes'); r.append(['#','id','roleURI','usedOn','definition']); r.append([1,'r','urn:role:a','presentationLink','Cash role'])
    p=w.create_sheet('Presentation Link')
    for role in ['urn:role:a','urn:role:a']:
        p.append([None,'LinkRole',None,role]); p.append([None,'Definition',None,'Cash role'])
        p.append(['#','prefix',None,'name','label','depth','order','priority','parent','arcrole','preferredLabel','systemid'])
        p.append([1,'ifrs',None,'Root','Root',0,1,0,None,'parent-child',None,'a.xsd'])
        p.append([2,'ifrs',None,'Cash','Cash',1,1,0,'ifrs:Root','parent-child','label','a.xsd'])
    l=w.create_sheet('Label Link'); l.append([None,'LinkRole',None,'urn:label']); l.append([])
    l.append([None,None,None,'ko',None,'en']); l.append(['#','prefix','name','label','verboseLabel','label'])
    l.append([1,'ifrs','Cash','현금','현금 설명','Cash'])
    r=w.create_sheet('Reference Link'); r.append([None,'LinkRole',None,'urn:ref']); r.append([]); r.append([])
    r.append(['#','prefix','name','label','verbose','role','Name','Number','Paragraph'])
    r.append([1,'ifrs','Cash',None,None,'disclosureRef','IAS',7,'6'])
    b=BytesIO(); w.save(b); return b.getvalue()


def parse(data=None, **changes):
    data=data or workbook()
    m=TaxonomyMetadata('2026','CURRENT_TAXONOMY','2026-12-31','2026-01-01','2026-12-31','public release notice',sha256(data).hexdigest())
    return parse_taxonomy_workbook(data,logical_uri='public:taxonomy.xlsx',metadata=replace(m,**changes),namespaces=(NamespaceBinding('ifrs','urn:ifrs:2026','release namespace declaration'),))


def test_current_workbook_has_provenance_and_partial_capability():
    s=parse(); assert s.current_applicable
    assert s.capabilities == 'PARTIAL' and s.dimensions_known is False
    cash=s.concepts[0]; assert cash.qname.key == '{urn:ifrs:2026}Cash'
    assert cash.period_type == 'instant' and cash.abstract is False
    assert s.concepts[1].abstract is True
    assert len(s.occurrences)==4 and len({o.occurrence_id for o in s.occurrences})==4
    assert s.occurrences[1].path == ('{urn:ifrs:2026}Root','{urn:ifrs:2026}Cash')
    assert {(l.language,l.text) for l in s.labels} == {('ko','현금'),('ko','현금 설명'),('en','Cash')}
    assert s.references[0].parts == (('Name','IAS'),('Number','7'),('Paragraph','6'))
    assert all(x.source.sheet and x.source.row for x in (*s.labels,*s.references,*s.occurrences))


@pytest.mark.parametrize('changes',[{'source_role':'PRIOR_COMPANY_XBRL'},{'applicability_evidence':''},{'applicable_to':'2025-12-31'},{'content_sha256':'0'*64},{'version':''}])
def test_current_claim_requires_bound_evidence(changes):
    assert not parse(**changes).current_applicable


def test_namespace_is_never_guessed():
    data=workbook(); s=parse_taxonomy_workbook(data,logical_uri='current.xlsx',metadata=parse(data).metadata)
    assert all(c.qname is None for c in s.concepts)
    assert 'UNRESOLVED_NAMESPACE' in {d.code for d in s.diagnostics}
    assert not s.current_applicable


@pytest.mark.parametrize('kwargs',[{'duplicate':True},{'malformed':True}])
def test_ambiguous_concepts_fail_closed(kwargs):
    s=parse(workbook(**kwargs)); assert not s.current_applicable
    assert any(d.severity == 'ERROR' for d in s.diagnostics)


def test_identity_binds_metadata_and_is_deterministic():
    data=workbook(); a=parse(data); assert a == parse(data)
    assert a.snapshot_id != parse(data,applicability_evidence='another notice').snapshot_id
    with pytest.raises(Exception): a.capabilities='FULL'


def test_malformed_zip_returns_diagnostic():
    s=parse(b'not an xlsx'); assert not s.current_applicable
    assert any(d.code == 'INVALID_WORKBOOK' for d in s.diagnostics)

def edited(change):
    from openpyxl import load_workbook
    b=BytesIO(workbook()); w=load_workbook(b); change(w)
    out=BytesIO(); w.save(out); w.close(); return out.getvalue()


def test_blank_abstract_remains_unknown_with_partial_warning():
    s=parse(edited(lambda w: setattr(w['Concepts']['G2'],'value',None)))
    assert s.concepts[0].abstract is None
    assert s.current_applicable
    assert any(d.code=='UNKNOWN_ABSTRACT' and d.severity=='WARNING' for d in s.diagnostics)


def test_missing_presentation_header_fails_closed():
    s=parse(edited(lambda w: w.remove(w['Presentation Link'])))
    assert not s.current_applicable


def test_duplicate_label_language_role_is_ambiguous():
    s=parse(edited(lambda w: setattr(w['Label Link']['E4'],'value','label')))
    assert not s.current_applicable
    assert any(d.code=='AMBIGUOUS_LABEL_HEADER' for d in s.diagnostics)


def test_presentation_without_role_marker_fails_closed():
    s=parse(edited(lambda w: w['Presentation Link'].delete_rows(1,2)))
    assert not s.current_applicable


def test_bad_parent_and_depth_fail_closed():
    s=parse(edited(lambda w: setattr(w['Presentation Link']['I5'],'value','ifrs:Cash')))
    assert not s.current_applicable
    s=parse(edited(lambda w: setattr(w['Presentation Link']['F5'],'value',9)))
    assert not s.current_applicable


def test_formula_is_not_executed_or_accepted():
    s=parse(edited(lambda w: setattr(w['Concepts']['C2'],'value','=1+1')))
    assert not s.current_applicable
    assert any(d.code=='FORMULA_UNSUPPORTED' for d in s.diagnostics)


def test_duplicate_namespace_mapping_not_last_writer_wins():
    data=workbook(); m=parse(data).metadata
    s=parse_taxonomy_workbook(data,logical_uri='a.xlsx',metadata=m,namespaces=(NamespaceBinding('ifrs','urn:a','a'),NamespaceBinding('ifrs','urn:b','b')))
    assert not s.current_applicable and all(c.qname is None for c in s.concepts)


def test_reference_missing_target_preserved_with_error():
    s=parse(edited(lambda w: setattr(w['Reference Link']['C5'],'value','Absent')))
    assert not s.current_applicable and s.references[0].qname is None
    assert any(d.code=='UNRESOLVED_CONCEPT' for d in s.diagnostics)

def test_optional_resources_missing_stay_partial():
    def remove(w):
        w.remove(w['Label Link']); w.remove(w['Reference Link'])
    s=parse(edited(remove)); assert s.current_applicable
    assert {'MISSING_LABELS','MISSING_REFERENCES'} <= {d.code for d in s.diagnostics}


def test_existing_label_sheet_without_header_is_not_silently_empty():
    s=parse(edited(lambda w: w['Label Link'].delete_rows(4,1)))
    assert not s.current_applicable


def test_presentation_without_any_occurrences_is_invalid():
    s=parse(edited(lambda w: w['Presentation Link'].delete_rows(1,20)))
    assert not s.current_applicable


def test_missing_nested_parent_is_unresolved():
    s=parse(edited(lambda w: setattr(w['Presentation Link']['I5'],'value',None)))
    assert not s.current_applicable

@pytest.mark.parametrize('sheet,cell',[
    ('Concepts','B2'),('Concepts','C2'),
    ('Presentation Link','B5'),('Presentation Link','D5'),
    ('Label Link','B5'),('Label Link','C5'),
    ('Reference Link','B5'),('Reference Link','C5'),
])
def test_nonempty_rows_missing_identity_are_diagnosed_and_preserved(sheet,cell):
    s=parse(edited(lambda w: setattr(w[sheet][cell],'value',None)))
    assert not s.current_applicable
    assert any(d.code=='MISSING_CONCEPT_IDENTITY' and d.source.sheet==sheet
               for d in s.diagnostics)
    records={'Concepts':s.concepts,'Presentation Link':s.occurrences,
             'Label Link':s.labels,'Reference Link':s.references}[sheet]
    assert any(record.qname is None for record in records)
    assert len(s.occurrences)==4 and len(s.labels)==3


@pytest.mark.parametrize('sheet', ['Concepts','Presentation Link','Label Link','Reference Link'])
def test_genuinely_empty_rows_do_not_create_identity_errors(sheet):
    def add_empty(w):
        w[sheet].append([None]*20)
    s=parse(edited(add_empty))
    assert s.current_applicable
    assert not any(d.code=='MISSING_CONCEPT_IDENTITY' for d in s.diagnostics)

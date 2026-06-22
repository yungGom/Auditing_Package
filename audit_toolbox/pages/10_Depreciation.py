"""유형자산 감가상각비 재계산"""
import streamlit as st
import pandas as pd
from datetime import date, datetime, timedelta
from dateutil.relativedelta import relativedelta
from io import BytesIO
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="감가상각 재계산", page_icon="🏗️", layout="wide")
st.title("🏗️ 유형자산 감가상각비 재계산")
st.caption("회사 원장 업로드 → 컬럼 매핑 → 재계산 → 차이 분석")

HF=PatternFill("solid",fgColor="002060"); HN=Font(name="맑은 고딕",bold=True,color="FFFFFF",size=10)
TF=Font(name="맑은 고딕",bold=True,size=14,color="002060")
SecF=PatternFill("solid",fgColor="D6DCE4"); NF=Font(name="맑은 고딕",size=10)
BF=Font(name="맑은 고딕",size=10,bold=True); HL=PatternFill("solid",fgColor="FFF2CC")
DiffP=Font(name="맑은 고딕",size=10,color="FF0000"); DiffM=Font(name="맑은 고딕",size=10,color="0000FF")
NM='#,##0'; DT='YYYY-MM-DD'
tb=Border(left=Side(style='thin',color='B0B0B0'),right=Side(style='thin',color='B0B0B0'),top=Side(style='thin',color='B0B0B0'),bottom=Side(style='thin',color='B0B0B0'))

def shdr(ws,r,cols):
    for c in range(1,cols+1):
        cl=ws.cell(row=r,column=c); cl.font=HN; cl.fill=HF; cl.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); cl.border=tb

def parse_dt(val):
    if pd.isna(val) or val is None: return None
    if isinstance(val,datetime): return val.date()
    if isinstance(val,date): return val
    s=str(val).strip()
    for f in ('%Y-%m-%d','%Y/%m/%d','%Y.%m.%d','%Y%m%d'):
        try: return datetime.strptime(s,f).date()
        except: continue
    return None

def dep_months_fn(d1,d2):
    if d1 is None or d2 is None or d2<d1: return 0
    return (d2.year-d1.year)*12+d2.month-d1.month+1

def calc_dep(row,base_date):
    acq=row.get("_취득일"); cost=int(row.get("_취득원가",0) or 0)
    ul=int(row.get("_내용연수월",0) or 0); resid=int(row.get("_잔존가치",0) or 0)
    cxd=row.get("_자본적지출일"); cxa=int(row.get("_자본적지출액",0) or 0)
    method=str(row.get("_상각방법","정액법") or "정액법")
    empty={"상각대상":0,"월상각비":0,"전기말누계":0,"당기상각비":0,"기준일누계":0,"장부금액":cost+cxa,"CAPEX":False}
    if acq is None or ul<=0 or cost<=0 or acq>base_date: return empty
    fys=date(base_date.year,1,1); prev=fys-timedelta(days=1)
    if "정액" in method:
        hcx=cxa>0 and cxd and cxd<=base_date
        if not hcx:
            db=cost-resid; md=db/ul
            bf=min(dep_months_fn(acq,prev),ul) if acq<fys else 0
            tb2=min(dep_months_fn(acq,base_date),ul); fy=tb2-bf
            ab=round(md*bf); fd=round(md*fy); at=ab+fd
            return {"상각대상":db,"월상각비":round(md),"전기말누계":ab,"당기상각비":fd,"기준일누계":at,"장부금액":max(0,cost-at),"CAPEX":False}
        else:
            db1=cost-resid; md1=db1/ul
            mp1=min((cxd.year-acq.year)*12+cxd.month-acq.month,ul); ac=round(md1*mp1)
            bv=cost-ac+cxa; rem=max(ul-mp1,1); db2=bv-resid; md2=db2/rem
            if acq>=fys: ab=0
            elif cxd>=fys: ab=round(md1*min(dep_months_fn(acq,prev),mp1))
            else: ab=ac+round(md2*min(dep_months_fn(cxd,prev),rem))
            p2b=min(dep_months_fn(cxd,base_date),rem); at=ac+round(md2*p2b); fd=at-ab
            return {"상각대상":db2,"월상각비":round(md2),"전기말누계":ab,"당기상각비":fd,"기준일누계":at,"장부금액":max(0,cost+cxa-at),"CAPEX":True}
    elif "정률" in method:
        rate=1-(resid/cost)**(12/ul) if cost>0 and resid>0 else 2/(ul/12)
        mmr=rate/12; tu=min(dep_months_fn(acq,base_date),ul)
        bal=cost; ab=0; fd=0
        for m in range(tu):
            d=max(round(bal*mmr),0)
            if bal-d<resid: d=max(0,round(bal-resid))
            bal-=d; cm=acq+relativedelta(months=m)
            if cm<fys: ab+=d
            else: fd+=d
        return {"상각대상":cost-resid,"월상각비":0,"전기말누계":ab,"당기상각비":fd,"기준일누계":ab+fd,"장부금액":max(0,cost+cxa-ab-fd),"CAPEX":False}
    return empty

def gen_dep_excel(rdf,bd,hc):
    wb=Workbook(); ws=wb.active; ws.title="감가상각 재계산"
    ws.merge_cells('A1:Q1'); ws.cell(row=1,column=1,value=f"유형자산 감가상각비 재계산표 (기준일: {bd})").font=TF
    ws.cell(row=1,column=1).alignment=Alignment(horizontal='center',vertical='center'); ws.row_dimensions[1].height=35
    row=3
    hdrs=["No.","자산명","취득일","취득원가","잔존가치","내용연수(월)","상각방법","CAPEX일","CAPEX액","상각대상","월상각비","전기말누계","당기상각비","기준일누계","장부금액"]
    if hc: hdrs+=["회사상각비","차이"]
    for ci,h in enumerate(hdrs,1): ws.cell(row=row,column=ci,value=h)
    shdr(ws,row,len(hdrs)); ws.row_dimensions[row].height=40; row+=1
    tf2=0;tc=0;td=0
    for idx,(_,r) in enumerate(rdf.iterrows(),1):
        co=r.get("_회사감가상각비",0) or 0; fd=r["당기상각비"]; diff=fd-co if hc else 0
        tf2+=fd;tc+=co;td+=diff
        vals=[idx,r.get("_자산명",""),r.get("_취득일"),r.get("_취득원가",0),r.get("_잔존가치",0),
              r.get("_내용연수월",0),r.get("_상각방법","정액법"),r.get("_자본적지출일"),r.get("_자본적지출액",0),
              r["상각대상"],r["월상각비"],r["전기말누계"],fd,r["기준일누계"],r["장부금액"]]
        if hc: vals+=[co,diff]
        for ci,v in enumerate(vals,1):
            c=ws.cell(row=row,column=ci,value=v); c.border=tb; c.font=NF
            if ci in [3,8]: c.number_format=DT; c.alignment=Alignment(horizontal='center')
            elif ci>=4: c.number_format=NM; c.alignment=Alignment(horizontal='right')
            if hc and ci==len(vals): c.font=DiffP if diff>0 else (DiffM if diff<0 else NF)
            if r.get("CAPEX"): c.fill=HL
        row+=1
    ws.cell(row=row,column=1,value="합계").font=BF; ws.cell(row=row,column=1).border=tb
    for c in range(1,len(hdrs)+1): ws.cell(row=row,column=c).fill=SecF; ws.cell(row=row,column=c).border=tb
    sums={4:rdf["_취득원가"].sum(),12:rdf["전기말누계"].sum(),13:tf2,14:rdf["기준일누계"].sum(),15:rdf["장부금액"].sum()}
    if hc: sums[16]=tc; sums[17]=td
    for ci,v in sums.items(): c=ws.cell(row=row,column=ci,value=v); c.font=BF; c.number_format=NM
    for ci,w in enumerate([6,24,14,16,14,12,10,14,16,16,14,18,18,18,18,18,14],1):
        if ci<=len(hdrs): ws.column_dimensions[get_column_letter(ci)].width=w
    buf=BytesIO(); wb.save(buf); buf.seek(0); return buf

SKIP="(매핑 안함)"
with st.expander("📖 사용 안내",expanded=False):
    st.markdown("회사 원장 엑셀을 그대로 업로드 → 컬럼 매핑 → 재계산. 정액법/정률법, CAPEX 반영 지원.")

f=st.file_uploader("📂 유형자산 원장 업로드",type=["xlsx","xls","csv"],key="du")
if f:
    if f.name.endswith('.csv'): df=pd.read_csv(f,encoding='utf-8-sig')
    else:
        xls=pd.ExcelFile(f); sn=st.selectbox("시트",xls.sheet_names) if len(xls.sheet_names)>1 else xls.sheet_names[0]
        df=pd.read_excel(f,sheet_name=sn)
    st.success(f"✅ {len(df)}행 × {len(df.columns)}열")
    with st.expander("미리보기"): st.dataframe(df.head(10),use_container_width=True)
    rc=[SKIP]+list(df.columns.astype(str))
    st.markdown("#### 컬럼 매핑")
    c1,c2=st.columns(2)
    with c1:
        st.markdown("**필수**"); mn=st.selectbox("자산명",rc,key="dn"); ma=st.selectbox("취득일",rc,key="da")
        mc=st.selectbox("취득원가",rc,key="dc"); mu=st.selectbox("내용연수",rc,key="dul")
    with c2:
        st.markdown("**선택**"); mr=st.selectbox("잔존가치",rc,key="d_resid"); mm=st.selectbox("상각방법",rc,key="dm")
        mco=st.selectbox("회사 감가상각비",rc,key="dco"); mcd=st.selectbox("자본적지출일",rc,key="dcd"); mca=st.selectbox("자본적지출액",rc,key="dca")
    lu=st.radio("내용연수 단위",["년","월"],horizontal=True,key="dlu"); bdt=st.date_input("기준일",value=date(2025,12,31),key="db")
    rok=all(m!=SKIP for m in [mn,ma,mc,mu])
    if not rok: st.warning("⚠️ 필수 매핑을 완료해주세요.")
    elif st.button("🔢 감가상각 재계산",type="primary",use_container_width=True,key="dep_run_btn"):
        with st.spinner("재계산 중..."):
            def gc(k): return k if k!=SKIP and k in df.columns else None
            mp=pd.DataFrame()
            mp["_자산명"]=df[mn].astype(str); mp["_취득일"]=df[ma].apply(parse_dt)
            mp["_취득원가"]=pd.to_numeric(df[mc],errors='coerce').fillna(0).astype(int)
            ulr=pd.to_numeric(df[mu],errors='coerce').fillna(0)
            mp["_내용연수월"]=(ulr*12).astype(int) if lu=="년" else ulr.astype(int)
            mp["_잔존가치"]=pd.to_numeric(df[gc(mr)],errors='coerce').fillna(0).astype(int) if gc(mr) else 0
            mp["_상각방법"]=df[gc(mm)].astype(str) if gc(mm) else "정액법"
            mp["_자본적지출일"]=df[gc(mcd)].apply(parse_dt) if gc(mcd) else None
            mp["_자본적지출액"]=pd.to_numeric(df[gc(mca)],errors='coerce').fillna(0).astype(int) if gc(mca) else 0
            hco=gc(mco) is not None
            mp["_회사감가상각비"]=pd.to_numeric(df[gc(mco)],errors='coerce').fillna(0).astype(int) if hco else 0
            for k in ["상각대상","월상각비","전기말누계","당기상각비","기준일누계","장부금액","CAPEX"]:
                mp[k]=[calc_dep(r,bdt)[k] for _,r in mp.iterrows()]
            st.session_state["dep_result"]=(mp,bdt,hco)
    if "dep_result" in st.session_state:
        mp,bd,hc=st.session_state["dep_result"]; st.divider(); st.subheader("📋 결과")
        c1,c2,c3=st.columns(3)
        c1.metric("자산 건수",f"{len(mp)}건"); c2.metric("당기 감가상각비",f"{mp['당기상각비'].sum():,.0f}"); c3.metric("기준일 장부금액",f"{mp['장부금액'].sum():,.0f}")
        if hc:
            tr=mp["당기상각비"].sum();tc2=mp["_회사감가상각비"].sum();d=tr-tc2
            st.info(f"📊 차이: 재계산 {tr:,.0f} - 회사 {tc2:,.0f} = **{d:,.0f}** ({'과소' if d>0 else '과대' if d<0 else '일치'})")
        disp=["_자산명","_취득일","_취득원가","_내용연수월","_상각방법","_자본적지출액","전기말누계","당기상각비","기준일누계","장부금액"]
        if hc: disp.append("_회사감가상각비")
        st.dataframe(mp[disp].style.format({k:"{:,.0f}" for k in ["_취득원가","_자본적지출액","전기말누계","당기상각비","기준일누계","장부금액","_회사감가상각비"]},na_rep="-"),use_container_width=True,height=400)
        buf=gen_dep_excel(mp,bd,hc)
        st.download_button("📥 엑셀 다운로드",data=buf,file_name=f"감가상각_재계산_{bd}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",type="primary",use_container_width=True)
else:
    st.info("👆 유형자산 원장 파일을 업로드하면 시작됩니다.")

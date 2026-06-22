"""리스 사용권자산 · 리스부채 재계산"""
import streamlit as st
import pandas as pd
from datetime import date, datetime, timedelta
from dateutil.relativedelta import relativedelta
from io import BytesIO
import calendar, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="리스 재계산", page_icon="🏢", layout="wide")
st.title("🏢 리스 사용권자산 · 리스부채 재계산")
st.caption("K-IFRS 16 | 리스부채 · 사용권자산 · 임차보증금 · 현재가치할인차금 · 복구충당부채")

HF=PatternFill("solid",fgColor="002060"); HN=Font(name="맑은 고딕",bold=True,color="FFFFFF",size=10)
TF=Font(name="맑은 고딕",bold=True,size=14,color="002060"); SF=Font(name="맑은 고딕",bold=True,size=11,color="002060")
SecF=PatternFill("solid",fgColor="D6DCE4"); SecN=Font(name="맑은 고딕",bold=True,size=10,color="002060")
NF=Font(name="맑은 고딕",size=10); IF=Font(name="맑은 고딕",size=10,color="0000FF")
BF=Font(name="맑은 고딕",size=10,bold=True); HL=PatternFill("solid",fgColor="FFF2CC")
DP=PatternFill("solid",fgColor="E8DAEF"); RF=PatternFill("solid",fgColor="E2EFDA")
NM='#,##0'; DT='YYYY-MM-DD'; PC='0.00%'
tb=Border(left=Side(style='thin',color='B0B0B0'),right=Side(style='thin',color='B0B0B0'),top=Side(style='thin',color='B0B0B0'),bottom=Side(style='thin',color='B0B0B0'))

def shdr(ws,r,cols):
    for c in range(1,cols+1):
        cl=ws.cell(row=r,column=c); cl.font=HN; cl.fill=HF; cl.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); cl.border=tb

def scell(ws,r,c,fmt=None,font=None,fill=None,align_h='right'):
    cl=ws.cell(row=r,column=c); cl.font=font or NF; cl.border=tb; cl.alignment=Alignment(horizontal=align_h,vertical='center')
    if fmt: cl.number_format=fmt
    if fill: cl.fill=fill

def parse_dt(val):
    if pd.isna(val) or val is None: return None
    if isinstance(val,datetime): return val.date()
    if isinstance(val,date): return val
    s=str(val).strip()
    for f in ('%Y-%m-%d','%Y/%m/%d','%Y.%m.%d','%Y%m%d'):
        try: return datetime.strptime(s,f).date()
        except: continue
    return None

def calc_lease(rd):
    start=rd["리스시작일"]; end=rd["리스종료일"]
    ar=rd["할인율"]/100; mr=ar/12
    pd_day=int(rd.get("월지급일",0) or 0) or start.day
    rent=int(rd.get("월임대료",0) or 0); dep=int(rd.get("보증금",0) or 0)
    idc=int(rd.get("초기직접원가",0) or 0); rn=int(rd.get("복구원가",0) or 0)
    er=(rd.get("인상률",0) or 0)/100; ep=int(rd.get("인상주기",12) or 12)
    ul=int(rd.get("내용연수",0) or 0); bd=rd["기준일"]
    prepaid=str(rd.get("지급방식","후불")).strip()=="선불"
    def mkpd(y,m,d): return date(y,m,min(d,calendar.monthrange(y,m)[1]))
    pays=[]; mi=0; cr=rent
    while True:
        y=start.year+(start.month-1+mi)//12; m=(start.month-1+mi)%12+1; p=mkpd(y,m,pd_day)
        if p>end: break
        if mi>0 and er>0 and mi%ep==0: cr*=(1+er)
        pays.append({"idx":mi,"date":p,"rent":round(cr)}); mi+=1
    tm=len(pays) if pays else max(1,(end.year-start.year)*12+end.month-start.month)
    pv=sum(p["rent"]/((1+mr)**(p["idx"]+(0 if prepaid else 1))) for p in pays)
    ll=round(pv); dpv=round(dep/((1+mr)**tm)) if dep>0 else 0; dd=dep-dpv
    rpv=round(rn/((1+ar)**(tm/12))) if rn>0 else 0
    rou=ll+dd+idc+rpv; dm=min(ul,tm) if ul>0 else tm; mdd=round(rou/dm) if dm>0 else 0
    sched=[]; lb=ll; rb=rou
    for i,p in enumerate(pays):
        intr=round(lb*mr); pri=p["rent"]-intr; d=(mdd if i<dm-1 else rb) if i<dm else 0
        sched.append({"회차":i+1,"일자":p["date"],"월임대료":p["rent"],"이자비용":intr,"원금상환":pri,
            "리스부채(기초)":lb,"리스부채(기말)":max(0,round(lb-pri)),"감가상각비":d,
            "사용권자산(기초)":rb,"사용권자산(기말)":max(0,round(rb-d))})
        lb=max(0,round(lb-pri)); rb=max(0,round(rb-d))
    fys=date(bd.year,1,1)
    fi=sum(s["이자비용"] for s in sched if fys<=s["일자"]<=bd)
    fd=sum(s["감가상각비"] for s in sched if fys<=s["일자"]<=bd)
    ba=None
    for s in sched:
        if s["일자"]<=bd: ba=s
    lat=ba["리스부채(기말)"] if ba else 0; oyl=bd+relativedelta(years=1)
    cl=min(sum(s["원금상환"] for s in sched if bd<s["일자"]<=oyl),lat)
    ds=[]
    if dd>0:
        b=dpv
        for yr in range(start.year,end.year+1):
            bb=b
            for mo in range(12):
                md2=date(yr,1,1)+relativedelta(months=mo)
                if md2<start or md2>end: continue
                b+=round(b*mr)
            ds.append({"연도":yr,"기초":bb,"이자수익":round(b-bb),"기말":b})
    rs=[]
    if rpv>0:
        b=rpv
        for yr in range(start.year,end.year+1):
            bb=b
            for mo in range(12):
                md2=date(yr,1,1)+relativedelta(months=mo)
                if md2<start or md2>end: continue
                b+=round(b*mr)
            rs.append({"연도":yr,"기초":bb,"전입액":round(b-bb),"기말":b})
    fdi=next((d["이자수익"] for d in ds if d["연도"]==bd.year),0)
    fru=next((r["전입액"] for r in rs if r["연도"]==bd.year),0)
    fde=next((d["기말"] for d in ds if d["연도"]==bd.year),dpv)
    return {"ll":ll,"dpv":dpv,"dd":dd,"rpv":rpv,"rou":rou,"md":mdd,"dm":dm,"tm":tm,
        "sched":sched,"ds":ds,"rs":rs,"fi":fi,"fd":fd,"fdi":fdi,"fru":fru,"fde":fde,
        "ba":ba,"cl":cl,"ncl":lat-cl,"dep":dep,"rn":rn,"idc":idc}

def gen_excel(data_list):
    wb=Workbook(); ws=wb.active; ws.title="리스 재계산"
    ws.merge_cells('A1:J1'); ws.cell(row=1,column=1,value="K-IFRS 16 리스 재계산표").font=TF
    ws.cell(row=1,column=1).alignment=Alignment(horizontal='center',vertical='center'); ws.row_dimensions[1].height=35
    row=3
    for ld,r in data_list:
        bd=ld["기준일"]; fy=bd.year
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=10)
        ws.cell(row=row,column=1,value=f"■ {ld['자산명']}").font=SF
        for c in range(1,11): ws.cell(row=row,column=c).fill=PatternFill("solid",fgColor="E8EFF7")
        row+=1
        def sec(t,cols=4):
            nonlocal row; ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=cols)
            ws.cell(row=row,column=1,value=t).font=SecN
            for c in range(1,cols+1): ws.cell(row=row,column=c).fill=SecF
            row+=1
        def ir(l,v,fmt=None,font=None):
            nonlocal row; ws.cell(row=row,column=1,value=l).font=NF; ws.cell(row=row,column=1).alignment=Alignment(horizontal='left'); ws.cell(row=row,column=1).border=tb
            cl=ws.cell(row=row,column=2,value=v); cl.font=font or IF; cl.border=tb; cl.alignment=Alignment(horizontal='right')
            if fmt: cl.number_format=fmt
            row+=1
        sec("1. 기본 정보")
        pday=int(ld.get("월지급일",0) or 0) or ld["리스시작일"].day
        for l,v,f in [("리스시작일",ld["리스시작일"],DT),("리스종료일",ld["리스종료일"],DT),("리스기간(월)",r["tm"],NM),
            ("월임대료",int(ld.get("월임대료",0) or 0),NM),("보증금(명목)",r["dep"],NM),("할인율(연)",ld["할인율"]/100,PC),
            ("지급방식",ld.get("지급방식","후불"),None),("월지급일",f"매월 {pday}일",None),("기준일",bd,DT)]: ir(l,v,f)
        row+=1
        sec("2. 초기 측정")
        for l,v in [("리스부채(리스료PV)",r["ll"]),("보증금PV",r["dpv"]),("보증금할인차금(현할차)",r["dd"]),
            ("복구충당부채PV",r["rpv"]),("",None),("리스부채",r["ll"]),("+보증금할인차금",r["dd"]),
            ("+초기직접원가",r["idc"]),("+복구충당부채PV",r["rpv"]),("사용권자산",r["rou"]),("",None),("월감가상각비",r["md"])]:
            if l=="" and v is None: row+=1; continue
            ir(l,v,NM,BF)
        sec("[개시일 분개]")
        for ci,h in enumerate(["차변","차변금액","대변","대변금액"],1):
            c=ws.cell(row=row,column=ci,value=h); c.font=HN; c.fill=HF; c.alignment=Alignment(horizontal='center',vertical='center'); c.border=tb
        row+=1
        dbs=[("사용권자산",r["rou"])]; crs=[]
        if r["dep"]>0: dbs.append(("임차보증금(명목)",r["dep"]))
        if r["ll"]>0: crs.append(("리스부채",r["ll"]))
        cash=r["dep"]+r["idc"]
        if cash>0: crs.append(("현금",cash))
        if r["dd"]>0: crs.append(("현재가치할인차금",r["dd"]))
        if r["rpv"]>0: crs.append(("복구충당부채",r["rpv"]))
        for ji in range(max(len(dbs),len(crs))):
            dn=dbs[ji][0] if ji<len(dbs) else ""; da=dbs[ji][1] if ji<len(dbs) else ""
            cn=crs[ji][0] if ji<len(crs) else ""; ca=crs[ji][1] if ji<len(crs) else ""
            for ci,v in enumerate([dn,da,cn,ca],1):
                c=ws.cell(row=row,column=ci,value=v); c.border=tb; c.font=NF
                c.alignment=Alignment(horizontal='left' if ci in [1,3] else 'right',vertical='center')
                if isinstance(v,(int,float)): c.number_format=NM
            row+=1
        row+=1
        ti=r["fi"]+r["fru"]
        sec(f"3. 당해연도 비용 ({fy}년)")
        for l,v in [("감가상각비",r["fd"]),("이자비용(리스부채)",r["fi"]),("전입액(복구충당부채)",r["fru"]),
            ("이자수익(현할차상각)",-r["fdi"] if r["fdi"]>0 else 0),("",None),("순비용합계",r["fd"]+ti-r["fdi"])]:
            if l=="" and v is None: row+=1; continue
            ir(l,v,NM,BF)
        row+=1
        if r["ba"]:
            sec(f"기준일({bd}) 잔액")
            for l,v in [("사용권자산",r["ba"]["사용권자산(기말)"]),("리스부채",r["ba"]["리스부채(기말)"]),("  유동",r["cl"]),("  비유동",r["ncl"])]: ir(l,v,NM,BF)
            if r["dep"]>0:
                pvdc=r["dep"]-r["fde"]
                for l,v in [("임차보증금(명목)",r["dep"]),("  현재가치할인차금(차감)",-pvdc if pvdc>0 else 0),("  순액",r["fde"])]: ir(l,v,NM,BF)
            if r["rpv"]>0:
                fre=next((x["기말"] for x in r["rs"] if x["연도"]==fy),r["rpv"]); ir("복구충당부채",fre,NM,BF)
        row+=1
        sec("4. 월별 상각 스케줄",10)
        for ci,h in enumerate(["회차","일자","월임대료","이자비용","원금상환","리스부채(기초)","리스부채(기말)","감가상각비","ROU(기초)","ROU(기말)"],1):
            ws.cell(row=row,column=ci,value=h)
        shdr(ws,row,10); row+=1
        for s in r["sched"]:
            for ci,k in enumerate(["회차","일자","월임대료","이자비용","원금상환","리스부채(기초)","리스부채(기말)","감가상각비","사용권자산(기초)","사용권자산(기말)"],1):
                ws.cell(row=row,column=ci,value=s[k]); fmt=DT if ci==2 else (None if ci==1 else NM)
                scell(ws,row,ci,fmt=fmt,align_h='center' if ci<=2 else 'right')
            if s["일자"].year==fy and s["일자"]<=bd:
                for ci in range(1,11): ws.cell(row=row,column=ci).fill=HL
            row+=1
        row+=1
        if r["ds"]:
            sec("5. 현재가치할인차금 스케줄",7)
            for ci,h in enumerate(["연도","보증금순액(기초)","이자수익","보증금순액(기말)","현할차(기초)","당기상각","현할차(기말)"],1):
                ws.cell(row=row,column=ci,value=h)
            shdr(ws,row,7); row+=1
            for ds in r["ds"]:
                pb=r["dep"]-ds["기초"]; pe=max(0,r["dep"]-ds["기말"])
                for ci,v in enumerate([ds["연도"],ds["기초"],ds["이자수익"],ds["기말"],pb,ds["이자수익"],pe],1):
                    ws.cell(row=row,column=ci,value=v); scell(ws,row,ci,fmt=None if ci==1 else NM,align_h='center' if ci==1 else 'right')
                if ds["연도"]==fy:
                    for ci in range(1,8): ws.cell(row=row,column=ci).fill=DP
                row+=1
            row+=1
        if r["rs"]:
            sec("6. 복구충당부채 스케줄",4)
            for ci,h in enumerate(["연도","기초","전입액","기말"],1): ws.cell(row=row,column=ci,value=h)
            shdr(ws,row,4); row+=1
            for rs2 in r["rs"]:
                for ci,v in enumerate([rs2["연도"],rs2["기초"],rs2["전입액"],rs2["기말"]],1):
                    ws.cell(row=row,column=ci,value=v); scell(ws,row,ci,fmt=None if ci==1 else NM,align_h='center' if ci==1 else 'right')
                if rs2["연도"]==fy:
                    for ci in range(1,5): ws.cell(row=row,column=ci).fill=RF
                row+=1
            row+=1
        row+=2
    for ci,w in enumerate([35,18,16,16,16,18,18,16,20,20],1): ws.column_dimensions[get_column_letter(ci)].width=w
    buf=BytesIO(); wb.save(buf); buf.seek(0); return buf

# ── UI ──
with st.expander("📖 사용 안내",expanded=False):
    st.markdown("화면에서 직접 입력하거나 기존 리스 입력 템플릿 엑셀을 업로드할 수 있습니다.")

method=st.radio("입력 방식",["직접 입력","엑셀 업로드"],horizontal=True,key="lm")
if method=="직접 입력":
    n=st.number_input("리스자산 건수",1,50,2,key="nl"); bdt=st.date_input("기준일",value=date(2025,12,31),key="lb")
    rows=[]
    for i in range(n):
        st.markdown(f"**리스자산 {i+1}**")
        c=st.columns([2,1,1,1,1,1,1,1])
        nm=c[0].text_input("자산명",f"자산{i+1}",key=f"n{i}"); s=c[1].date_input("시작",date(2024,1,1),key=f"s{i}")
        e=c[2].date_input("종료",date(2028,12,31),key=f"e{i}"); r=c[3].number_input("월임대료",value=5000000,step=100000,key=f"r{i}")
        d=c[4].number_input("보증금",value=50000000,step=1000000,key=f"d{i}"); rt=c[5].number_input("할인율%",value=4.5,step=0.1,key=f"rt{i}")
        ul=c[6].number_input("내용연수(월)",value=60,step=1,key=f"u{i}"); pd2=c[7].number_input("지급일",value=1,min_value=1,max_value=31,key=f"p{i}")
        with st.expander(f"자산 {i+1} 추가옵션"):
            c2=st.columns(4); tm=c2[0].selectbox("지급방식",["후불","선불"],key=f"t{i}")
            esc=c2[1].number_input("인상률%",0.0,key=f"esc{i}"); idc2=c2[2].number_input("초기직접원가",0,key=f"idc{i}")
            rest=c2[3].number_input("복구원가(명목)",0,key=f"rest{i}")
        rows.append({"자산명":nm,"리스시작일":s,"리스종료일":e,"월임대료":r,"보증금":d,"할인율":rt,"지급방식":tm,"월지급일":pd2,"인상률":esc,"인상주기":12,"초기직접원가":idc2,"복구원가":rest,"내용연수":ul,"기준일":bdt})
    if st.button("🔢 재계산 실행",type="primary",use_container_width=True,key="lr_btn"):
        with st.spinner("재계산 중..."): st.session_state["lease_result"]=[(ld,calc_lease(ld)) for ld in rows]
else:
    up=st.file_uploader("리스 입력 템플릿",type=["xlsx"],key="lu")
    if up:
        wb=openpyxl.load_workbook(up,data_only=True); ws=wb["리스정보입력"]; rows=[]
        for row in ws.iter_rows(min_row=4,max_col=15,values_only=True):
            if not row[0]: continue
            nm=str(row[0]).strip()
            if not nm: continue
            rows.append({"자산명":nm,"리스시작일":parse_dt(row[1]),"리스종료일":parse_dt(row[2]),
                "월임대료":int(float(row[3] or 0)),"보증금":int(float(row[4] or 0)),"할인율":float(row[5] or 0),
                "지급방식":str(row[6] or "후불"),"월지급일":int(float(row[7] or 0)) if row[7] else 0,
                "인상률":float(row[8] or 0),"인상주기":int(float(row[9] or 12)),"초기직접원가":int(float(row[10] or 0)),
                "복구원가":int(float(row[11] or 0)),"내용연수":int(float(row[12] or 0)),"기준일":parse_dt(row[13])})
        wb.close(); st.success(f"✅ {len(rows)}건 로드")
        if st.button("🔢 재계산 실행",type="primary",use_container_width=True,key="lr_btn2"):
            with st.spinner("재계산 중..."): st.session_state["lease_result"]=[(ld,calc_lease(ld)) for ld in rows]

if "lease_result" in st.session_state:
    st.divider(); st.subheader("📋 결과")
    for ld,r in st.session_state["lease_result"]:
        with st.expander(f"**{ld['자산명']}** — 리스부채 {r['ll']:,} | ROU {r['rou']:,}",expanded=True):
            c1,c2,c3,c4=st.columns(4)
            c1.metric("리스부채",f"{r['ll']:,}"); c2.metric("사용권자산",f"{r['rou']:,}")
            c3.metric("당기 감가상각비",f"{r['fd']:,}"); c4.metric("당기 이자비용",f"{r['fi']:,}")
            if r["dd"]>0:
                c5,c6=st.columns(2); c5.metric("보증금할인차금",f"{r['dd']:,}"); c6.metric("이자수익(현할차)",f"{r['fdi']:,}")
            st.dataframe(pd.DataFrame(r["sched"]).style.format({k:"{:,.0f}" for k in ["월임대료","이자비용","원금상환","리스부채(기초)","리스부채(기말)","감가상각비","사용권자산(기초)","사용권자산(기말)"]}),use_container_width=True,height=300)
    buf=gen_excel(st.session_state["lease_result"])
    st.download_button("📥 엑셀 다운로드",data=buf,file_name="리스_재계산_결과.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",type="primary",use_container_width=True)

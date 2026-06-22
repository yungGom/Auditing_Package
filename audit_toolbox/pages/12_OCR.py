"""
============================================================
Audit Toolbox - 의사록 OCR 정리 (페이지)
============================================================
Streamlit 멀티페이지 - 의사록 OCR 통합 도구
- 탭1: PDF 파일명 일괄 수정
- 탭2: OCR 실행 → 엑셀 생성/다운로드

외부 API/웹 연동 없음 (로컬 처리)

[필요 패키지]
streamlit, easyocr, pdf2image, openpyxl, Pillow, tqdm

[추가 설치 (PDF 처리용)]
Windows: winget install oschwartz10612.Poppler

[보안 메모]
- OCR 중간 산출물(PDF→PNG 변환 이미지)은 OS 임시 폴더에 생성되며
  처리 종료 시(예외 발생 포함) 자동 삭제됩니다.
- 탭1 파일명 변경 시 경로 구분자/위험 문자를 제거하여
  의도치 않은 폴더 이동을 방지합니다.
"""

import os
import re
import shutil
import tempfile
import warnings
from datetime import datetime
from pathlib import Path
from io import BytesIO

import streamlit as st
import pandas as pd

warnings.filterwarnings("ignore", message=".*pin_memory.*")

# ============================================================
# 설정
# ============================================================
CONFIG = {
    "input_dir": r"",
    "output_dir": r"",

    "languages": ["ko", "en"],
    "gpu": False,
    "dpi": 300,
    "poppler_path": None,
}

SUPPORTED_EXT = {".pdf", ".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp"}


# ============================================================
# Streamlit 페이지 설정
# ============================================================
st.set_page_config(page_title="의사록 OCR 정리", page_icon="📋", layout="wide")
st.title("📋 의사록 OCR 정리")
st.caption("PDF 파일명 수정 → OCR → 엑셀 정리 (로컬 처리, 외부 전송 없음)")


# ============================================================
# 의존성 체크
# ============================================================
@st.cache_resource
def check_deps():
    """필요 패키지 확인"""
    missing = []
    try:
        import easyocr  # noqa: F401
    except ImportError:
        missing.append("easyocr")
    try:
        from pdf2image import convert_from_path  # noqa: F401
    except ImportError:
        missing.append("pdf2image")
    try:
        from openpyxl import Workbook  # noqa: F401
    except ImportError:
        missing.append("openpyxl")
    try:
        from PIL import Image  # noqa: F401
    except ImportError:
        missing.append("Pillow")

    if not shutil.which("pdftoppm"):
        missing.append("poppler (winget install oschwartz10612.Poppler)")

    return missing


missing_deps = check_deps()
if missing_deps:
    st.error(f"다음 패키지가 설치되지 않았습니다: {', '.join(missing_deps)}")
    st.code("pip install easyocr pdf2image openpyxl Pillow tqdm", language="bash")
    st.info(
        "Poppler가 누락된 경우 명령 프롬프트에서 다음을 실행하세요:\n"
        "winget install oschwartz10612.Poppler"
    )
    st.stop()

import easyocr
from pdf2image import convert_from_path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from PIL import Image  # noqa: F401  (PIL은 pdf2image가 내부에서 사용)


# ============================================================
# 유틸리티 함수
# ============================================================
def get_files_in_dir(input_dir):
    """폴더에서 지원 파일 목록 반환"""
    input_path = Path(input_dir)
    if not input_path.exists():
        input_path.mkdir(parents=True, exist_ok=True)
        return []
    files = []
    for f in sorted(input_path.iterdir()):
        if f.suffix.lower() in SUPPORTED_EXT:
            files.append(f)
    return files


def sanitize_filename(name):
    """
    파일명에서 경로 구분자 및 위험 문자를 제거한다.
    - '..\\..\\foo.pdf' 또는 '../bar/foo.pdf' 같은 입력이 들어와도
      디렉터리 부분을 모두 떼어내고 순수 파일명만 남긴다.
    - Windows 금지 문자(< > : " | ? *)는 '_'로 치환한다.
    - 앞뒤 점/공백을 제거해 '..' 같은 위험 입력을 무력화한다.
    """
    name = str(name).strip()
    # 역슬래시를 슬래시로 통일한 뒤 마지막 경로 요소만 취함 → 디렉터리 부분 제거
    name = name.replace("\\", "/").split("/")[-1]
    # Windows 파일명 금지 문자 치환
    name = re.sub(r'[<>:"|?*]', "_", name)
    # 앞뒤 점/공백 제거 ('..', '.' 등 무력화)
    name = name.strip(". ")
    return name


def parse_filename(filename):
    """
    파일명에서 회차/유형/안건 파싱
    규칙: 2026-01-주총의사록-정기주주총회.pdf
          2026-01-이사회의사록-배당결의.pdf
          2025-01_자기주식취득결정.pdf (기존 형식도 호환)
    """
    stem = Path(filename).stem

    # 새 형식: 2026-01-주총의사록-제목 또는 2026-01-이사회의사록-제목
    match_new = re.match(
        r"^(\d{4}-?\d{0,2})[-_\s]*(주총의사록|이사회의사록|주주총회의사록)[-_\s]*(.*)",
        stem
    )
    if match_new:
        return {
            "회차": match_new.group(1).strip(),
            "유형": match_new.group(2).strip(),
            "안건명": match_new.group(3).strip(),
        }

    # 기존 형식: 2025-01_자기주식취득결정
    match_old = re.match(r"^(\d{4}-?\d{0,2})[_\s]*(.*)", stem)
    if match_old:
        return {
            "회차": match_old.group(1).strip(),
            "유형": "",  # 유형 없음 → OCR 원문에서도 판단하지 않고 빈칸
            "안건명": match_old.group(2).strip(),
        }

    return {"회차": stem, "유형": "", "안건명": ""}


def classify_by_filename(filename):
    """파일명 기반 유형 분류 (OCR 내용 무시)"""
    info = parse_filename(filename)
    doc_type = info["유형"]
    if "주총" in doc_type or "주주총회" in doc_type:
        return "주주총회의사록"
    elif "이사회" in doc_type:
        return "이사회의사록"
    return "의사록(유형미상)"


# ============================================================
# OCR 엔진 (캐시)
# ============================================================
@st.cache_resource
def load_ocr_reader():
    """EasyOCR 리더 초기화 (앱 최초 1회)"""
    return easyocr.Reader(CONFIG["languages"], gpu=CONFIG["gpu"], verbose=False)


# ============================================================
# OCR 처리 함수
# ============================================================
def pdf_to_images(pdf_path, temp_dir, file_idx):
    """PDF → 이미지 변환 (영문 임시경로)"""
    temp_path = Path(temp_dir)
    temp_path.mkdir(parents=True, exist_ok=True)

    kwargs = {"dpi": CONFIG["dpi"]}
    if CONFIG["poppler_path"]:
        kwargs["poppler_path"] = CONFIG["poppler_path"]

    try:
        images = convert_from_path(str(pdf_path), **kwargs)
    except Exception as e:
        return [], str(e)

    image_paths = []
    for i, img in enumerate(images):
        img_path = temp_path / f"doc{file_idx:03d}_page_{i+1:03d}.png"
        img.save(str(img_path), "PNG")
        image_paths.append(img_path)

    return image_paths, None


def ocr_image(reader, image_path):
    """이미지 OCR"""
    try:
        results = reader.readtext(str(image_path))
        results.sort(key=lambda x: (x[0][0][1], x[0][0][0]))
        return results, None
    except Exception as e:
        return [], str(e)


def merge_ocr_lines(ocr_results, line_threshold=15):
    """OCR 결과 줄 병합"""
    if not ocr_results:
        return []
    lines = []
    current_line = []
    current_y = ocr_results[0][0][0][1]

    for bbox, text, conf in ocr_results:
        y = bbox[0][1]
        if abs(y - current_y) > line_threshold:
            if current_line:
                current_line.sort(key=lambda x: x[0])
                lines.append(" ".join([t for _, t in current_line]).strip())
            current_line = [(bbox[0][0], text)]
            current_y = y
        else:
            current_line.append((bbox[0][0], text))

    if current_line:
        current_line.sort(key=lambda x: x[0])
        lines.append(" ".join([t for _, t in current_line]).strip())
    return lines


def extract_date_from_text(full_text):
    """텍스트에서 일시 추출"""
    date_patterns = [
        r"\d{4}\s*[년.]\s*\d{1,2}\s*[월.]\s*\d{1,2}\s*일?[^~\n]*(?:오전|오후)?\s*\d{0,2}\s*시?\s*\d{0,2}\s*분?",
        r"\d{4}\s*[년.]\s*\d{1,2}\s*[월.]\s*\d{1,2}\s*일?",
        r"\d{4}[./\-]\d{1,2}[./\-]\d{1,2}",
    ]
    for line in full_text.split("\n"):
        line_s = line.strip()
        if not line_s:
            continue
        for pattern in date_patterns:
            match = re.search(pattern, line_s)
            if match:
                return match.group(0).strip()
    return ""


def process_file(reader, file_path, file_idx, temp_dir):
    """
    파일 1개 OCR 처리
    temp_dir: 호출자가 관리하는 임시 폴더 경로.
              (탭2에서 tempfile.TemporaryDirectory()로 생성 → 처리 후 자동 삭제)
    """
    file_path = Path(file_path)
    pages_text = []

    if file_path.suffix.lower() == ".pdf":
        image_paths, err = pdf_to_images(file_path, temp_dir, file_idx)
        if err:
            return None, f"PDF 변환 실패: {err}"
        if not image_paths:
            return None, "PDF 변환 결과 없음"

        for img_path in image_paths:
            page_num = int(img_path.stem.split("_")[-1])
            results, err = ocr_image(reader, img_path)
            lines = merge_ocr_lines(results)
            pages_text.append({
                "page": page_num,
                "full_text": "\n".join(lines),
                "text_count": len("".join(lines)),
            })
    else:
        # 이미지: 영문 임시경로 복사
        temp_path = Path(temp_dir)
        temp_path.mkdir(parents=True, exist_ok=True)
        safe_path = temp_path / f"doc{file_idx:03d}_img{file_path.suffix.lower()}"
        shutil.copy2(str(file_path), str(safe_path))

        results, err = ocr_image(reader, safe_path)
        lines = merge_ocr_lines(results)
        pages_text.append({
            "page": 1,
            "full_text": "\n".join(lines),
            "text_count": len("".join(lines)),
        })

    if not pages_text:
        return None, "텍스트 추출 실패"

    full_text = "\n".join([p["full_text"] for p in pages_text])
    total_chars = sum(p["text_count"] for p in pages_text)

    # 파일명 기반 분류 (OCR 내용으로 판단하지 않음)
    file_info = parse_filename(file_path.name)
    doc_type = classify_by_filename(file_path.name)

    result = {
        "filename": file_path.name,
        "pages_text": pages_text,
        "parsed_info": {
            "유형": doc_type,
            "회차": file_info["회차"],
            "안건명": file_info["안건명"],
            "일시": extract_date_from_text(full_text),
            "원문": full_text.strip(),
        },
        "total_chars": total_chars,
    }

    return result, None


# ============================================================
# 엑셀 생성
# ============================================================
def create_excel(all_results):
    """결과를 엑셀로 생성, BytesIO 반환"""
    wb = Workbook()
    ws = wb.active
    ws.title = "의사록 검토"

    # 스타일
    title_font = Font(name="맑은 고딕", bold=True, size=12)
    header_font = Font(name="맑은 고딕", bold=True, size=10)
    header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    normal_font = Font(name="맑은 고딕", size=9)
    normal_align = Alignment(vertical="top", wrap_text=True, horizontal="left")
    center_align = Alignment(vertical="center", wrap_text=True, horizontal="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # 열 너비
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 30
    for col_letter in ["E", "F", "G", "H", "I", "J", "K", "L", "M"]:
        ws.column_dimensions[col_letter].width = 8
    ws.column_dimensions["N"].width = 12
    ws.column_dimensions["O"].width = 14

    # 분류
    shareholder_items = [r for r in all_results if "주주총회" in r["parsed_info"]["유형"]]
    board_items = [r for r in all_results if "이사회" in r["parsed_info"]["유형"]]
    other_items = [r for r in all_results if r not in shareholder_items and r not in board_items]

    ROWS_PER_ENTRY = 8

    def apply_border_range(ws, sr, er, sc, ec):
        for r in range(sr, er + 1):
            for c in range(sc, ec + 1):
                ws.cell(row=r, column=c).border = thin_border

    def write_section_header(ws, row, title):
        ws.cell(row=row, column=2, value=title).font = title_font
        h_row = row + 2
        for col_letter, text in {"B": "회차", "C": "일자", "D": "장소 및 출석인원",
                                  "E": "안건", "N": "승인여부", "O": "관련 재무제표"}.items():
            cell = ws.cell(row=h_row, column=ord(col_letter) - ord("A") + 1, value=text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        ws.merge_cells(f"E{h_row}:M{h_row}")
        for c in range(5, 14):
            ws.cell(row=h_row, column=c).border = thin_border
        return h_row + 1

    def write_entry(ws, sr, result, rows=ROWS_PER_ENTRY):
        info = result["parsed_info"]
        er = sr + rows - 1

        ws.cell(row=sr, column=2, value=info["회차"]).font = normal_font
        ws.cell(row=sr, column=2).alignment = center_align
        ws.merge_cells(f"B{sr}:B{er}")

        ws.cell(row=sr, column=3, value=info["일시"]).font = normal_font
        ws.cell(row=sr, column=3).alignment = center_align
        ws.merge_cells(f"C{sr}:C{er}")

        ws.cell(row=sr, column=4, value="").font = normal_font
        ws.merge_cells(f"D{sr}:D{er}")

        ws.cell(row=sr, column=5, value=info["원문"]).font = normal_font
        ws.cell(row=sr, column=5).alignment = normal_align
        ws.merge_cells(f"E{sr}:M{er}")

        ws.cell(row=sr, column=14, value="").font = normal_font
        ws.cell(row=sr, column=14).alignment = center_align
        ws.merge_cells(f"N{sr}:N{er}")

        ws.cell(row=sr, column=15, value="").font = normal_font
        ws.cell(row=sr, column=15).alignment = center_align

        apply_border_range(ws, sr, er, 2, 15)
        return er + 1

    def write_empty(ws, sr, rows=ROWS_PER_ENTRY):
        er = sr + rows - 1
        ws.merge_cells(f"E{sr}:M{er}")
        apply_border_range(ws, sr, er, 2, 15)
        return er + 1

    # 섹션 작성
    cur = 2

    ds = write_section_header(ws, cur, "2. 주총의사록의 내용 검토")
    if shareholder_items:
        for r in shareholder_items:
            ds = write_entry(ws, ds, r)
    else:
        ds = write_empty(ws, ds)
    cur = ds + 2

    ds = write_section_header(ws, cur, "3. 이사회의사록의 내용 검토")
    if board_items:
        for r in board_items:
            ds = write_entry(ws, ds, r)
    else:
        ds = write_empty(ws, ds)
    cur = ds + 2

    if other_items:
        ds = write_section_header(ws, cur, "4. 기타 의사록")
        for r in other_items:
            ds = write_entry(ws, ds, r)

    # 원문 시트
    sub_font = Font(name="맑은 고딕", bold=True, size=10)
    sub_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

    for idx, result in enumerate(all_results, 1):
        sn = f"{idx}_{result['filename'][:25]}"
        sn = re.sub(r'[\\/*?\[\]:]', '_', sn)[:31]
        ws_raw = wb.create_sheet(title=sn)
        ws_raw.column_dimensions["A"].width = 10
        ws_raw.column_dimensions["B"].width = 100

        ws_raw.cell(row=1, column=1, value="파일명").font = sub_font
        ws_raw.cell(row=1, column=1).fill = sub_fill
        ws_raw.cell(row=1, column=2, value=result["filename"]).font = normal_font

        ws_raw.cell(row=2, column=1, value="유형").font = sub_font
        ws_raw.cell(row=2, column=1).fill = sub_fill
        ws_raw.cell(row=2, column=2, value=result["parsed_info"]["유형"]).font = normal_font

        ws_raw.cell(row=3, column=1, value="일시").font = sub_font
        ws_raw.cell(row=3, column=1).fill = sub_fill
        ws_raw.cell(row=3, column=2, value=result["parsed_info"]["일시"]).font = normal_font

        ws_raw.cell(row=5, column=1, value="페이지").font = sub_font
        ws_raw.cell(row=5, column=1).fill = sub_fill
        ws_raw.cell(row=5, column=2, value="OCR 추출 원문").font = sub_font
        ws_raw.cell(row=5, column=2).fill = sub_fill

        row = 6
        for pd_item in result["pages_text"]:
            ws_raw.cell(row=row, column=1, value=f"p.{pd_item['page']}").font = normal_font
            ws_raw.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="top")
            ws_raw.cell(row=row, column=2, value=pd_item["full_text"]).font = normal_font
            ws_raw.cell(row=row, column=2).alignment = normal_align
            row += 1

    # BytesIO로 반환
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============================================================
# 탭 구성
# ============================================================
tab1, tab2 = st.tabs(["① 파일명 수정", "② OCR 실행"])

# ============================================================
# 탭1: 파일명 수정
# ============================================================
with tab1:
    st.subheader("📝 PDF 파일명 일괄 수정")
    st.markdown(
        "원본 폴더의 PDF 파일명을 수정합니다.  \n"
        "**파일명 규칙**: `2026-01-주총의사록-제목` 또는 `2026-01-이사회의사록-제목`  \n"
        "이 규칙대로 작성하면 탭②에서 자동으로 주총/이사회 분류됩니다."
    )

    # 폴더 경로 표시/수정
    input_dir = st.text_input(
        "원본 폴더 경로",
        value=CONFIG["input_dir"],
        key="input_dir_tab1",
    )
    CONFIG["input_dir"] = input_dir

    if st.button("📂 파일 목록 새로고침", key="refresh_tab1"):
        st.rerun()

    files = get_files_in_dir(CONFIG["input_dir"])

    if not files:
        st.warning(f"'{CONFIG['input_dir']}' 폴더에 파일이 없습니다.")
        st.info("지원 형식: PDF, PNG, JPG, JPEG, TIFF, BMP")
    else:
        st.success(f"{len(files)}개 파일 발견")

        # 파일 목록 데이터프레임
        df = pd.DataFrame({
            "BEFORE (현재 파일명)": [f.name for f in files],
            "AFTER (변경할 파일명)": [f.name for f in files],
        })

        st.markdown("**AFTER 열을 편집하세요** (엑셀에서 Ctrl+V 붙여넣기 가능)")
        edited = st.data_editor(
            df,
            use_container_width=True,
            num_rows="fixed",
            key="rename_editor",
            column_config={
                "BEFORE (현재 파일명)": st.column_config.TextColumn(disabled=True),
                "AFTER (변경할 파일명)": st.column_config.TextColumn(),
            },
        )

        # 변경사항 미리보기
        # 입력값은 sanitize_filename()을 거친 결과로 비교 → 미리보기와 실제 변경값이 일치
        changes = []
        sanitized_warn = []
        for i in range(len(files)):
            before = df.iloc[i]["BEFORE (현재 파일명)"]
            raw_after = str(edited.iloc[i]["AFTER (변경할 파일명)"]).strip()
            after = sanitize_filename(raw_after)
            # 입력값이 정제 과정에서 바뀌었으면(경로 구분자/위험 문자 포함) 사용자에게 알림
            if after != raw_after:
                sanitized_warn.append((raw_after, after))
            if before != after and after:
                changes.append((i, before, after))

        if sanitized_warn:
            st.warning(
                "⚠️ 일부 입력에서 경로 구분자 또는 사용할 수 없는 문자가 제거되었습니다 "
                "(폴더 이동 방지):\n"
                + "\n".join([f"- `{raw}` → `{clean}`" for raw, clean in sanitized_warn])
            )

        if changes:
            st.markdown(f"**변경 예정: {len(changes)}건**")
            for _, before, after in changes:
                info = parse_filename(after)
                type_tag = f" → **{classify_by_filename(after)}**" if info["유형"] else ""
                st.markdown(f"- `{before}` → `{after}`{type_tag}")

        col1, col2 = st.columns(2)
        with col1:
            if st.button("✅ 파일명 변경 실행", type="primary", use_container_width=True, key="rename_run"):
                if not changes:
                    st.info("변경할 파일이 없습니다.")
                else:
                    errors = []
                    success = 0
                    for idx, before, after in changes:
                        # 확장자 보정
                        if not any(after.lower().endswith(ext) for ext in SUPPORTED_EXT):
                            ext = Path(before).suffix
                            after = after + ext

                        old_path = files[idx]
                        new_path = old_path.parent / after

                        # 안전장치: 정제 후에도 new_path가 원본 폴더를 벗어나면 건너뜀
                        try:
                            if new_path.parent.resolve() != old_path.parent.resolve():
                                errors.append(f"{before}: 원본 폴더를 벗어나는 경로 — 건너뜀")
                                continue
                            old_path.rename(new_path)
                            success += 1
                        except Exception as e:
                            errors.append(f"{before}: {e}")

                    if success:
                        st.success(f"✅ {success}건 파일명 변경 완료")
                    if errors:
                        for err in errors:
                            st.error(f"❌ {err}")
                    st.rerun()

        with col2:
            if st.button("🔄 초기화", use_container_width=True, key="rename_reset"):
                st.rerun()


# ============================================================
# 탭2: OCR 실행
# ============================================================
with tab2:
    st.subheader("🔍 OCR 실행 → 엑셀 생성")
    st.markdown(
        "원본 폴더의 파일을 OCR 처리하여 엑셀로 정리합니다.  \n"
        "**파일명에 `주총의사록` 또는 `이사회의사록`이 포함되어 있으면** 자동 분류됩니다."
    )

    # 폴더 경로
    input_dir2 = st.text_input(
        "원본 폴더 경로",
        value=CONFIG["input_dir"],
        key="input_dir_tab2",
    )
    CONFIG["input_dir"] = input_dir2

    output_dir2 = st.text_input(
        "출력 폴더 경로",
        value=CONFIG["output_dir"],
        key="output_dir_tab2",
    )
    CONFIG["output_dir"] = output_dir2

    files = get_files_in_dir(CONFIG["input_dir"])

    if not files:
        st.warning(f"'{CONFIG['input_dir']}' 폴더에 파일이 없습니다.")
        st.info("탭①에서 파일명을 먼저 수정해주세요.")
    else:
        # 파일 목록 + 파싱 결과 미리보기
        preview_data = []
        for f in files:
            info = parse_filename(f.name)
            doc_type = classify_by_filename(f.name)
            size_mb = f.stat().st_size / (1024 * 1024)
            preview_data.append({
                "파일명": f.name,
                "회차": info["회차"],
                "유형": doc_type,
                "안건명": info["안건명"],
                "크기(MB)": f"{size_mb:.1f}",
            })

        st.dataframe(
            pd.DataFrame(preview_data),
            use_container_width=True,
            hide_index=True,
        )

        # 유형미상 경고
        unknown = [p for p in preview_data if "미상" in p["유형"]]
        if unknown:
            st.warning(
                f"⚠️ {len(unknown)}개 파일의 유형을 파악할 수 없습니다. "
                "파일명에 `주총의사록` 또는 `이사회의사록`을 포함시켜주세요. "
                "(탭①에서 수정 가능)"
            )

        # OCR 실행
        if st.button("🚀 OCR 실행", type="primary", use_container_width=True, key="ocr_run"):
            os.makedirs(CONFIG["output_dir"], exist_ok=True)

            # OCR 엔진 로딩
            with st.spinner("OCR 엔진 로딩 중... (최초 실행 시 모델 다운로드로 시간이 걸립니다)"):
                reader = load_ocr_reader()

            all_results = []
            progress_bar = st.progress(0, text="OCR 처리 중...")
            status_area = st.empty()

            # 임시 이미지(PDF→PNG 변환물)는 OS 임시 폴더에 생성한다.
            # with 블록을 벗어나면 — 정상 종료든 예외 발생이든 — 폴더째 자동 삭제되어
            # 피감 의사록 이미지가 디스크에 잔존하지 않는다.
            with tempfile.TemporaryDirectory(prefix="audit_ocr_") as tmp_dir:
                total = len(files)
                for file_idx, file_path in enumerate(files, 1):
                    progress = file_idx / total
                    progress_bar.progress(progress, text=f"[{file_idx}/{total}] {file_path.name}")
                    status_area.info(f"처리 중: {file_path.name}")

                    result, err = process_file(reader, file_path, file_idx, tmp_dir)

                    if err:
                        status_area.warning(f"⚠️ {file_path.name}: {err}")
                    elif result:
                        all_results.append(result)

                progress_bar.progress(1.0, text="OCR 완료!")
            # ── tmp_dir 자동 삭제 완료 (이하 처리는 메모리상의 all_results만 사용) ──

            if all_results:
                # 결과 요약
                status_area.empty()
                st.success(f"✅ {len(all_results)}건 OCR 처리 완료")

                # 분류 결과 표시
                sh_count = sum(1 for r in all_results if "주주총회" in r["parsed_info"]["유형"])
                bd_count = sum(1 for r in all_results if "이사회" in r["parsed_info"]["유형"])
                ot_count = len(all_results) - sh_count - bd_count

                col1, col2, col3 = st.columns(3)
                col1.metric("주총의사록", f"{sh_count}건")
                col2.metric("이사회의사록", f"{bd_count}건")
                if ot_count > 0:
                    col3.metric("기타", f"{ot_count}건")

                # 상세 결과 테이블
                result_data = []
                for r in all_results:
                    result_data.append({
                        "파일명": r["filename"],
                        "유형": r["parsed_info"]["유형"],
                        "회차": r["parsed_info"]["회차"],
                        "일시": r["parsed_info"]["일시"] or "(추출 실패)",
                        "글자수": f"{r['total_chars']:,}",
                    })

                st.dataframe(
                    pd.DataFrame(result_data),
                    use_container_width=True,
                    hide_index=True,
                )

                # 엑셀 생성 + 다운로드
                with st.spinner("엑셀 파일 생성 중..."):
                    excel_buf = create_excel(all_results)

                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                excel_filename = f"의사록_OCR정리_{timestamp}.xlsx"

                # 파일로도 저장
                output_path = Path(CONFIG["output_dir"]) / excel_filename
                try:
                    with open(str(output_path), "wb") as f:
                        f.write(excel_buf.getvalue())
                    st.info(f"📁 저장됨: {output_path}")
                except Exception as e:
                    st.warning(f"파일 저장 실패: {e}")

                # 다운로드 버튼
                st.download_button(
                    "📥 엑셀 다운로드",
                    data=excel_buf.getvalue(),
                    file_name=excel_filename,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )

                # 원문 미리보기 (expander)
                with st.expander("📄 OCR 원문 미리보기 (클릭해서 펼치기)"):
                    for r in all_results:
                        st.markdown(f"**{r['filename']}** ({r['parsed_info']['유형']})")
                        st.text_area(
                            f"원문 - {r['filename']}",
                            value=r["parsed_info"]["원문"],
                            height=200,
                            key=f"preview_{r['filename']}",
                        )
                        st.divider()

            else:
                status_area.error("처리된 파일이 없습니다.")


# ============================================================
# 사이드바 정보
# ============================================================
with st.sidebar:
    st.markdown("### ⚙️ 설정")
    st.markdown(f"**원본 폴더**: `{CONFIG['input_dir']}`")
    st.markdown(f"**출력 폴더**: `{CONFIG['output_dir']}`")
    st.divider()

    st.markdown("### 📌 파일명 규칙")
    st.code("2026-01-주총의사록-제목.pdf\n2026-01-이사회의사록-제목.pdf", language=None)
    st.markdown(
        "- `2026-01`: 회차 (엑셀 회차 칸)\n"
        "- `주총의사록`/`이사회의사록`: 섹션 자동 분류\n"
        "- `제목`: 안건명 (참고용)"
    )
    st.divider()

    st.markdown("### 🔒 보안")
    st.markdown(
        "- 모든 처리는 **로컬 PC**에서만 수행\n"
        "- 외부 서버 전송 **없음**\n"
        "- 인터넷 끊고 실행해도 동작\n"
        "- (단, EasyOCR 모델 최초 다운로드 1회만 인터넷 필요)\n"
        "- OCR 중간 이미지는 OS 임시 폴더에 생성 후 **자동 삭제**"
    )

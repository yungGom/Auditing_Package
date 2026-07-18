"""단기리스 · 소액자산 면제 판단 (K-IFRS 1116호 문단 5~8, B3~B8)"""
import streamlit as st
import pandas as pd
from datetime import date
from io import BytesIO
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="면제 판단", page_icon="🏷️", layout="wide")
st.title("🏷️ 단기리스 · 소액자산 면제 판단")
st.caption("K-IFRS 1116호 문단 5~8 | 면제 가능 여부 판단 + 비용 처리액 계산")

with st.expander("📖 면제 기준 요약", expanded=False):
    st.markdown("""
    **단기리스 면제 (문단 5(a), 6~8)**
    - 리스기간 ≤ 12개월 (연장선택권 포함)
    - 매수선택권이 **없어야** 함
    - 자산 유형별로 일괄 선택 (개별 선택 불가)
    - 면제 시: 리스료를 정액 또는 다른 체계적 기준으로 비용 인식

    **소액자산 면제 (문단 5(b), B3~B8)**
    - **새 자산 기준** 가치가 소액 (IASB 기준 약 USD 5,000, 원화 약 650~700만원)
    - 다른 자산에 대한 의존성/상호관련성이 높지 않을 것
    - 개별 리스 단위로 선택 가능
    - 예시: 태블릿, 개인용 컴퓨터, 소형 사무가구, 전화기
    - 반례: 자동차 (일반적으로 소액 아님)
    """)

st.divider()
st.subheader("리스 정보 입력")

n = st.number_input("판단할 리스 건수", 1, 50, 3, key="n_exempt")
threshold = st.number_input("소액자산 기준금액 (원)", value=7000000, step=100000, format="%d",
                            help="IASB 적용사례에서 USD 5,000 제시. 원화 환산 시 약 650~700만원")

items = []
for i in range(n):
    st.markdown(f"**리스 {i+1}**")
    c = st.columns([2, 1, 1, 1, 1, 1])
    name = c[0].text_input("자산명", f"자산{i+1}", key=f"en{i}")
    term = c[1].number_input("리스기간(월)", value=12, min_value=1, key=f"et{i}")
    purchase_opt = c[2].selectbox("매수선택권", ["없음", "있음"], key=f"ep{i}")
    new_value = c[3].number_input("새자산 가치(원)", value=5000000, step=100000, key=f"ev{i}")
    interdep = c[4].selectbox("자산 상호의존성", ["낮음", "높음"], key=f"ei{i}")
    monthly_rent = c[5].number_input("월 리스료", value=500000, step=10000, key=f"er{i}")

    items.append({
        "자산명": name, "리스기간(월)": term,
        "매수선택권": purchase_opt, "새자산가치": new_value,
        "상호의존성": interdep, "월리스료": monthly_rent,
    })

base_dt = st.date_input("기준일", value=date(2025, 12, 31), key="exempt_base")

if st.button("🔢 면제 판단 실행", type="primary", use_container_width=True, key="exempt_run"):
    results = []
    for item in items:
        # 단기리스 판단
        short_term = item["리스기간(월)"] <= 12 and item["매수선택권"] == "없음"
        # 소액자산 판단
        low_value = item["새자산가치"] <= threshold and item["상호의존성"] == "낮음"

        if short_term and low_value:
            exempt_type = "단기+소액 (둘 다 해당)"
        elif short_term:
            exempt_type = "단기리스"
        elif low_value:
            exempt_type = "소액자산"
        else:
            exempt_type = "면제불가"

        is_exempt = short_term or low_value

        # 비용 계산 (정액)
        fy_start = date(base_dt.year, 1, 1)
        total_rent = item["월리스료"] * item["리스기간(월)"]
        monthly_expense = total_rent / item["리스기간(월)"]  # 정액
        fy_months = min(item["리스기간(월)"], 12)  # 간이 계산
        fy_expense = round(monthly_expense * fy_months)

        results.append({
            "자산명": item["자산명"],
            "리스기간": f"{item['리스기간(월)']}개월",
            "매수선택권": item["매수선택권"],
            "새자산가치": item["새자산가치"],
            "상호의존성": item["상호의존성"],
            "월리스료": item["월리스료"],
            "단기리스": "O" if short_term else "X",
            "소액자산": "O" if low_value else "X",
            "면제유형": exempt_type,
            "면제가능": is_exempt,
            "총리스료": total_rent,
            "당기비용(정액)": fy_expense,
        })

    st.session_state["exempt_results"] = results

if "exempt_results" in st.session_state:
    results = st.session_state["exempt_results"]
    st.divider()
    st.subheader("📋 판단 결과")

    exempt_count = sum(1 for r in results if r["면제가능"])
    non_exempt = sum(1 for r in results if not r["면제가능"])

    c1, c2, c3 = st.columns(3)
    c1.metric("전체", f"{len(results)}건")
    c2.metric("면제 가능", f"{exempt_count}건", delta="비용처리" if exempt_count > 0 else None)
    c3.metric("면제 불가", f"{non_exempt}건", delta="K-IFRS 16 적용" if non_exempt > 0 else None, delta_color="inverse")

    df = pd.DataFrame(results)
    display_cols = ["자산명", "리스기간", "매수선택권", "새자산가치", "상호의존성",
                    "단기리스", "소액자산", "면제유형", "월리스료", "당기비용(정액)"]
    st.dataframe(df[display_cols].style.format({
        "새자산가치": "{:,.0f}", "월리스료": "{:,.0f}", "당기비용(정액)": "{:,.0f}",
    }).apply(lambda row: ['background-color: #E8F5E9' if row.name in [i for i, r in enumerate(results) if r["면제가능"]]
                          else 'background-color: #FCE4EC'] * len(row), axis=1),
        use_container_width=True, hide_index=True)

    # 면제 항목 비용 합계
    total_exempt_expense = sum(r["당기비용(정액)"] for r in results if r["면제가능"])
    if total_exempt_expense > 0:
        st.info(f"💰 면제 대상 리스의 당기 비용 합계: **{total_exempt_expense:,.0f}원** (정액 기준)")

    # 면제 불가 항목 안내
    non_exempt_items = [r["자산명"] for r in results if not r["면제가능"]]
    if non_exempt_items:
        st.warning(f"⚠️ 면제 불가 자산: **{', '.join(non_exempt_items)}** → 리스이용자 기본 모듈에서 재계산 필요")

    # 분개 예시
    st.markdown("#### 면제 리스 분개 (매기)")
    st.code("Dr. 리스료 (비용)    xxx\n    Cr. 현금          xxx", language=None)

    # 엑셀 다운로드
    wb = Workbook()
    ws = wb.active; ws.title = "면제판단"
    HF2 = PatternFill("solid", fgColor="002060")
    HN2 = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10)
    NF2 = Font(name="맑은 고딕", size=10)
    tb2 = Border(left=Side(style='thin', color='B0B0B0'), right=Side(style='thin', color='B0B0B0'),
                 top=Side(style='thin', color='B0B0B0'), bottom=Side(style='thin', color='B0B0B0'))

    headers = ["자산명", "리스기간", "매수선택권", "새자산가치", "상호의존성",
               "단기리스", "소액자산", "면제유형", "월리스료", "총리스료", "당기비용(정액)"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = HN2; c.fill = HF2; c.alignment = Alignment(horizontal='center', vertical='center'); c.border = tb2
    for ri, r in enumerate(results, 2):
        for ci, k in enumerate(headers, 1):
            v = r.get(k, "")
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = NF2; c.border = tb2
            if k in ["새자산가치", "월리스료", "총리스료", "당기비용(정액)"]:
                c.number_format = '#,##0'

    buf = BytesIO(); wb.save(buf); buf.seek(0)
    st.download_button("📥 엑셀 다운로드", data=buf, file_name=f"면제판단_{base_dt}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       type="primary", use_container_width=True)

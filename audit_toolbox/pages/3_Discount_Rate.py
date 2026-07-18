"""할인율 산정 도우미 (K-IFRS 1116호 문단 26, 부록 A)"""
import streamlit as st
import pandas as pd
from datetime import date
from io import BytesIO
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

st.set_page_config(page_title="할인율 산정", page_icon="📐", layout="wide")
st.title("📐 할인율 산정 도우미")
st.caption("K-IFRS 1116호 문단 26 | 내재이자율 또는 증분차입이자율 산출 근거 문서화")

with st.expander("📖 할인율 산정 기준", expanded=False):
    st.markdown("""
    **K-IFRS 1116호 문단 26:**
    > 리스의 내재이자율을 쉽게 산정할 수 있는 경우 그 이자율로, 아닌 경우 증분차입이자율로 리스료를 할인

    **내재이자율:** 리스제공자가 받는 리스료 + 무보증잔존가치의 현재가치 = 자산 공정가치 + 초기직접원가

    **증분차입이자율:** 리스이용자가 비슷한 경제적 환경에서 비슷한 기간에 걸쳐 비슷한 담보로
    사용권자산과 비슷한 가치의 자산 취득에 필요한 자금을 차입한다면 지급해야 할 이자율

    **증분차입이자율 산출 = 기준금리 + 신용스프레드 ± 리스 조정**
    """)

st.divider()

# ── 이자율 유형 선택 ──
rate_type = st.radio("이자율 유형", ["내재이자율 (알고 있는 경우)", "증분차입이자율 (산출 필요)"], horizontal=True)

if "내재이자율" in rate_type:
    st.subheader("내재이자율 입력")
    c1, c2 = st.columns(2)
    asset_name = c1.text_input("리스자산명", "본사 사무실", key="ir_name")
    implicit_rate = c2.number_input("내재이자율 (%)", value=10.0, step=0.1, key="ir_rate")
    source = st.text_input("근거 자료", "리스계약서 상 명시 / 리스제공자 제공", key="ir_source")

    if st.button("📋 산정 근거 정리", type="primary", use_container_width=True, key="ir_btn"):
        st.session_state["rate_result"] = {
            "type": "내재이자율",
            "asset": asset_name,
            "rate": implicit_rate,
            "source": source,
            "buildup": None,
        }

else:
    st.subheader("증분차입이자율 산출")

    asset_name = st.text_input("리스자산명", "본사 사무실", key="ibr_name")

    st.markdown("#### Step 1: 기준금리")
    c1, c2, c3 = st.columns(3)
    base_rate_source = c1.selectbox("기준금리 출처", [
        "국고채 수익률", "통안채 수익률", "금융채(AAA) 수익률", "CD 금리", "기타"
    ], key="ibr_base_src")
    base_rate_tenor = c2.selectbox("만기", ["1년", "2년", "3년", "5년", "10년", "기타"], index=2, key="ibr_tenor")
    base_rate = c3.number_input("기준금리 (%)", value=3.50, step=0.01, key="ibr_base")
    base_rate_date = st.date_input("조회일", value=date(2025, 12, 31), key="ibr_base_date")
    base_rate_ref = st.text_input("출처 URL/문서", "한국은행 경제통계시스템 (ECOS) / 금융투자협회 채권정보센터", key="ibr_base_ref")

    st.markdown("#### Step 2: 신용스프레드")
    c1, c2, c3 = st.columns(3)
    credit_rating = c1.text_input("신용등급", "BBB+", key="ibr_cr")
    credit_source = c2.selectbox("등급 출처", [
        "NICE신용평가", "한국기업평가(한기평)", "한국신용평가(KIS)", "NICEBIZLINE", "기타"
    ], key="ibr_cr_src")
    credit_spread = c3.number_input("신용스프레드 (%)", value=1.50, step=0.01, key="ibr_spread")

    spread_method = st.selectbox("스프레드 산출 방법", [
        "해당 등급 회사채 수익률 - 국고채 수익률",
        "KOFIA 회사채 수익률 직접 조회",
        "한기평/NICE 신용등급별 스프레드표 참조",
        "실제 차입금리에서 역산",
        "기타"
    ], key="ibr_spread_method")

    if "등급별" in spread_method or "KOFIA" in spread_method:
        st.info("💡 이전에 NICEBIZLINE → 한기평 스프레드 → 노치 가속비율 외삽법으로 B- 수익률을 산출하신 적이 있었죠. 같은 방법론을 여기에 기록하시면 됩니다.")

    spread_ref = st.text_input("스프레드 근거 자료", "금융투자협회 KOFIA 채권수익률 / NICEBIZLINE 신용등급 보고서", key="ibr_spread_ref")

    st.markdown("#### Step 3: 리스 조정 (선택)")
    c1, c2 = st.columns(2)
    collateral_adj = c1.number_input("담보 조정 (%)", value=0.0, step=0.1, key="ibr_coll",
                                     help="유담보 차입 → 무담보 대비 금리 낮음. 마이너스(-)로 입력")
    other_adj = c2.number_input("기타 조정 (%)", value=0.0, step=0.1, key="ibr_other",
                                help="통화 조정, 기간 조정 등")

    st.markdown("#### Step 4: 참고 정보 (선택)")
    c1, c2 = st.columns(2)
    ref_borrow_rate = c1.number_input("실제 차입금리 참고 (%)", value=0.0, step=0.1, key="ibr_ref_borrow",
                                      help="회사의 실제 은행 차입금리가 있으면 참고용으로 기록")
    kofia_yield = c2.number_input("KOFIA 회사채 수익률 참고 (%)", value=0.0, step=0.1, key="ibr_kofia")
    ref_borrow_source = st.text_input("참고 출처", "은행 대출약정서 / portal.kfb.or.kr 기업대출금리", key="ibr_ref_src")

    final_rate = base_rate + credit_spread + collateral_adj + other_adj

    st.markdown("---")
    st.markdown(f"### 📊 산출 결과: **{final_rate:.2f}%**")

    buildup_data = [
        ("기준금리", f"{base_rate_source} {base_rate_tenor}", f"{base_rate:.2f}%"),
        ("(+) 신용스프레드", f"{credit_rating} 등급", f"{credit_spread:.2f}%"),
    ]
    if collateral_adj != 0:
        buildup_data.append(("(±) 담보 조정", "", f"{collateral_adj:+.2f}%"))
    if other_adj != 0:
        buildup_data.append(("(±) 기타 조정", "", f"{other_adj:+.2f}%"))
    buildup_data.append(("**= 증분차입이자율**", "", f"**{final_rate:.2f}%**"))

    st.table(pd.DataFrame(buildup_data, columns=["구분", "내용", "이자율"]))

    if ref_borrow_rate > 0:
        diff = final_rate - ref_borrow_rate
        st.info(f"📌 실제 차입금리({ref_borrow_rate:.2f}%)와의 차이: {diff:+.2f}%p — "
                f"{'합리적 범위' if abs(diff) < 2 else '차이 큼, 검토 필요'}")

    if st.button("📋 산정 근거 정리", type="primary", use_container_width=True, key="ibr_btn"):
        st.session_state["rate_result"] = {
            "type": "증분차입이자율",
            "asset": asset_name,
            "rate": final_rate,
            "source": f"{base_rate_ref} / {spread_ref}",
            "buildup": {
                "base_rate": base_rate, "base_source": f"{base_rate_source} {base_rate_tenor}",
                "base_date": str(base_rate_date), "base_ref": base_rate_ref,
                "credit_rating": credit_rating, "credit_source": credit_source,
                "credit_spread": credit_spread, "spread_method": spread_method, "spread_ref": spread_ref,
                "collateral_adj": collateral_adj, "other_adj": other_adj,
                "ref_borrow_rate": ref_borrow_rate, "kofia_yield": kofia_yield,
            },
        }

# ── 결과 출력 ──
if "rate_result" in st.session_state:
    r = st.session_state["rate_result"]
    st.divider()
    st.subheader("📋 산정 근거 문서")

    doc = f"""### 리스 할인율 산정 근거서

**리스자산:** {r['asset']}
**이자율 유형:** {r['type']}
**적용 할인율:** {r['rate']:.2f}%

"""
    if r['buildup']:
        b = r['buildup']
        doc += f"""#### 산출 내역

| 구분 | 내용 | 비율 |
|------|------|------|
| 기준금리 | {b['base_source']} ({b['base_date']}) | {b['base_rate']:.2f}% |
| 신용스프레드 | {b['credit_rating']} ({b['credit_source']}) | {b['credit_spread']:.2f}% |
"""
        if b['collateral_adj'] != 0:
            doc += f"| 담보 조정 | | {b['collateral_adj']:+.2f}% |\n"
        if b['other_adj'] != 0:
            doc += f"| 기타 조정 | | {b['other_adj']:+.2f}% |\n"
        doc += f"| **증분차입이자율** | | **{r['rate']:.2f}%** |\n"

        doc += f"""
#### 근거 자료
- 기준금리: {b['base_ref']}
- 신용스프레드: {b['spread_ref']} ({b['spread_method']})
"""
        if b['ref_borrow_rate'] > 0:
            doc += f"- 참고(실제 차입금리): {b['ref_borrow_rate']:.2f}%\n"
        if b['kofia_yield'] > 0:
            doc += f"- 참고(KOFIA 회사채): {b['kofia_yield']:.2f}%\n"
    else:
        doc += f"#### 근거: {r['source']}\n"

    st.markdown(doc)

    # 엑셀 다운로드
    wb = Workbook()
    ws = wb.active; ws.title = "할인율 산정"
    HF2 = PatternFill("solid", fgColor="002060")
    HN2 = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10)
    TF2 = Font(name="맑은 고딕", bold=True, size=14, color="002060")
    NF2 = Font(name="맑은 고딕", size=10)
    BF2 = Font(name="맑은 고딕", size=10, bold=True)
    IF2 = Font(name="맑은 고딕", size=10, color="0000FF")
    tb2 = Border(left=Side(style='thin', color='B0B0B0'), right=Side(style='thin', color='B0B0B0'),
                 top=Side(style='thin', color='B0B0B0'), bottom=Side(style='thin', color='B0B0B0'))

    ws.merge_cells('A1:C1')
    ws.cell(row=1, column=1, value="리스 할인율 산정 근거서").font = TF2
    ws.cell(row=1, column=1).alignment = Alignment(horizontal='center'); ws.row_dimensions[1].height = 30

    row = [3]
    def add_row(label, val, fmt=None):
        ws.cell(row=row[0], column=1, value=label).font = NF2; ws.cell(row=row[0], column=1).border = tb2
        c = ws.cell(row=row[0], column=2, value=val); c.font = IF2; c.border = tb2
        if fmt: c.number_format = fmt
        row[0] += 1

    add_row("리스자산", r['asset'])
    add_row("이자율 유형", r['type'])
    add_row("적용 할인율", r['rate'] / 100, '0.00%')
    row[0] += 1

    if r['buildup']:
        b = r['buildup']
        ws.cell(row=row[0], column=1, value="산출 내역").font = BF2; row[0] += 1
        for ci, h in enumerate(["구분", "내용", "비율"], 1):
            c = ws.cell(row=row[0], column=ci, value=h); c.font = HN2; c.fill = HF2; c.border = tb2
            c.alignment = Alignment(horizontal='center')
        row[0] += 1

        items = [("기준금리", b['base_source'], b['base_rate']),
                 ("신용스프레드", f"{b['credit_rating']} ({b['credit_source']})", b['credit_spread'])]
        if b['collateral_adj'] != 0: items.append(("담보 조정", "", b['collateral_adj']))
        if b['other_adj'] != 0: items.append(("기타 조정", "", b['other_adj']))
        items.append(("증분차입이자율", "", r['rate']))

        for label, desc, val in items:
            ws.cell(row=row[0], column=1, value=label).font = NF2; ws.cell(row=row[0], column=1).border = tb2
            ws.cell(row=row[0], column=2, value=desc).font = NF2; ws.cell(row=row[0], column=2).border = tb2
            ws.cell(row=row[0], column=3, value=val / 100).font = BF2; ws.cell(row=row[0], column=3).border = tb2
            ws.cell(row=row[0], column=3).number_format = '0.00%'
            row[0] += 1

    ws.column_dimensions['A'].width = 20; ws.column_dimensions['B'].width = 40; ws.column_dimensions['C'].width = 15

    buf = BytesIO(); wb.save(buf); buf.seek(0)
    st.download_button("📥 엑셀 다운로드", data=buf, file_name=f"할인율_산정_{r['asset']}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       type="primary", use_container_width=True)

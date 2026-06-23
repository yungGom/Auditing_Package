"""
금융기관 탐지 엔진 (모집단 완전성 검토용)
- 외부 통신 없음. 전부 로컬 처리.
- 판정이 아니라 '후보 추출 + 근거 제시'가 목적.
"""
import os
import re
import unicodedata

# ─────────────────────────────────────────────────────────────
# 레이어 1: 금융기관 명칭 사전 (카테고리별)
# 회사별로 자유롭게 추가/삭제 가능한 고정 사전
# ─────────────────────────────────────────────────────────────
FI_DICT = {
    "은행": [
        "국민은행", "KB국민", "신한은행", "우리은행", "하나은행", "농협은행", "NH농협",
        "기업은행", "IBK", "산업은행", "KDB", "수협은행", "수협", "SC제일", "한국씨티", "씨티은행",
        "대구은행", "iM뱅크", "부산은행", "경남은행", "광주은행", "전북은행", "제주은행",
        "케이뱅크", "카카오뱅크", "토스뱅크", "외환은행", "씨티",
    ],
    "증권": [
        "미래에셋증권", "삼성증권", "한국투자증권", "NH투자증권", "KB증권", "키움증권",
        "대신증권", "메리츠증권", "하나증권", "신한투자증권", "유안타", "교보증권",
        "한화투자증권", "현대차증권", "DB금융투자", "이베스트", "SK증권",
    ],
    "보험": [
        "삼성생명", "한화생명", "교보생명", "신한라이프", "NH농협생명", "동양생명",
        "삼성화재", "DB손해보험", "현대해상", "KB손해보험", "메리츠화재", "한화손해보험",
        "흥국화재", "롯데손해보험", "MG손해보험",
    ],
    "보증기관": ["신용보증기금", "기술보증기금", "서울보증보험", "신용보증재단", "무역보험공사", "주택금융공사"],
    "카드": ["삼성카드", "신한카드", "현대카드", "KB국민카드", "롯데카드", "하나카드", "우리카드", "BC카드"],
    "저축은행": ["저축은행", "OK저축", "SBI저축", "웰컴저축", "페퍼저축"],
    "캐피탈/기타": ["캐피탈", "자산운용", "신탁", "할부금융", "리스금융", "팩토링"],
}

# 기관유형(사전 카테고리) → 조회서 양식 매핑
CONFIRM_FORM = {
    "은행": "은행", "증권": "증권사", "보험": "보험회사", "보증기관": "보증기관",
    "카드": "여신전문", "저축은행": "저축은행", "캐피탈/기타": "여신전문",
    "기타(접미어)": "기타(검토필요)",
}

# 사전에 없어도 아래 일반 접미어가 거래처명에 있으면 후보로 승격
GENERIC_SUFFIX = ["은행", "Bank", "증권", "Securities", "생명", "화재", "해상", "손해보험",
                  "보험", "캐피탈", "저축은행", "자산운용", "카드"]

# ─────────────────────────────────────────────────────────────
# 레이어 2: 금융기관 거래가 반드시 흔적을 남기는 계정과목
#   거래처명이 비거나 약칭이어도 계정으로 후보를 건진다 (완전성 보강)
# ─────────────────────────────────────────────────────────────
FI_ACCOUNTS = {
    "예금/현금성": ["보통예금", "당좌예금", "정기예금", "정기적금", "현금성자산", "단기금융상품", "기타예금"],
    "차입": ["단기차입금", "장기차입금", "유동성장기부채", "사채", "당좌차월", "외화차입금"],
    "손익": ["이자수익", "이자비용", "지급수수료"],
    "유가증권/파생": ["당기손익공정가치측정금융자산", "기타포괄손익공정가치측정금융자산",
                  "상각후원가측정금융자산", "파생상품자산", "파생상품부채", "매도가능증권", "단기매매증권"],
}
FI_ACCOUNT_FLAT = {a: cat for cat, lst in FI_ACCOUNTS.items() for a in lst}

# 적요(摘要)에서 잡는 보조 키워드 (계정이 두루뭉술할 때)
MEMO_HINTS = ["이자", "예금이자", "차입", "대출이자", "지급이자", "수수료", "외화환산", "배당"]


def _norm(s):
    """거래처명 정규화: 법인격 표기·공백·지점명 제거 → 기관 단위로 롤업"""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s)).strip()
    s = re.sub(r"\(주\)|㈜|주식회사|\(유\)|유한회사|\(주\)", "", s)
    # 지점/영업점 표기 제거 (롤업용 키 생성)
    s = re.sub(r"\s*(제?\d+|[가-힣A-Za-z]+)\s*(지점|영업점|출장소|센터|지店)\s*$", "", s)
    s = re.sub(r"\s+", "", s)
    return s


def _match_dict(name):
    """레이어1: 사전 매칭. (카테고리, 매칭키워드) 반환, 없으면 None"""
    if not name:
        return None
    nm = unicodedata.normalize("NFKC", str(name))
    for cat, kws in FI_DICT.items():
        for kw in kws:
            if kw in nm:
                return (cat, kw)
    # 일반 접미어 승격
    for suf in GENERIC_SUFFIX:
        if suf in nm:
            return ("기타(접미어)", suf)
    return None


def _pick_vendor(pairs):
    """
    거래처 후보 (컬럼, 값) 목록에서 채택할 거래처를 고른다.
    우선순위: 사전매칭(실제 기관 사전) > 접미어매칭. 둘 다 없으면 None.
    반환: (값, 컬럼, _match_dict결과) or None
    """
    suffix_fallback = None
    for col, val in pairs:
        if not val or not str(val).strip():
            continue
        m = _match_dict(val)
        if not m:
            continue
        if m[0] != "기타(접미어)":
            return (str(val).strip(), col, m)        # 사전매칭 즉시 채택
        if suffix_fallback is None:
            suffix_fallback = (str(val).strip(), col, m)  # 접미어는 후순위 보관
    return suffix_fallback


def scan(records):
    """
    records: list of dict, 각 행은
      {row_no, source, date, slip_no, account, vendor, memo, amount, ...}
    거래처는 단일 vendor 컬럼뿐 아니라 vendor_candidates(여러 컬럼의 값)도 스캔한다
    (더존 관리항목1~6 분산 대응). vendor_cand_cols 가 있으면 매칭된 컬럼을 기록한다.
    하위호환: vendor_candidates 가 없으면 기존처럼 vendor 단일값을 쓴다.
    반환: 후보별 집계 + 원천 추적 리스트
    """
    hits = []  # 원천 추적용: 행 단위 매칭 결과
    for r in records:
        # 거래처 후보 (컬럼, 값) 구성 — 다중 컬럼 + 하위호환 단일 vendor
        cand_vals = list(r.get("vendor_candidates") or [])
        cand_cols = list(r.get("vendor_cand_cols") or [])
        if len(cand_cols) != len(cand_vals):
            cand_cols = [None] * len(cand_vals)
        v_single = r.get("vendor") or ""
        if v_single and v_single not in cand_vals:
            cand_vals.append(v_single)
            cand_cols.append(r.get("vendor_col") or "vendor")
        pairs = list(zip(cand_cols, cand_vals))

        # 후보 중 사전/접미어 매칭되는 거래처 채택, 없으면 첫 비공란 후보
        picked = _pick_vendor(pairs)
        if picked:
            vendor, vendor_col, dict_hit = picked
        else:
            dict_hit = None
            vendor = next((str(v).strip() for _, v in pairs if v and str(v).strip()), "")
            vendor_col = None

        account = (r.get("account") or "").strip()
        memo = r.get("memo") or ""
        acct_hit = FI_ACCOUNT_FLAT.get(account)
        memo_hit = next((h for h in MEMO_HINTS if h in memo), None)

        if not (dict_hit or acct_hit or memo_hit):
            continue

        if dict_hit:
            basis, category, kw = "사전매칭", dict_hit[0], dict_hit[1]
            form = CONFIRM_FORM.get(category, "기타(검토필요)")
        elif acct_hit:
            basis, category, kw = "계정추적", acct_hit, account
            form = None  # 기관유형 미상 → 회계사 판단
        else:
            basis, category, kw = "적요키워드", "추정", memo_hit
            form = None

        hits.append({
            **r,
            "vendor": vendor,                 # 채택된 거래처(대표)
            "norm_vendor": _norm(vendor) or "(거래처미상)",
            "basis": basis,
            "category": category,
            "form": form,
            "matched": kw,
            "vendor_col": vendor_col,         # 매칭된 컬럼(서면 검토 패널용)
        })
    return hits


# ─────────────────────────────────────────────────────────────
# PATCH 8: 분개장 검토 스텝 (계정 슬라이서 + 드릴다운 + 일괄동작)
#   자동 탐지는 1차 추천(체크 기본값)으로만 쓰고 최종 판단은 사람이 한다.
# ─────────────────────────────────────────────────────────────
def _detect_row(r):
    """
    한 행의 거래처 채택 + 탐지 분류(검토 드릴다운용).
    반환 dict:
      vendor/vendor_col, basis/category/form/matched(집계용),
      auto(자동추천=체크 기본 ON), tag('자동추천'|'적요확인필요'|'비금융추정'),
      highlight(적요 확인 필요 강조)
    """
    cand_vals = list(r.get("vendor_candidates") or [])
    cand_cols = list(r.get("vendor_cand_cols") or [])
    if len(cand_cols) != len(cand_vals):
        cand_cols = [None] * len(cand_vals)
    v_single = r.get("vendor") or ""
    if v_single and v_single not in cand_vals:
        cand_vals.append(v_single)
        cand_cols.append(r.get("vendor_col") or "vendor")
    pairs = list(zip(cand_cols, cand_vals))
    first_vendor, first_col = "", None
    for c, v in pairs:
        if v and str(v).strip():
            first_vendor, first_col = str(v).strip(), c
            break

    account = (r.get("account") or "").strip()
    memo = r.get("memo") or ""
    acct_hit = FI_ACCOUNT_FLAT.get(account)

    picked = _pick_vendor(pairs)
    if picked:  # 레이어1: 거래처(관리항목 포함) 사전/접미어 매칭 → 자동추천
        vendor, vcol, dh = picked
        return dict(vendor=vendor, vendor_col=vcol, basis="사전매칭", category=dh[0],
                    form=CONFIRM_FORM.get(dh[0], "기타(검토필요)"), matched=dh[1],
                    auto=True, tag="자동추천", highlight=False)
    if acct_hit:  # 레이어2: 금융 계정 → 자동추천(거래처 미상이면 강조)
        return dict(vendor=first_vendor, vendor_col=first_col, basis="계정추적",
                    category=acct_hit, form=None, matched=account,
                    auto=True, tag="자동추천", highlight=(first_vendor == ""))
    memo_bank = _match_dict(memo)
    memo_kw = next((h for h in MEMO_HINTS if h in memo), None)
    if memo_bank and memo_bank[0] != "기타(접미어)":  # 적요에 진짜 기관명 → 사람이 확인
        return dict(vendor=memo_bank[1], vendor_col="적요", basis="적요매칭",
                    category=memo_bank[0], form=CONFIRM_FORM.get(memo_bank[0], "기타(검토필요)"),
                    matched=memo_bank[1], auto=False, tag="적요확인필요", highlight=True)
    if memo_kw or memo_bank:  # 약한 단서(키워드/접미어) → 사람이 확인
        return dict(vendor=first_vendor or (memo_bank[1] if memo_bank else ""),
                    vendor_col=("적요" if not first_vendor else first_col),
                    basis="적요키워드", category="추정",
                    form=("기타(검토필요)" if memo_bank else None),
                    matched=(memo_kw or (memo_bank[1] if memo_bank else "")),
                    auto=False, tag="적요확인필요", highlight=True)
    return dict(vendor=first_vendor, vendor_col=first_col, basis=None, category=None,
                form=None, matched="", auto=False, tag="비금융추정", highlight=False)


def account_summary(records):
    """분개장 계정과목 distinct 목록 + 건수 + 금융계정 여부(기본 체크용).
       정렬: 금융계정 먼저, 그 다음 건수 많은 순."""
    counts = {}
    for r in records:
        a = (r.get("account") or "").strip()
        if not a:
            continue
        counts[a] = counts.get(a, 0) + 1
    out = [{"account": a, "count": n, "is_fi": a in FI_ACCOUNT_FLAT}
           for a, n in counts.items()]
    out.sort(key=lambda x: (not x["is_fi"], -x["count"], x["account"]))
    return out


def review_journal(records, selected_accounts=None):
    """
    분개장 행을 검토용으로 분류(드릴다운/일괄동작용).
    selected_accounts(set) 가 주어지면 해당 계정 행만 반환.
    각 행: row_no/date/account/memo/amount/vendor_display/vendor_col/
           auto_include/highlight/tag/_detect/_rec/raw/raw_cols
    """
    out = []
    for r in records:
        account = (r.get("account") or "").strip()
        if selected_accounts is not None and account not in selected_accounts:
            continue
        d = _detect_row(r)
        out.append({
            "row_no": r.get("row_no"), "source": r.get("source"), "kind": r.get("kind"),
            "date": r.get("date") or "", "account": account,
            "memo": r.get("memo") or "", "amount": r.get("amount") or 0,
            "vendor_display": d["vendor"] or "(공란)", "vendor_col": d["vendor_col"],
            "auto_include": d["auto"], "highlight": d["highlight"], "tag": d["tag"],
            "_detect": d, "_rec": r,
            "raw": r.get("raw"), "raw_cols": r.get("raw_cols"),
        })
    return out


def hits_from_included(review_rows, include_map=None):
    """
    포함 확정된 검토행 → scan 형식 hits 생성 (aggregate 입력용).
    include 결정: include_map(row_no→bool) 우선, 없으면 각 행 auto_include.
    포함된 행이 비탐지(basis None)면 '수기포함'으로 후보화.
    _hits 추적용으로 review_account/included_via 기록.
    """
    hits = []
    for rr in review_rows:
        if include_map is not None:
            inc = include_map.get(rr["row_no"], rr["auto_include"])
        else:
            inc = rr["auto_include"]
        if not inc:
            continue
        d = rr["_detect"]
        vendor = d["vendor"] or ""
        rec = dict(rr.get("_rec") or {})
        rec.update({
            "vendor": vendor,
            "norm_vendor": _norm(vendor) or "(거래처미상)",
            "basis": d["basis"] or "수기포함",
            "category": d["category"] if d["category"] is not None else "수기",
            "form": d["form"],
            "matched": d["matched"] or vendor or "수기",
            "vendor_col": d["vendor_col"],
            "review_account": rr["account"],
            "included_via": rr["tag"],
        })
        hits.append(rec)
    return hits


def aggregate(hits):
    """레이어3: 정규화 기관 단위로 집계. 조회 대상 마스터 생성."""
    agg = {}
    for h in hits:
        key = h["norm_vendor"]
        if key not in agg:
            agg[key] = {
                "기관(정규화)": key,
                "조회서양식": None,
                "근거": set(),
                "매칭값": set(),
                "발견건수": 0,
                "원천행": [],
                "_hits": [],
            }
        a = agg[key]
        a["근거"].add(h["basis"])
        a["매칭값"].add(str(h["matched"]))
        a["발견건수"] += 1
        a["원천행"].append(h["row_no"])
        a["_hits"].append(h)
        # 사전매칭으로 기관유형이 확정되면 채택
        if h["form"] and a["조회서양식"] is None:
            a["조회서양식"] = h["form"]
    for a in agg.values():
        a["사전매칭여부"] = "Y" if "사전매칭" in a["근거"] else "N(검토필요)"
        if a["조회서양식"] is None:
            a["조회서양식"] = "미분류(검토필요)"
        a["근거"] = ", ".join(sorted(a["근거"]))
        a["매칭값"] = ", ".join(sorted(a["매칭값"]))
        a["원천행"] = ", ".join(str(x) for x in a["원천행"])
        a["포함근거"] = "당기 탐지"
        # 온라인 조회 가능 여부 + 고객센터 자동기입
        o = lookup_online(a["기관(정규화)"], a["조회서양식"])
        a["조회방법"] = "온라인 조회" if o["online"] else "서면 조회(확인필요)"
        a["전화번호"] = o["전화번호"]
        a["온라인정식명"] = o["정식명"]
    return list(agg.values())


def inst_key(name, form=None):
    """기관 동일성 판단 키. 온라인 참가 기관은 정식명으로 통일(명칭 변형 흡수)."""
    o = lookup_online(name, form)
    if o["online"]:
        return ("ON", o["정식명"])
    return ("OFF", _canon(name))


def merge_prior(candidates, prior_names):
    """
    전기 발송 조회처(담당자 직접 입력)를 당기 탐지 결과에 병합.
    - 명칭 변형(KEB하나은행=하나은행 등)은 inst_key의 온라인 정식명 통일로 흡수.
    - 당기에 이미 있으면 '포함근거'에 전기 표시, 없으면 '전기보유분'으로 추가.
    반환: (병합된 candidates, 신규 추가된 전기보유분 수)
    """
    idx = {}
    for c in candidates:
        c.setdefault("포함근거", "당기 탐지")
        k = inst_key(c.get("온라인정식명") or c["기관(정규화)"], c["조회서양식"])
        idx[k] = c
    added = 0
    for raw in prior_names:
        nm = str(raw).strip()
        if not nm:
            continue
        dh = _match_dict(nm)
        form = CONFIRM_FORM.get(dh[0], None) if dh else None
        o = lookup_online(nm, form)
        if not form and o["유형"] in ("은행", "증권사", "보험회사", "보증기관"):
            form = o["유형"]
        k = inst_key(nm, form)
        if k in idx:
            c = idx[k]
            if "전기" not in c["포함근거"]:
                c["포함근거"] = c["포함근거"] + " + 전기 조회처"
        else:
            newc = {
                "기관(정규화)": o["정식명"] or _norm(nm) or nm,
                "온라인정식명": o["정식명"],
                "조회서양식": form or "미분류(검토필요)",
                "사전매칭여부": "N(전기)",
                "근거": "전기 조회처",
                "매칭값": nm,
                "발견건수": 0,
                "원천행": "-",
                "_hits": [],
                "조회방법": "온라인 조회" if o["online"] else "서면 조회(확인필요)",
                "전화번호": o["전화번호"],
                "포함근거": "전기 조회처(당기 미탐지)",
            }
            candidates.append(newc)
            idx[k] = newc
            added += 1
    return candidates, added


# ─────────────────────────────────────────────────────────────
# 온라인 조회 가능 기관 목록 (전자조회 참가기관)
#   탐지된 기관이 여기 있으면 → "온라인 조회" + 고객센터 자동기입
#   type 표기는 CONFIRM_FORM 기준(은행/증권사/보험회사/보증기관/기타)
# ─────────────────────────────────────────────────────────────
ONLINE_FI = [
    # 은행
    ("HSBC", "은행", "2019-09-23", "02)2004-0298"),
    ("KEB하나은행", "은행", "2014-12-01", "1599-1111"),
    ("SC제일은행", "은행", "2015-06-15", "1588-1599"),
    ("경남은행", "은행", "2013-12-02", "1600-8585"),
    ("광주은행", "은행", "2013-10-31", "1588-3388"),
    ("국민은행", "은행", "2013-12-18", "1599-9999"),
    ("기업은행", "은행", "2013-12-16", "1566-2566"),
    ("농협은행", "은행", "2013-11-11", "1588-2100"),
    ("미즈호은행", "은행", "2025-03-10", "02-3782-6094"),
    ("부산은행", "은행", "2013-11-20", "1588-6200"),
    ("산업은행", "은행", "2015-12-01", "1588-1500"),
    ("수협은행", "은행", "2013-12-02", "1588-1515"),
    ("신한은행", "은행", "2013-12-02", "1599-8000"),
    ("씨티은행", "은행", "2013-12-27", "1588-7000"),
    ("아이엠뱅크(구 대구)", "은행", "2013-12-02", "1588-5050"),
    ("우리은행", "은행", "2014-01-03", "1588-5000, 1599-5000"),
    ("전북은행", "은행", "2014-06-01", "1588-4477"),
    ("제이피모간체이스은행", "은행", "2024-11-06", "02-758-5026"),
    ("제주은행", "은행", "2013-11-11", "1588-0079"),
    ("중소벤처기업진흥공단", "기타", "2026-06-01", "055-751-9607"),
    ("카카오뱅크", "은행", "2017-12-19", "1599-3333"),
    ("케이뱅크", "은행", "2018-12-10", "1522-1000"),
    ("토스뱅크", "은행", "2023-11-20", "1661-7654"),
    ("한국수출입은행", "은행", "2024-01-02", "02)3779-6114"),
    # 증권
    ("KB증권", "증권사", "2020-01-03", "1588-6611"),
    ("NH투자증권", "증권사", "2019-12-20", "1544-0000"),
    ("SK증권", "증권사", "2023-12-27", "1544-8245"),
    ("교보증권", "증권사", "2017-12-21", "1544-0900"),
    ("대신증권", "증권사", "2024-11-08", "1544-4488"),
    ("메리츠증권", "증권사", "2023-11-03", "1588-3400"),
    ("미래에셋증권", "증권사", "2020-12-28", "1588-6800"),
    ("삼성증권", "증권사", "2017-12-21", "1588-2323"),
    ("신한투자증권", "증권사", "2020-12-31", "1588-0365"),
    ("아이엠증권", "증권사", "2018-05-11", "1588-7171"),
    ("키움증권", "증권사", "2020-01-28", "02)1544-9000"),
    ("하나증권", "증권사", "2020-10-16", "1588-3111"),
    ("한국증권금융", "증권사", "2020-01-13", "02)3770-8800"),
    ("한국투자증권", "증권사", "2019-08-05", "1544-5000, 1588-0023"),
    ("한화투자증권", "증권사", "2020-12-28", "02)3772-7727"),
    ("현대차증권", "증권사", "2020-12-18", "1588-6655"),
    # 보험
    ("DB손해보험", "보험회사", "2021-12-31", "1588-0100"),
    ("KB손해보험", "보험회사", "2021-11-01", "02)6900-2556"),
    ("NH농협생명보험", "보험회사", "2022-12-13", "02)3786-8156"),
    ("NH농협손해보험", "보험회사", "2024-05-10", "1644-9000"),
    ("교보생명보험", "보험회사", "2021-11-22", "02)721-2707"),
    ("롯데손해보험", "보험회사", "2022-01-14", "02)3455-3621"),
    ("메리츠화재해상보험", "보험회사", "2018-01-15", "1566-7711"),
    ("미래에셋생명", "보험회사", "2023-06-16", "02-3271-4948"),
    ("삼성생명보험", "보험회사", "2022-08-16", "1588-3114"),
    ("삼성화재해상보험", "보험회사", "2020-01-10", "02)758-4634"),
    ("신한라이프생명보험", "보험회사", "2023-01-06", "02)3455-4741"),
    ("에이아이지손해보험", "보험회사", "2024-11-18", "02-316-5982"),
    ("한화생명보험", "보험회사", "2025-01-13", "1588-6363"),
    ("한화손해보험", "보험회사", "2020-12-31", "02)316-0476"),
    ("현대해상화재보험", "보험회사", "2019-08-19", "1588-5656"),
    ("흥국생명보험", "보험회사", "2025-01-13", "1588-2288"),
    # 보증기관
    ("기술보증기금", "보증기관", "2020-12-31", "1544-1120"),
    ("서울보증보험", "보증기관", "2017-12-21", "1670-7000"),
    ("신용보증기금", "보증기관", "2024-12-30", "1588-6565"),
]

# 토큰 정규화로 안 잡히는 표기 차이 보정 (분개장 표기 → 온라인목록 표준 표기)
ONLINE_ALIAS = {
    "대구은행": "아이엠뱅크", "iM뱅크": "아이엠뱅크", "디지비대구은행": "아이엠뱅크",
    "하이투자증권": "아이엠증권", "iM증권": "아이엠증권",
    "한국씨티은행": "씨티은행", "씨티": "씨티은행",
    "중소기업은행": "기업은행", "IBK기업은행": "기업은행",
    "제이피모간": "제이피모간체이스", "JP모간": "제이피모간체이스",
}


def _canon(name):
    """온라인 매칭 비교용: 법인격·지점·괄호·공백 제거 (영문약어는 식별자라 보존)"""
    s = unicodedata.normalize("NFKC", str(name or "")).upper().strip()
    s = re.sub(r"\(.*?\)", "", s)  # 괄호 내용 제거
    s = re.sub(r"주식회사|㈜|\(주\)|\(유\)", "", s)
    s = re.sub(r"\s*(제?\d+|[가-힣A-Z]+)\s*(지점|영업점|출장소|센터)\s*$", "", s)
    s = re.sub(r"\s+", "", s)
    return s


# 사전 인덱싱
_ONLINE_INDEX = [(_canon(ONLINE_ALIAS.get(n, n)), n, t, d, tel) for (n, t, d, tel) in ONLINE_FI]


def lookup_online(vendor_norm, form):
    """
    탐지 기관(vendor_norm) + 기관유형(form)을 온라인 목록과 대조.
    반환: {"online":bool, "정식명":str, "전화번호":str, "서비스일자":str} or 미매칭시 online=False
    """
    target = _canon(ONLINE_ALIAS.get(vendor_norm, vendor_norm))
    if not target:
        return {"online": False, "정식명": "", "전화번호": "", "서비스일자": "", "유형": ""}
    type_known = form in ("은행", "증권사", "보험회사", "보증기관")
    matches = []
    for canon, name, typ, date, tel in _ONLINE_INDEX:
        # 유형이 확정된 경우 같은 유형만, 미확정이면 전체 대상
        if type_known and typ != form and typ != "기타":
            continue
        if target == canon or target in canon or canon in target:
            matches.append((name, typ, date, tel, len(canon)))
    if not matches:
        return {"online": False, "정식명": "", "전화번호": "", "서비스일자": "", "유형": ""}
    # 유형 정확히 일치하는 것 우선, 그 다음 이름 길이 가까운 순(가장 구체적 매칭)
    matches.sort(key=lambda m: (m[1] != form, abs(m[4] - len(target))))
    name, typ, date, tel, _ = matches[0]
    return {"online": True, "정식명": name, "전화번호": tel, "서비스일자": date, "유형": typ}


# ─────────────────────────────────────────────────────────────
# PATCH 5: 사전·온라인목록 외부 엑셀 분리 (회계사가 코드 없이 갱신)
#   같은 폴더에 아래 파일이 있으면 내장 기본값 대신 사용한다:
#     금융기관사전.xlsx     (열: 카테고리, 키워드)
#     온라인조회목록.xlsx   (열: 기관명, 유형, 서비스일자, 전화번호)
#   파일이 없거나 읽기 실패 시 내장 기본값으로 자동 복귀(동작 불변).
#   외부 통신 없음. 로컬 엑셀만 읽는다.
# ─────────────────────────────────────────────────────────────
_BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DICT_XLSX_NAME = "금융기관사전.xlsx"
ONLINE_XLSX_NAME = "온라인조회목록.xlsx"

# 내장 기본값 스냅샷 (외부 파일 없을 때 복귀용)
_DEFAULT_FI_DICT = {k: list(v) for k, v in FI_DICT.items()}
_DEFAULT_ONLINE_FI = list(ONLINE_FI)

EXTERNAL_DICT_LOADED = False
EXTERNAL_ONLINE_LOADED = False


def _load_external_dict(path):
    """금융기관사전.xlsx → {카테고리:[키워드,...]} 또는 None(없음/실패)."""
    if not path or not os.path.exists(path):
        return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        d = {}
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row:
                continue  # 헤더 / 빈 행
            cat = str(row[0]).strip() if row[0] is not None else ""
            kw = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            if cat and kw:
                d.setdefault(cat, []).append(kw)
        wb.close()
        return d or None
    except Exception:
        return None


def _load_external_online(path):
    """온라인조회목록.xlsx → [(기관명, 유형, 서비스일자, 전화번호),...] 또는 None."""
    if not path or not os.path.exists(path):
        return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        out = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row:
                continue
            name = str(row[0]).strip() if row[0] is not None else ""
            typ = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            svc = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            tel = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
            if name and typ:
                out.append((name, typ, svc, tel))
        wb.close()
        return out or None
    except Exception:
        return None


def reload_external(base_dir=None):
    """외부 엑셀이 있으면 사전/온라인목록을 교체, 없으면 내장 기본값으로 복귀."""
    global FI_DICT, ONLINE_FI, _ONLINE_INDEX
    global EXTERNAL_DICT_LOADED, EXTERNAL_ONLINE_LOADED
    base = base_dir or _BASE_DIR
    d = _load_external_dict(os.path.join(base, DICT_XLSX_NAME))
    o = _load_external_online(os.path.join(base, ONLINE_XLSX_NAME))
    FI_DICT = d if d else {k: list(v) for k, v in _DEFAULT_FI_DICT.items()}
    ONLINE_FI = o if o else list(_DEFAULT_ONLINE_FI)
    EXTERNAL_DICT_LOADED = bool(d)
    EXTERNAL_ONLINE_LOADED = bool(o)
    _ONLINE_INDEX = [(_canon(ONLINE_ALIAS.get(n, n)), n, t, dt, tel)
                     for (n, t, dt, tel) in ONLINE_FI]
    return EXTERNAL_DICT_LOADED, EXTERNAL_ONLINE_LOADED


def export_templates(dirpath):
    """현재 사전/온라인목록을 엑셀 템플릿 2개로 내보낸다(회계사 편집용)."""
    from openpyxl import Workbook
    os.makedirs(dirpath, exist_ok=True)
    dict_path = os.path.join(dirpath, DICT_XLSX_NAME)
    online_path = os.path.join(dirpath, ONLINE_XLSX_NAME)

    wb = Workbook(); ws = wb.active; ws.title = "금융기관사전"
    ws.append(["카테고리", "키워드"])
    for cat, kws in FI_DICT.items():
        for kw in kws:
            ws.append([cat, kw])
    wb.save(dict_path)

    wb2 = Workbook(); ws2 = wb2.active; ws2.title = "온라인조회목록"
    ws2.append(["기관명", "유형", "서비스일자", "전화번호"])
    for (n, t, dt, tel) in ONLINE_FI:
        ws2.append([n, t, dt, tel])
    wb2.save(online_path)
    return dict_path, online_path


# 임포트 시 외부 파일 자동 반영 (없으면 내장 기본값 유지 → 동작 불변)
reload_external()

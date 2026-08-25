# -*- coding: utf-8 -*-
"""l2_probe — L2 표간 대사 CLI 진입점. l2_extract + l2_labels → 콘솔 요약 + xlsx.

이번 라운드 범위(설계안 확정): 단독 실행만. marks.json 편입 안 함, 렌더러·게이트
불가침. 완전 오프라인.
"""
import argparse, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import l2_extract
import l2_labels


def _company_from_filename(pdf_path):
    """samples/[회사명]문서종류(...).pdf 규약에서 회사명 추출. 못 찾으면 None
    (labels/{회사}.json이 없으면 l2_labels.load가 공통 사전만으로 동작)."""
    base = os.path.basename(pdf_path)
    if base.startswith("["):
        end = base.find("]")
        if end > 0:
            return base[1:end]
    return None


def _dict_applies(labels_path, pdf_path):
    """사전의 scope는 주석 번호에 묶여 있어 문서(보고서 종류·시점)가 다르면 무효다
    (2026-08-25 실측: 조선내화 반기연결용 사전을 연차보고서에 그대로 돌렸더니
    주석 번호 우연일치로 CONFIRMED 7건이 가짜로 나옴). applies_to에 파일명이
    명시돼 있으면 정확히 일치할 때만 사용 — table_aliases+fingerprint(TABLE_SHIFT)
    전면 구현 전까지의 안전장치."""
    import json
    with open(labels_path, encoding="utf-8") as f:
        doc = json.load(f)
    allow = doc.get("applies_to")
    if not allow:
        return True
    return os.path.basename(pdf_path) in allow


def run(pdf_path, company=None, out=None, quiet=False):
    company = company or _company_from_filename(pdf_path)
    labels_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "labels", f"{company}.json") \
        if company else None
    dict_exists = bool(labels_path and os.path.isfile(labels_path))
    dict_mismatch = dict_exists and not _dict_applies(labels_path, pdf_path)
    has_dict = dict_exists and not dict_mismatch
    warning = None
    if dict_mismatch:
        # 침묵 탈락 금지 — 사전이 있는데도 안 걸리면(파일명 불일치) UNMAPPED이
        # 대량 발생해도 이유를 알 수 없다. --quiet와 무관하게 항상 알린다
        # (2026-08-25 조건부 승인: 사전 미적용은 조용히 넘어가면 안 됨).
        warning = (f"사전 labels/{company}.json이 이 문서에 적용되지 않았습니다 "
                   f"(applies_to 불일치) — 공통 사전만 적용됩니다. "
                   f"table_aliases+fingerprint(TABLE_SHIFT) 정식 구현 전 임시방편(CLAUDE.md 참고).")
        print(f"[L2] *** 경고: {warning} ***")
    use_company = company if has_dict else None

    raw = l2_extract.extract(pdf_path)
    res = l2_labels.build(pdf_path, use_company, raw)

    if not quiet:
        mapped_raw = {t["raw_label"] for t in raw} - {t["raw_label"] for t in res["unmapped"]}
        unmapped_raw = {t["raw_label"] for t in res["unmapped"]}
        print(f"L2 표간 대사 → 매칭 라벨 {len(mapped_raw)}종 / UNMAPPED {len(unmapped_raw)}종"
              f"({len(res['unmapped'])}건)"
              + ("" if has_dict else f"  (※ labels/{company}.json 없음 — 공통 사전만 적용, 전량에 가깝게 UNMAPPED 정상)"))
        print(f"  CONFIRMED  {len(res['confirmed']):3d}건 (assert_equal 위반 — 검토 대상)")
        print(f"  UNDECLARED {len(res['undeclared']):3d}건 (assert_equal도 exclude_pairs도 없음 — 등록 필요 여부 판단)")
        print(f"  EXCLUDED   {len(res['excluded']):3d}건 (exclude_pairs로 억제됨, 참고용)")
        print(f"  경고: WEAK_COL {len(res['weak_col'])}건")
        for r in res["confirmed"]:
            print(f"    [CONFIRMED] {r['canonical_label']:20s} {r['table_a']:>10s} {r['won_a']:>16,.0f} "
                  f"vs {r['table_b']:>10s} {r['won_b']:>16,.0f}  차이 {r['diff']:+,.0f}원"
                  f"  (p{r['page_a']}/p{r['page_b']})")

    base = os.path.splitext(os.path.basename(pdf_path))[0]
    outdir = out or os.path.dirname(os.path.abspath(pdf_path))
    xlsx_path = os.path.join(outdir, base + "_L2대사.xlsx")
    sugg_path = os.path.join(outdir, base + "_undeclared_suggestions.json")
    _save_xlsx(xlsx_path, res, warning)
    _save_suggestions(sugg_path, res, pdf_path)
    if not quiet:
        print(f"산출물: {xlsx_path}")
        if res["undeclared"]:
            print(f"산출물: {sugg_path} (UNDECLARED {len(res['undeclared'])}건의 assert_equal 후보)")
    return res


def _sheet(wb, name, hdr, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(hdr)
    for cc in ws[1]:
        cc.font = Font(bold=True, color="FFFFFF")
        cc.fill = PatternFill("solid", fgColor="404040")
        cc.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _sign_flag(r):
    """부호반전 열 표시 — 어느 쪽 값에 sign_overrides가 적용됐는지(회계사가 나중에
    "여기 부호 뒤집었구나"를 바로 확인할 수 있어야 한다, 2026-08-25 조건부 승인)."""
    tags = []
    if r.get("sign_a", 1) != 1:
        tags.append("A")
    if r.get("sign_b", 1) != 1:
        tags.append("B")
    return ",".join(tags)


def _finding_row(r):
    return [r["canonical_label"], r["column_key"] or "", r["period_key"] or "",
            r["table_a"], r["table_b"], r["won_a"], r["won_b"], r["diff"],
            r["page_a"], r["page_b"], _sign_flag(r)]


def _save_xlsx(path, res, warning=None):
    wb = openpyxl.Workbook()

    if warning:
        # 사전 미적용 경고 — xlsx를 열자마자 보이도록 맨 앞 시트(2026-08-25 조건부 승인,
        # 침묵 탈락 금지: 사전이 안 걸린 걸 회계사가 UNMAPPED 폭증만 보고는 모른다).
        ws0 = wb.create_sheet("⚠경고")
        ws0.append(["사전이 적용되지 않았습니다"])
        ws0["A1"].font = Font(bold=True, size=14, color="C00000")
        ws0.append([warning])
        ws0.column_dimensions["A"].width = 100
        ws0["A2"].alignment = Alignment(wrap_text=True)

    hdr = ["항목", "열", "기간", "표A", "표B", "값A(원)", "값B(원)", "차이(원)", "페이지A", "페이지B", "부호반전"]
    widths = [22, 18, 10, 12, 12, 18, 18, 16, 8, 8, 10]

    ws = _sheet(wb, "CONFIRMED", hdr, [_finding_row(r) for r in res["confirmed"]], widths)
    for row in ws.iter_rows(min_row=2):
        for cc in row[5:8]:
            cc.number_format = "#,##0"
        row[7].fill = PatternFill("solid", fgColor="FFC7CE")

    # UNDECLARED — 요청된 컬럼명(항목A/항목B/값A/값B/차이/출처 페이지, 2026-08-25 확정)
    # + 부호반전(같은 원칙 적용)
    und_hdr = ["항목A", "항목B", "값A", "값B", "차이", "출처 페이지", "부호반전"]
    und_rows = [[f"{r['canonical_label']} ({r['table_a']})", f"{r['canonical_label']} ({r['table_b']})",
                 r["won_a"], r["won_b"], r["diff"], f"p{r['page_a']}/p{r['page_b']}", _sign_flag(r)]
                for r in res["undeclared"]]
    ws2 = _sheet(wb, "UNDECLARED", und_hdr, und_rows, [26, 26, 18, 18, 16, 12, 10])
    for row in ws2.iter_rows(min_row=2):
        for cc in row[2:4]:
            cc.number_format = "#,##0"
        row[4].fill = PatternFill("solid", fgColor="FFEB9C")

    ws3 = _sheet(wb, "EXCLUDED", hdr + ["사유"],
                 [_finding_row(r) + [r.get("reason", "")] for r in res["excluded"]], widths + [30])
    for row in ws3.iter_rows(min_row=2):
        for cc in row[5:8]:
            cc.number_format = "#,##0"

    unmapped_rows = [[t["raw_label"], t["table_id"], t["page"]] for t in res["unmapped"]]
    _sheet(wb, "UNMAPPED", ["원문 라벨", "표", "페이지"], unmapped_rows, [30, 14, 8])

    warn_rows = [["WEAK_COL", t["table_id"], f"p{t['page']} c{t['col']} — 헤더 없음, 위치로만 판정"]
                 for t in res["weak_col"]]
    _sheet(wb, "WARNINGS", ["유형", "표", "상세"], warn_rows, [14, 14, 50])

    wb.remove(wb["Sheet"])
    wb.active = 0
    wb.save(path)


def _save_suggestions(path, res, pdf_path):
    import json
    suggestions = []
    seen = set()
    for r in res["undeclared"]:
        key = (r["canonical_label"], r["table_a"], r["table_b"])
        if key in seen:
            continue
        seen.add(key)
        suggestions.append({
            "kind": "assert_equal",
            "label": r["canonical_label"],
            "_evidence": {"table_id_a": r["table_a"], "table_id_b": r["table_b"],
                          "value_a": r["won_a"], "value_b": r["won_b"],
                          "page_a": r["page_a"], "page_b": r["page_b"]},
        })
    doc = {"schema": "dsd-l2-undeclared-suggestions/1",
           "generated_from": os.path.basename(pdf_path), "suggestions": suggestions}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=2)
        f.write("\n")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(prog="l2_probe", description="L2 표간 대사 — 단독 실행(콘솔+xlsx). marks.json 편입 안 함")
    ap.add_argument("pdf")
    ap.add_argument("--company", default=None, help="labels/{회사}.json 이름 (기본: 파일명 [회사명] 자동 추출)")
    ap.add_argument("--out", default=None)
    ap.add_argument("--quiet", action="store_true")
    a = ap.parse_args()
    if not os.path.isfile(a.pdf):
        print(f"[오류] 파일이 없습니다: {a.pdf}"); sys.exit(2)
    run(a.pdf, company=a.company, out=a.out, quiet=a.quiet)

# -*- coding: utf-8 -*-
"""DSD 풋팅 엔진 — 통합 실행 (A1·A2·A3·A5·A7·B·C7·F1). 완전 오프라인.

Phase 1(판정/렌더 분리, 2026-08-14): 이 파일은 분석만 한다. 마크는 픽셀이 아니라
marks.py가 직렬화하는 데이터(<원본>_marks.json)이고, 실제 그리기는 render.py가
전담한다. `python foot.py`만 쓰면 현행과 동일하게 동작한다(내부적으로 marks.json도
함께 생성된다). 편집 후 재렌더는 `python render.py --marks ...`를 직접 쓴다.
"""
import argparse, collections, datetime, os, re, subprocess, sys, time
import pdfplumber
from core import (check_table, verdict, find_unit, grid_info, mixed_currency, is_total_label,
                  norm, SKIP_REASON_TEXT)
from statements import foot_hier, foot_a5, stmt_type, APPLY
import tieout, notes, prose, refmap, consist, docmeta
import marks as marksio
import render as renderer

# 검증 종류 → 화면 설명 문구(rev.2 formula). 코드값을 그대로 내보내지 않는다.
FORMULA_TEXT = {"A1": "세로 합계", "A2": "가로 합계", "A3": "계층 합계",
                "A5": "가감 관계식", "C": "본표↔주석 금액 대사"}


def fmt_like(val, sample):
    """재계산 금액을 공시 표기와 같은 규약으로 찍는다. rev.2가 shown_value와 자리를
    맞대어 어긋난 자리만 강조하도록 요구하므로, 음수 표기(괄호 vs 마이너스)가 다르면
    자리 비교가 어긋난다 — 원문이 괄호를 쓰면 괄호로 맞춘다."""
    s = f"{abs(val):,.0f}"
    if val >= 0:
        return s
    return f"({s})" if (sample or "").strip().startswith("(") else f"-{s}"

_ap = argparse.ArgumentParser(prog="foot",
    description="DSD 풋팅 — 표시 수치 정합성 검증 → <원본>_틱마크.pdf / <원본>_예외색인.xlsx "
                "(완전 오프라인 · 자동은 추천, 확정은 회계사 — candidate only)")
_ap.add_argument("pdf", nargs="?", help="감사보고서 PDF 경로")
_ap.add_argument("tol_pos", nargs="?", type=float, help=argparse.SUPPRESS)  # 구형: final.py 보고서.pdf 0
_ap.add_argument("--tol", type=float, default=None, help="A7 허용오차 (기본 0)")
_ap.add_argument("--round-steps", dest="round_steps", type=int, default=1,
                 help="표 단위 스텝 허용 — 백만원 표=1백만원, 천원 표=1천원 (기본 1, 2026-08-09 확정)")
_ap.add_argument("--min-won", dest="min_won", type=float, default=refmap.MIN_WON,
                 help="C 대사 소액 제외, 원 환산 절대금액 (기본 1억)")
_ap.add_argument("--out", default=None, help="산출물 폴더 (기본: 입력 PDF와 같은 폴더)")
_ap.add_argument("--terms", default=None,
                 help="F2 기준 표현 CSV (한 줄에 표현 하나, # 주석). 미지정 시 빈도 기준")
_ap.add_argument("--quiet", action="store_true", help="요약 출력 생략")
_cli, _ = _ap.parse_known_args()                 # --update-gates 등은 gates.py가 처리
if not _cli.pdf:
    _ap.print_help(); sys.exit(2)

PDF = _cli.pdf
TOL = _cli.tol if _cli.tol is not None else (_cli.tol_pos if _cli.tol_pos is not None else 0.0)
RSTEPS = _cli.round_steps
MINWON = _cli.min_won
QUIET = _cli.quiet

# ── 입력 검증 ──
if not os.path.isfile(PDF):
    print(f"[오류] 파일이 없습니다: {PDF}"); sys.exit(2)
if not PDF.lower().endswith(".pdf"):
    print(f"[오류] PDF 파일이 아닙니다: {PDF}"); sys.exit(2)
with pdfplumber.open(PDF) as _p:
    _chars = sum(len(pg.extract_text() or "") for pg in _p.pages[:5])
if _chars < 100:
    print("[오류] 텍스트 레이어가 없습니다(스캔본 추정). OCR은 조용히 틀리므로 설계상 "
          "사용하지 않습니다 — 이 파일은 처리를 거부합니다."); sys.exit(2)

_base = os.path.splitext(os.path.basename(PDF))[0]
_outdir = _cli.out or (os.path.dirname(os.path.abspath(PDF)) or ".")
os.makedirs(_outdir, exist_ok=True)
OUT_PDF = os.path.join(_outdir, _base + "_틱마크.pdf")
OUT_XLSX = os.path.join(_outdir, _base + "_예외색인.xlsx")
MARKS_JSON = os.path.join(_outdir, _base + "_marks.json")

# ── 버전 스탬프 (감사조서 추적성) ──
try:
    VER = subprocess.run(["git", "rev-parse", "--short", "HEAD"], capture_output=True,
                         text=True, cwd=os.path.dirname(os.path.abspath(__file__))
                         ).stdout.strip() or "nogit"
except Exception:
    VER = "nogit"
RUN_TS = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

# ── F2 기준 표현 (고정 템플릿, 선택) ──
# 회계법인/조서 관행상 기준 표현이 정해져 있으면 빈도 대신 그 표현을 주된 표기로 쓴다.
# 파일이 없거나 미지정이면 현행(빈도 기준) 동작 그대로.
TERMS = set()
if _cli.terms:
    try:
        with open(_cli.terms, encoding="utf-8-sig") as _tf:
            for _ln in _tf:
                _t = _ln.split(",")[0].strip()
                if _t and not _t.startswith("#"):
                    TERMS.add(_t.replace(" ", ""))
    except OSError:
        print(f"[경고] --terms 파일을 열 수 없습니다: {_cli.terms} — 빈도 기준으로 진행")

FCON = consist.report(PDF)
CIR, TAGS, RDIFF, RLINKS, RUN, RMAIN, REXCL = refmap.marks(PDF, TOL, MINWON)
TUNIT = refmap.unit_texts(PDF)                 # 표 귀속 단위 원문 — 복합 통화 A2 제외용
cirmap=collections.defaultdict(list); dmap=collections.defaultdict(list)
for x in CIR: cirmap[(x["page"], x["table"])].append(x)
for x in RDIFF: dmap[(x["page"], x["table"])].append(x)
tagmap=collections.defaultdict(list)
for (pg_,tb_,rw_,cl_), lbls in TAGS.items(): tagmap[(pg_,tb_)].append((rw_,cl_,lbls))

exc=[]; stat=collections.Counter(); nounit=[]; pros=[]; npara=0
SUBLOG=[]                              # 하위항목(소계 귀속) 판정 로그 — 경로별 검증용
SPLIT_SUSPECT=[]                       # 분할 의심 페이지 전환 — 참고 정보 (판정 무영향)
MARKS=[]                               # 마크 데이터 (픽셀이 아니라 데이터 — Phase 1)
A2_OCC=collections.defaultdict(int)    # (page,table,row)별 발생 순번 — A2는 col 필드가
                                        # 없어(이중 합계열 시 같은 행에 여러 A2 가능) id
                                        # 계산에만 쓰는 보조 카운터. bbox 산출 로직은 불변.

def ends_open(tb):
    """표의 마지막 데이터 행이 합계행(계·총계류)이 아니면 True — 분할 의심 신호.
    3축 전수 계측(2026-08): 표준 분할 5건은 carry가 전부 병합. 이 신호는 3축에 없는
    조판(헤더 반복형 등)이 실전에서 나타날 때 조용히 넘어가지 않게 하는 탐지기다."""
    try:
        G, K, V, hdr, ncol, nrow, numcols = grid_info(tb)
        if not numcols: return False
        first_num = min(numcols)
        for i in range(nrow - 1, hdr - 1, -1):
            if any(K[i][j] == "NUM" for j in numcols):
                lab = next((G[i][j] for j in range(first_num)
                            if G[i][j] and K[i][j] == "TEXT"), "")
                return not is_total_label(lab)
        return False
    except Exception:
        return False

_T0 = time.time()
carry=None
STMT_PAGES={}                          # page → 본표 유형 (목차 depth 0 생성용)
TABLE_LABEL={}                         # (page, table) → 사람이 읽는 표 이름 (없으면 None)
UNRESOLVED_LABEL=[]                    # 표 이름 추출 실패 — 빈 문자열로 흘리지 않는다
PAGE_SEQ={}                            # page → 그 페이지에서 쓴 seq 개수 (L2 이어붙이기용)
PREV_CAP=None                          # 직전 표 제목 — 페이지 넘어 이어지는 표에 물려준다
with pdfplumber.open(PDF) as pdf:
    _HEADS = docmeta.headings(pdf)
    _TITLE, _TITLE_FALLBACK = docmeta.title(pdf)
    if _TITLE_FALLBACK:
        _TITLE = _base
        if not QUIET:
            print(f"[문서] 표지에서 제목을 읽지 못해 파일명으로 대체합니다: {_TITLE}")
    for pi,page in enumerate(pdf.pages,1):
        SEQ=0                                        # 페이지별 원 그리기 순서 — R-1의 핵심
        txt=page.extract_text() or ""
        unit=find_unit(txt); st=stmt_type(txt)
        if st: STMT_PAGES[pi]=st
        tobjs=page.find_tables()
        _plines = page.extract_text_lines() if tobjs else []
        for _ti,_t in enumerate(tobjs,1):
            _cap = docmeta.table_caption(_plines, _t.bbox)
            if _cap is None and _ti == 1 and _t.bbox[1] < 100 and PREV_CAP:
                # 페이지 첫 표가 지면 맨 위에서 시작하면(위에 글줄이 아예 없다) 앞 페이지
                # 표의 이어짐이다 — 제목은 앞 페이지에 있다. 이 신호가 없을 때 물려받으면
                # 남의 제목을 붙이므로 '지면 최상단 + 첫 표'로만 한정한다.
                _cap = PREV_CAP
            if _cap is None and st is None:
                UNRESOLVED_LABEL.append({"page": pi, "table": _ti})
            if _cap: PREV_CAP = _cap
            TABLE_LABEL[(pi,_ti)] = _cap
        pr_d=prose.check_page(page); pr_f6=prose.check_leading_space(page)
        pr=pr_d+pr_f6; paras=prose.paragraphs(page)
        if not tobjs and not pr and not paras: carry=None; continue
        if not unit: nounit.append(pi)
        for _gi,gp in enumerate([g for g in FCON["F4_다중공백"] if g["page"]==pi]):
            x0,t0,x1,b0 = gp["bbox"]
            MARKS.append(dict(id=marksio.mid("gapx","F4",pi,extra=_gi), seq=SEQ, page=pi, kind="gapx",
                box=marksio.box4(x0,t0,x1,b0), text=None,
                source=dict(check="F4", table=None, row=None, col=None, label=None), verdict=None,
                evidence=dict(before=gp["before"], after=gp["after"]),
                origin="tool", status="pending", note=None)); SEQ+=1
        for _pidx,pg_ in enumerate(paras):
            x_end,ptop,pbot = pg_["bbox"]
            MARKS.append(dict(id=marksio.mid("slash","D",pi,extra=_pidx), seq=SEQ, page=pi, kind="slash",
                box=marksio.box4(x_end,ptop,x_end,pbot), text=None,
                source=dict(check="D", table=None, row=None, col=None, label=None), verdict=None,
                evidence=dict(nline=pg_["nline"], preview=pg_["text"]),
                origin="tool", status="pending", note=None)); SEQ+=1
            npara += 1
        for _fi,pz in enumerate(pr):
            _chk = "D" if _fi < len(pr_d) else "F6"
            MARKS.append(dict(id=marksio.mid("flagword",_chk,pi,extra=_fi), seq=SEQ, page=pi, kind="flagword",
                box=marksio.box4(*pz["bbox"]), text=None,
                source=dict(check=_chk, table=None, row=None, col=None, label=pz["text"]), verdict=None,
                evidence=dict(msg=pz["msg"], level=pz["level"]),
                origin="tool", status="pending", note=None)); SEQ+=1
            pros.append([pi, pz["text"], pz["msg"], pz["level"]])
        for ti,t in enumerate(tobjs,1):
            data=t.extract()
            if not data or len(data)<2: continue
            off=0
            if ti==1 and carry is not None:
                _why=None
                try:
                    _,_,_,h0,nc0,_,_=grid_info(data)
                    if h0==0 and nc0==carry[1]: off=len(carry[0]); data=carry[0]+data
                    elif h0!=0: _why=f"헤더 반복(h0={h0})"
                    else: _why=f"열수 불일치 {carry[1]}→{nc0}"
                except Exception: _why="grid_info 실패"
                # 분할 의심: 앞 페이지 마지막 표가 합계행 없이 끝났는데 병합되지 않음.
                # 참고 정보만 — A/B/C 지표·판정에 영향 없음. 오탐 다수(다음 주석의 새 표).
                if _why and len(carry)>3 and carry[2]:
                    SPLIT_SUSPECT.append(dict(prev=carry[3], page=pi, reason=_why, seq=SEQ)); SEQ+=1
            x0s=[]                                       # 첫 열 라벨 x0 — 들여쓰기 하위항목 판정용
            pwords = page.extract_words()
            for rr in t.rows:
                bb = rr.cells[0] if rr.cells else None
                xw = None
                if bb:
                    ws_=[w["x0"] for w in pwords
                         if bb[0]-1 <= w["x0"] and w["x1"] <= bb[2]+1
                         and bb[1]-1 <= w["top"] and w["bottom"] <= bb[3]+1]
                    xw = min(ws_) if ws_ else None
                x0s.append(xw)
            if off: x0s = [None]*off + x0s               # carry 병합 행은 좌표 미상
            rs = check_table(data, x0s=x0s, sublog=SUBLOG, ctx=(pi,ti),
                             excl_a2=mixed_currency(TUNIT.get((pi,ti))))  # A1·A2 (+구형 A3 제거됨)
            rs = [r for r in rs if r["kind"] in ("A1","A2")]
            if "A3" in APPLY.get(st, ()): rs += foot_hier(data)
            if "A5" in APPLY.get(st,()):      rs += foot_a5(data)
            def cellbox(gi, cj):
                gi -= off
                if not (0 <= gi < len(t.rows)): return None
                cells = t.rows[gi].cells
                if cj is not None and cj < len(cells) and cells[cj]: return cells[cj]
                for cc in reversed(cells):
                    if cc: return cc
                return None
            for x in cirmap.get((pi,ti), []):
                bb = cellbox(x["row"]+off, x["ncol"])
                if not bb: continue
                want = {str(n) for n in x["notes"]}
                for wd in page.extract_words():
                    if not (bb[0]-1 <= wd["x0"] and wd["x1"] <= bb[2]+1
                            and bb[1]-1 <= wd["top"] and wd["bottom"] <= bb[3]+1): continue
                    txt = wd["text"].strip()
                    note_num = txt.rstrip(",")
                    if note_num in want:
                        x1w = wd["x1"]
                        if txt.endswith(","):            # 뒤 쉼표는 원 밖으로
                            x1w -= (wd["x1"]-wd["x0"]) / max(len(txt),1) * 0.75
                        MARKS.append(dict(
                            id=marksio.mid("circle","C",pi,table=ti,row=x["row"],extra=f"n{note_num}"),
                            seq=SEQ, page=pi, kind="circle",
                            box=marksio.box4(wd["x0"], wd["top"], x1w, wd["bottom"]), text=None,
                            source=dict(check="C", table=ti, row=x["row"], col=x["ncol"], label=note_num),
                            verdict=None, evidence=dict(notes=sorted(x["notes"])),
                            origin="tool", status="pending", note=None)); SEQ+=1
            for rw_, cl_, lbls in tagmap.get((pi,ti), []):
                bb = cellbox(rw_+off, cl_)
                if bb:
                    seen=[]; [seen.append(l) for l in lbls if l not in seen]
                    _txt=" ".join(seen[:2])
                    MARKS.append(dict(id=marksio.mid("reftag","C",pi,table=ti,row=rw_,col=cl_),
                        seq=SEQ, page=pi, kind="reftag",
                        box=marksio.box4(*bb), text=_txt,
                        source=dict(check="C", table=ti, row=rw_, col=cl_, label=None),
                        verdict=None, evidence=dict(labels=seen),
                        origin="tool", status="pending", note=None)); SEQ+=1
            for x in dmap.get((pi,ti), []):
                bb = cellbox(x["row"]+off, x["col"])
                if bb:
                    _txt=f"{x['diff']:+,.0f} {x['where']}"
                    # C 레퍼 미성립은 '금액이 틀렸다'가 아니라 '주석에서 동일 금액을 찾지
                    # 못했다'다 → type=unverified (marks.type_of, 회계사 판단 2026-08-25).
                    MARKS.append(dict(id=marksio.mid("cross","C",pi,table=ti,row=x["row"],col=x["col"]),
                        seq=SEQ, page=pi, kind="cross",
                        box=marksio.box4(*bb), text=_txt,
                        source=dict(check="C", table=ti, row=x["row"], col=x["col"], label=x.get("label")),
                        verdict=None, evidence=dict(val=x.get("val"), other=x.get("other"),
                                                    diff=x["diff"], where=x["where"],
                                                    reason="REF_NOT_FOUND"),
                        account=x.get("label"), level="L1",
                        shown_value=(f"{x['val']:,.0f}" if x.get("val") is not None else None),
                        computed_value=None,
                        delta=SKIP_REASON_TEXT["REF_NOT_FOUND"],
                        formula=f"{FORMULA_TEXT['C']} — 가장 가까운 금액 {x['where']}",
                        operands=[], counterparts=[],
                        l2_class=None, column_key=None, paper_no=None,
                        comment=None, verified_at=RUN_TS, reviewed_at=None, reviewed_by=None,
                        origin="tool", status="pending", note=None)); SEQ+=1
            for r in rs:
                v=verdict(r,TOL,RSTEPS); stat[v]+=1
                # '1원차이' 태그 — 원 단위 표에서 |차이|<=1이면 일괄 확인용 표시.
                # 판정은 바꾸지 않는다 (원 단위는 반올림이 없어 흡수 금지 — 회계사 확인 대상)
                tg = ("부호규약" if v == "SIGN" else
                      "1원차이" if v == "DIFF" and abs(r["calc"]-r["disp"]) <= 1
                      and (TUNIT.get((pi,ti)) or "").strip() == "원" else "")
                gi=r.get("row"); bbox=None
                if gi is not None:
                    gi-=off
                    if 0<=gi<len(t.rows):
                        cells=t.rows[gi].cells; cj=r.get("col")
                        if cj is not None and cj<len(cells) and cells[cj]: bbox=cells[cj]
                        else:
                            for cc in reversed(cells):
                                if cc: bbox=cc; break
                if bbox:
                    x0,t0,x1,b0=bbox
                    _kind = {"OK":"tick","ROUND":"tick","DIFF":"cross"}.get(v, "question")
                    _text = f"{r['calc']-r['disp']:+,.0f}" if v == "DIFF" else None
                    _col_for_id = r.get("col")
                    if _col_for_id is None:
                        _a2key = (pi, ti, r.get("row"))
                        _col_for_id = f"x{A2_OCC[_a2key]}"; A2_OCC[_a2key] += 1
                    # ── rev.2 화면 필드 ────────────────────────────────────
                    # 원문 표기: A2는 col이 없어(마크 위치는 행 마지막 셀) tcol을 쓴다.
                    _vc = r.get("col") if r.get("col") is not None else r.get("tcol")
                    _shown = None
                    if _vc is not None and 0 <= r.get("row", -1) < len(data):
                        _row_cells = data[r["row"]]
                        if _vc < len(_row_cells): _shown = norm(_row_cells[_vc])
                    # 미검증 사유 — 코드가 없으면 조용히 넘기지 않고 표면화한다
                    if v in ("SKIP","SIGN"):
                        _rc = "SIGN_CONVENTION" if v == "SIGN" else r.get("reason")
                        _why = SKIP_REASON_TEXT.get(_rc) or f"UNRESOLVED_REASON({_rc})"
                        # 사유에 구체적인 행·항목 이름을 붙인다. "산식에 포함되지 않은
                        # 행이 있습니다"만으로는 회계사가 무엇을 볼지 알 수 없다.
                        if r.get("reason_detail"): _why = f"{_why}: {r['reason_detail']}"
                    else:
                        _rc = _why = None
                    _fml = FORMULA_TEXT.get(r["kind"], r["kind"])
                    if v in ("SKIP","SIGN"): _fml = f"{_fml} — {_why}"
                    elif r.get("n"):         _fml = f"{_fml} — 구성 항목 {r['n']}개"
                    # operands: 합계에 들어간 셀. ★마크가 아니라 좌표다 — 성분 셀에는
                    # 마크가 없고, 만들면 새 드로잉이라 R-1이 깨진다(설계안 4절 승인).
                    _ops=[]
                    for _o in (r.get("operands") or []):
                        # ★ +off를 붙이지 않는다. operands는 본체 마크(r["row"])와 같은
                        # 병합 기준 인덱스인데, cellbox가 안에서 -off를 하므로 +off를
                        # 붙이면 환산이 상쇄되어 지면 범위를 벗어난다. 분할 병합 표
                        # (off>0)에서 성분이 조용히 버려지고(조선내화 p62: 21→7) 살아남은
                        # 것도 엉뚱한 행을 가리켰다. cirmap/tagmap은 지면 기준 인덱스라
                        # +off가 맞지만 여기는 아니다.
                        _ob = cellbox(_o["row"], _o.get("col"))
                        # 좌표가 없어도 버리지 않는다. 분할 병합 표에서 앞 페이지에 있는
                        # 성분은 이 지면에 셀이 없을 뿐, 계산에는 들어간 항목이다 —
                        # 버리면 성분 합이 재계산금액을 재현하지 못해 회계사가 판정할 수
                        # 없다(조선내화 p62: 21개 중 14개가 앞 페이지). bbox=null로 두고
                        # 화면은 그 항목의 하이라이트만 건너뛴다.
                        _otxt = None
                        if 0 <= _o["row"] < len(data):
                            _orow = data[_o["row"]]
                            if _o.get("col") is not None and _o["col"] < len(_orow):
                                _otxt = norm(_orow[_o["col"]])
                        _olab = None
                        if 0 <= _o["row"] < len(data):
                            _olab = next((norm(c) for c in data[_o["row"]] if norm(c)), None)
                        _ops.append(dict(label=_olab, value=_otxt, sign=_o.get("sign", "+"),
                                         bbox=(marksio.bbox4(marksio.box4(*_ob)) if _ob else None)))
                    MARKS.append(dict(
                        id=marksio.mid(_kind, r["kind"], pi, table=ti, row=r.get("row"), col=_col_for_id),
                        seq=SEQ, page=pi, kind=_kind,
                        box=marksio.box4(x0,t0,x1,b0), text=_text,
                        source=dict(check=r["kind"], table=ti, row=r.get("row"), col=r.get("col"),
                                   label=r.get("label")),
                        verdict=v,
                        evidence=dict(disp=r["disp"], calc=r["calc"], diff=r["calc"]-r["disp"], n=r["n"],
                                     unit=unit or "미표기", tag=(tg or None), reason=_rc),
                        account=r.get("label"), level="L1",
                        shown_value=_shown,
                        computed_value=(None if v in ("SKIP","SIGN") else fmt_like(r["calc"], _shown)),
                        delta=(_why if v in ("SKIP","SIGN") else f"{r['calc']-r['disp']:+,.0f}"),
                        formula=_fml, operands=_ops,
                        counterparts=[],           # A계열은 대사 대상 없음 확정([]≠null)
                        l2_class=None, column_key=None, paper_no=None,
                        comment=None, verified_at=RUN_TS, reviewed_at=None, reviewed_by=None,
                        origin="tool", status="pending", note=None)); SEQ+=1
                if v in ("DIFF","ROUND","SKIP","SIGN"):
                    exc.append([pi,ti,r["kind"],r["label"],unit or "미표기",
                                r["disp"],r["calc"],r["calc"]-r["disp"],r["n"],
                                {"DIFF":"차이","ROUND":"단수차이","SKIP":"미검증","SIGN":"미검증"}[v],tg])
        try:
            _last=tobjs[-1].extract()
            _,_,_,_,ncL,_,_=grid_info(_last)
            carry=(_last,ncL,ends_open(_last),pi)
        except Exception: carry=None
        PAGE_SEQ[pi]=SEQ                 # L2 마크를 이 페이지 뒤에 이어 붙일 때 쓴다
    _NPAGES = pi
    # ── 목차 + 마크 귀속 (판정 이후 후처리 — 판정에 영향 없음) ──
    SECTIONS = docmeta.sections(pdf, _HEADS, STMT_PAGES)
    for _m in MARKS:
        _m["section_id"] = docmeta.section_of(SECTIONS, _m["page"], _m["box"]["top"])
        _tb = (_m.get("source") or {}).get("table")
        if _tb is None:
            _m["table_label"] = None
        elif STMT_PAGES.get(_m["page"]):
            # 본표는 표 위에 제목 문장이 없다(있는 건 회사명·단위 머리글이다) — 캡션
            # 추출을 아예 쓰지 않고 목차의 본표 이름을 쓴다.
            _m["table_label"] = next((s["label"] for s in SECTIONS
                                      if s["id"] == STMT_PAGES[_m["page"]]), None)
        else:
            _lab = TABLE_LABEL.get((_m["page"], _tb))
            _sid = _m.get("section_id") or ""
            _m["table_label"] = (docmeta.with_note(_lab, _sid[1:])
                                 if (_lab and _sid.startswith("n")) else _lab)

    # ── L2 표간 대사 편입 ──────────────────────────────────────────────
    # 회사 사전이 이 문서에 적용될 때만 돌린다 — 적용 안 되면 결과가 전량 UNMAPPED이라
    # PDF를 한 번 더 파싱하는 비용만 버린다(등록 4축이 여기 해당해 실행시간·드로잉 불변).
    L2_MARKS=[]; L2_RES=None
    try:
        import l2_probe, l2_labels, l2_extract, l2_marks
        _co = l2_probe._company_from_filename(PDF)
        _lp = os.path.join(os.path.dirname(os.path.abspath(__file__)), "labels", f"{_co}.json") if _co else None
        if _lp and os.path.isfile(_lp) and l2_probe._dict_applies(_lp, PDF):
            L2_RES = l2_labels.build(PDF, _co, l2_extract.extract(PDF))
            def _l2_label(side):
                # 본표는 table_id 자체가 유형이다 — 페이지로 찾으면 분할 2면차(p7 등)에서
                # stmt_type이 None이라 이름을 놓친다.
                _st = l2_marks.stmt_of(side["table_id"])
                if _st:
                    return next((s["label"] for s in SECTIONS if s["id"] == _st), None)
                _l = (TABLE_LABEL.get((side["page"], side["table_seq"]))
                      if side.get("table_seq") is not None else None)
                _mn = re.match(r"n(\d+)", side["table_id"] or "")
                return docmeta.with_note(_l, _mn.group(1)) if (_l and _mn) else _l
            for _cls, _rows in (("confirmed", L2_RES["confirmed"]),
                                ("undeclared", L2_RES["undeclared"])):
                L2_MARKS += l2_marks.build(_rows, _cls, _l2_label, PAGE_SEQ, RUN_TS,
                                           marksio.mid, marksio.box4, marksio.bbox4)
            for _m in L2_MARKS:
                _m["section_id"] = docmeta.section_of(SECTIONS, _m["page"], _m["box"]["top"])
            MARKS += L2_MARKS
    except Exception as _e:
        print(f"[L2] 편입 실패(분석 결과에는 영향 없음): {_e}")

_T_ANALYZE = time.time() - _T0

B,cons = tieout.run(PDF)
decl,refs,miss,unref,gap = notes.run(PDF)

RUN_META = {"commit": VER, "at": RUN_TS,
            "params": {"tol": TOL, "round_steps": RSTEPS, "min_won": MINWON,
                       "terms": (_cli.terms or None)}}
DOCUMENT = {"title": _TITLE, "title_from_filename": _TITLE_FALLBACK,
            "last_verified_at": RUN_TS,
            "sections": [{k: s[k] for k in ("id","label","page","depth")} for s in SECTIONS],
            # 표 이름을 못 뽑은 표 — 빈 문자열로 흘리지 않고 여기로 표면화한다
            "unresolved_labels": UNRESOLVED_LABEL}
MARKS_DOC = marksio.save(MARKS_JSON, PDF, MARKS, SPLIT_SUSPECT, RUN_META, _NPAGES, DOCUMENT)
renderer.render_all(PDF, MARKS_DOC, OUT_PDF, quiet=QUIET)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
wb=openpyxl.Workbook()
def sheet(name, hdr, rows, widths):
    ws=wb.create_sheet(name); ws.append(hdr)
    for cc in ws[1]:
        cc.font=Font(bold=True,color="FFFFFF"); cc.fill=PatternFill("solid",fgColor="404040")
        cc.alignment=Alignment(horizontal="center")
    for r in rows: ws.append(r)
    for i,wd in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=wd
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    return ws
wb.remove(wb.active)
ws=sheet("A_예외색인",["페이지","표","유형","항목","단위","표시금액","계산금액","차이","항목수","판정","태그"],
         exc,[8,6,8,30,10,18,18,16,8,10,10])
for row in ws.iter_rows(min_row=2):
    for cc in row[5:8]: cc.number_format="#,##0"
    if row[9].value=="차이": row[9].fill=PatternFill("solid",fgColor="FFC7CE")
    elif row[9].value=="미검증": row[9].fill=PatternFill("solid",fgColor="FFEB9C")
wsb=sheet("B_본표간연계",["번호","검증항목","값1","값2","판정"],
          [[a,b,c_,d_,e] for a,b,c_,d_,e,_ in B],[8,52,18,18,10])
for row in wsb.iter_rows(min_row=2):
    for cc in row[2:4]: cc.number_format="#,##0"
    if row[4].value=="차이": row[4].fill=PatternFill("solid",fgColor="FFC7CE")
    elif row[4].value=="미검증": row[4].fill=PatternFill("solid",fgColor="FFEB9C")
sheet("C_본표주석레퍼",["구분","페이지","항목","금액","대상","비고"],
      [["성립",m["page"],m["label"],m["val"],
        " ".join(sorted({f"FN{n}" for c_ in cs for n in (c_["notes"] & m["refs"])})),
        f"주석 p{sorted({c_['page'] for c_ in cs})}"] for m,cs in RLINKS]
    + [["미성립",m["page"],m["label"],m["val"],
        " ".join(f"FN{n}" for n in sorted(m["refs"])),"주석에서 동일 금액 미발견"] for m,_ in RUN]
    + [["단위제외",m["page"],m["label"],m["val"],
        " ".join(f"FN{n}" for n in sorted(m["refs"])),"복합·판독 불가 단위 — 대사 미수행(?)"] for m in REXCL],
      [10,8,28,18,16,30])
_frows=[]
for p_ in FCON["F1_단위누락"]:
    _frows.append(["F1 단위누락", p_, "-", "금액 표가 있으나 (단위: ) 표기 없음", ""])
for g_, forms in FCON["F2_표현불일치"]:
    _pref = [f_ for f_ in forms if f_ in TERMS]       # 기준 표현 지정 시 빈도보다 우선
    if _pref:
        dom = (_pref[0], forms[_pref[0]]); _why = f"기준 표현 '{dom[0]}'(지정)과 불일치"
    else:
        dom = max(forms.items(), key=lambda x: x[1][0])
        _why = f"주된 표현 '{dom[0]}'({dom[1][0]}회)과 불일치"
    for f_, (n_, ps_) in sorted(forms.items(), key=lambda x: -x[1][0]):
        if f_ == dom[0]: continue
        _frows.append(["F2 표현불일치", ps_[0] if ps_ else "-", f_,
                       f"{_why} — {n_}회",
                       " ".join(f"p{x}" for x in ps_)])
for k_, forms, tot in FCON["F3_라벨불일치"]:
    dom = max(forms.items(), key=lambda x: x[1][0])
    for f_, (n_, ps_) in sorted(forms.items(), key=lambda x: -x[1][0]):
        if f_ == dom[0]: continue
        _frows.append(["F3 라벨불일치", ps_[0] if ps_ else "-", f_,
                       f"주된 표기 '{dom[0]}'({dom[1][0]}회)과 불일치 — {n_}회",
                       " ".join(f"p{x}" for x in ps_)])
for g_ in FCON["F4_다중공백"]:
    _frows.append(["F4 다중공백", g_["page"], f"{g_['before']} | {g_['after']}",
                   "어절 간격이 통상보다 넓음", ""])
sheet("F_일관성",["유형","페이지","대상","내용","출현위치"],_frows,[14,8,26,44,30])
sheet("D_줄글표기",["페이지","지적어절","내용","신뢰도"],pros,[8,20,40,10])
sheet("C7_주석참조",["구분","내용"],
      [["선언된 주석 수",len(decl)],["본표 참조 수",len(refs)],
       ["참조됐으나 주석 없음",str(miss or "없음")],
       ["주석 있으나 본표 미참조",str(unref or "없음")],
       ["주석번호 결번",str(gap or "없음")]],[26,60])
tot=sum(stat.values())
sheet("요약",["항목","값"],
      [["A 산술검증 총건수",tot],["  OK",stat['OK']],["  단수차이(ROUND)",stat['ROUND']],
       ["  차이(DIFF)",stat['DIFF']],["  미검증(SKIP)",stat['SKIP']],
       ["  OK 비율",f"{stat['OK']/tot*100:.1f}%" if tot else "-"],
       ["B 연계검증 총건수",len(B)],["  OK",sum(1 for r in B if r[4]=='OK')],
       ["  차이",sum(1 for r in B if r[4]=='차이')],["  미검증",sum(1 for r in B if r[4]=='미검증')],
       ["D 검토완료(/) 문단 수",npara],["D 표기 지적 건수",len(pros)],
       ["허용오차(A7)",TOL],["round_steps(표 단위 스텝)",RSTEPS],
       ["min_won(C 소액 제외, 원)",MINWON],["연결 구조 감지",str(cons)],
       ["단위 미표기 페이지",str(nounit or "없음")],
       ["도구 버전(커밋)",VER],["실행 일시",RUN_TS],
       ["분할 의심 페이지 전환(참고)",
        (f"{len(SPLIT_SUSPECT)}건 — " +
         ", ".join(f"p{s['prev']}→p{s['page']}({s['reason']})" for s in SPLIT_SUSPECT))
        if SPLIT_SUSPECT else "없음"],
       ["범위","표시 수치 상호 정합성 한정. 원장·조서 대사는 별도 절차."]],[30,60])
wb.save(OUT_XLSX)
if not QUIET:
    print(f"A 산술 {tot}건 → OK {stat['OK']} ({stat['OK']/tot*100:.1f}%) / ROUND {stat['ROUND']} / DIFF {stat['DIFF']} / SKIP {stat['SKIP']} / SIGN {stat['SIGN']}")
    print(f"B 연계 {len(B)}건 → OK {sum(1 for r in B if r[4]=='OK')} / 차이 {sum(1 for r in B if r[4]=='차이')} / 미검증 {sum(1 for r in B if r[4]=='미검증')}")
    print(f"C 레퍼 → 성립 {len(RLINKS)} / 미성립 {len(RUN)} / 차이 {len(RDIFF)}" +
          (f" · 단위제외 {len(REXCL)}건" if REXCL else ""))
    print(f"F 일관성 → 단위누락 {len(FCON['F1_단위누락'])}p / 표현불일치 {len(FCON['F2_표현불일치'])}그룹 / 라벨불일치 {len(FCON['F3_라벨불일치'])}건 / 다중공백 {len(FCON['F4_다중공백'])}건")
    print(f"D 줄글 → 검토완료(/) {npara}문단 / 표기 지적 {len(pros)}건")
    print(f"C7 주석 → 결번 {gap or '없음'} / 참조무주석 {miss or '없음'}")
    _sub_pre = sum(1 for e in SUBLOG if e["path"]=="접두")
    _sub_ind = [e for e in SUBLOG if e["path"].startswith("들여쓰기")]
    print(f"하위항목 인식 → 총 {len(SUBLOG)}건 (접두 {_sub_pre} / 들여쓰기 {len(_sub_ind)})")
    for e in _sub_ind:
        print(f"  [들여쓰기] p{e['ctx'][0]} 표{e['ctx'][1]} 행{e['row']} '{e['label']}' ← 상위 '{e['parent']}' ({e['path']})")
    print(f"연결 감지 {cons} · 단위 미표기 {nounit or '없음'} · 허용오차 ±{TOL:g}")
    print(f"분할 의심 페이지 전환(참고) → {len(SPLIT_SUSPECT)}건")
    print(f"분석 소요: {_T_ANALYZE:.2f}s (마크 {len(MARKS)}개, marks.json: {MARKS_JSON})")
print(f"산출물: {OUT_PDF} · {OUT_XLSX}")

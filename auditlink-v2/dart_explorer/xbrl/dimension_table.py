"""D-2: 차원 표 렌더러 (이미지1 형식 — DART 뷰어식 배치).

- 열: 축(axis)의 member 조합 (축 2개면 2단 중첩 헤더) + 마지막 [합계]
  = 차원 미지정(default member) 팩트
- 행: LineItems (pre.xml order, 들여쓰기)
- 기간별 블록 분리 "(기간: 당기)/(전기)", 음수는 빨간 괄호 (숫자서식)
- def.xml: all → hypercube-dimension → dimension-domain → domain-member
  아크롤 체인 (dimension-member 표기 수용)

DART 구조 반영 (삼성전자 실측):
- 연결/별도 축(ConsolidatedAndSeparateFinancialStatementsAxis)은 열이 아니라
  role별 필터 — 정의 문자열의 "- 연결/- 별도"로 member를 선택해 팩트를 거른다
- 주석 role 하나가 def에서 a/b/c… 서브롤(하이퍼큐브 단위)로 분할됨
  → 한 시트에 하이퍼큐브별 섹션 테이블을 연속 렌더
스코프 제외: typed dimension.
"""
import collections
import os
import re

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .taxonomy import LB, XL, TaxonomyPackage, _pkg_glob, _sheet_name

XI = "http://www.xbrl.org/2003/instance"
XD = "http://xbrl.org/2006/xbrldi"

FILTER_AXIS = "ifrs-full_ConsolidatedAndSeparateFinancialStatementsAxis"
_MEMBER_CONsolidated = "ifrs-full_ConsolidatedMember"
_MEMBER_SEPARATE = "ifrs-full_SeparateMember"

_BOLD = Font(bold=True)
_HDR_FILL = PatternFill("solid", start_color="D9E1F2")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_NUMFMT = "#,##0;[RED](#,##0)"              # 음수 = 빨간 괄호
_NUMFMT_DEC = "#,##0.00;[RED](#,##0.00)"

_PERIOD_NAMES = ["당기", "전기", "전전기"]


def _qid(text):
    return (text or "").strip().replace(":", "_")


def _role_code(uri):
    m = re.search(r"role[-_]([A-Za-z0-9]+)$", uri or "")
    return m.group(1) if m else (uri or "")


def _base_code(code):
    return re.sub(r"[a-z]+$", "", code)


def _try_num(val):
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# 인스턴스
# ---------------------------------------------------------------------------

class XbrlInstance:
    """팩트·컨텍스트 (explicit dimension만)."""

    def __init__(self, folder):
        path = _pkg_glob(folder, "*.xbrl")
        if not path:
            raise FileNotFoundError(f".xbrl 인스턴스 없음: {folder}")
        root = etree.parse(path[0]).getroot()
        self.contexts = {}
        for c in root.findall(f"{{{XI}}}context"):
            p = c.find(f"{{{XI}}}period")
            inst = p.find(f"{{{XI}}}instant")
            if inst is not None:
                period = ("instant", inst.text, inst.text)
            else:
                s = p.find(f"{{{XI}}}startDate")
                e = p.find(f"{{{XI}}}endDate")
                period = ("duration",
                          s.text if s is not None else "",
                          e.text if e is not None else "")
            dims = {}
            for m in c.findall(f".//{{{XD}}}explicitMember"):
                dims[_qid(m.get("dimension"))] = _qid(m.text)
            self.contexts[c.get("id")] = {
                "type": period[0], "start": period[1], "end": period[2],
                "dims": dims,
            }
        self.facts = collections.defaultdict(list)
        for el in root:
            if not isinstance(el.tag, str):
                continue
            cref = el.get("contextRef")
            if cref is None or cref not in self.contexts:
                continue
            val = (el.text or "").strip()
            if not val:
                continue
            qn = etree.QName(el.tag)
            cid = f"{el.prefix}_{qn.localname}" if el.prefix else qn.localname
            self.facts[cid].append({
                "ctx": self.contexts[cref], "value": val,
                "decimals": el.get("decimals", ""),
            })


# ---------------------------------------------------------------------------
# def.xml 하이퍼큐브
# ---------------------------------------------------------------------------

class HypercubeDef:
    """definitionLink별 하이퍼큐브. 주석 서브롤(a/b/c…)을 기본 코드로 묶는다."""

    def __init__(self, folder):
        self.by_base = collections.defaultdict(list)
        paths = _pkg_glob(folder, "*_def.xml")
        if not paths:
            return
        root = etree.parse(paths[0]).getroot()
        for dl in root.findall(f"{{{LB}}}definitionLink"):
            role_uri = dl.get(f"{{{XL}}}role")
            loc = {l.get(f"{{{XL}}}label"):
                   (l.get(f"{{{XL}}}href") or "").split("#")[-1]
                   for l in dl.findall(f"{{{LB}}}loc")}
            arcs = collections.defaultdict(list)
            for a in dl.findall(f"{{{LB}}}definitionArc"):
                kind = (a.get(f"{{{XL}}}arcrole") or "").rsplit("/", 1)[-1]
                if kind == "dimension-member":
                    kind = "domain-member"
                arcs[kind].append((loc.get(a.get(f"{{{XL}}}from")),
                                   loc.get(a.get(f"{{{XL}}}to")),
                                   float(a.get("order", "0") or 0)))
            if not arcs["all"]:
                continue
            children = collections.defaultdict(list)
            for src, dst, order in arcs["domain-member"]:
                children[src].append((order, dst))

            def closure(node):
                out, stack = [], [node]
                seen = {node}
                while stack:
                    cur = stack.pop(0)
                    out.append(cur)
                    for _, ch in sorted(children.get(cur, [])):
                        if ch not in seen:
                            seen.add(ch)
                            stack.append(ch)
                return out

            for li_node, table, _ in sorted(arcs["all"], key=lambda t: t[2]):
                axes = []
                for src, axis, order in sorted(arcs["hypercube-dimension"],
                                               key=lambda t: t[2]):
                    if src != table:
                        continue
                    domains = [d for s, d, _ in arcs["dimension-domain"]
                               if s == axis]
                    default = next((d for s, d, _ in
                                    arcs["dimension-default"] if s == axis),
                                   None)
                    members = []
                    for dom in domains:
                        members += [m for m in closure(dom)[1:]
                                    if m not in members]
                    axes.append({"axis": axis, "domains": domains,
                                 "default": default, "members": members})
                code = _role_code(role_uri)
                self.by_base[_base_code(code)].append({
                    "role_uri": role_uri, "code": code, "axes": axes,
                    "primary": set(closure(li_node)),
                })


# ---------------------------------------------------------------------------
# Role 테이블 모델 (섹션 = 하이퍼큐브 1개)
# ---------------------------------------------------------------------------

def _structural_concepts(cubes):
    out = set()
    for cube in cubes:
        for ax in cube["axes"]:
            out.add(ax["axis"])
            out.update(ax["domains"])
            out.update(ax["members"])
            if ax["default"]:
                out.add(ax["default"])
    return out


class RoleTable:
    def __init__(self, pkg: TaxonomyPackage, inst: XbrlInstance,
                 cube: HypercubeDef, role_uri, definition, pl):
        self.pkg = pkg
        self.definition = definition
        self.cubes = cube.by_base.get(_base_code(_role_code(role_uri)), [])
        structural = _structural_concepts(self.cubes)

        pre_rows = []
        for r in pkg.tree_rows(pl):
            cid = f"{r['prefix']}_{r['id']}" if r["prefix"] else r["id"]
            if cid.endswith(("Table", "Axis")) or cid in structural:
                continue
            pre_rows.append((r, cid))

        # 연결/별도 필터 member (role 정의 문자열 기준)
        self.filter_member = None
        if "별도" in definition or "Separate" in definition:
            self.filter_member = _MEMBER_SEPARATE
        elif "연결" in definition or "Consolidated" in definition:
            self.filter_member = _MEMBER_CONsolidated

        def keep(f):
            m = f["ctx"]["dims"].get(FILTER_AXIS)
            return m is None or self.filter_member is None or \
                m == self.filter_member

        facts = {cid: [f for f in inst.facts.get(cid, []) if keep(f)]
                 for _, cid in pre_rows}

        # 섹션 구성: 하이퍼큐브별 (없으면 전체 1섹션·축0)
        self.sections = []
        if self.cubes:
            for cb in self.cubes:
                rows = [(r, cid) for r, cid in pre_rows
                        if cid in cb["primary"]]
                if not rows:
                    continue
                sec = _Section(self, rows, facts, cb)
                if sec.has_facts:
                    self.sections.append(sec)
        if not self.sections:
            sec = _Section(self, pre_rows, facts, None)
            if sec.has_facts:
                self.sections.append(sec)

    @property
    def has_facts(self):
        return bool(self.sections)

    def member_label(self, member):
        return self.pkg._label(self.pkg.labels_ko, member)


class _Section:
    """하이퍼큐브 1개 = 표 1개."""

    def __init__(self, parent: RoleTable, rows, all_facts, cube):
        self.parent = parent
        self.rows = rows
        self.facts = {cid: all_facts.get(cid, []) for _, cid in rows}
        self.has_facts = any(self.facts.values())

        axes_def = [ax for ax in (cube["axes"] if cube else [])
                    if ax["axis"] != FILTER_AXIS]
        self.axes = []
        for ax in axes_def:
            observed = {f["ctx"]["dims"].get(ax["axis"])
                        for fl in self.facts.values() for f in fl}
            observed.discard(None)
            used = [m for m in ax["members"] if m in observed]
            if not ax["members"]:
                # 큐브가 member 미선언 시에만 관측값 사용
                # (선언된 경우 타 하이퍼큐브 member 유입 차단)
                used = sorted(observed)
            if used:
                self.axes.append({**ax, "members": used})
        self.axis_ids = [ax["axis"] for ax in self.axes]

        periods = {}
        for fl in self.facts.values():
            for f in fl:
                c = f["ctx"]
                periods[(c["type"], c["start"], c["end"])] = True
        durs = sorted([p for p in periods if p[0] == "duration"],
                      key=lambda p: p[1])
        durs.sort(key=lambda p: p[2], reverse=True)
        insts = sorted([p for p in periods if p[0] == "instant"],
                       key=lambda p: p[2], reverse=True)
        self.block_kind = "duration" if durs else "instant"
        self.blocks = durs if durs else insts
        self.inst_cols = insts
        if self.axes:
            def has_member_fact(block):
                return any(
                    (f["ctx"]["type"], f["ctx"]["start"],
                     f["ctx"]["end"]) == block
                    and any(f["ctx"]["dims"].get(a) for a in self.axis_ids)
                    for fl in self.facts.values() for f in fl)
            self.blocks = [b for b in self.blocks if has_member_fact(b)]

    # --- 값 조회 -----------------------------------------------------------
    def _dims_match(self, ctx, combo):
        proj = {a: ctx["dims"].get(a) for a in self.axis_ids}
        want = {a: None for a in self.axis_ids}
        want.update(combo)
        return proj == want

    def _in_block(self, ctx, block):
        btype, bstart, bend = block
        if ctx["type"] == btype and (ctx["start"], ctx["end"]) == \
                (bstart, bend):
            return True
        if btype == "duration" and ctx["type"] == "instant":
            return ctx["end"] == bend
        return False

    def cell(self, concept, block, combo):
        for f in self.facts.get(concept, []):
            if self._in_block(f["ctx"], block) and \
                    self._dims_match(f["ctx"], combo):
                return f
        return None

    def block_label(self, block):
        ends = sorted({b[2] for b in self.blocks}, reverse=True)
        rank = ends.index(block[2])
        name = _PERIOD_NAMES[rank] if rank < 3 else f"{rank + 1}기 전"
        if block[0] == "duration":
            return f"(기간: {name} {block[1]} ~ {block[2]})"
        return f"(기준일: {name}말 {block[2]})"


# ---------------------------------------------------------------------------
# 엑셀 렌더링
# ---------------------------------------------------------------------------

def _write_cell(ws, r, c, fact):
    if fact is None:
        return False
    num = _try_num(fact["value"])
    cell = ws.cell(row=r, column=c)
    if num is not None:
        cell.value = num
        dec = fact.get("decimals", "")
        try:
            frac = dec not in ("", "0", "INF") and int(float(dec)) > 0
        except ValueError:
            frac = False
        cell.number_format = _NUMFMT_DEC if frac else _NUMFMT
    else:
        text = fact["value"]
        cell.value = text[:100] + ("…" if len(text) > 100 else "")
    return True


def _hdr(ws, r, c, value):
    cell = ws.cell(row=r, column=c, value=value)
    cell.font = _BOLD
    cell.fill = _HDR_FILL
    cell.alignment = _CENTER
    return cell


def _render_section_flat(ws, sec: _Section, r):
    """축 0개: 행=element, 열=기간."""
    cols = list(dict.fromkeys(
        ([b for b in sec.blocks] if sec.block_kind == "duration" else [])
        + sec.inst_cols))
    labels = []
    di = ii = 0
    for c in cols:
        if c[0] == "duration":
            nm = _PERIOD_NAMES[di] if di < 3 else f"{di + 1}기 전"
            labels.append(f"{nm}\n{c[1]}~{c[2]}")
            di += 1
        else:
            nm = (_PERIOD_NAMES[ii] + "말") if ii < 3 else ""
            labels.append(f"{nm}\n{c[2]}")
            ii += 1
    _hdr(ws, r, 1, "계정과목")
    for j, lab in enumerate(labels):
        _hdr(ws, r, 2 + j, lab)
    r += 1
    for row, concept in sec.rows:
        ws.cell(row=r, column=1, value="    " * row["depth"] + row["ko"])
        for j, c in enumerate(cols):
            _write_cell(ws, r, 2 + j, sec.cell(concept, c, {}))
        r += 1
    return r + 1


def _render_section_dim(ws, sec: _Section, r):
    """축 1~2개: 기간 블록 × (member 조합 + [합계])."""
    for block in sec.blocks:
        ws.cell(row=r, column=1, value=sec.block_label(block)).font = _BOLD
        r += 1
        combos = [{}]
        for ax in sec.axes:
            combos = [dict(c, **{ax["axis"]: m}) for c in combos
                      for m in ax["members"]]
        combos = [c for c in combos
                  if any(sec.cell(concept, block, c)
                         for _, concept in sec.rows)]
        if not combos:
            r += 1
            continue
        nhdr = len(sec.axes)
        for ai, ax in enumerate(sec.axes):
            ws.cell(row=r + ai, column=1,
                    value="계정과목" if ai == nhdr - 1 else "")
            if ai == nhdr - 1:
                ws.cell(row=r + ai, column=1).font = _BOLD
            col = 2
            j = 0
            while j < len(combos):
                m = combos[j][ax["axis"]]
                span = 1
                while j + span < len(combos) and \
                        combos[j + span][ax["axis"]] == m and \
                        all(combos[j + span][a["axis"]] ==
                            combos[j][a["axis"]]
                            for a in sec.axes[:ai]):
                    span += 1
                _hdr(ws, r + ai, col, sec.parent.member_label(m))
                if span > 1:
                    ws.merge_cells(start_row=r + ai, start_column=col,
                                   end_row=r + ai, end_column=col + span - 1)
                col += span
                j += span
            if ai == nhdr - 1:
                _hdr(ws, r + ai, 2 + len(combos), "[합계]")
        r += nhdr
        for row, concept in sec.rows:
            ws.cell(row=r, column=1,
                    value="    " * row["depth"] + row["ko"])
            for j, combo in enumerate(combos):
                _write_cell(ws, r, 2 + j, sec.cell(concept, block, combo))
            _write_cell(ws, r, 2 + len(combos),
                        sec.cell(concept, block, {}))
            r += 1
        r += 1
    return r


def render_dimension_tables(folder, out_path=None, role_filter=None):
    """패키지 폴더 → Role별 차원 표 엑셀. 요약 dict 반환."""
    pkg = TaxonomyPackage(folder)
    inst = XbrlInstance(folder)
    cube = HypercubeDef(folder)
    wb = Workbook()
    wb.remove(wb.active)
    rendered = []
    used_names = set()
    for role_uri, definition, pl in pkg.roles():
        if role_filter and role_filter not in definition and \
                role_filter not in role_uri:
            continue
        table = RoleTable(pkg, inst, cube, role_uri, definition, pl)
        if not table.has_facts:
            continue
        ws = wb.create_sheet(_sheet_name(definition, used_names))
        ws.append([definition])
        ws.cell(row=1, column=1).font = _BOLD
        r = 3
        for sec in table.sections:
            if sec.axes:
                r = _render_section_dim(ws, sec, r)
            else:
                r = _render_section_flat(ws, sec, r)
        ws.column_dimensions["A"].width = 52
        for ci in range(2, 14):
            ws.column_dimensions[get_column_letter(ci)].width = 17
        ws.freeze_panes = "B2"
        rendered.append({
            "definition": definition,
            "sections": len(table.sections),
            "axes": max((len(s.axes) for s in table.sections), default=0),
            "rows": sum(len(s.rows) for s in table.sections),
        })
    if not rendered:
        raise ValueError("렌더링할 Role이 없습니다 (필터 확인).")
    if out_path is None:
        out_path = os.path.join(folder, "차원표.xlsx")
    wb.save(out_path)
    return {"out_path": out_path, "roles": rendered}

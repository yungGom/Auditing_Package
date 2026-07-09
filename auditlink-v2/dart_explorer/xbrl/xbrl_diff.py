"""
XBRL 변화사항(버전 diff) 추출기  (로컬 실행 전용)
사용법:
  (1) 기간 비교 :  python xbrl_diff.py <당기.xbrl>
        └ 파일 안의 당기 vs 전기를 비교 (재무상태표=당기말vs전기말, 손익=당반기vs전년동기)
  (2) 버전 비교 :  python xbrl_diff.py <전기.xbrl> <당기.xbrl>
        └ 두 파일의 각 당기값을 요소별로 대사 (재작성·정정 대사, 신규/삭제 계정 탐지)
필요: pip install lxml pandas openpyxl  (같은 폴더에 xbrl_extract.py 필요)
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xbrl_extract import extract

def key(r):  # 계정 식별 키 (날짜 제외)
    return (r['요소ID'], r['연결별도'], r['차원'], r['기간구분'])

def pct(prev, cur):
    if prev in (None, 0) or cur is None: return None
    return round((cur - prev) / abs(prev) * 100, 1)

def period_compare(rows):
    """한 파일 안에서 기간구분별로 당기(최신종료일) vs 전기(차순위) 비교"""
    out = []
    for pt in ('instant', 'duration'):
        sub = [r for r in rows if r['기간구분'] == pt and r['금액'] is not None]
        dates = sorted({r['종료일'] for r in sub}, reverse=True)
        if len(dates) < 2:
            continue
        cur_d, prev_d = dates[0], dates[1]
        cur = {key(r): r for r in sub if r['종료일'] == cur_d}
        prev = {key(r): r for r in sub if r['종료일'] == prev_d}
        for k in cur.keys() | prev.keys():
            c, p = cur.get(k), prev.get(k)
            base = c or p
            cv = c['금액'] if c else None
            pv = p['금액'] if p else None
            note = '' if (c and p) else ('신규' if c else '제거')
            if '(사용중단)' in (base['한글명칭'] or ''): note = (note + ' 사용중단').strip()
            out.append(dict(한글명칭=base['한글명칭'], 요소ID=base['요소ID'],
                연결별도=base['연결별도'], 차원=base['차원'], 기간구분=pt,
                전기=pv, 당기=cv,
                증감액=((cv or 0) - (pv or 0)) if (c and p) else None,
                증감율_pct=pct(pv, cv) if (c and p) else None, 비고=note))
    return out

def _current_only(rows):
    """버전비교용: 각 파일에서 기간구분별 최신 종료일 값만 추출"""
    res = {}
    for pt in ('instant', 'duration'):
        sub = [r for r in rows if r['기간구분'] == pt and r['금액'] is not None]
        if not sub: continue
        cur_d = max(r['종료일'] for r in sub)
        for r in sub:
            if r['종료일'] == cur_d:
                res[key(r)] = r
    return res

def version_compare(rows_old, rows_new):
    a, b = _current_only(rows_old), _current_only(rows_new)
    out = []
    for k in a.keys() | b.keys():
        ra, rb = a.get(k), b.get(k)
        base = rb or ra
        av = ra['금액'] if ra else None
        bv = rb['금액'] if rb else None
        note = '' if (ra and rb) else ('당기파일 신규' if rb else '당기파일 삭제')
        if ra and rb and av != bv: note = '값 변동'
        if '(사용중단)' in (base['한글명칭'] or ''): note = (note + ' 사용중단').strip()
        out.append(dict(한글명칭=base['한글명칭'], 요소ID=base['요소ID'],
            연결별도=base['연결별도'], 차원=base['차원'], 기간구분=base['기간구분'],
            전기파일=av, 당기파일=bv,
            증감액=((bv or 0) - (av or 0)) if (ra and rb) else None,
            증감율_pct=pct(av, bv) if (ra and rb) else None, 비고=note))
    return out

def write_excel(rows, out, valcols):
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    df = pd.DataFrame(rows)
    # 변동 큰 순 + 비고 있는 것 우선 정렬
    df['_abs'] = df['증감액'].abs().fillna(-1)
    df = df.sort_values(['비고', '_abs'], ascending=[False, False]).drop(columns='_abs')
    changed = df[(df['비고'] != '') | (df['증감액'].fillna(0) != 0)]
    with pd.ExcelWriter(out, engine='openpyxl') as w:
        changed.to_excel(w, sheet_name='변화사항', index=False)
        df.to_excel(w, sheet_name='전체대사', index=False)
    wb = load_workbook(out)
    hdr = PatternFill('solid', start_color='1F4E78'); hf = Font(name='맑은 고딕', bold=True, color='FFFFFF')
    warn = PatternFill('solid', start_color='FCE4D6'); newf = PatternFill('solid', start_color='FFF2CC')
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        cols = {c.value: c.column_letter for c in ws[1]}
        for c in ws[1]:
            c.fill = hdr; c.font = hf; c.alignment = Alignment(horizontal='center')
        for col in ws.columns:
            L = col[0].column_letter
            mx = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[L].width = min(max(mx + 2, 10), 44)
            for c in col[1:]:
                c.font = Font(name='맑은 고딕')
                if isinstance(c.value, (int, float)): c.number_format = '#,##0;(#,##0)'
        # 비고 강조
        bcol = cols.get('비고')
        if bcol:
            for row in ws.iter_rows(min_row=2):
                v = ws[f'{bcol}{row[0].row}'].value or ''
                if '사용중단' in v or '삭제' in v or '제거' in v:
                    for c in row: c.fill = warn
                elif '신규' in v:
                    for c in row: c.fill = newf
    wb.save(out)

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    if len(sys.argv) == 2:
        rows = period_compare(extract(sys.argv[1]))
        out = sys.argv[1].rsplit('.', 1)[0] + '_당기전기비교.xlsx'
        mode = '기간비교(당기 vs 전기)'
    else:
        rows = version_compare(extract(sys.argv[1]), extract(sys.argv[2]))
        out = sys.argv[2].rsplit('.', 1)[0] + '_버전비교.xlsx'
        mode = '버전비교(전기파일 vs 당기파일)'
    write_excel(rows, out, None)
    print(f'[{mode}] 완료: {len(rows)}개 계정 -> {out}')

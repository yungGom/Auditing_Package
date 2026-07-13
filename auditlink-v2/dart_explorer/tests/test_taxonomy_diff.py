"""D-4 게이트: 택사노미 버전 관리 (감지·대조·리포트).

1. 실파일 2세대 taxdiff → 기대값(신설 1,334 / 폐지 156 / 유지 7,652) 근사 일치
2. 홀드아웃 회사 1곳(삼성전자 전기 XBRL) taxcheck → 폐지 검출 + 대체 후보 제시
3. D-4d: 코퍼스 확장 중 신설 표준과 고유사도 매칭 사례 5건 이상
"""
import os

import pytest

from dart_explorer.xbrl.taxonomy_diff import (TAXONOMY_DIR, detect_version,
                                              get_promotions,
                                              parse_presentation_concepts,
                                              resolve_version_dir,
                                              run_taxcheck, taxdiff)

_HAS_2GEN = (os.path.isdir(os.path.join(TAXONOMY_DIR, "2024-06-30")) and
            os.path.isdir(os.path.join(TAXONOMY_DIR, "2026-01-31")))
_SAMSUNG_PRIOR = os.path.join(
    os.path.dirname(__file__), "..", "cache", "xbrl", "00126380",
    "20250311001085_11011")


# ---------------------------------------------------------------------------
# 헤더 탐지 파서 — 세대 간 열 배치 차이 회귀 방지
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not _HAS_2GEN, reason="택사노미 2세대 미배치")
def test_detect_version():
    assert detect_version(resolve_version_dir("2024-06-30")) == "2024-06-30"
    assert detect_version(resolve_version_dir("2026-01-31")) == "2026-01-31"


@pytest.mark.skipif(not _HAS_2GEN, reason="택사노미 2세대 미배치")
def test_parse_concepts_ko_label_not_english():
    """한글 라벨 열 판정이 값(한글 포함 여부) 기준임을 보증 — 열 위치 함정 회귀."""
    concepts = parse_presentation_concepts(resolve_version_dir("2024-06-30"))
    c = concepts["ifrs-full_CurrentAssets"]
    assert c["label_ko"] == "유동자산"
    assert c["label_en"] and "current" in c["label_en"].lower()
    assert not any(ch.isascii() and ch.isalpha() for ch in c["label_ko"])


# ---------------------------------------------------------------------------
# 게이트 1: taxdiff 기대값 근사 일치
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not _HAS_2GEN, reason="택사노미 2세대 미배치")
def test_gate1_taxdiff_expected_counts(tmp_path):
    res = taxdiff("2024-06-30", "2026-01-31", str(tmp_path / "diff.xlsx"))
    assert res["old_version"] == "2024-06-30"
    assert res["new_version"] == "2026-01-31"
    assert len(res["added"]) == 1334          # 정확 일치 (실측)
    assert len(res["removed"]) == 156          # 정확 일치 (실측)
    assert abs(res["kept"] - 7652) <= 200       # 근사 일치 (±소량 오차 허용)
    assert os.path.exists(res["out_path"])

    from openpyxl import load_workbook
    wb = load_workbook(res["out_path"])
    assert set(wb.sheetnames) == {"신설", "폐지", "변경"}
    assert wb["신설"].max_row - 1 == 1334
    assert wb["폐지"].max_row - 1 == 156


# ---------------------------------------------------------------------------
# 게이트 2: 홀드아웃 회사(삼성전자) 전기 XBRL taxcheck
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not (_HAS_2GEN and os.path.isdir(_SAMSUNG_PRIOR)),
                    reason="택사노미 2세대 또는 전기 XBRL 캐시 미확보")
def test_gate2_taxcheck_deprecation_and_replacement(tmp_path):
    res = run_taxcheck(_SAMSUNG_PRIOR, "2026-01-31",
                       out_path=str(tmp_path / "taxcheck.xlsx"))
    assert res["yellow"] > 0, "폐지 element 미검출"
    assert res["green"] > res["yellow"]         # 대부분은 그대로 사용 가능

    yellow_rows = [r for r in res["rows"] if r["status"] == "노랑"]
    assert all(r["candidates"] for r in yellow_rows), \
        "폐지 element에 대체 후보 미부여"
    assert all(len(r["candidates"]) <= 4 for r in yellow_rows)

    from openpyxl import load_workbook
    wb = load_workbook(res["out_path"])
    assert "체크리스트" in wb.sheetnames


# ---------------------------------------------------------------------------
# 게이트 3: D-4d 확장→표준 승격 감지 (5건 이상 고유사도 매칭)
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not _HAS_2GEN, reason="택사노미 2세대 미배치")
def test_gate3_promotion_detection(tmp_path):
    from dsd_tool.mapping import MappingCorpus
    corpus = MappingCorpus()
    if not corpus.extensions:
        pytest.skip("코퍼스 확장 데이터 없음")
    promos = get_promotions("2024-06-30", "2026-01-31", corpus,
                            min_similarity=0.55)
    assert len(promos) >= 5, f"고유사도 매칭 {len(promos)}건 (목표 5건+)"
    for p in promos[:5]:
        assert p["similarity"] >= 0.55
        assert p["new_element_id"] and p["ext_label"]


@pytest.mark.skipif(not (_HAS_2GEN and os.path.isdir(_SAMSUNG_PRIOR)),
                    reason="택사노미 2세대 또는 전기 XBRL 캐시 미확보")
def test_taxcheck_integrates_promotions(tmp_path):
    """D-4c 리포트에 D-4d 승격후보가 통합되는지 (taxonomies/에 2세대뿐일 때
    자동으로 구버전을 채택)."""
    res = run_taxcheck(_SAMSUNG_PRIOR, "2026-01-31",
                       out_path=str(tmp_path / "taxcheck2.xlsx"))
    assert res["promotions"], "승격 후보가 체크리스트에 통합되지 않음"
    from openpyxl import load_workbook
    wb = load_workbook(res["out_path"])
    assert "확장→표준 승격후보" in wb.sheetnames


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not _HAS_2GEN, reason="택사노미 2세대 미배치")
def test_taxdiff_cli(tmp_path, capsys):
    from dart_explorer.__main__ import main
    out = str(tmp_path / "cli_diff.xlsx")
    assert main(["taxdiff", "2024-06-30", "2026-01-31", "-o", out]) == 0
    text = capsys.readouterr().out
    assert "신설 1334" in text and "폐지 156" in text

"""A-5b 게이트: 전기대사 (삼성전자 FY2025 ↔ FY2024 실파일 페어).

1. 본문 전기열 전수 대사 — 대부분 TRUE 기대
2. 주석 제목 매칭률 + 표 대사 (최초 실행 = 벤치마크)
3. 변조 테스트: 전기 사본 당기값 1개 +1 → 해당 셀만 FALSE
4. 모드 ② 호환성 판정은 RECON_모드2_호환성.md 문서로 회신
"""
import os
import re
import shutil

import pytest
from openpyxl import load_workbook

from dsd_tool.excel_out import extract
from dsd_tool.recon import GUIDE, recon

_REAL = os.path.join(os.path.dirname(__file__), "..", "fixtures", "real")
_CUR = os.path.join(_REAL, "[삼성전자(주)]_2025_[감사보고서].dsd")
_PRI = os.path.join(_REAL, "[삼성전자(주)]_2024_[감사보고서].dsd")

pytestmark = pytest.mark.skipif(
    not (os.path.exists(_CUR) and os.path.exists(_PRI)),
    reason="삼성 FY2025/FY2024 페어 없음")


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("recon")
    return recon(_CUR, _PRI, out_path=str(tmp / "전기대사.xlsx"))


# ---------------------------------------------------------------------------
# 게이트 1: 본문 전수 대사 — 삼성은 재작성 없음 → 전부 TRUE
# ---------------------------------------------------------------------------

def test_gate1_statements_all_true(result):
    s = result["stmt"]
    assert s["n"] >= 90                        # BS/PL/PL1/CE/CF 전수(행 단위)
    assert s["false"] == 0, [
        (sheet, r["label"], r["note"])
        for sheet, rows in result["stmt_results"].items()
        for r in rows if not r["true"]]
    # CE 행 단위(전기 블록 행 × 자본항목 열 전체 일치) 대사 포함 확인
    ce = result["stmt_results"].get("CE", [])
    assert len(ce) >= 5
    assert all(r["true"] for r in ce)
    assert all(r.get("pri_row") for r in ce)   # 전기 파일 행과 실제 매칭됨


# ---------------------------------------------------------------------------
# 게이트 2: 주석 제목 매칭 + 표 대사 벤치마크
# ---------------------------------------------------------------------------

def test_gate2_note_false_breakdown(result):
    """확정 기준(2026-07-10): TRUE율은 게이트가 아니다 — 정당한 FALSE를
    찾는 게 목적. 게이트 = 값 상이 중 매칭 오류 오탐 0건 + FALSE 전수 분해."""
    assert result["note_matched"] / result["note_total"] >= 0.9
    n = result["notes"]
    assert n["n"] >= 100                       # 표 셀 대조가 실제 수행됨

    bd = result["breakdown"]
    assert bd["mismatch_flags"] == 0           # ★ 게이트: 매칭 오류 오탐 0건
    # ① 항목 없음: 목록화 (삼성 실측 — FY2025 신규 공시 행)
    assert bd["missing"] >= 1
    # ③ 표 쌍 매칭 실패: 원인별 분류 존재
    assert bd["unpaired"] >= 1
    # ④ 제목 매칭 실패 1건: FY2024에 없는 신규 주석('공정가치 측정')
    assert bd["title_misses"] == result["note_total"] - \
        result["note_matched"]

    # FALSE분해 시트: 값 상이 건은 당기/전기 셀 주소·값·열 헤더 병기
    wb = load_workbook(result["out_path"])
    assert "FALSE분해" in wb.sheetnames
    ws = wb["FALSE분해"]
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("항목 없음" in t for t in text)
    assert any("표 쌍 매칭 실패" in t for t in text)
    assert any("제목 매칭 실패" in t for t in text)
    assert any("당기 셀" in t for t in text) and \
        any("전기 열 헤더" in t for t in text)


# ---------------------------------------------------------------------------
# 게이트 3: 변조 테스트 — 전기 사본 당기값 1개 +1 → 해당 셀만 FALSE
# ---------------------------------------------------------------------------

def test_gate3_tamper_detection(result, tmp_path):
    cur_x = str(tmp_path / "당기.xlsx")
    pri_x = str(tmp_path / "전기.xlsx")
    extract(_CUR, cur_x)
    extract(_PRI, pri_x)

    base = recon(cur_x, pri_x, out_path=str(tmp_path / "기준.xlsx"))
    assert base["stmt"]["false"] == 0

    # 전기 파일 BS 당기 열에서 '현금및현금성자산' 값 +1 변조
    wb = load_workbook(pri_x)
    ws = wb["BS"]
    target = None
    for row in ws.iter_rows():
        label = str(row[0].value or "")
        if "현금및현금성자산" in label:
            for c in row[1:]:
                if isinstance(c.value, (int, float)):
                    c.value = c.value + 1      # 당기 열(첫 숫자 열)
                    target = label
                    break
            break
    assert target
    wb.save(pri_x)

    tampered = recon(cur_x, pri_x, out_path=str(tmp_path / "변조.xlsx"))
    false_rows = [(sheet, r["label"], r["note"])
                  for sheet, rows in tampered["stmt_results"].items()
                  for r in rows if not r["true"]]
    assert len(false_rows) == 1                # 해당 셀만 FALSE
    sheet, label, note = false_rows[0]
    assert sheet == "BS" and "현금및현금성자산" in label
    assert "값 상이" in note and "±1" in note

    # 출력 엑셀: 요약 카운트 + 안내문 (상세는 실무 양식 병렬 시트)
    wb2 = load_workbook(tampered["out_path"])
    ws0 = wb2["요약"]
    text = "\n".join(str(c.value) for row in ws0.iter_rows()
                     for c in row if c.value)
    assert GUIDE in text
    assert "FALSE 1" in text.replace("FALSE  ", "FALSE ")


# ---------------------------------------------------------------------------
# 출력 형식 (A-5a 톤 통일)
# ---------------------------------------------------------------------------

def test_output_format(result):
    """실무 양식(문용.xlsb 배치): 당기 원문 | 판정 수식 | 전기 원문 병렬.

    판정 규칙(확정): 셀 참조 수식 =A1=B1 (복수 셀 =AND(...)), 전기
    대응 행 없으면 리터럴 "FALSE" — 파이썬 bool 금지.
    """
    wb = load_workbook(result["out_path"])
    assert wb.sheetnames[0] == "요약"
    # 시트 구성 = 실무 양식 (원문 배치 그대로 BS/PL/…/주석 번호)
    assert {"BS", "PL", "CE", "CF"} <= set(wb.sheetnames)
    assert any(n.isdigit() for n in wb.sheetnames)

    pat = re.compile(
        r"^=(?:AND\()?[A-Z]+\d+=[A-Z]+\d+(?:,[A-Z]+\d+=[A-Z]+\d+)*\)?$")
    for sheet in ("BS", "PL", "CE", "CF"):
        ws = wb[sheet]
        vcol = next((c.column for c in ws[1] if c.value == "판정"), None)
        assert vcol, f"{sheet}: 판정 열 없음"
        vals = [row[0].value for row in
                ws.iter_rows(min_row=2, min_col=vcol, max_col=vcol)
                if row[0].value is not None]
        assert vals, f"{sheet}: 판정 값 없음"
        assert all(isinstance(v, str) and (pat.match(v) or v == "FALSE")
                   for v in vals), \
            f"{sheet}: 수식/FALSE 외 판정 값 {vals[:3]}"
        assert any(pat.match(str(v)) for v in vals), f"{sheet}: 수식 없음"

    # 좌측 당기 원문 + 우측 전기 원문 병렬 (같은 행에 같은 계정 라벨)
    bs = wb["BS"]
    found_parallel = False
    for row in bs.iter_rows(min_row=5):
        vals = [(c.column, str(c.value)) for c in row if c.value is not None]
        labels = [v for _, v in vals if "현금및현금성자산" in v]
        if len(labels) >= 2:
            found_parallel = True
            break
    assert found_parallel, "좌우 병렬 원문 배치 아님"

    # 서식 규격 (문용.xlsb 실측): 숫자 음수 괄호 서식 + 데이터 영역
    # thin 테두리 + 표 헤더 회색(DCDCDC) — TRUE는 무강조, FALSE만 빨강
    num_cells = [c for row in bs.iter_rows() for c in row
                 if isinstance(c.value, (int, float))]
    assert num_cells
    assert all(c.number_format == "#,##0;(#,##0)" for c in num_cells)
    assert all(c.border.left.style == "thin" for c in num_cells)
    fills = {str(c.fill.start_color.rgb) for row in bs.iter_rows()
             for c in row if c.fill and c.fill.patternType == "solid"}
    assert any(f.endswith("DCDCDC") for f in fills), "표 헤더 회색 없음"
    from dsd_tool.recon import _eval_verdict
    for row in bs.iter_rows():
        for c in row:
            v = _eval_verdict(bs, c.value) if c.value is not None else None
            if v is True:
                assert c.fill.patternType is None, "TRUE는 무강조"
            elif v is False:
                assert str(c.fill.start_color.rgb).endswith("FFC7CE")

    # 요약 → 시트 하이퍼링크 + FALSE분해 시트
    ws0 = wb["요약"]
    links = [str(c.value) for row in ws0.iter_rows() for c in row
             if c.value and str(c.value).startswith("=HYPERLINK")]
    assert any('#\'BS\'' in ln or "#BS" in ln for ln in links)
    assert "FALSE분해" in wb.sheetnames


# ---------------------------------------------------------------------------
# 회귀: 요약 카운트 == 시트별 판정 실측 카운트 (전 시트, 재발 방지)
# ---------------------------------------------------------------------------

def test_summary_equals_rendered_verdicts(result):
    """생성된 엑셀을 다시 읽어 요약 표의 (대사, TRUE, FALSE)가 각 시트에
    실제 렌더된 판정 열 실측(수식 참조 셀 평가)과 일치하는지 전 시트 검증."""
    wb = load_workbook(result["out_path"])
    ws0 = wb["요약"]

    # 요약 표 파싱: 헤더 [구분, 시트, 대사, TRUE, FALSE, ...] 이후 행들
    summary_rows = {}
    in_table = False
    for row in ws0.iter_rows(values_only=True):
        if row[:2] == ("구분", "시트"):
            in_table = True
            continue
        if in_table:
            if not row[0]:
                break
            summary_rows[str(row[1])] = (row[2], row[3], row[4])
    assert summary_rows, "요약 표 없음"
    assert {"BS", "PL", "CE", "CF"} <= set(summary_rows)

    from dsd_tool.recon import count_rendered_verdicts
    for sheet, (n, t, f) in summary_rows.items():
        rn, rt, rf = count_rendered_verdicts(wb[sheet])
        assert (n, t, f) == (rn, rt, rf), \
            f"{sheet}: 요약 {(n, t, f)} != 실측 {(rn, rt, rf)}"

    # 요약 상단 전체 판정 문구의 합계도 실측 총합과 일치
    total = [str(c.value) for row in ws0.iter_rows() for c in row
             if c.value and "전체 판정" in str(c.value)][0]
    m = re.search(r"본문 대사 (\d+) · TRUE (\d+) · FALSE (\d+) / "
                  r"주석 대사 (\d+) · TRUE (\d+) · FALSE (\d+)", total)
    assert m
    sn, st, sf = (sum(v[i] for s, v in summary_rows.items()
                      if s in ("BS", "PL", "PL1", "CE", "CF"))
                  for i in range(3))
    assert (int(m.group(1)), int(m.group(2)), int(m.group(3))) == \
        (sn, st, sf)

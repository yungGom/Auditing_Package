"""DART 편집기 버전 관리 게이트 (패치 A-3).

게이트: fixtures/real 전체에서 editver 정확 추출,
가짜 버전 픽스처로 미확인 경고 발동 확인.
"""
import glob
import os

import pytest
from openpyxl import load_workbook

from dsd_tool.excel_out import extract
from dsd_tool.version import (is_known, known_versions, read_version_info,
                              update_known_versions)

from .fixture import build_dsd

_REAL_DIR = os.path.join(os.path.dirname(__file__), "..", "fixtures", "real")
_REAL_FILES = sorted(glob.glob(os.path.join(_REAL_DIR, "*.dsd")))


# ---------------------------------------------------------------------------
# 게이트 1: 실파일 전체에서 editver 정확 추출
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not _REAL_FILES, reason="실제 DSD 없음")
def test_editver_extraction_all_real_files():
    for path in _REAL_FILES:
        with open(path, "rb") as f:
            info = read_version_info(f.read())
        assert info["editver"] == "5.049", os.path.basename(path)
        assert info["schema"] == "dart4.xsd"
        assert info["docver"] in ("4.1", "3.5")


@pytest.mark.skipif(not _REAL_FILES, reason="실제 DSD 없음")
def test_editver_in_extract_info_and_meta_sheet(tmp_path):
    xlsx = str(tmp_path / "v.xlsx")
    info = extract(_REAL_FILES[0], xlsx)
    assert info["editver"] == "5.049"
    assert info["editver_known"] is True   # KNOWN_VERSIONS.md에 등재됨
    meta = {r[0]: r[1] for r in
            load_workbook(xlsx)["_META"].iter_rows(values_only=True)}
    assert meta["editver"] == "5.049"
    assert meta["schema"] == "dart4.xsd"


# ---------------------------------------------------------------------------
# 게이트 2: 미확인 editver → 경고 발동
# ---------------------------------------------------------------------------

def test_unknown_editver_warning(tmp_path, capsys):
    fake = build_dsd(str(tmp_path / "가짜버전.dsd"), editver="9.999")
    info = extract(fake, str(tmp_path / "가짜버전.xlsx"))
    assert info["editver"] == "9.999"
    assert info["editver_known"] is False

    from dsd_tool.cli import main
    fake2 = build_dsd(str(tmp_path / "가짜버전2.dsd"), editver="9.999")
    assert main(["extract", fake2, "-o", str(tmp_path / "f2.xlsx")]) == 0
    out = capsys.readouterr().out
    assert "미확인 DART 편집기 버전" in out
    assert "batch_validate" in out          # 갱신 안내 문구
    assert "KNOWN_VERSIONS.md" in out


def test_known_editver_no_warning(tmp_path, capsys):
    dsd = build_dsd(str(tmp_path / "정상.dsd"))     # editver=5.049
    from dsd_tool.cli import main
    assert main(["extract", dsd, "-o", str(tmp_path / "n.xlsx")]) == 0
    out = capsys.readouterr().out
    assert "editver=5.049" in out
    assert "미확인" not in out


def test_meta_missing(tmp_path):
    """meta.xml 없는 DSD: editver None + 미확인 처리 (크래시 없음)."""
    import io
    import zipfile
    from .fixture import CONTENTS_XML
    p = str(tmp_path / "nometa.dsd")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("contents.xml", CONTENTS_XML.encode("utf-8"))
    open(p, "wb").write(buf.getvalue())
    info = extract(p, str(tmp_path / "nometa.xlsx"))
    assert info["editver"] is None
    assert info["editver_known"] is False


# ---------------------------------------------------------------------------
# KNOWN_VERSIONS.md 파싱·갱신
# ---------------------------------------------------------------------------

def test_known_versions_registry():
    vers = known_versions()
    assert "5.049" in vers                  # batch_validate가 등재함
    assert is_known("5.049") and not is_known("9.999") and not is_known(None)


def test_update_known_versions_merge(tmp_path, monkeypatch):
    import dsd_tool.version as V
    fake_md = str(tmp_path / "KNOWN_VERSIONS.md")
    monkeypatch.setattr(V, "known_versions_path", lambda: fake_md)
    V.update_known_versions({"5.049": {"files": 12, "g2_pass": True}})
    V.update_known_versions({"5.106": {"files": 1, "g2_pass": False}})
    doc = open(fake_md, encoding="utf-8").read()
    assert "| 5.049 | 12 | PASS |" in doc   # 기존 행 유지
    assert "| 5.106 | 1 | FAIL |" in doc    # 신규 병합
    assert V.known_versions() == {"5.049", "5.106"}

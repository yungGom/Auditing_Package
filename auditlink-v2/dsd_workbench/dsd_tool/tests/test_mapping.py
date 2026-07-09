"""D-3b: 계정과목 → element 매핑 추천 테스트.

블라인드 20개 적중률 게이트(80%+)는 전량 코퍼스 완료 후 별도 측정 —
여기서는 스코어 요소·판정 규칙·입출력 파이프라인을 검증한다.
"""
import os
import sqlite3

import pytest

from dsd_tool import mapping as M


# ---------------------------------------------------------------------------
# 유사도 단위 테스트 (한글 특화 규칙)
# ---------------------------------------------------------------------------

def test_jamo_similarity():
    assert M.jamo_similarity("매출채권", "매출채권") == 1.0
    assert M.jamo_similarity("매출채권", "매출체권") > 0.85   # 오타 흡수
    assert M.jamo_similarity("매출채권", "현금") < 0.3


def test_antonym_penalty():
    """단기차입금 vs 장기차입금: 자모 1글자 차이지만 반의 개념 — 페널티."""
    anto = M.similarity(M.normalize("단기차입금"), "단기차입금", "장기차입금")
    syno = M.similarity(M.normalize("단기차입금"), "단기차입금",
                        "유동 차입금(사채 포함)")
    assert syno > anto                          # 동의 변형이 반의어를 이김
    assert anto < 0.45


def test_synonym_variants():
    assert "유동차입금" in M._query_variants("단기차입금")
    assert "비유동차입금" in M._query_variants("장기차입금")
    # 비유동 안의 '유동'을 잘못 치환하지 않음
    assert "단기부채" not in M._query_variants("비유동부채")


def test_normalize():
    assert M.normalize("Ⅰ. 유동자산 (당기)") == "유동자산"
    assert M.normalize("자본금 [구성요소]") == "자본금"


def test_role_score():
    assert M._role_score("BS", {"D210005"}) == 1.0
    assert M._role_score("CF", {"D210005"}) == 0.0
    assert M._role_score(None, {"D210005"}) == 0.5
    assert M._role_score("주석", {"D822100"}) == 1.0


# ---------------------------------------------------------------------------
# 가짜 코퍼스로 조회·판정·출력 검증
# ---------------------------------------------------------------------------

@pytest.fixture()
def fake_db(tmp_path):
    db = str(tmp_path / "corpus.sqlite")
    con = sqlite3.connect(db)
    con.executescript("""
    CREATE TABLE companies(corp_code TEXT PRIMARY KEY, corp_name TEXT,
      stock_code TEXT, induty_code TEXT, year INT, rcept_no TEXT,
      report_nm TEXT, status TEXT, error TEXT, fetched_at TEXT);
    CREATE TABLE usages(corp_code TEXT, element_id TEXT, is_ext INT,
      label_ko TEXT, roles TEXT, dims TEXT, n_facts INT,
      PRIMARY KEY(corp_code, element_id));
    CREATE TABLE extensions(corp_code TEXT, element_id TEXT, label_ko TEXT,
      roles TEXT, dims TEXT, n_facts INT,
      PRIMARY KEY(corp_code, element_id));
    CREATE TABLE standard_labels(element_id TEXT PRIMARY KEY,
      label_ko TEXT, label_en TEXT);
    """)
    for i in range(1, 11):                     # 10사: 매출채권 (업종 264는 3사)
        code = f"{i:08d}"
        con.execute("INSERT INTO companies VALUES(?,?,?,?,?,?,?,?,?,?)",
                    (code, f"회사{i}", "0", "264" if i <= 3 else "111",
                     2025, None, None, "ok", None, ""))
        con.execute("INSERT INTO usages VALUES(?,?,?,?,?,?,?)",
                    (code, "ifrs-full_TradeReceivables", 0, "매출채권",
                     "D210005", "", 2))
    con.execute("INSERT INTO usages VALUES(?,?,?,?,?,?,?)",
                ("00000001", "dart-gcd_AuthorName", 0, "매출채권",
                 "D000000", "", 1))            # gcd 메타 — 제외돼야 함
    con.execute("INSERT INTO usages VALUES(?,?,?,?,?,?,?)",
                ("00000001", "ifrs-full_CashAndCashEquivalents", 0,
                 "현금및현금성자산", "D210005", "", 2))
    con.execute("INSERT INTO extensions VALUES(?,?,?,?,?,?)",
                ("00000001", "entity1_SpecialBondGain",
                 "전환사채 발행 특별이익", "D210005", "", 1))
    con.execute("INSERT INTO standard_labels VALUES(?,?,?)",
                ("ifrs-full_TradeReceivables", "매출채권", "Trade receivables"))
    con.commit()
    con.close()
    return db


def test_suggest_and_verdict(fake_db):
    corpus = M.MappingCorpus(fake_db, refs_path="(없음)")
    assert "dart-gcd_AuthorName" not in corpus.elements   # 메타 제외

    r = M.suggest(corpus, "매출채권", "BS", induty="264")
    assert r["verdict"] == "후보"
    top = r["candidates"][0]
    assert top["element_id"] == "ifrs-full_TradeReceivables"
    assert top["sim"] == 1.0
    assert top["n_same_induty"] == 3            # 동업종 가중 근거
    assert "동업종 3사" in top["evidence"]

    # 무관 계정 → 적합 표준 없음 + 유사 확장 사례
    r2 = M.suggest(corpus, "전환사채발행특별이익", "BS")
    assert r2["verdict"] == "적합 표준 없음"
    assert r2["similar_extensions"]
    assert r2["similar_extensions"][0]["label"] == "전환사채 발행 특별이익"


def test_map_accounts_excel(fake_db, tmp_path):
    from openpyxl import load_workbook
    tpl = str(tmp_path / "입력.xlsx")
    M.write_template(tpl)
    wb = load_workbook(tpl)
    ws = wb.active
    ws.delete_rows(2)                           # 예시 행 제거
    ws.append(["매출채권", "BS", ""])
    ws.append(["이상한계정명", "주석", ""])
    wb.save(tpl)

    res = M.map_accounts(tpl, out_path=str(tmp_path / "out.xlsx"),
                         db_path=fake_db)
    assert res["items"] == 2 and res["no_match"] == 1
    out = load_workbook(res["out_path"]).active
    rows = list(out.iter_rows(values_only=True))
    assert rows[0][:3] == ("계정과목명", "구분", "판정")
    body = [r for r in rows[1:] if any(r)]
    assert any(r[4] == "ifrs-full_TradeReceivables" and r[9] == "☐"
               for r in body)                   # 선택 열은 빈 체크박스
    assert any(r[2] == "적합 표준 없음" for r in body)


def test_map_cli(fake_db, tmp_path, capsys):
    from dsd_tool.cli import main
    tpl = str(tmp_path / "t.xlsx")
    assert main(["map", "--template", tpl]) == 0
    assert os.path.exists(tpl)
    assert main(["map", tpl, "--corpus", fake_db,
                 "-o", str(tmp_path / "o.xlsx")]) == 0
    out = capsys.readouterr().out
    assert "매핑 후보 생성" in out and "자동 확정 없음" in out

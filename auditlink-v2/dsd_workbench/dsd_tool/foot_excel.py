"""A-5a: Footing 결과 엑셀 — AI_Footing 형식 재현.

예시 파일 실측 (AI_Footing 5종):
- 시트 구성: 총괄표 | 원문 | BS/PL/CE/CF | 1..N (주석)
- 총괄표: 좌측 [시트|제목] (하이퍼링크+원문 제목), 우측
  [시트|시트 링크|푸팅 오류|크로스 오류|총합(=I{r}+J{r} 수식)]
- 각 시트 M2에 총괄표 복귀 하이퍼링크
- 오류 셀 표시: 노란 채우기 (예시 5종 중 2종에서 실측 — 비일관하므로
  노랑 컨벤션만 계승, 위치는 A-4 검증 결과의 정확 좌표 사용, 상세는 메모)

용어 매핑: 푸팅 오류 = A-4 합계검증(단수차+불일치),
크로스 오류 = 주석대사 미매칭(못찾음).
"""
import os
import re

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .excel_out import MAP_SHEET, META_SHEET
from .foot import FOOT_SHEET, FUZZY, MISMATCH

_BOLD = Font(bold=True)
_HDR_FILL = PatternFill("solid", start_color="D9E1F2")
_YELLOW = PatternFill("solid", start_color="FFFFFF00")   # 예시 실측 색
_LINK_FONT = Font(color="FF0563C1", underline="single")

_SUMMARY_LEFT_HDR = ["시트", "제목"]
_SUMMARY_RIGHT_HDR = ["시트", "시트 링크", "푸팅 오류", "크로스 오류", "총합"]


def _sheet_title(ws):
    """시트 원문 제목: FS는 A1 제목행, 주석은 'N. 제목'의 제목부."""
    v = str(ws.cell(1, 1).value or "").strip()
    m = re.match(r"^\d+\.\s*(.+)$", v)
    return m.group(1).strip() if m else v


def _quote(name):
    return f"'{name}'" if not name.isalnum() else name


def write_ai_footing(xlsx_path, foot_result, out_path=None,
                     company_name=None):
    """foot() 결과 → AI_Footing_{회사명}.xlsx 독립 파일. 요약 dict 반환."""
    wb = load_workbook(xlsx_path)
    data_sheets = [s for s in wb.sheetnames
                   if s not in (MAP_SHEET, META_SHEET, FOOT_SHEET)]

    # --- 오류 집계 (시트별) -------------------------------------------------
    foot_err = {}
    for r in foot_result["foot"]:
        if r["verdict"] in (FUZZY, MISMATCH):
            foot_err.setdefault(r["sheet"], []).append(r)
    cross_err = {}
    for r in foot_result["notes"]:
        if not r["found"]:
            cross_err.setdefault(r["sheet"], []).append(r)

    # --- 오류 셀 표시 (노란 채우기 + 메모 — A-4 정확 좌표) ------------------
    for sheet, errs in foot_err.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for e in errs:
            m = re.match(r"R(\d+)", str(e["loc"]))
            if not m:
                continue
            row = int(m.group(1))
            cell = ws.cell(row=row, column=1)
            cell.fill = _YELLOW
            note = (f"[푸팅 {e['verdict']}] Σ자식 {e['expected']:,.0f} vs "
                    f"기재 {e['actual']:,.0f} (차이 {e['diff']:,.0f}) — "
                    f"{e['scope']}/{e['direction']}")
            cell.comment = Comment(note, "dsd_tool foot")
    for sheet, errs in cross_err.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for e in errs:
            cell = ws.cell(row=e["row"], column=1)
            cell.fill = _YELLOW
            cell.comment = Comment(
                f"[크로스 미매칭] 주석 {e['refs']} 에서 값 "
                f"{e['value']:,.0f} ({e['period']}) 미발견", "dsd_tool foot")

    # --- 각 시트 복귀 링크 (M2 — 예시 위치 그대로) ---------------------------
    for sheet in data_sheets:
        ws = wb[sheet]
        c = ws.cell(row=2, column=13)
        c.value = '=HYPERLINK("#총괄표!A1","총괄표로 이동")'
        c.font = _LINK_FONT

    # --- 원문 시트 (전체 통합 뷰 — 시트 내용 세로 연결) ----------------------
    src_order = list(data_sheets)
    ws_all = wb.create_sheet("원문", 0)
    out_row = 1
    for sheet in src_order:
        ws = wb[sheet]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or cell.column > 12:
                    continue
                tgt = ws_all.cell(row=out_row + cell.row - 1,
                                  column=cell.column)
                tgt.value = cell.value
                if cell.fill and cell.fill.patternType:
                    tgt.fill = PatternFill(
                        cell.fill.patternType,
                        start_color=cell.fill.start_color,
                        end_color=cell.fill.end_color)
        out_row += ws.max_row + 2

    # --- 총괄표 -------------------------------------------------------------
    ws0 = wb.create_sheet("총괄표", 0)
    ws0.cell(row=2, column=2, value="시트정보").font = _BOLD
    ws0.cell(row=2, column=7, value="푸팅/크로스 에러 사항").font = _BOLD
    for j, h in enumerate(_SUMMARY_LEFT_HDR):
        c = ws0.cell(row=4, column=2 + j, value=h)
        c.font = _BOLD
        c.fill = _HDR_FILL
    for j, h in enumerate(_SUMMARY_RIGHT_HDR):
        c = ws0.cell(row=4, column=7 + j, value=h)
        c.font = _BOLD
        c.fill = _HDR_FILL

    left = [("원문", _sheet_title(wb[src_order[0]]) if src_order else "")]
    left += [(s, _sheet_title(wb[s])) for s in src_order]
    for i, (name, title) in enumerate(left):
        r = 5 + i
        c = ws0.cell(row=r, column=2,
                     value=f'=HYPERLINK("#{_quote(name)}!A1","{name}")')
        c.font = _LINK_FONT
        ws0.cell(row=r, column=3, value=title)

    total_foot = total_cross = 0
    for i, name in enumerate(src_order):
        r = 5 + i
        nf = len(foot_err.get(name, []))
        nc = len(cross_err.get(name, []))
        total_foot += nf
        total_cross += nc
        ws0.cell(row=r, column=7, value=name)
        c = ws0.cell(row=r, column=8,
                     value=f'=HYPERLINK("#{_quote(name)}!$A$1",'
                           f'"시트 바로가기")')
        c.font = _LINK_FONT
        ws0.cell(row=r, column=9, value=float(nf))
        ws0.cell(row=r, column=10, value=float(nc))
        ws0.cell(row=r, column=11, value=f"=총괄표!I{r}+총괄표!J{r}")
    ws0.column_dimensions["C"].width = 40
    for col in ("B", "G", "H"):
        ws0.column_dimensions[col].width = 16

    # --- 저장 ---------------------------------------------------------------
    if MAP_SHEET in wb.sheetnames:
        del wb[MAP_SHEET]
    if META_SHEET in wb.sheetnames:
        del wb[META_SHEET]
    if FOOT_SHEET in wb.sheetnames:
        del wb[FOOT_SHEET]
    if company_name is None:
        company_name = re.sub(r"\.xlsx$", "", os.path.basename(xlsx_path))
    if out_path is None:
        out_path = os.path.join(os.path.dirname(os.path.abspath(xlsx_path)),
                                f"AI_Footing_{company_name}.xlsx")
    wb.save(out_path)
    return {"out_path": out_path, "foot_errors": total_foot,
            "cross_errors": total_cross,
            "sheets": ["총괄표", "원문"] + src_order}

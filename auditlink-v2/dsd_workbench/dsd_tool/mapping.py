"""D-3b: 계정과목 → XBRL element 매핑 추천 (완전 로컬).

고객사 계정명은 외부로 나가지 않는다 — dart_explorer가 만들어 둔
mapping_corpus.sqlite(공개 데이터 집계 파일)를 읽기 전용으로 조회할 뿐이다.
LLM/외부 AI 불사용: 실증 빈도 + 규칙 스코어 (재현성·감사조서 근거).

입력 템플릿 (고정 3열): 계정과목명 | 구분(BS/PL/CE/CF/주석) | 비고
출력: 계정과목명당 후보 Top 4 —
  element ID | 표준 한글라벨 | 스코어 | 근거 | 기준서 Reference | 선택☐
자동 확정 없음 — 선택 열은 회계사가 기입한다.

스코어 = 가중합:
- 실증 빈도: 해당 element 사용 회사 수 (동업종 ×2 가중)
- 문자열 유사도: 자모 분해 편집거리 + 토큰 자카드 (한글 계정명 특화)
- Role 정합: 구분(BS 등)과 element 소속 Role(D210000계열 등) 일치
최고 스코어가 임계 미만이면 "적합 표준 없음" + 유사 확장(extensions) 사례 제시.
dart-gcd_*(문서개황 메타)는 후보에서 제외.
"""
import collections
import json
import math
import os
import re
import sqlite3

# 기본 코퍼스 위치 (파일 기반 연결 — dart_explorer 산출물)
_REPO = os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__))))
DEFAULT_CORPUS = os.path.join(_REPO, "dart_explorer", "corpus",
                              "mapping_corpus.sqlite")
DEFAULT_REFS = os.path.join(_REPO, "dart_explorer", "corpus",
                            "standard_refs.json")

WEIGHTS = {"freq": 0.35, "sim": 0.45, "role": 0.20}
SIM_MIX = {"jamo": 0.7, "jaccard": 0.3}
DEFAULT_THRESHOLD = 0.45       # 스코어 임계 — 미만 → "적합 표준 없음"
SIM_FLOOR = 0.35               # 유사도 하한 — 빈도·Role만으로 후보 판정 방지

TEMPLATE_HEADER = ["계정과목명", "구분(BS/PL/CE/CF/주석)", "비고"]
ROLE_PREFIX = {
    "BS": ("D2",), "PL": ("D3", "D4"), "CE": ("D6",), "CF": ("D5",),
    "주석": ("D8",),
}


# ---------------------------------------------------------------------------
# 한글 유사도 (자모 분해 편집거리 + 토큰 자카드)
# ---------------------------------------------------------------------------

_CHO = "ㄱㄲㄴㄷㄸㄹㅁㅂㅃㅅㅆㅇㅈㅉㅊㅋㅌㅍㅎ"
_JUNG = "ㅏㅐㅑㅒㅓㅔㅕㅖㅗㅘㅙㅚㅛㅜㅝㅞㅟㅠㅡㅢㅣ"
_JONG = ("", "ㄱ", "ㄲ", "ㄳ", "ㄴ", "ㄵ", "ㄶ", "ㄷ", "ㄹ", "ㄺ", "ㄻ", "ㄼ",
         "ㄽ", "ㄾ", "ㄿ", "ㅀ", "ㅁ", "ㅂ", "ㅄ", "ㅅ", "ㅆ", "ㅇ", "ㅈ",
         "ㅊ", "ㅋ", "ㅌ", "ㅍ", "ㅎ")


def normalize(name):
    """계정명 정규화: 공백·괄호 표기·기호 제거."""
    s = str(name or "")
    s = re.sub(r"\[.*?\]", "", s)              # [개요], [구성요소] 등
    s = re.sub(r"\(.*?\)", "", s)              # (당기), (단위 : 원) 등
    s = re.sub(r"[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+\.?", "", s)
    s = re.sub(r"[^가-힣A-Za-z0-9]", "", s)
    return s.lower()


def jamo(s):
    out = []
    for ch in s:
        code = ord(ch)
        if 0xAC00 <= code <= 0xD7A3:
            idx = code - 0xAC00
            out.append(_CHO[idx // 588])
            out.append(_JUNG[(idx % 588) // 28])
            if idx % 28:
                out.append(_JONG[idx % 28])
        else:
            out.append(ch)
    return "".join(out)


def _levenshtein(a, b):
    if not a or not b:
        return max(len(a), len(b))
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[-1] + 1,
                           prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def jamo_similarity(a, b):
    ja, jb = jamo(a), jamo(b)
    if not ja and not jb:
        return 0.0
    return 1.0 - _levenshtein(ja, jb) / max(len(ja), len(jb))


def _tokens(name):
    return set(re.findall(r"[가-힣]{2,}|[A-Za-z]+|\d+", str(name or "")))


def token_jaccard(a, b):
    ta, tb = _tokens(a), _tokens(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)


# 재무 계정명 접두 처리:
# - 동의 변형: 단기↔유동, 장기↔비유동 (관행상 혼용) → 변형 후 매칭 시도
# - 반의 쌍: 자모 1~2자 차이로 유사도가 높게 나오는 반대 개념 → 페널티
_SYNONYM_SUBS = [("단기", "유동"), ("유동", "단기"),
                 ("장기", "비유동"), ("비유동", "장기")]
_ANTONYM_PAIRS = [("단기", "장기"), ("유동", "비유동"), ("당기", "전기"),
                  ("기초", "기말"), ("증가", "감소"), ("유입", "유출")]
_ANTONYM_PENALTY = 0.5


def _prefix_flags(s):
    flags = set()
    t = s
    if "비유동" in t:
        flags.add("비유동")
        t = t.replace("비유동", "")
    for w in ("유동", "단기", "장기", "당기", "전기", "기초", "기말",
              "증가", "감소", "유입", "유출"):
        if w in t:
            flags.add(w)
    return flags


def _antonym_penalty(qn, ln):
    fq, fl = _prefix_flags(qn), _prefix_flags(ln)
    for a, b in _ANTONYM_PAIRS:
        if (a in fq and b in fl and a not in fl) or \
                (b in fq and a in fl and b not in fl):
            return _ANTONYM_PENALTY
    return 1.0


def _query_variants(qn):
    vs = {qn}
    for a, b in _SYNONYM_SUBS:
        if a == "유동" and "비유동" in qn:
            continue
        if a in qn:
            vs.add(qn.replace(a, b, 1))
    return vs


def similarity(query_norm, query_raw, label):
    ln = normalize(label)
    if not ln:
        return 0.0
    best_jamo = max(jamo_similarity(v, ln)
                    for v in _query_variants(query_norm))
    sim = (SIM_MIX["jamo"] * best_jamo
           + SIM_MIX["jaccard"] * token_jaccard(query_raw, label))
    return sim * _antonym_penalty(query_norm, ln)


# ---------------------------------------------------------------------------
# 코퍼스 로드
# ---------------------------------------------------------------------------

class MappingCorpus:
    """mapping_corpus.sqlite 읽기 전용 스냅샷 (조회기 시작 시 메모리 적재)."""

    def __init__(self, db_path=None, refs_path=None):
        path = db_path or DEFAULT_CORPUS
        if not os.path.exists(path):
            raise FileNotFoundError(
                f"매핑 코퍼스 없음: {path} — dart_explorer에서 "
                "`corpus build`로 먼저 생성하세요.")
        con = sqlite3.connect(f"file:{path}?mode=ro", uri=True, timeout=30)
        self.induty = dict(con.execute(
            "SELECT corp_code, induty_code FROM companies "
            "WHERE status = 'ok'"))
        self.total_companies = len(self.induty)

        # element별: 사용 회사 목록, 라벨 변형, Role 코드
        self.elements = {}
        rows = con.execute(
            "SELECT element_id, corp_code, label_ko, roles FROM usages "
            "WHERE is_ext = 0 AND element_id NOT LIKE 'dart-gcd%'")
        for eid, corp, label, roles in rows:
            e = self.elements.setdefault(
                eid, {"corps": set(), "labels": set(), "roles": set()})
            e["corps"].add(corp)
            if label:
                e["labels"].add(label)
            for r in (roles or "").split(","):
                if r:
                    e["roles"].add(r)

        self.std_labels = dict(
            (eid, ko) for eid, ko, _ in
            con.execute("SELECT element_id, label_ko, label_en "
                        "FROM standard_labels"))

        # 확장 사례 (적합 표준 없음일 때 제시)
        self.extensions = collections.defaultdict(set)
        for eid, corp, label in con.execute(
                "SELECT element_id, corp_code, label_ko FROM extensions"):
            if label:
                self.extensions[label].add(corp)
        con.close()

        self.refs = {}
        refs_path = refs_path or DEFAULT_REFS
        if os.path.exists(refs_path):
            with open(refs_path, encoding="utf-8") as f:
                self.refs = json.load(f)


# ---------------------------------------------------------------------------
# 후보 스코어링
# ---------------------------------------------------------------------------

def _role_score(category, roles):
    cat = (category or "").strip()
    prefixes = ROLE_PREFIX.get(cat.upper()) or ROLE_PREFIX.get(cat)
    if not prefixes:
        return 0.5                              # 구분 미지정 → 중립
    return 1.0 if any(r.startswith(prefixes) for r in roles) else 0.0


def suggest(corpus: MappingCorpus, name, category=None, induty=None,
            top=4, threshold=DEFAULT_THRESHOLD):
    """계정과목명 1건 → 후보 Top N. dict 반환.

    {"name", "verdict": "후보"|"적합 표준 없음", "candidates": [...],
     "similar_extensions": [...]}
    """
    qn = normalize(name)
    denom = math.log1p(2 * max(1, corpus.total_companies))
    scored = []
    for eid, e in corpus.elements.items():
        labels = set(e["labels"])
        if eid in corpus.std_labels:
            labels.add(corpus.std_labels[eid])
        sim = max((similarity(qn, name, lb) for lb in labels), default=0.0)
        if sim <= 0.15:                         # 무관 후보 조기 제외
            continue
        n = len(e["corps"])
        n_same = sum(1 for c in e["corps"]
                     if induty and corpus.induty.get(c) == str(induty))
        freq = math.log1p(n + n_same) / denom   # 동업종 ×2 (한 번 더 가산)
        role = _role_score(category, e["roles"])
        score = (WEIGHTS["sim"] * sim + WEIGHTS["freq"] * freq
                 + WEIGHTS["role"] * role)
        best_label = max(labels, key=lambda lb: similarity(qn, name, lb))
        evidence = []
        if induty and n_same:
            evidence.append(f"동업종 {n_same}사 사용")
        evidence.append(f"전체 {n}사 사용")
        evidence.append(f"라벨 '{best_label}' 유사도 {sim:.2f}")
        if role == 1.0:
            evidence.append(f"Role 정합({','.join(sorted(e['roles'])[:3])})")
        scored.append({
            "element_id": eid,
            "std_label": corpus.std_labels.get(
                eid, next(iter(e["labels"]), "")),
            "score": round(score, 4), "sim": round(sim, 3),
            "n_companies": n, "n_same_induty": n_same,
            "evidence": ", ".join(evidence),
            "reference": "; ".join(corpus.refs.get(eid, [])[:3]),
        })
    scored.sort(key=lambda c: c["score"], reverse=True)
    cands = scored[:top]

    result = {"name": name, "category": category, "candidates": cands,
              "verdict": "후보", "similar_extensions": []}
    if not cands or cands[0]["score"] < threshold or \
            cands[0]["sim"] < SIM_FLOOR:
        result["verdict"] = "적합 표준 없음"
        exts = []
        for label, corps in corpus.extensions.items():
            sim = similarity(qn, name, label)
            if sim > 0.3:
                exts.append({"label": label, "n_companies": len(corps),
                             "sim": round(sim, 3)})
        exts.sort(key=lambda x: (x["sim"], x["n_companies"]), reverse=True)
        result["similar_extensions"] = exts[:3]
    return result


# ---------------------------------------------------------------------------
# 엑셀 입출력
# ---------------------------------------------------------------------------

def write_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "계정과목"
    ws.append(TEMPLATE_HEADER)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", start_color="D9E1F2")
    ws.append(["매출채권", "BS", "예시 행 — 지우고 사용"])
    for col, w in (("A", 34), ("B", 22), ("C", 30)):
        ws.column_dimensions[col].width = w
    wb.save(path)
    return path


def read_input(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True).active
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 and row and str(row[0] or "").strip() == "계정과목명":
            continue
        if not row or not str(row[0] or "").strip():
            continue
        rows.append({"name": str(row[0]).strip(),
                     "category": str(row[1] or "").strip() or None,
                     "note": str(row[2] or "").strip()
                     if len(row) > 2 and row[2] else ""})
    return rows


def map_accounts(input_path, out_path=None, db_path=None, induty=None,
                 top=4, threshold=DEFAULT_THRESHOLD):
    """입력 템플릿 → 후보 Top4 엑셀. 요약 dict 반환."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    corpus = MappingCorpus(db_path)
    items = read_input(input_path)
    if not items:
        raise ValueError("입력에 계정과목이 없습니다 (1열: 계정과목명).")

    wb = Workbook()
    ws = wb.active
    ws.title = "매핑후보"
    header = ["계정과목명", "구분", "판정", "후보#", "element ID",
              "표준 한글라벨", "스코어", "근거", "기준서 Reference", "선택☐"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", start_color="D9E1F2")

    results = []
    for item in items:
        res = suggest(corpus, item["name"], item["category"],
                      induty=induty, top=top, threshold=threshold)
        results.append(res)
        if res["verdict"] == "적합 표준 없음":
            ext_txt = "; ".join(
                f"확장사례 '{e['label']}' ({e['n_companies']}사)"
                for e in res["similar_extensions"]) or "유사 확장 사례 없음"
            ws.append([item["name"], item["category"], "적합 표준 없음",
                       "", "", "", "", ext_txt, "", ""])
            ws.cell(row=ws.max_row, column=3).font = Font(
                color="FFCC0000", bold=True)
        for rank, c in enumerate(res["candidates"], 1):
            if res["verdict"] == "적합 표준 없음" and rank > 2:
                break                            # 참고 후보 2개만
            ws.append([item["name"] if rank == 1 and
                       res["verdict"] == "후보" else "",
                       item["category"] if rank == 1 and
                       res["verdict"] == "후보" else "",
                       "" if res["verdict"] == "후보" else "(참고)",
                       rank, c["element_id"], c["std_label"], c["score"],
                       c["evidence"], c["reference"], "☐"])
    widths = [26, 8, 14, 5, 46, 28, 8, 46, 30, 6]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    if out_path is None:
        out_path = re.sub(r"\.xlsx$", "", input_path) + "_매핑후보.xlsx"
    wb.save(out_path)
    return {"out_path": out_path, "items": len(items),
            "no_match": sum(1 for r in results
                            if r["verdict"] == "적합 표준 없음"),
            "results": results}
